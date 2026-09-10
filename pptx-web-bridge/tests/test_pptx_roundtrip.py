"""PPTX → JSON → HTML / PPTX の往復試験。"""
import base64
import io

from PIL import Image
from pptx import Presentation

from app.pptx_generator import generate_pptx
from app.pptx_parser import parse_pptx
from app.quality_check import check_presentation
from app.validate import validate
from app.web_renderer import build_bundle, render_html


def test_parse_sample_deck(sample_pptx_bytes):
    p = parse_pptx(sample_pptx_bytes, "sample_deck.pptx")
    assert validate(p) == []
    assert len(p["slides"]) == 6
    assert p["canvas"]["aspect"] == "16:9"
    assert p["slides"][0]["layout"] == "title"
    assert p["slides"][0]["title"] == "業務改善提案（サンプル）"
    types = {el["type"] for s in p["slides"] for el in s["elements"]}
    assert {"text", "shape", "line", "table", "image"} <= types
    assert p["slides"][1]["notes"].startswith("発表者ノート")
    # 箇条書きの階層
    body = [el for el in p["slides"][1]["elements"] if el["role"] == "body"][0]
    assert body["paragraphs"][1]["level"] == 1 and body["paragraphs"][1]["bullet"] == "bullet"
    # 画像資産と縦横比
    assert len(p["assets"]) == 1
    asset = next(iter(p["assets"].values()))
    assert asset["width_px"] == 800 and asset["height_px"] == 450


def test_group_children_are_flattened_with_parent_coordinates(sample_pptx_bytes):
    p = parse_pptx(sample_pptx_bytes, "sample_deck.pptx")
    s5 = p["slides"][4]
    shapes = [el for el in s5["elements"] if el["type"] == "shape"]
    assert shapes, "グループ内の矩形が平坦化されていること"
    b = shapes[0]["bbox"]
    assert 0 <= b["x"] <= p["canvas"]["width_pt"] and 0 <= b["y"] <= p["canvas"]["height_pt"]


def test_render_html_contains_all_slides(sample_pptx_bytes):
    p = parse_pptx(sample_pptx_bytes, "sample_deck.pptx")
    html = render_html(p, inline_assets=True, inline_viewer=True)
    for s in p["slides"]:
        assert f'id="{s["id"]}"' in html
    assert "data:image/png;base64" in html
    assert "<table>" in html
    assert "<ul>" in html
    assert "<script" in html and "viewer-header" in html


def test_bundle_files(sample_pptx_bytes):
    p = parse_pptx(sample_pptx_bytes, "sample_deck.pptx")
    files = build_bundle(p)
    assert {"index.html", "viewer.css", "viewer.js", "presentation.json"} <= set(files)
    assert any(k.startswith("assets/") for k in files)
    assert b"data_base64" not in files["presentation.json"]


def test_pptx_roundtrip_keeps_structure(sample_pptx_bytes):
    p = parse_pptx(sample_pptx_bytes, "sample_deck.pptx")
    data, warns = generate_pptx(p, "editable")
    assert warns == []
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == 6
    assert abs(prs.slide_width / prs.slide_height - 16 / 9) < 0.01
    again = parse_pptx(data, "rt.pptx")
    assert [s["title"] for s in again["slides"]] == [s["title"] for s in p["slides"]]
    assert again["slides"][3]["layout"] == "table"
    tbl = [el for el in again["slides"][3]["elements"] if el["type"] == "table"][0]
    assert tbl["rows"][0][0]["text"] == "区分"
    assert again["slides"][1]["notes"].startswith("発表者ノート")
    # 画像の縦横比が保たれる
    img = [el for el in again["slides"][4]["elements"] if el["type"] == "image"][0]
    b = img["bbox"]
    assert abs(b["w"] / b["h"] - 800 / 450) < 0.02
    assert not [i for i in check_presentation(again) if i["severity"] == "error"]


def test_generated_pptx_has_no_template_chrome_after_reread(sample_pptx_bytes):
    """部品付きテンプレート（表紙背景・ロゴ・帯・ページ番号）は図形として出力され、再読込では本文に混ざらない。"""
    from app.pipeline import prepare

    p = parse_pptx(sample_pptx_bytes, "sample_deck.pptx", template_id="corporate_standard")
    p, _e, _f = prepare(p)
    data, _ = generate_pptx(p, "editable")
    prs = Presentation(io.BytesIO(data))
    cover_names = {sh.name for sh in prs.slides[0].shapes}
    content_names = {sh.name for sh in prs.slides[1].shapes}
    assert "cover_background" in cover_names and "logo" in cover_names and "footer" in cover_names
    assert {"logo", "bar", "page_number"} <= content_names
    # 表紙の題名はテンプレート定義の位置・サイズ・色（白 32pt）のテキストシェイプになっている
    title_shape = next(sh for sh in prs.slides[0].shapes if sh.name.startswith("title:"))
    run = title_shape.text_frame.paragraphs[0].runs[0]
    assert run.font.size.pt == 32 and str(run.font.color.rgb) == "FFFFFF"
    again = parse_pptx(data, "rt.pptx")
    assert all(el.get("role") != "footer" for el in again["slides"][0]["elements"])
    assert [el["type"] for el in again["slides"][1]["elements"]] == ["text", "text"]


def test_closing_slide_layout():
    from app.template_kit import make_closing_slide
    from app.model import new_presentation
    from app.pipeline import prepare

    p = new_presentation("t", template_id="corporate_standard")
    p["slides"].append(make_closing_slide("s001", 0, "ご清聴ありがとうございました（サンプル）"))
    p, errors, _ = prepare(p)
    assert errors == []
    assert p["slides"][0]["elements"][0]["bbox"]["y"] > p["canvas"]["height_pt"] * 0.5
    data, warns = generate_pptx(p, "editable")
    names = {sh.name for sh in Presentation(io.BytesIO(data)).slides[0].shapes}
    assert "logo" in names and warns == []


def test_unsupported_element_becomes_box_not_error():
    from app.model import bbox, new_presentation, new_slide, unsupported_element

    p = new_presentation("t")
    s = new_slide("s001", 0)
    s["elements"].append(unsupported_element("e1", "chart", bbox(10, 10, 200, 100), "グラフ（未対応）"))
    p["slides"].append(s)
    data, warns = generate_pptx(p, "editable")
    assert any(w["code"] == "UNSUPPORTED_AS_BOX" for w in warns)
    assert len(Presentation(io.BytesIO(data)).slides) == 1
