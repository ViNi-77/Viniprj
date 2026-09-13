"""レイアウト方式（Phase N）: スライドマスター / レイアウトを直接読む。

推測ではなく、テンプレートが明示的に持っている設計を読めていることを確かめる。
資料は `samples/make_template_matrix.py` が実行時に生成する（リポジトリには置かない）。
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

import pytest
from pptx import Presentation

from app import pptx_parser, template_from_layout as tfl

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture(scope="module")
def matrix(tmp_path_factory) -> dict[str, bytes]:
    sys.path.insert(0, str(ROOT / "samples"))
    from make_template_matrix import build  # type: ignore

    return {p.stem: p.read_bytes() for p in build(tmp_path_factory.mktemp("layouts"))}


def _chrome(pres: dict) -> list[dict]:
    """装飾（帯・ロゴ）にあたる要素。プレースホルダ由来の文字は除く。"""
    return [e for e in pres["slides"][0]["elements"] if e["type"] in ("shape", "image")]


def _first_with_chrome(info: dict) -> dict:
    return next(l for l in info["layouts"] if l["chrome_count"] + l["chrome_from_master"] > 0)


def test_chrome_in_layout_is_found_and_rendered(matrix):
    """装飾をレイアウトに置いたテンプレート: 帯とロゴがレイアウトから出る。"""
    data = matrix["chrome_on_layout"]
    info = tfl.enumerate_layouts(data, "chrome_on_layout.pptx")
    assert info["available"] is True and not info["warnings"]
    target = _first_with_chrome(info)
    assert target["chrome_count"] == 2

    pres = tfl.parse_layout(Presentation(io.BytesIO(data)), target["master"], target["index"])
    chrome = _chrome(pres)
    assert len(chrome) == 2
    bar = next(e for e in chrome if e["type"] == "shape")
    assert bar["fill"] == "#003087" and abs(bar["bbox"]["y"] - 490) < 2
    assert any(e["type"] == "image" for e in chrome), "ロゴが出ていない"


def test_chrome_on_master_is_inherited(matrix):
    """装飾がマスターにある場合も、レイアウトを見れば出る（showMasterSp を尊重）。"""
    data = matrix["chrome_on_master"]
    info = tfl.enumerate_layouts(data, "chrome_on_master.pptx")
    assert info["available"] is True
    assert info["masters"][0]["chrome_count"] == 2
    target = _first_with_chrome(info)
    assert target["chrome_count"] == 0 and target["chrome_from_master"] == 2, "マスターの分は別勘定にする"

    chrome = _chrome(tfl.parse_layout(Presentation(io.BytesIO(data)), target["master"], target["index"]))
    assert len(chrome) == 2, "マスターの装飾がレイアウトの描画に出ていない"


def test_chrome_on_slide_falls_back_with_a_reason(matrix):
    """装飾がスライド側にあるファイルでは使えない。理由を必ず出す（黙って 0 個にしない）。"""
    info = tfl.enumerate_layouts(matrix["chrome_on_slide_wide_logo"], "chrome_on_slide_wide_logo.pptx")
    assert info["available"] is False
    w = next(x for x in info["warnings"] if x["code"] == "LAYOUT_HAS_NO_CHROME")
    assert "レイアウト" in w["message"] and "推測" in w["message"]
    assert all(l["chrome_count"] == 0 for l in info["layouts"])


def test_placeholders_are_kept_as_labelled_boxes(matrix):
    """レイアウトのフッター・日付・ページ番号は「どこに何が入るか」なので捨てない。"""
    data = matrix["chrome_on_layout"]
    info = tfl.enumerate_layouts(data, "chrome_on_layout.pptx")
    target = _first_with_chrome(info)
    pres = tfl.parse_layout(Presentation(io.BytesIO(data)), target["master"], target["index"])
    texts = ["".join(r.get("text", "") for p in e.get("paragraphs", []) for r in p.get("runs", []))
             for e in pres["slides"][0]["elements"] if e["type"] == "text"]
    assert any("フッター" in t for t in texts), "空のフッタープレースホルダが落ちている"
    assert len(texts) >= 4, f"題名・本文・フッター類が揃っていない: {texts}"


def test_region_boxes_come_from_real_placeholders(matrix):
    """本文の流し込み先は推測せず、本物の本文プレースホルダの枠を使う。"""
    data = matrix["chrome_on_layout"]
    info = tfl.enumerate_layouts(data, "chrome_on_layout.pptx")
    target = _first_with_chrome(info)
    boxes = tfl.region_boxes(Presentation(io.BytesIO(data)), target["master"], target["index"])
    assert "title" in boxes and "body" in boxes
    assert boxes["body"]["y"] > boxes["title"]["y"], "本文が題名より上にある"
    assert all(v > 0 for v in boxes["body"].values())


def test_all_masters_are_walked(matrix):
    """`prs.slide_layouts` はマスター 0 だけなので使わない。全マスターを走査する。"""
    data = matrix["three_slides"]
    prs = Presentation(io.BytesIO(data))
    info = tfl.enumerate_layouts(data, "three_slides.pptx")
    assert len(info["layouts"]) == sum(len(m.slide_layouts) for m in prs.slide_masters)


def test_layout_mode_does_not_change_normal_parsing(matrix):
    """通常のスライド解析の挙動は変えない（フッター類は従来どおり捨てる）。"""
    from app.pptx_parser import parse_pptx

    pres = parse_pptx(matrix["three_slides"], "three_slides.pptx")
    texts = ["".join(r.get("text", "") for p in e.get("paragraphs", []) for r in p.get("runs", []))
             for s in pres["slides"] for e in s["elements"] if e["type"] == "text"]
    assert not any("フッター" == t for t in texts), "レイアウトモードの印が通常解析に漏れている"
    assert pptx_parser._layout_mode.get() is False


def test_bad_index_raises_a_readable_error(matrix):
    prs = Presentation(io.BytesIO(matrix["three_slides"]))
    with pytest.raises(ValueError, match="レイアウト"):
        tfl.parse_layout(prs, 0, 999)
    with pytest.raises(ValueError, match="マスター"):
        tfl.parse_layout(prs, 9, 0)


# ---------------------------------------------------------------- 保存・描画・出力
@pytest.fixture(scope="module")
def layout_template(matrix) -> dict:
    data = matrix["chrome_on_layout"]
    info = tfl.enumerate_layouts(data, "chrome_on_layout.pptx")
    hit = _first_with_chrome(info)
    ref = {"master": hit["master"], "index": hit["index"]}
    return tfl.build_template(data, "chrome_on_layout.pptx", {"cover": ref, "content": ref}, template_id="t_phase_n")["proposal"]


def test_build_template_externalizes_images(layout_template):
    """画像をテンプレートの資産フォルダへ出す。飛ばすと装飾が画面から消える。"""
    from app.config import resource_path

    part = layout_template["content"]
    assert layout_template["mode"] == "layout"
    assert layout_template["layout_map"]["content"]["name"]
    refs = part["chrome_assets"]
    assert refs, "画像が外部化されていない"
    for aid, ref in refs.items():
        assert Path(resource_path(ref["path"])).exists(), f"{aid} のファイルが無い"
        assert any(e.get("asset_id") == aid for e in part["chrome_elements"])


def test_saved_template_passes_validation(layout_template):
    from app import template_store

    assert template_store.validate_template(layout_template) == []
    bad = {**layout_template, "layout_map": {"content": {"master": 0}}}
    assert any("layout_map" in e for e in template_store.validate_template(bad))
    assert any("mode" in e for e in template_store.validate_template({**layout_template, "mode": "nonsense"}))


def test_chrome_is_drawn_behind_the_content(layout_template):
    """装飾は資料の要素に混ぜず、背面に敷くだけ（選択・削除・保存の対象にしない）。"""
    from app import pipeline, template_kit
    from app.template_from_pptx import sample_presentation
    from app.web_renderer import slide_html

    with template_kit.use_template(layout_template):
        pres, _e, _f = pipeline.prepare(sample_presentation(layout_template))
        spec = template_kit.chrome_spec(pres["slides"][1], pres)
        assert len(spec["elements"]) == 2 and spec["element_assets"]
        assert not spec["bars"] and not spec["images"], "レイアウト方式では推測の部品を作らない"
        html = slide_html(pres["slides"][1], pres, inline_assets=True, template=layout_template, with_notes=False)
    assert html.count('class="tpl tpl-el"') == 2
    assert "data:image/png;base64," in html, "ロゴが画面に出ていない"
    assert not any(e.get("asset_id", "").startswith("img") for s in pres["slides"] for e in s["elements"]), \
        "装飾が資料の要素に混ざっている"


def test_export_uses_the_mapped_layout_without_double_drawing(layout_template):
    """土台 PPTX ありでは保存したレイアウトを使い、装飾は PowerPoint に継承させる（二重に描かない）。"""
    from pptx import Presentation as _P

    from app import pipeline, template_kit
    from app.template_from_pptx import sample_presentation

    with template_kit.use_template(layout_template):
        data, _pres, warns = pipeline.export_pptx(sample_presentation(layout_template), "editable", use_base_pptx=None)
    slide = _P(io.BytesIO(data)).slides[1]
    assert slide.slide_layout.name == layout_template["layout_map"]["content"]["name"]
    assert len([s for s in slide.slide_layout.shapes if not s.is_placeholder]) == 2, "レイアウト側に装飾が無い"
    assert not [s for s in slide.shapes if s.name.startswith(("bar", "logo"))], "装飾が二重に描かれている"
    assert not [w for w in warns if w["code"] == "LAYOUT_MAP_OUT_OF_RANGE"]


def test_export_without_base_draws_the_chrome(layout_template):
    """土台 PPTX を使わないときは、装飾を図形として置く（白紙にしない）。"""
    from pptx import Presentation as _P

    from app import pipeline, template_kit
    from app.template_from_pptx import sample_presentation

    t = {k: v for k, v in layout_template.items() if k != "base_pptx"}
    with template_kit.use_template(t):
        data, _pres, _w = pipeline.export_pptx(sample_presentation(t), "editable", use_base_pptx=False)
    shapes = _P(io.BytesIO(data)).slides[1].shapes
    assert len([s for s in shapes if s.shape_type is not None]) >= 4, "装飾が置かれていない"


def test_out_of_range_layout_map_warns_and_falls_back(layout_template):
    """土台 PowerPoint を差し替えるとレイアウト番号がずれる。落ちずに警告して自動選択へ戻す。"""
    from app import pipeline, template_kit
    from app.template_from_pptx import sample_presentation

    t = {**layout_template, "layout_map": {**layout_template["layout_map"], "content": {"master": 0, "index": 99}}}
    with template_kit.use_template(t):
        _data, _pres, warns = pipeline.export_pptx(sample_presentation(t), "editable", use_base_pptx=None)
    assert any(w["code"] == "LAYOUT_MAP_OUT_OF_RANGE" for w in warns)


def test_unrenderable_image_shows_a_labelled_box():
    """EMF / WMF はレイアウト方式でも「画面では表示できない形式」の枠を出す（白くしない）。"""
    from app.web_renderer import element_html

    pres = {"assets": {"a1": {"mime": "image/emf", "filename": "a1.emf", "data_base64": "", "unrenderable": True}}, "canvas": {"width_pt": 960, "height_pt": 540}}
    el = {"id": "e1", "type": "image", "asset_id": "a1", "bbox": {"x": 0, "y": 0, "w": 100, "h": 40}, "alt": "ロゴ"}
    html = element_html(el, pres, inline_assets=True)
    assert "unrenderable" in html and "画面では表示できない形式" in html


# ---------------------------------------------------------------- API
@pytest.fixture(scope="module")
def client():
    from fastapi.testclient import TestClient

    from app.main import app

    return TestClient(app)


def _upload(client, matrix, name: str, tid: str) -> dict:
    r = client.post("/api/templates/from-pptx", files={"file": (f"{name}.pptx", matrix[name])}, data={"template_id": tid})
    assert r.status_code == 200, r.text
    return r.json()


def test_api_reports_layout_availability_in_one_round_trip(client, matrix):
    """1 往復で「レイアウト方式が使えるか」まで返す（画面が両方式を出せるように）。"""
    d = _upload(client, matrix, "chrome_on_layout", "api_l1")
    layouts = d["layouts"]
    assert layouts["available"] is True and len(layouts["layouts"]) == 11
    assert any(l["chrome_count"] > 0 for l in layouts["layouts"])
    assert d["proposal"]["id"] == "api_l1"  # 推測方式の提案も同時に返る


def test_api_reports_the_reason_when_layouts_are_empty(client, matrix):
    d = _upload(client, matrix, "chrome_on_slide_wide_logo", "api_l2")
    layouts = d["layouts"]
    assert layouts["available"] is False
    assert any("レイアウト" in str(w.get("message", w)) for w in layouts["warnings"])
    assert all(l["chrome_count"] == 0 for l in layouts["layouts"]), "図形数を表に出せる形で返す"


def test_api_layout_preview_and_build(client, matrix):
    d = _upload(client, matrix, "chrome_on_layout", "api_l3")
    hit = next(l for l in d["layouts"]["layouts"] if l["chrome_count"] + l["chrome_from_master"] > 0)
    ref = {"master": hit["master"], "index": hit["index"]}

    # 選択 UI のサムネイル: レイアウトそのものを描くので、装飾は通常の要素として出る
    r = client.post("/api/templates/layout-preview", json={"template_id": "api_l3", **ref})
    assert r.status_code == 200
    html = r.json()["html"]
    assert "#003087" in html.upper() or "003087" in html.lower(), "帯が描かれていない"
    assert "el-image" in html, "ロゴが描かれていない"

    r2 = client.post("/api/templates/from-layout", json={"template_id": "api_l3", "filename": "chrome_on_layout.pptx", "layout_map": {"content": ref}})
    assert r2.status_code == 200, r2.text
    built = r2.json()
    assert built["proposal"]["mode"] == "layout"
    assert built["proposal"]["layout_map"]["content"]["name"] == hit["name"]
    assert "content" in built["previews"]


def test_api_rejects_unknown_source_and_empty_map(client, matrix):
    assert client.post("/api/templates/layout-preview", json={"template_id": "never_uploaded", "master": 0, "index": 0}).status_code == 409
    _upload(client, matrix, "chrome_on_layout", "api_l4")
    assert client.post("/api/templates/from-layout", json={"template_id": "api_l4", "layout_map": {}}).status_code == 422
    assert client.post("/api/templates/layout-preview", json={"template_id": "api_l4", "master": 0, "index": 999}).status_code == 422
