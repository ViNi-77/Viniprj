"""レイアウト品質（Phase A）: 文字サイズの一本化、画像の実寸配置と横並び、分割の抑制、HTML/PPTX 取込の修正。"""
from __future__ import annotations

import io

from PIL import Image

from app import typography
from app.html_parser import parse_html
from app.layout import layout_presentation, layout_slide
from app.model import bbox, image_element, new_presentation, new_slide, paragraph, run, text_element
from app.pptx_generator import generate_pptx
from app.pptx_parser import parse_pptx
from app.quality_check import _overlap_area, check_presentation
from app.report import element_rows
from app.validate import validate_and_repair
from app.web_renderer import render_html


def _png(w: int, h: int) -> bytes:
    im = Image.new("RGB", (w, h), (30, 80, 160))
    buf = io.BytesIO()
    im.save(buf, format="PNG")
    return buf.getvalue()


def _pres_with(elements: list[dict], assets: dict | None = None, layout: str = "title_body") -> dict:
    p = new_presentation("t")
    p["assets"] = assets or {}
    s = new_slide("s001", 0, layout=layout)
    s["elements"] = elements
    p["slides"] = [s]
    return p


def _asset(w: int, h: int) -> dict:
    import base64

    return {"mime": "image/png", "filename": "a.png", "data_base64": base64.b64encode(_png(w, h)).decode("ascii"), "width_px": w, "height_px": h}


# ---------------------------------------------------------------- A1 文字サイズ
def test_band_clamps_inherited_but_keeps_explicit():
    el = text_element("e1", [paragraph([run("明示", size_pt=32.0), run("継承", size_pt=6.0, inherited=["size_pt"])]), paragraph([run("大きい継承", size_pt=60.0, inherited=["size_pt"])])], role="body")
    el["paragraphs"][0]["runs"][1]["inherited"] = ["size_pt"]
    el["paragraphs"][1]["runs"][0]["inherited"] = ["size_pt"]
    typography.normalize_element(el)
    runs = [r for p in el["paragraphs"] for r in p["runs"]]
    assert runs[0]["size_pt"] == 32.0  # 明示値はそのまま
    assert runs[1]["size_pt"] == 12.0 and runs[2]["size_pt"] == 20.0  # 継承値は本文の帯域 12〜20 へ
    assert el["font_pt"] == 32.0


def test_effective_size_uses_element_font_and_scale():
    el = text_element("e1", [paragraph([run("a")])], role="title")
    typography.normalize_element(el)
    assert el["font_pt"] == 28.0
    el["font_scale"] = 0.9
    assert abs(typography.effective_size({"text": "a"}, el) - 25.2) < 1e-6
    assert abs(typography.effective_size({"text": "a", "size_pt": 40}, el) - 36.0) < 1e-6


def test_renderer_report_and_pptx_agree_on_size():
    el = text_element("e1", [paragraph([run("見出し")])], role="title", box=bbox(36, 36, 888, 64))
    p = _pres_with([el])
    p["slides"][0]["elements"][0]["font_scale"] = 0.9
    typography.normalize_presentation(p)
    html = render_html(p, inline_assets=True, inline_viewer=True)
    assert "font-size:25.2px" in html
    rows = element_rows(p)
    assert rows[0]["font_pt_max"] == 25.2
    data, _w = generate_pptx(p, "editable")
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    shp = next(s for s in prs.slides[0].shapes if s.has_text_frame and s.text_frame.text == "見出し")
    assert abs(shp.text_frame.paragraphs[0].runs[0].font.size.pt - 25.2) < 0.01  # 縮小後の実サイズを書く


# ---------------------------------------------------------------- A2 画像
def test_image_keeps_natural_size_and_pairs_with_text():
    html = b"<html><body><section><h2>t</h2><p>" + "本文です。".encode("utf-8") * 20 + b"</p><img src='s.png' alt='x'></section></body></html>"
    p = parse_html(html, "t.html", {"s.png": _png(400, 300)}, computed_style=False)
    els = p["slides"][0]["elements"]
    img = next(e for e in els if e["type"] == "image")
    txt = next(e for e in els if e["type"] == "text" and e.get("role") == "body")
    assert img["layout_hint"]["row"] == txt["layout_hint"]["row"] and img["layout_hint"]["side"] == "right" and txt["layout_hint"]["side"] == "left"
    lp = layout_presentation(p)
    els = lp["slides"][0]["elements"]
    img = next(e for e in els if e["type"] == "image")
    txt = next(e for e in els if e["type"] == "text" and e.get("role") == "body")
    assert img["bbox"]["y"] == txt["bbox"]["y"]  # 同じ行
    assert abs(img["bbox"]["w"] - 300) < 0.01 and abs(img["bbox"]["h"] - 225) < 0.01  # 400px×0.75、縦横比 3:4 を保つ
    assert img["bbox"]["x"] > txt["bbox"]["x"] + txt["bbox"]["w"]  # 右側
    assert img["bbox"]["h"] >= 120


def test_image_goes_to_next_page_instead_of_shrinking():
    long_text = text_element("t1", [paragraph([run("あ" * 700)])], role="body")  # 残り高さが最小画像高さ 120pt を下回る量
    img = image_element("i1", "a1", None)
    p = _pres_with([text_element("h", [paragraph([run("題名")])], role="title"), long_text, img], {"a1": _asset(1600, 900)})
    lp = layout_presentation(p)
    placed = [(s["id"], e["id"], e["bbox"]) for s in lp["slides"] for e in s["elements"] if e["id"] == "i1"]
    assert placed and placed[0][0].endswith("_p2")  # 縮めず次ページへ
    assert placed[0][2]["h"] >= 120


def test_wide_image_capped_by_height_ratio_and_centered():
    img = image_element("i1", "a1", None)
    p = _pres_with([text_element("h", [paragraph([run("題名")])], role="title"), img], {"a1": _asset(3000, 2000)})
    lp = layout_presentation(p)
    b = next(e for e in lp["slides"][0]["elements"] if e["id"] == "i1")["bbox"]
    assert abs(b["h"] - 540 * 0.45) < 0.01
    assert abs((b["x"] + b["w"] / 2) - 480) < 0.01  # 左右中央


# ---------------------------------------------------------------- A3 分割の抑制
def test_autofit_before_split_sets_font_scale():
    paras = [paragraph([run("い" * 90)]) for _ in range(9)]  # 1 枚をわずかに超える量
    p = _pres_with([text_element("h", [paragraph([run("題名")])], role="title"), text_element("b", paras, role="body")])
    lp = layout_presentation(p)
    assert len(lp["slides"]) == 1
    body = next(e for e in lp["slides"][0]["elements"] if e["id"] == "b")
    assert 0.85 <= body["font_scale"] < 1.0
    assert not any(w["code"] == "SLIDE_SPLIT" for w in lp["warnings"])


def test_split_marks_continuation():
    paras = [paragraph([run("う" * 90)]) for _ in range(30)]
    p = _pres_with([text_element("h", [paragraph([run("題名")])], role="title"), text_element("b", paras, role="body")])
    lp = layout_presentation(p)
    assert len(lp["slides"]) >= 2
    s2 = lp["slides"][1]
    assert s2["continuation_of"] == "s001" and s2["continuation_index"] == 1
    assert any(w["code"] == "SLIDE_SPLIT" for w in s2["warnings"])
    title = next(e for e in s2["elements"] if e.get("role") == "title")
    assert title["paragraphs"][0]["runs"][-1]["text"].endswith("（続き）")


def test_small_remaining_space_pushes_element_instead_of_splitting():
    filler = text_element("f", [paragraph([run("え" * 90)]) for _ in range(6)], role="body")
    tail = text_element("t", [paragraph([run("お" * 90)]) for _ in range(3)], role="body")
    p = _pres_with([text_element("h", [paragraph([run("題名")])], role="title"), filler, tail])
    lp = layout_presentation(p)
    ids = [[e["id"] for e in s["elements"]] for s in lp["slides"]]
    assert "t" not in ids[0] and "t_cont" not in ids[0]  # 残りが少ないので分割せず次ページへ
    assert any("t" in page for page in ids[1:])


# ---------------------------------------------------------------- A4 HTML 取込
def test_hero_background_becomes_slide_background_with_text_color():
    html = b"<html><body><section style='background:#0B3D91;color:#fff'><h1>Hero</h1><p>sub</p></section><section><h2>b</h2><p>x</p></section></body></html>"
    p = parse_html(html, "t.html", {}, computed_style=False)
    assert p["slides"][0]["background"] == {"color": "#0B3D91", "text_color": "#FFFFFF"}
    assert p["slides"][1].get("background") is None
    html_out = render_html(layout_presentation(p), inline_assets=True, inline_viewer=True)
    assert "background:#0B3D91;color:#FFFFFF" in html_out


def test_hr_becomes_line_and_band_has_small_height():
    html = "<html><body><section><h2>t</h2><div class='card' style='background:#E60A3C'><p>短い帯</p></div><hr><p>x</p></section></body></html>".encode("utf-8")
    p = parse_html(html, "t.html", {}, computed_style=False)
    els = p["slides"][0]["elements"]
    band = next(e for e in els if e["type"] == "shape")
    assert band["layout_hint"].get("band") is True and band["fill"] == "#E60A3C"
    assert any(e["type"] == "line" for e in els)
    lp = layout_presentation(p)
    els = lp["slides"][0]["elements"]
    band = next(e for e in els if e["type"] == "shape")
    line = next(e for e in els if e["type"] == "line")
    assert band["bbox"]["h"] < 48 and band["bbox"]["h"] >= 14 * 1.6
    assert line["bbox"]["h"] <= 1.0 and line["points"][0][1] == line["points"][1][1]


def test_unresolved_image_is_placeholder_not_error():
    html = b"<html><body><section><h2>t</h2><img src='https://example.invalid/a.png' alt='ext'></section></body></html>"
    p = parse_html(html, "t.html", {}, computed_style=False)
    lp = layout_presentation(p)
    img = next(e for e in lp["slides"][0]["elements"] if e["type"] == "image")
    assert img["placeholder"] is True and img["asset_id"] is None
    assert abs(img["bbox"]["w"] - 888 * 0.4) < 0.01
    issues = check_presentation(lp)
    assert not any(i["severity"] == "error" for i in issues)
    assert any(i["code"] == "IMAGE_PLACEHOLDER" for i in issues)
    html_out = render_html(lp, inline_assets=True, inline_viewer=True)
    assert "el-image placeholder" in html_out and 'class="el el-unsupported"' not in html_out
    data, warns = generate_pptx(lp, "editable")
    assert data[:2] == b"PK" and not any(w["code"] == "ASSET_MISSING" for w in warns)


# ---------------------------------------------------------------- A5 PPTX 取込
def test_flow_elements_avoid_fixed_ones():
    fixed_title = text_element("ft", [paragraph([run("固定題名")])], role="title", box=bbox(36, 36, 888, 64))
    fixed_logo = image_element("fl", "a1", bbox(700, 420, 200, 100))
    body = text_element("b", [paragraph([run("本文" * 20)])], role="body")
    p = _pres_with([fixed_title, fixed_logo, body], {"a1": _asset(200, 100)})
    lp = layout_presentation(p)
    els = lp["slides"][0]["elements"]
    b = next(e for e in els if e["id"] == "b")["bbox"]
    for fid in ("ft", "fl"):
        assert _overlap_area(b, next(e for e in els if e["id"] == fid)["bbox"]) == 0
    assert b["y"] >= 36 + 64


def test_continuation_keeps_fixed_title():
    fixed_title = text_element("ft", [paragraph([run("固定題名")])], role="title", box=bbox(36, 36, 888, 64))
    body = text_element("b", [paragraph([run("か" * 90)]) for _ in range(30)], role="body")
    lp = layout_presentation(_pres_with([fixed_title, body]))
    assert len(lp["slides"]) >= 2
    t2 = next(e for e in lp["slides"][1]["elements"] if e.get("role") == "title")
    assert t2["paragraphs"][0]["runs"][-1]["text"].endswith("（続き）") and t2["bbox"]["y"] == 36


def test_pptx_crop_is_baked_and_rotation_kept(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    png = tmp_path / "p.png"
    png.write_bytes(_png(400, 200))
    pic = s.shapes.add_picture(str(png), Inches(1), Inches(1), Inches(4), Inches(2))
    pic.crop_left = 0.25
    pic.crop_right = 0.25
    tb = s.shapes.add_textbox(Inches(1), Inches(4), Inches(3), Inches(1))
    tb.text_frame.text = "回転"
    tb.rotation = 15.0
    buf = io.BytesIO()
    prs.save(buf)
    p = parse_pptx(buf.getvalue(), "r.pptx")
    els = p["slides"][0]["elements"]
    img = next(e for e in els if e["type"] == "image")
    asset = p["assets"][img["asset_id"]]
    assert asset["width_px"] == 200 and asset["height_px"] == 200  # 左右 25% ずつ切った実画像
    assert img["crop"]["left"] == 0.25
    assert any(w["code"] == "IMAGE_CROP_BAKED" for w in p["warnings"])
    txt = next(e for e in els if e["type"] == "text")
    assert txt["rotation_deg"] == 15.0
    html = render_html(p, inline_assets=True, inline_viewer=True)
    assert "transform:rotate(15deg)" in html
    data, _w = generate_pptx(p, "editable")
    out = Presentation(io.BytesIO(data))
    rotated = [sh for sh in out.slides[0].shapes if abs(sh.rotation - 15.0) < 0.01]
    assert rotated


# ---------------------------------------------------------------- A6 スキーマ
def test_schema_migration_and_new_fields_validate():
    p = new_presentation("t")
    p["schema_version"] = "1.0"
    s = new_slide("s001", 0)
    s["continuation_of"] = "missing"
    el = text_element("e1", [paragraph([run("a")])], box=bbox(0, 0, 100, 20), font_pt=16.0, font_scale=1.7, rotation_deg=10.0)
    s["elements"] = [el]
    p["slides"] = [s]
    fixed, errors, fixes = validate_and_repair(p)
    assert not errors
    codes = {f["code"] for f in fixes}
    assert {"SCHEMA_MIGRATED", "FONT_SCALE_FIXED", "CONTINUATION_FIXED"} <= codes or ("SCHEMA_MIGRATED" in codes and "CONTINUATION_FIXED" in codes and fixed["slides"][0]["elements"][0]["font_scale"] == 1.0)
    assert fixed["schema_version"] == "1.1" and "continuation_of" not in fixed["slides"][0]


def test_layout_slide_keeps_columns_images_readable():
    assets = {"a1": _asset(800, 600)}
    els = [text_element("h", [paragraph([run("題名")])], role="title")]
    for c in range(3):
        els.append(text_element(f"c{c}", [paragraph([run("カード" * 10)])], role="card", layout_hint={"column": c, "columns": 3}))
        els.append(image_element(f"i{c}", "a1", None, layout_hint={"column": c, "columns": 3}))
    p = _pres_with(els, assets)
    slides = layout_slide(p["slides"][0], p["canvas"], p["assets"])
    imgs = [e for s in slides for e in s["elements"] if e["type"] == "image"]
    assert len(imgs) == 3 and all(abs(i["bbox"]["w"] / i["bbox"]["h"] - 800 / 600) < 0.02 for i in imgs)
    assert all(i["bbox"]["h"] >= 120 for i in imgs)
