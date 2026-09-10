"""模擬 12 枚 PPTX の往復、テンプレート部品、要素判別レポート、API 追加分の試験。"""
import io
import json
from pathlib import Path

from fastapi.testclient import TestClient
from pptx import Presentation

from app import pipeline
from app.main import app
from app.pptx_generator import generate_pptx
from app.pptx_parser import parse_pptx
from app.quality_check import check_presentation
from app.report import build_report, to_markdown

ROOT = Path(__file__).resolve().parents[1]
MOCK = ROOT / "mock-pptx" / "output" / "mock_bidirectional_conversion_12slides.pptx"
client = TestClient(app)


def _mock_bytes() -> bytes:
    if not MOCK.exists():
        import subprocess
        import sys

        subprocess.run([sys.executable, str(ROOT / "mock-pptx" / "src" / "generate_mock_pptx.py")], check=True)
    return MOCK.read_bytes()


def test_mock_deck_parses_without_warnings():
    p = parse_pptx(_mock_bytes(), "mock.pptx")
    assert len(p["slides"]) == 12
    assert p["warnings"] == []
    manifest = json.loads((MOCK.parent / "mock_bidirectional_conversion_manifest.json").read_text(encoding="utf-8"))
    assert [s["title"] for s in p["slides"]] == [m["title"] for m in manifest["slides"]]
    assert not check_presentation(p)
    types = {el["type"] for s in p["slides"] for el in s["elements"]}
    assert {"text", "shape", "line", "table", "image"} <= types
    shapes = {el.get("shape") for s in p["slides"] for el in s["elements"] if el["type"] == "shape"}
    assert {"arrow_right", "triangle", "diamond", "ellipse", "rounded_rect"} <= shapes
    # 表ヘッダーの白文字がラン書式として残る
    tbl = next(el for s in p["slides"] for el in s["elements"] if el["type"] == "table")
    assert tbl["rows"][0][0]["paragraphs"][0]["runs"][0]["color"] == "#FFFFFF"


def test_mock_deck_roundtrip_editable():
    p = parse_pptx(_mock_bytes(), "mock.pptx")
    data, warns = generate_pptx(p, "editable")
    assert warns == []
    again = parse_pptx(data, "rt.pptx")
    assert [s["title"] for s in again["slides"]] == [s["title"] for s in p["slides"]]
    assert sum(1 for s in again["slides"] for el in s["elements"] if el["type"] == "image") == 2
    assert not [i for i in check_presentation(again) if i["severity"] == "error"]


def test_report_lists_text_as_shapes_with_numbers():
    p = parse_pptx(_mock_bytes(), "mock.pptx")
    rep = build_report(p)
    assert rep["summary"]["text_as_shapes"] == rep["summary"]["text_with_bbox"] > 50
    assert rep["summary"]["images"] == 2 and rep["summary"]["editable_ratio"] == 1.0
    row = next(r for r in rep["rows"] if r["type"] == "text" and r["role"] == "title")
    assert all(isinstance(row[k], float) for k in ("x", "y", "w", "h")) and row["font_pt_max"] >= 28
    md = to_markdown(p)
    assert "| 種別 |" in md and "画像化しない" in md


def test_template_parts_in_web_bundle_and_pptx():
    from app.web_renderer import build_bundle

    p = parse_pptx(_mock_bytes(), "mock.pptx", template_id="corporate_standard")
    p, _e, _f = pipeline.prepare(p)
    files = build_bundle(p)
    assert "assets/template_logo_horizontal_blue.png" in files and "assets/template_cover_background.jpg" in files
    assert "conversion_report.md" in files
    html = files["index.html"].decode("utf-8")
    assert "tpl-bar" in html and "kind-cover" in html
    data, _ = generate_pptx(p, "editable")
    prs = Presentation(io.BytesIO(data))
    assert any(sh.name == "cover_background" for sh in prs.slides[0].shapes)
    assert any(sh.name == "bar" for sh in prs.slides[1].shapes)


def test_api_report_and_closing_and_loglevel():
    pres = client.post("/api/import/pptx", files={"file": ("mock.pptx", _mock_bytes())}, data={"template_id": "corporate_standard"}).json()["presentation"]
    r = client.post("/api/report", json={"presentation": pres})
    assert r.status_code == 200 and r.json()["summary"]["elements"] > 50
    r = client.post("/api/closing-slide", json={"presentation": pres, "name": "サンプル（架空）"})
    assert r.status_code == 200 and r.json()["added"] is True
    slides = r.json()["presentation"]["slides"]
    assert slides[-1]["layout"] == "closing" and slides[-1]["elements"][0]["bbox"]
    r2 = client.post("/api/closing-slide", json={"presentation": r.json()["presentation"]})
    assert r2.json()["added"] is False
    assert client.post("/api/loglevel", json={"level": "DEBUG"}).json()["level"] == "DEBUG"
    assert client.get("/api/loglevel").json()["level"] == "DEBUG"
    assert client.post("/api/loglevel", json={"level": "INFO"}).status_code == 200
    assert client.post("/api/loglevel", json={"level": "bogus"}).status_code == 400


def test_closing_and_cover_survive_roundtrip():
    """Issue #8: 表紙・最終ページの種別が PPTX 往復で保たれ、マーカーは要素として現れない。"""
    from app.template_kit import make_closing_slide

    pres = client.post("/api/import/pptx", files={"file": ("mock.pptx", _mock_bytes())}, data={"template_id": "corporate_standard"}).json()["presentation"]
    pres["slides"].append(make_closing_slide("s_end", len(pres["slides"]), "ご清聴ありがとうございました（サンプル）"))
    pres, _e, _f = pipeline.prepare(pres)
    data, _ = generate_pptx(pres, "editable")
    again = parse_pptx(data, "rt.pptx")
    assert again["slides"][0]["layout"] == "title"
    assert again["slides"][-1]["layout"] == "closing"
    assert [el["type"] for el in again["slides"][-1]["elements"]] == ["text"]
    # 2 周目でも同じ
    data2, _ = generate_pptx(pipeline.prepare(again)[0], "editable")
    again2 = parse_pptx(data2, "rt2.pptx")
    assert again2["slides"][-1]["layout"] == "closing" and again2["slides"][0]["layout"] == "title"
