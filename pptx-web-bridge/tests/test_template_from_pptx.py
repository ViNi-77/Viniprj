"""PPTX からテンプレートを作る（Phase C）: 部品の推定、保存・一覧・削除、未保存テンプレートでのプレビュー、土台 PPTX を使った出力。"""
from __future__ import annotations

import io
import sys
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from app import pipeline, template_kit, template_store
from app.config import get_config
from app.main import app
from app.pptx_generator import generate_pptx
from app.pptx_parser import parse_pptx
from app.template_from_pptx import analyze, classify_slides, parts_of, sample_presentation, set_part_box

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture(scope="module")
def brand_pptx() -> bytes:
    path = ROOT / "samples" / "brand_template.pptx"
    if not path.exists():
        sys.path.insert(0, str(ROOT / "samples"))
        from make_brand_template_pptx import build  # type: ignore

        build(path)
    return path.read_bytes()


@pytest.fixture(scope="module")
def client() -> TestClient:
    return TestClient(app)


@pytest.fixture(scope="module")
def analyzed(brand_pptx: bytes) -> dict:
    return analyze(brand_pptx, "brand_template.pptx")


def test_slide_roles_first_cover_last_closing():
    assert classify_slides(3, [50, 40, 10]) == ["cover", "content", "closing"]
    assert classify_slides(2, [50, 10]) == ["cover", "closing"]
    assert classify_slides(2, [50, 200]) == ["cover", "content"]
    assert classify_slides(4, [1, 1, 1, 1], {1: "skip", 3: "content"}) == ["cover", "skip", "content", "content"]


def test_cover_parts_background_logo_title_footer(analyzed: dict):
    cover = analyzed["proposal"]["cover"]
    assert cover["background_image"].endswith("cover_background.jpg") and cover["background_source"] == "slide"
    assert cover["logo"]["image"].endswith("cover_logo.png")
    # 左上（0.5in, 0.3in）のロゴ: 正規化後 36×21.6pt 付近、幅 1.8in → 129.6pt
    assert abs(cover["logo"]["x"] - 36) < 2 and abs(cover["logo"]["y"] - 21.6) < 2 and abs(cover["logo"]["w"] - 129.6) < 2
    assert cover["title"]["size_pt"] == 36 and cover["title"]["color"] == "#FFFFFF" and cover["title"]["bold"] is True
    assert cover["subtitle"]["size_pt"] == 20
    assert cover["footer"]["text"] == "{date} / © Sample Brand" and cover["footer"]["align"] == "right"
    assert cover["date_format"] == "%Y年%-m月%-d日"
    assert "page_number" not in cover  # レイアウトのページ番号プレースホルダはスライドに置かれていないので拾わない


def test_content_parts_bar_logo_body_page_number(analyzed: dict):
    content = analyzed["proposal"]["content"]
    bar = content["bar"]
    assert bar["color"] == "#001A72" and bar["slant_pt"] > 5 and bar["w"] > 700 and bar["h"] < 30 and bar["y"] > 500
    assert content["logo"]["y"] > 500 and content["logo"]["x"] < 40
    assert content["title"]["color"] == "#001A72" and content["title"]["size_pt"] == 28
    body = content["body"]
    assert body["y"] > content["title"]["y"] + content["title"]["h"] - 1 and body["h"] > 300 and body["source"] == "slide"
    pn = content["page_number"]
    assert pn["text"] == "{page}" and pn["align"] == "right" and pn["size_pt"] == 9 and pn["x"] > 800
    assert "footer" not in content


def test_closing_parts_logo_and_message(analyzed: dict):
    closing = analyzed["proposal"]["closing"]
    assert closing["logo"]["x"] > 300 and closing["logo"]["w"] > 250
    assert closing["message"]["align"] == "center" and closing["message"]["size_pt"] == 18
    assert "title" not in closing  # マスターの題名プレースホルダは拾わない


def test_theme_colors_fonts_and_assets(analyzed: dict):
    p = analyzed["proposal"]
    assert p["source"] == "user" and p["id"] == "brand_template"
    assert p["colors"]["background"] == "#FFFFFF" and all(v.startswith("#") and len(v) == 7 for v in p["colors"].values())
    assert set(p["fonts"]) == {"heading", "body"}
    assert p["base_pptx"].endswith("base.pptx") and abs(p["base_canvas"]["width_pt"] - 960) < 0.1
    d = get_config().path("user_template_assets_dir") / "brand_template"
    names = sorted(x.name for x in d.iterdir())
    assert "base.pptx" in names and "cover_background.jpg" in names and "cover_logo.png" in names and "content_logo.png" in names
    assert [s["role"] for s in analyzed["slides"]] == ["cover", "content", "closing"]


def test_parts_list_and_set_part_box(analyzed: dict):
    parts = analyzed["parts"]
    keys = {(p["part"], p["key"]) for p in parts}
    assert {("cover", "background_image"), ("cover", "logo"), ("cover", "title"), ("content", "bar"), ("content", "body"), ("closing", "message")} <= keys
    logo = next(p for p in parts if p["part"] == "cover" and p["key"] == "logo")
    assert logo["box"]["h"] > 0 and logo["type"] == "image"
    bg = next(p for p in parts if p["key"] == "background_image")
    assert bg["box"] is None
    moved = set_part_box(analyzed["proposal"], "content", "bar", {"x": 100, "y": 500, "w": 800, "h": 20})
    assert moved["content"]["bar"]["x"] == 100 and analyzed["proposal"]["content"]["bar"]["x"] != 100  # 元は変えない
    moved2 = set_part_box(analyzed["proposal"], "cover", "logo", {"x": 10, "y": 10, "w": 100, "h": 999})
    assert moved2["cover"]["logo"]["w"] == 100 and "h" not in moved2["cover"]["logo"]


def test_role_override_makes_content_from_second_slide(brand_pptx: bytes):
    r = analyze(brand_pptx, "brand_template.pptx", template_id="brand_override", roles={2: "content"})
    assert [s["role"] for s in r["slides"]] == ["cover", "content", "content"]
    assert "closing" not in r["proposal"]
    template_store.delete_template("brand_override")


def test_preview_uses_unsaved_template_and_body_area(analyzed: dict):
    t = analyzed["proposal"]
    with template_kit.use_template(t):
        pres, _e, _f = pipeline.prepare(sample_presentation(t))
        assert template_kit.template_for(pres) is t
        content = pres["slides"][1]
        body = next(e for e in content["elements"] if e["role"] == "body")
        area = template_kit.content_area(t, pres["canvas"])
        assert abs(body["bbox"]["x"] - area["x"]) < 0.01 and body["bbox"]["y"] >= area["y"] - 0.01 and body["bbox"]["w"] <= area["w"] + 0.01
        title = next(e for e in content["elements"] if e["role"] == "title")
        assert abs(title["bbox"]["y"] - t["content"]["title"]["y"]) < 0.01
        closing = pres["slides"][2]
        msg = closing["elements"][0]
        assert abs(msg["bbox"]["y"] - t["closing"]["message"]["y"]) < 0.01
        from app.web_renderer import slide_html

        html = slide_html(content, pres, inline_assets=True, template=t, with_notes=False)
        assert 'data-tpl="bar"' in html and 'data-tpl="logo"' in html and 'data-tpl="page_number"' in html
    assert template_kit.template_for(pres)["id"] != t["id"]  # with を抜けると通常の解決に戻る


def test_api_from_pptx_preview_save_use_delete(client: TestClient, brand_pptx: bytes, sample_pptx_bytes: bytes):
    r = client.post("/api/templates/from-pptx", files={"file": ("brand_template.pptx", brand_pptx)}, data={"roles": "{}"})
    assert r.status_code == 200
    d = r.json()
    assert set(d["previews"]) == {"cover", "content", "closing"} and len(d["thumbs"]) == 3 and d["canvas"]["width_pt"] == 960
    proposal = d["proposal"]
    r2 = client.post("/api/templates/preview", json={"template": proposal})
    assert r2.status_code == 200 and len(r2.json()["parts"]) == len(d["parts"])
    # ID を変えて保存 → 画像フォルダが移り、一覧に source=user で載る
    r3 = client.put("/api/templates/My Brand!", json={"template": proposal, "previous_id": proposal["id"]})
    assert r3.status_code == 200, r3.text
    saved = r3.json()["template"]
    assert saved["id"] == "my_brand" and saved["source"] == "user" and "my_brand" in saved["content"]["logo"]["image"]
    listed = [t for t in r3.json()["templates"] if t["id"] == "my_brand"]
    assert listed and listed[0]["source"] == "user" and listed[0]["has_base_pptx"], r3.json()["templates"]
    assert not (get_config().path("user_template_assets_dir") / "brand_template").exists()
    assert client.get("/api/config").json()["templates"][0]["source"] == "builtin"
    # 保存したテンプレートで取込・描画
    pres = client.post("/api/import/pptx", files={"file": ("sample_deck.pptx", sample_pptx_bytes)}, data={"template_id": "my_brand"}).json()["presentation"]
    assert pres["theme"]["template_id"] == "my_brand"
    html = client.post("/api/render/slides", json={"presentation": pres, "indices": [1]}).json()["slides"]["1"]
    assert "tpl-bar" in html and "tpl-logo" in html
    # 組込と同じ ID・不正な色は拒否
    bad = dict(proposal, colors={"primary": "blue"})
    assert client.put("/api/templates/plain", json={"template": bad}).status_code == 400
    assert client.put("/api/templates/x_bad", json={"template": bad}).status_code == 400
    assert client.delete("/api/templates/plain").status_code == 400
    r4 = client.delete("/api/templates/my_brand")
    assert r4.status_code == 200 and r4.json()["deleted"] is True
    assert not (get_config().path("user_template_assets_dir") / "my_brand").exists()
    assert all(t["id"] != "my_brand" for t in r4.json()["templates"])
    assert client.post("/api/templates/from-pptx", files={"file": ("x.pptx", b"not a pptx")}).status_code == 422


def test_export_with_base_pptx_keeps_master_and_reparses(brand_pptx: bytes, sample_pptx_bytes: bytes):
    # 前の試験が brand_template の画像フォルダを移動・削除しているため、この試験専用の ID で解析し直す
    t = template_store.save_template(dict(analyze(brand_pptx, "brand_template.pptx", template_id="brand_base")["proposal"]))
    try:
        pres = parse_pptx(sample_pptx_bytes, "sample_deck.pptx", template_id=t["id"])
        data, _pres, warns = pipeline.export_pptx(pres, "editable", use_base_pptx=True)
        assert not [w for w in warns if w["code"].startswith("BASE_PPTX")]
        prs = Presentation(io.BytesIO(data))
        assert len(prs.slides) == len(pres["slides"])
        assert "blank" in prs.slide_layout_names[0].lower() if hasattr(prs, "slide_layout_names") else True
        assert all(len(list(s.slide_layout.placeholders)) <= 3 for s in prs.slides)  # 白紙レイアウト（フッター類のみ）
        names = [sh.name for sh in prs.slides[1].shapes]
        assert "bar" in names and "logo" in names and "page_number" in names  # スライド由来の部品は描く
        again = parse_pptx(data, "rt.pptx")
        assert [s["title"] for s in again["slides"]][:3] == [s["title"] for s in pres["slides"]][:3]
        assert not any(el.get("alt") in ("bar", "logo") for s in again["slides"] for el in s["elements"])
        # 土台を使わない従来出力も同じテンプレートで動く
        data2, warns2 = generate_pptx(_pres, "editable", use_base_pptx=False)
        assert data2 and not warns2
    finally:
        template_store.delete_template(t["id"])
