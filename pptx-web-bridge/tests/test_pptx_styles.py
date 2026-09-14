"""Issue #4（テーマ色）/ #5（継承スタイル）の試験。期待値はサンプル PPTX のテーマ XML から読む。"""
import io

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

from app.pptx_parser import parse_pptx
from app.pptx_styles import ThemeInfo, apply_brightness, apply_color_transforms

NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main", "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}


def _theme_xml(data: bytes):
    prs = Presentation(io.BytesIO(data))
    master = prs.slide_masters[0]
    return etree.fromstring(master.part.part_related_by(RT.THEME).blob), master


def _scheme(root, key: str) -> str:
    el = root.find(f".//a:clrScheme/a:{key}", NS)[0]
    return "#" + (el.get("lastClr") or el.get("val")).upper()  # sysClr は lastClr が実色


def _master_size(master, style: str, level: int) -> float:
    tx = master._element.find("p:txStyles", NS)
    return int(tx.find(f"p:{style}", NS).find(f"a:lvl{level}pPr", NS).find("a:defRPr", NS).get("sz")) / 100


def test_theme_info_reads_scheme_and_fonts(sample_pptx_bytes):
    root, master = _theme_xml(sample_pptx_bytes)
    t = ThemeInfo(master)
    assert t.scheme_hex("accent1") == _scheme(root, "accent1")
    assert t.scheme_hex("tx1") == _scheme(root, "dk1")  # clrMap 経由
    assert t.resolve_font("+mn-lt") == root.find(".//a:fontScheme/a:minorFont/a:latin", NS).get("typeface")
    assert t.resolve_font("Meiryo") == "Meiryo"


def test_apply_brightness():
    assert apply_brightness("#000000", 0.5) == "#808080"
    assert apply_brightness("#FFFFFF", -0.5) == "#808080"
    assert apply_brightness("#4F81BD", 0) == "#4F81BD"


def _transformed(base: str, inner: str) -> str:
    el = etree.fromstring(f'<a:srgbClr xmlns:a="{NS["a"]}" val="{base}">{inner}</a:srgbClr>')
    return apply_color_transforms("#" + base.upper(), el)


def test_lum_transforms_match_powerpoint_palette():
    """PowerPoint の色パレット（青 アクセント 1 = 4472C4）の「明るく / 暗く」と一致する。"""
    assert _transformed("4472C4", '<a:lumMod val="60000"/><a:lumOff val="40000"/>') == "#8FAADC"  # 明るく 40%
    assert _transformed("4472C4", '<a:lumMod val="75000"/>') == "#2F5597"  # 暗く 25%
    assert _transformed("4472C4", '<a:lumMod val="50000"/>') == "#203864"  # 暗く 50%
    assert _transformed("4472C4", "") == "#4472C4"


def test_shade_and_tint_use_linear_gamma():
    """shade / tint はリニアガンマ空間。ECMA-376 の例では 00FF00 に shade 50% で 00BC00。"""
    assert _transformed("00FF00", '<a:shade val="50000"/>') == "#00BC00"
    assert _transformed("FFFFFF", '<a:shade val="0"/>') == "#000000"
    assert _transformed("000000", '<a:tint val="0"/>') == "#FFFFFF"
    # 以前は shade / tint を無視して元の色のまま返していた
    assert _transformed("4472C4", '<a:shade val="50000"/>') != "#4472C4"


def test_lum_mod_and_off_combine():
    """lumMod と lumOff が両方あるとき、輝度は L*lumMod + lumOff になる（片方だけ見ない）。"""
    assert _transformed("808080", '<a:lumMod val="0"/><a:lumOff val="100000"/>') == "#FFFFFF"
    assert _transformed("808080", '<a:lumMod val="0"/>') == "#000000"


def test_shape_fill_keeps_shade(sample_pptx_bytes):
    """図形の塗りの schemeClr + shade が反映される（python-pptx は shade を落とす）。"""
    prs = Presentation(io.BytesIO(sample_pptx_bytes))
    badge = next(sh for sh in prs.slides[4].shapes if sh.has_text_frame and "accent2" in sh.text_frame.text)
    sch = badge.fill._xPr.find("a:solidFill", NS).find("a:schemeClr", NS)
    for child in list(sch):
        sch.remove(child)
    etree.SubElement(sch, "{%s}shade" % NS["a"]).set("val", "50000")
    buf = io.BytesIO()
    prs.save(buf)
    root, master = _theme_xml(sample_pptx_bytes)
    accent2 = ThemeInfo(master).scheme_hex("accent2")
    p = parse_pptx(buf.getvalue(), "shaded.pptx")
    shaded = next(el for el in p["slides"][4]["elements"] if el["type"] == "shape" and "accent2" in el["paragraphs"][0]["runs"][0]["text"])
    assert shaded["fill"] == _transformed(accent2[1:], '<a:shade val="50000"/>')
    assert shaded["fill"] != _scheme(root, "accent2")


def test_theme_colors_resolved_in_parse(sample_pptx_bytes):
    root, _m = _theme_xml(sample_pptx_bytes)
    p = parse_pptx(sample_pptx_bytes, "sample_deck.pptx")
    assert not any(w["code"] == "THEME_COLOR_UNRESOLVED" for w in p["warnings"])
    s5 = p["slides"][4]
    caption = next(el for el in s5["elements"] if el["type"] == "text" and "月別" in el["paragraphs"][0]["runs"][0]["text"])
    assert caption["paragraphs"][0]["runs"][0]["color"] == _scheme(root, "accent1")
    badge = next(el for el in s5["elements"] if el["type"] == "shape" and "accent2" in el["paragraphs"][0]["runs"][0]["text"])
    assert badge["fill"] == apply_brightness(_scheme(root, "accent2"), 0.4)
    assert badge["paragraphs"][0]["runs"][0]["color"] == _scheme(root, "dk2")  # tx2 → dk2


def test_inherited_sizes_and_bullets(sample_pptx_bytes):
    _root, master = _theme_xml(sample_pptx_bytes)
    p = parse_pptx(sample_pptx_bytes, "sample_deck.pptx")
    title = next(el for el in p["slides"][1]["elements"] if el["role"] == "title")
    body = next(el for el in p["slides"][1]["elements"] if el["role"] == "body")
    assert title["paragraphs"][0]["runs"][0]["size_pt"] == _master_size(master, "titleStyle", 1)
    assert title["paragraphs"][0]["bullet"] is None
    assert body["paragraphs"][0]["runs"][0]["size_pt"] == _master_size(master, "bodyStyle", 1)
    assert body["paragraphs"][1]["level"] == 1 and body["paragraphs"][1]["runs"][0]["size_pt"] == _master_size(master, "bodyStyle", 2)
    assert body["paragraphs"][0]["bullet"] == "bullet"
    # 素のテキストボックスは otherStyle（既定 18pt）だが、明示サイズ 16pt がある run は明示値が勝つ
    note = next(el for el in p["slides"][2]["elements"] if el["type"] == "text" and el["role"] == "body")
    assert note["paragraphs"][0]["runs"][0]["size_pt"] == 16.0
    # フォントは +mn-lt → テーマの minor latin
    assert body["paragraphs"][0]["runs"][0]["font"] == ThemeInfo(master).resolve_font("+mn-lt")


def test_font_scale_autofit(sample_pptx_bytes):
    """normAutofit fontScale が run サイズに掛かる。"""
    prs = Presentation(io.BytesIO(sample_pptx_bytes))
    body = prs.slides[1].placeholders[1]
    bodypr = body.text_frame._txBody.find("a:bodyPr", NS)
    for child in list(bodypr):
        bodypr.remove(child)
    auto = etree.SubElement(bodypr, "{%s}normAutofit" % NS["a"])
    auto.set("fontScale", "50000")
    buf = io.BytesIO()
    prs.save(buf)
    p = parse_pptx(buf.getvalue(), "scaled.pptx")
    body_el = next(el for el in p["slides"][1]["elements"] if el["role"] == "body")
    assert body_el["paragraphs"][0]["runs"][0]["size_pt"] == round(_master_size(Presentation(io.BytesIO(sample_pptx_bytes)).slide_masters[0], "bodyStyle", 1) * 0.5, 1)


def test_inherited_values_are_marked_and_explicit_ones_are_not(sample_pptx_bytes):
    """継承で補った値には inherited 印が付き、明示された値には付かない。

    印の使い道は変わった（かつてはテンプレート適用の上書き判定に使っていた）。
    いまは「この色・大きさは本当にその資料が指定したものか」を区別するために要る。
    """
    p = parse_pptx(sample_pptx_bytes, "sample_deck.pptx")
    body = next(el for el in p["slides"][1]["elements"] if el["role"] == "body")
    assert {"size_pt", "color", "font"} <= set(body["paragraphs"][0]["runs"][0]["inherited"])
    # 明示 12pt の注記には size_pt の印が付かない
    cap = next(el for el in p["slides"][4]["elements"] if el["type"] == "text" and "月別" in el["paragraphs"][0]["runs"][0]["text"])
    r = cap["paragraphs"][0]["runs"][0]
    assert r["size_pt"] == 12.0 and "size_pt" not in r.get("inherited", [])




def test_clr_map_override_resolves_inverted_slide(sample_pptx_bytes):
    """Issue #14: スライドの clrMapOvr（bg1↔dk1, tx1↔lt1）で、継承文字色 tx1 が lt1（白）に解決される。"""
    prs = Presentation(io.BytesIO(sample_pptx_bytes))
    slide = prs.slides[1]
    # python-pptx のスライドは既定で <p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr> を持つので、その中身を差し替える
    ovr = slide._element.find("p:clrMapOvr", NS)
    if ovr is None:
        ovr = etree.SubElement(slide._element, "{%s}clrMapOvr" % NS["p"])
    for child in list(ovr):
        ovr.remove(child)
    mapping = etree.SubElement(ovr, "{%s}overrideClrMapping" % NS["a"])
    for k, v in {"bg1": "dk1", "tx1": "lt1", "bg2": "dk2", "tx2": "lt2", "accent1": "accent1", "accent2": "accent2", "accent3": "accent3", "accent4": "accent4", "accent5": "accent5", "accent6": "accent6", "hlink": "hlink", "folHlink": "folHlink"}.items():
        mapping.set(k, v)
    buf = io.BytesIO()
    prs.save(buf)
    root, _m = _theme_xml(sample_pptx_bytes)
    p = parse_pptx(buf.getvalue(), "inverted.pptx")
    assert not any(w["code"].startswith("CLRMAP") for w in p["warnings"])
    inverted_title = next(el for el in p["slides"][1]["elements"] if el["role"] == "title")
    normal_title = next(el for el in p["slides"][2]["elements"] if el["role"] == "title")
    assert inverted_title["paragraphs"][0]["runs"][0]["color"] == _scheme(root, "lt1")
    assert normal_title["paragraphs"][0]["runs"][0]["color"] == _scheme(root, "dk1")
