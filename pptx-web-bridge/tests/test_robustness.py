"""探索的デバッグで見つかった弱点の回帰試験（境界入力・修復・分割）。"""
import io
import json

from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches

from app import pipeline
from app.html_parser import parse_html
from app.layout import layout_presentation
from app.main import app
from app.model import new_presentation, new_slide, paragraph, run, table_element, text_element, cell
from app.validate import validate, validate_and_repair

client = TestClient(app)


def test_repair_coerces_paragraphs_runs_and_cells():
    obj = {"slides": [{"id": "s", "elements": [
        {"id": "e", "type": "text", "paragraphs": [{"runs": "文字列"}, "junk", {"runs": [5, None, {"text": 3, "size_pt": -1}]}, "そのまま段落"]},
        {"id": "t", "type": "table", "rows": [["a", 1], "x", [{"text": None, "colspan": "2"}]]},
    ]}]}
    fixed, errors, fixes = validate_and_repair(obj)
    assert errors == []
    paras = fixed["slides"][0]["elements"][0]["paragraphs"]
    assert paras[0]["runs"][0]["text"] == "文字列" and paras[-1]["runs"][0]["text"] == "そのまま段落"
    assert all(isinstance(r["text"], str) for p in paras for r in p["runs"])
    rows = fixed["slides"][0]["elements"][1]["rows"]
    assert rows[0][0]["text"] == "a" and rows[0][1]["text"] == "1" and rows[1][0]["colspan"] == 2
    assert {f["code"] for f in fixes} >= {"PARAGRAPH_FIXED", "TABLE_FIXED"}


def test_repair_clamps_huge_bbox_and_font_then_exports():
    obj = {"slides": [{"id": "s", "elements": [{"id": "e", "type": "text", "bbox": {"x": 1e308, "y": -1e308, "w": 1e308, "h": float("nan") if False else 1}, "paragraphs": [{"runs": [{"text": "x", "size_pt": 99999}]}]}]}]}
    fixed, errors, fixes = validate_and_repair(obj)
    assert errors == [] and any(f["code"] == "BBOX_CLAMPED" for f in fixes)
    b = fixed["slides"][0]["elements"][0]["bbox"]
    assert b["x"] <= 2 * fixed["canvas"]["width_pt"] and b["y"] >= -fixed["canvas"]["height_pt"]
    assert fixed["slides"][0]["elements"][0]["paragraphs"][0]["runs"][0]["size_pt"] is None
    data, _p, warns = pipeline.export_pptx(fixed, "editable")
    assert not any(w["code"] == "ELEMENT_WRITE_FAILED" for w in warns)
    assert len(Presentation(io.BytesIO(data)).slides) == 1


def test_grid_with_many_cards_wraps_into_columns():
    html = ('<html><body><section><h2>t</h2><div class="grid">' + "".join(f'<div class="card"><h3>c{i}</h3><p>x</p></div>' for i in range(10)) + "</div></section></body></html>").encode()
    p = parse_html(html, "g.html", {}, computed_style=False)
    cards = [e for e in p["slides"][0]["elements"] if e.get("role") == "card"]
    assert len(cards) == 10 and all(e["layout_hint"]["columns"] == 3 for e in cards)
    assert any(w["code"] == "GRID_COLUMNS_LIMITED" for w in p["warnings"])
    lp = layout_presentation(p)
    first = [e for e in lp["slides"][0]["elements"] if e.get("role") == "card"]
    assert len({round(e["bbox"]["x"]) for e in first}) == 3  # 3 列に並ぶ
    assert len(first) >= 6


def test_table_is_split_across_slides_with_header():
    p = new_presentation("t")
    s = new_slide("s001", 0, title="表")
    s["elements"].append(text_element("t", [paragraph([run("表")])], role="title"))
    rows = [[cell("列A", bold=True), cell("列B", bold=True)]] + [[cell(f"r{i}"), cell("x")] for i in range(60)]
    s["elements"].append(table_element("tbl", rows, None, header_rows=1))
    p["slides"].append(s)
    lp = layout_presentation(p)
    assert len(lp["slides"]) >= 3
    tables = [e for sl in lp["slides"] for e in sl["elements"] if e["type"] == "table"]
    assert sum(len(t["rows"]) - 1 for t in tables) == 60  # ヘッダーを除く行数は保存される
    assert all(t["rows"][0][0]["text"] == "列A" for t in tables)  # 続きの表にもヘッダー
    assert not any(w["code"] == "ELEMENT_TRUNCATED" for w in lp["warnings"])


def test_long_single_paragraph_is_split_by_sentence():
    p = new_presentation("t")
    s = new_slide("s001", 0, title="長文")
    s["elements"].append(text_element("t", [paragraph([run("長文")])], role="title"))
    s["elements"].append(text_element("b", [paragraph([run("これは文です。" * 300)])], role="body"))
    p["slides"].append(s)
    lp = layout_presentation(p)
    assert len(lp["slides"]) >= 2
    assert not any(w["code"] == "ELEMENT_TRUNCATED" for w in lp["warnings"])
    joined = "".join(r["text"] for sl in lp["slides"] for e in sl["elements"] if e.get("role") == "body" for pp in e["paragraphs"] for r in pp["runs"])
    assert joined == "これは文です。" * 300


def test_empty_html_warns_no_content():
    p = parse_html(b"<html></html>", "empty.html", {}, computed_style=False)
    assert p["slides"] == [] and any(w["code"] == "NO_CONTENT" for w in p["warnings"])


def test_empty_placeholders_are_dropped():
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[1])  # 未入力のタイトル・本文プレースホルダ
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text_frame.text = ""
    s.shapes.add_textbox(Inches(1), Inches(3), Inches(3), Inches(1)).text_frame.text = "本文あり"
    buf = io.BytesIO()
    prs.save(buf)
    p = pipeline.import_pptx(buf.getvalue(), "ph.pptx")["presentation"]
    assert p["slides"][0]["elements"] == []
    assert [e["type"] for e in p["slides"][1]["elements"]] == ["text"]


def test_wrong_format_messages():
    r = client.post("/api/import/html", files={"file": ("x.html", open("samples/sample_deck.pptx", "rb").read())})
    assert r.status_code == 422 and "PPTX" in r.json()["detail"]
    r = client.post("/api/import/pptx", files={"file": ("x.pptx", b"<!doctype html><html></html>")})
    assert r.status_code == 422 and "HTML" in r.json()["detail"]


def test_unknown_mode_warns():
    p = pipeline.import_pptx(open("samples/sample_deck.pptx", "rb").read(), "d.pptx")["presentation"]
    r = client.post("/api/export/pptx", json={"presentation": p, "mode": "bogus"})
    assert r.status_code == 200 and "MODE_UNKNOWN" in r.headers["x-warnings"]
