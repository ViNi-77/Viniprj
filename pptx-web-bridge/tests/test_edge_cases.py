"""Issue #11: 手動だった 3 項目の自動化（見た目優先フォールバック、4:3 資料、上限・大容量）。基本設計基準書 9.2 の例外系・境界値。"""
import io
import json
import time

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches

from app import rasterize
from app.config import get_config
from app.layout import layout_presentation
from app.main import app
from app.model import new_presentation, new_slide, paragraph, run, text_element
from app.pptx_generator import generate_pptx
from app.pptx_parser import parse_pptx

client = TestClient(app)


def test_visual_mode_falls_back_to_editable_when_raster_unavailable(sample_pptx_bytes, monkeypatch):
    """Playwright/Chromium が使えない環境で visual を選ぶと、警告付きで編集性優先として出力される。"""

    def boom(*_a, **_k):
        raise RuntimeError("Chromium を起動できません（試験用に強制）")

    monkeypatch.setattr(rasterize, "render_slide_images", boom)
    p = parse_pptx(sample_pptx_bytes, "sample_deck.pptx")
    data, warns = generate_pptx(p, "visual")
    assert any(w["code"] == "VISUAL_MODE_UNAVAILABLE" for w in warns)
    prs = Presentation(io.BytesIO(data))
    # 文字がテキストシェイプとして残っている（画像化されていない）
    assert any(sh.has_text_frame and sh.text_frame.text for sh in prs.slides[1].shapes)


def test_hybrid_without_raster_uses_boxes(monkeypatch):
    from app.model import bbox, unsupported_element

    monkeypatch.setattr(rasterize, "render_slide_images", lambda *_a, **_k: (_ for _ in ()).throw(RuntimeError("no chromium")))
    p = new_presentation("t")
    s = new_slide("s001", 0)
    s["elements"].append(unsupported_element("e1", "chart", bbox(10, 10, 200, 100), "グラフ（未対応）"))
    p["slides"].append(s)
    data, warns = generate_pptx(p, "hybrid")
    codes = {w["code"] for w in warns}
    assert "HYBRID_RASTER_UNAVAILABLE" in codes and "UNSUPPORTED_AS_BOX" in codes
    assert len(Presentation(io.BytesIO(data)).slides) == 1


def _deck_4_3() -> bytes:
    prs = Presentation()  # python-pptx 既定は 10in × 7.5in = 4:3
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "4:3 の資料（サンプル）"
    s.placeholders[1].text = "比率を保ったまま変換する"
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    s2.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(1)).text_frame.text = "2 枚目"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_4_3_deck_keeps_ratio_and_warns():
    p = parse_pptx(_deck_4_3(), "deck43.pptx")
    assert p["canvas"]["aspect"] == "4:3"
    assert any(w["code"] == "ASPECT_NOT_16_9" for w in p["warnings"])
    assert abs(p["canvas"]["width_pt"] / p["canvas"]["height_pt"] - 4 / 3) < 0.01
    data, warns = generate_pptx(p, "editable")
    prs = Presentation(io.BytesIO(data))
    assert abs(prs.slide_width / prs.slide_height - 4 / 3) < 0.01
    r = client.post("/api/preview/html", json={"presentation": p})
    assert r.status_code == 200 and 'data-canvas-w="720"' in r.text
    # テンプレート部品も 4:3 に拡縮される（帯が右端で切れない）
    from app.template_kit import chrome_spec

    p["theme"]["template_id"] = "corporate_standard"
    p["slides"][1]["index"] = 1
    spec = chrome_spec(p["slides"][1], p)
    bar = spec["bars"][0]
    assert bar["x"] + bar["w"] <= p["canvas"]["width_pt"] + 0.5


def test_upload_limit_returns_413(monkeypatch):
    from app import main as main_mod

    monkeypatch.setattr(main_mod, "_MAX_UPLOAD", 1024)
    r = client.post("/api/import/pptx", files={"file": ("big.pptx", b"x" * 2048)})
    assert r.status_code == 413


def test_large_presentation_layout_completes():
    """100 枚超・各 20 段落の資料でもレイアウトと品質検査が完走する（時間を記録）。"""
    p = new_presentation("大量")
    for i in range(120):
        s = new_slide(f"s{i:03d}", i, title=f"スライド {i + 1}")
        s["elements"].append(text_element(f"t{i}", [paragraph([run(f"題名 {i + 1}")])], role="title"))
        s["elements"].append(text_element(f"b{i}", [paragraph([run("本文 " + "あ" * 30)], bullet="bullet") for _ in range(20)], role="body"))
        p["slides"].append(s)
    t0 = time.time()
    lp = layout_presentation(p)
    r = client.post("/api/quality", json={"presentation": lp})
    elapsed = time.time() - t0
    assert r.status_code == 200
    assert len(lp["slides"]) >= 120 and lp["slides"][-1]["index"] == len(lp["slides"]) - 1
    assert elapsed < 60, f"処理時間 {elapsed:.1f}s（上限 60s）"
