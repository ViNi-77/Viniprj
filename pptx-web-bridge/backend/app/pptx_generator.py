"""Presentation JSON → PPTX 生成器（python-pptx）。

出力モード:
- editable: すべてネイティブ要素（テキスト、図形、画像、表、線）。再編集を優先。
- visual  : スライド全体を画像化して貼る。見た目を優先。Playwright が無ければ editable へフォールバック。
- hybrid  : テキスト・表・図形はネイティブ、未対応要素だけ画像（画像化できなければ枠）。
"""
from __future__ import annotations

import base64
import io
import re
from typing import Any

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from . import diagrams, template_kit
from .config import get_config
from .logging_setup import get_logger
from .model import add_warning, warning
from .typography import effective_size, element_font_pt, font_scale, role_default_pt
from .units import pt_to_emu

log = get_logger("pptx_generator")

_SHAPES = {
    "rect": MSO_SHAPE.RECTANGLE,
    "rounded_rect": MSO_SHAPE.ROUNDED_RECTANGLE,
    "ellipse": MSO_SHAPE.OVAL,
    "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE,
    "diamond": MSO_SHAPE.DIAMOND,
    "arrow_right": MSO_SHAPE.RIGHT_ARROW,
    "parallelogram": MSO_SHAPE.PARALLELOGRAM,
}
_ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}
_ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}


def _rgb(color: str | None) -> RGBColor | None:
    if not color or not str(color).startswith("#") or len(color) != 7:
        return None
    try:
        return RGBColor.from_string(color[1:])
    except ValueError:
        return None


def _is_dark(color: str | None) -> bool:
    c = _rgb(color)
    if c is None:
        return False
    r, g, b = c[0], c[1], c[2]
    return (0.299 * r + 0.587 * g + 0.114 * b) < 128


def _resolve_font(name: str | None, fallback: dict, default: str) -> str:
    """フォント代替表に従い、最初の候補名を返す（実際の有無は PowerPoint 側で解決される）。"""
    if not name:
        return default
    aliases = fallback.get("aliases", {})
    if name in aliases and aliases[name]:
        return aliases[name][0]
    return name


class _Gen:
    def __init__(self, presentation: dict, mode: str, use_base_pptx: bool = False):
        self.p = presentation
        self.mode = mode
        self.cfg = get_config()
        self.template = template_kit.template_for(presentation)
        # 読み込んだ PPTX を土台にする（ユーザーテンプレート）: マスター・レイアウト由来の部品は土台が持つので描かない
        self.base_path = None
        if use_base_pptx and self.template.get("base_pptx"):
            from .config import resource_path

            bp = resource_path(self.template["base_pptx"])
            self.base_path = bp if bp.exists() else None
        self.skip_inherited_chrome = self.base_path is not None
        self.placeholder_title: dict | None = None
        self.fonts = self.cfg.font_fallback()
        self.theme_fonts = presentation.get("theme", {}).get("fonts", {})
        self.colors = presentation.get("theme", {}).get("colors", {})
        self.warnings: list[dict] = []
        self.images: list[bytes | None] = []
        self.slide_text_color: str | None = None  # 暗い背景のスライドで色未指定の文字に使う

    # --- 共通 ---
    def role_size(self, role: str | None) -> float:
        return role_default_pt(role)

    def font_for(self, role: str | None, run_font: str | None) -> str:
        default = self.theme_fonts.get("heading" if role == "title" else "body") or self.fonts.get("default_body", "Meiryo")
        return _resolve_font(run_font, self.fonts, default)

    def fill_paragraphs(self, text_frame: Any, paragraphs: list[dict], role: str | None, default_color: str | None = None, el: dict | None = None) -> None:
        """段落を書き込む。サイズは typography の実効サイズ（自動縮小込み）を実値で書き、normAutofit は使わない。"""
        text_frame.word_wrap = True
        first = True
        base_size = element_font_pt(el) * font_scale(el) if el is not None else self.role_size(role)
        default_color = default_color or self.slide_text_color
        for para in paragraphs:
            p = text_frame.paragraphs[0] if first else text_frame.add_paragraph()
            first = False
            level = int(para.get("level", 0) or 0)
            p.level = min(level, 8)
            if para.get("align") in _ALIGN:
                p.alignment = _ALIGN[para["align"]]
            bullet = para.get("bullet")
            self._set_bullet(p, bullet, level, base_size)
            for r in para.get("runs", []):
                run = p.add_run()
                run.text = r.get("text", "")
                f = run.font
                f.size = Pt(effective_size(r, el) if el is not None else float(r.get("size_pt") or base_size))
                f.bold = bool(r.get("bold")) or (role == "title" and r.get("bold") is None)
                f.italic = bool(r.get("italic"))
                if r.get("underline"):
                    f.underline = True
                # 継承で補ったフォント（テーマの latin 名など）は書かず、テンプレート既定に任せる
                f.name = self.font_for(role, None if "font" in (r.get("inherited") or []) else r.get("font"))
                color = _rgb(r.get("color")) or _rgb(default_color) or _rgb(self.colors.get("primary") if role == "title" else self.colors.get("text"))
                if color is not None:
                    f.color.rgb = color
                href = r.get("href")
                if href and str(href).lower().startswith(("http://", "https://", "mailto:")):
                    try:
                        run.hyperlink.address = href
                    except Exception:  # noqa: BLE001
                        pass
            if not para.get("runs"):
                p.add_run().text = ""

    @staticmethod
    def _set_bullet(p: Any, bullet: str | None, level: int, size: float) -> None:
        """段落 XML に箇条書き設定を書く（python-pptx に API が無いため直接編集）。"""
        pPr = p._p.get_or_add_pPr()
        for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buFont"):
            for e in pPr.findall(qn(tag)):
                pPr.remove(e)
        if bullet in ("bullet", "number"):
            indent = int(size * 1.2 * 12700)
            pPr.set("marL", str(indent * (level + 1)))
            pPr.set("indent", str(-indent))
            if bullet == "bullet":
                bu = etree.SubElement(pPr, qn("a:buChar"))
                bu.set("char", "•" if level == 0 else "–")
            else:
                bu = etree.SubElement(pPr, qn("a:buAutoNum"))
                bu.set("type", "arabicPeriod")
        else:
            etree.SubElement(pPr, qn("a:buNone"))
            if level:
                pPr.set("marL", str(int(size * 1.2 * 12700 * level)))

    # --- 要素 ---
    def add_text(self, slide: Any, el: dict) -> None:
        b = el["bbox"]
        tb = slide.shapes.add_textbox(Emu(pt_to_emu(b["x"])), Emu(pt_to_emu(b["y"])), Emu(pt_to_emu(b["w"])), Emu(pt_to_emu(b["h"])))
        tb.name = f"{el.get('role') or 'body'}:{el.get('id', 'text')}"
        if el.get("fill") and _rgb(el["fill"]):
            tb.fill.solid()
            tb.fill.fore_color.rgb = _rgb(el["fill"])
        tf = tb.text_frame
        tf.vertical_anchor = _ANCHOR.get(el.get("vertical_align") or "top", MSO_ANCHOR.TOP)
        self.fill_paragraphs(tf, el.get("paragraphs", []), el.get("role"), el=el)
        self._apply_rotation(tb, el)

    @staticmethod
    def _apply_rotation(shape: Any, el: dict) -> None:
        rot = el.get("rotation_deg")
        if rot:
            try:
                shape.rotation = float(rot)
            except (TypeError, ValueError):
                pass

    def add_shape(self, slide: Any, el: dict) -> None:
        b = el["bbox"]
        shp = slide.shapes.add_shape(_SHAPES.get(el.get("shape") or "rect", MSO_SHAPE.RECTANGLE), Emu(pt_to_emu(b["x"])), Emu(pt_to_emu(b["y"])), Emu(pt_to_emu(b["w"])), Emu(pt_to_emu(b["h"])))
        shp.name = f"{el.get('role') or 'shape'}:{el.get('id', 'shape')}"
        fill = _rgb(el.get("fill"))
        if fill is not None:
            shp.fill.solid()
            shp.fill.fore_color.rgb = fill
        else:
            shp.fill.background()
        stroke = _rgb(el.get("stroke"))
        if stroke is not None:
            shp.line.color.rgb = stroke
            shp.line.width = Pt(float(el.get("stroke_width_pt") or 1.0))
        else:
            shp.line.fill.background()
        shp.shadow.inherit = False
        tf = shp.text_frame
        tf.vertical_anchor = _ANCHOR.get(el.get("vertical_align") or "middle", MSO_ANCHOR.MIDDLE)
        tf.margin_left = tf.margin_right = Pt(8)
        tf.margin_top = tf.margin_bottom = Pt(6)
        default_color = "#FFFFFF" if _is_dark(el.get("fill")) else ("#222222" if el.get("fill") else None)
        paras = el.get("paragraphs", [])
        if paras:
            self.fill_paragraphs(tf, paras, el.get("role") if el.get("role") in ("title", "subtitle", "caption") else "body", default_color, el=el)
        self._apply_rotation(shp, el)

    def add_image(self, slide: Any, el: dict, slide_dict: dict) -> None:
        b = el["bbox"]
        asset = self.p.get("assets", {}).get(el.get("asset_id") or "")
        if not asset:
            if el.get("placeholder"):
                self.add_placeholder_box(slide, b, str(el.get("alt") or "画像"))
            else:
                self.warn("ASSET_MISSING", "画像資産が見つからないため枠のみ出力します。", slide_dict, el)
                self.add_placeholder_box(slide, b, "画像なし")
            return
        try:
            blob = base64.b64decode(asset.get("data_base64", ""))
            stream = io.BytesIO(blob)
            x, y, w, h = b["x"], b["y"], b["w"], b["h"]
            fit = el.get("fit") or "contain"
            nat_w, nat_h = asset.get("width_px"), asset.get("height_px")
            if fit == "contain" and nat_w and nat_h and w > 0 and h > 0:
                # 枠内に縦横比を保って収め、中央に寄せる
                ratio = float(nat_w) / float(nat_h)
                if w / h > ratio:
                    new_w = h * ratio
                    x += (w - new_w) / 2
                    w = new_w
                else:
                    new_h = w / ratio
                    y += (h - new_h) / 2
                    h = new_h
            pic = slide.shapes.add_picture(stream, Emu(pt_to_emu(x)), Emu(pt_to_emu(y)), Emu(pt_to_emu(w)), Emu(pt_to_emu(h)))
            pic.name = el.get("id", "image")
            self._apply_rotation(pic, el)
            if fit == "cover" and nat_w and nat_h and w > 0 and h > 0:
                ratio_box, ratio_img = w / h, float(nat_w) / float(nat_h)
                if ratio_img > ratio_box:
                    crop = (1 - ratio_box / ratio_img) / 2
                    pic.crop_left = pic.crop_right = crop
                else:
                    crop = (1 - ratio_img / ratio_box) / 2
                    pic.crop_top = pic.crop_bottom = crop
        except Exception as e:  # noqa: BLE001
            self.warn("IMAGE_WRITE_FAILED", f"画像を書き出せません: {e}", slide_dict, el)
            self.add_placeholder_box(slide, b, "画像エラー")

    def add_line(self, slide: Any, el: dict) -> None:
        pts = el.get("points") or []
        b = el.get("bbox") or {"x": 0, "y": 0, "w": 0, "h": 0}
        if len(pts) == 2:
            (x1, y1), (x2, y2) = pts
        else:
            x1, y1, x2, y2 = b["x"], b["y"], b["x"] + b["w"], b["y"] + b["h"]
        con = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(pt_to_emu(x1)), Emu(pt_to_emu(y1)), Emu(pt_to_emu(x2)), Emu(pt_to_emu(y2)))
        con.name = el.get("id", "line")
        stroke = _rgb(el.get("stroke")) or _rgb(self.colors.get("line")) or RGBColor(0x66, 0x66, 0x66)
        con.line.color.rgb = stroke
        con.line.width = Pt(float(el.get("stroke_width_pt") or 1.0))

    def add_table(self, slide: Any, el: dict, slide_dict: dict) -> None:
        rows = el.get("rows") or []
        if not rows:
            return
        b = el["bbox"]
        # 論理列数は colspan を考慮して求める
        ncols = max(sum(int(c.get("colspan", 1) or 1) for c in r) for r in rows)
        nrows = len(rows)
        gf = slide.shapes.add_table(nrows, ncols, Emu(pt_to_emu(b["x"])), Emu(pt_to_emu(b["y"])), Emu(pt_to_emu(b["w"])), Emu(pt_to_emu(b["h"])))
        gf.name = el.get("id", "table")
        tbl = gf.table
        widths = el.get("col_widths_pt") or []
        if len(widths) == ncols and sum(widths) > 0:
            scale = b["w"] / sum(widths)
            for i, w in enumerate(widths):
                tbl.columns[i].width = Emu(pt_to_emu(w * scale))
        header_rows = int(el.get("header_rows", 1) or 0)
        size = float(self.cfg.get("layout.caption_font_pt", 12)) + 2
        occupied: set[tuple[int, int]] = set()
        for ri, row in enumerate(rows):
            ci = 0
            for c in row:
                while (ri, ci) in occupied:
                    ci += 1
                if ci >= ncols:
                    break
                cs, rs = int(c.get("colspan", 1) or 1), int(c.get("rowspan", 1) or 1)
                target = tbl.cell(ri, ci)
                if (cs > 1 or rs > 1) and ri + rs - 1 < nrows and ci + cs - 1 < ncols:
                    try:
                        target.merge(tbl.cell(ri + rs - 1, ci + cs - 1))
                    except Exception:  # noqa: BLE001
                        pass
                    for rr in range(ri, ri + rs):
                        for cc in range(ci, ci + cs):
                            occupied.add((rr, cc))
                else:
                    occupied.add((ri, ci))
                tf = target.text_frame
                paras = c.get("paragraphs") or [{"runs": [{"text": line}], "level": 0, "bullet": None} for line in str(c.get("text", "")).split("\n")]
                if c.get("bold") or ri < header_rows:
                    for p_ in paras:
                        for r_ in p_["runs"]:
                            r_ = r_.setdefault("bold", True)
                self.fill_paragraphs(tf, paras, "caption")
                for p_ in tf.paragraphs:
                    for r_ in p_.runs:
                        r_.font.size = Pt(size)
                    if c.get("align") in _ALIGN:
                        p_.alignment = _ALIGN[c["align"]]
                fill = _rgb(c.get("fill")) or (_rgb(self.colors.get("surface")) if ri < header_rows else None)
                if fill is not None:
                    target.fill.solid()
                    target.fill.fore_color.rgb = fill
                elif ri >= header_rows:
                    target.fill.background()
                ci += cs

    def add_placeholder_box(self, slide: Any, b: dict, label: str) -> None:
        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(pt_to_emu(b["x"])), Emu(pt_to_emu(b["y"])), Emu(pt_to_emu(max(b["w"], 10))), Emu(pt_to_emu(max(b["h"], 10))))
        shp.fill.background()
        shp.line.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
        shp.line.dash_style = 4  # MSO_LINE.DASH
        tf = shp.text_frame
        tf.text = label
        for p in tf.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)

    def add_cropped_image(self, slide: Any, el: dict, slide_index: int) -> bool:
        """スライド画像から要素の領域を切り出して貼る（hybrid の未対応要素向け）。"""
        img = self.images[slide_index] if slide_index < len(self.images) else None
        b = el.get("bbox")
        if not img or not b or b["w"] <= 0 or b["h"] <= 0:
            return False
        try:
            from PIL import Image

            cw, ch = float(self.p["canvas"]["width_pt"]), float(self.p["canvas"]["height_pt"])
            with Image.open(io.BytesIO(img)) as im:
                sx, sy = im.width / cw, im.height / ch
                box = (int(b["x"] * sx), int(b["y"] * sy), int((b["x"] + b["w"]) * sx), int((b["y"] + b["h"]) * sy))
                crop = im.crop(box)
                buf = io.BytesIO()
                crop.save(buf, format="PNG")
                buf.seek(0)
            pic = slide.shapes.add_picture(buf, Emu(pt_to_emu(b["x"])), Emu(pt_to_emu(b["y"])), Emu(pt_to_emu(b["w"])), Emu(pt_to_emu(b["h"])))
            pic.name = el.get("id", "image")
            return True
        except Exception as e:  # noqa: BLE001
            log.warning("切り出し画像の生成に失敗: %s", e)
            return False

    def _inherited(self, item: dict) -> bool:
        return self.skip_inherited_chrome and item.get("source") in ("layout", "master")

    def add_template_background(self, slide: Any, sd: dict, prs: Any) -> None:
        """テンプレートの背景色・背景画像（本文より先に追加して背面にする）。"""
        chrome = template_kit.chrome_spec(sd, self.p)
        bg = (sd.get("background") or {}).get("color") or chrome.get("background_color")
        if _rgb(bg) and not (self.skip_inherited_chrome and not (sd.get("background") or {}).get("color")):
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = _rgb(bg)
        if chrome.get("background_image") and not self._inherited({"source": chrome.get("background_source")}):
            data = template_kit.image_data(chrome["background_image"])
            if data:
                pic = slide.shapes.add_picture(io.BytesIO(base64.b64decode(data[1])), 0, 0, prs.slide_width, prs.slide_height)
                pic.name = "cover_background"

    def add_template_chrome(self, slide: Any, slide_dict: dict) -> None:
        """テンプレート部品: 帯・ロゴ・フッター・ページ番号・機密表示。図形名で識別し再読込時に除外する。"""
        chrome = template_kit.chrome_spec(slide_dict, self.p)
        for bar in chrome["bars"]:
            if self._inherited(bar):
                continue
            if bar.get("slant_pt"):
                shp = slide.shapes.add_shape(MSO_SHAPE.PARALLELOGRAM, Emu(pt_to_emu(bar["x"])), Emu(pt_to_emu(bar["y"])), Emu(pt_to_emu(bar["w"])), Emu(pt_to_emu(bar["h"])))
                try:
                    # プリセット parallelogram の傾き = min(w, h) × adj。上限 adj は w / min(w, h)
                    ss = max(1.0, min(bar["w"], bar["h"]))
                    shp.adjustments[0] = min(bar["w"] / ss, bar["slant_pt"] / ss)
                except Exception:  # noqa: BLE001
                    pass
            else:
                shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(pt_to_emu(bar["x"])), Emu(pt_to_emu(bar["y"])), Emu(pt_to_emu(bar["w"])), Emu(pt_to_emu(bar["h"])))
            shp.name = "bar"
            shp.fill.solid()
            shp.fill.fore_color.rgb = _rgb(bar["color"]) or RGBColor(0, 0, 0)
            shp.line.fill.background()
            shp.shadow.inherit = False
        for img in chrome["images"]:
            data = template_kit.image_data(img["image"])
            if not data or self._inherited(img):
                continue
            pic = slide.shapes.add_picture(io.BytesIO(base64.b64decode(data[1])), Emu(pt_to_emu(img["x"])), Emu(pt_to_emu(img["y"])), Emu(pt_to_emu(img["w"])), Emu(pt_to_emu(img["h"])))
            pic.name = img.get("name", "logo")
        for t in chrome["texts"]:
            if self._inherited(t):
                continue
            tb = slide.shapes.add_textbox(Emu(pt_to_emu(t["x"])), Emu(pt_to_emu(t["y"])), Emu(pt_to_emu(t["w"])), Emu(pt_to_emu(t["h"])))
            tb.name = t.get("name", "footer")
            tf = tb.text_frame
            tf.word_wrap = False
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            tf.text = t["text"]
            para = tf.paragraphs[0]
            para.alignment = _ALIGN.get(t.get("align", "left"), PP_ALIGN.LEFT)
            for r in para.runs:
                r.font.size = Pt(float(t["size_pt"]))
                r.font.name = self.font_for("caption", None)
                color = _rgb(t.get("color"))
                if color is not None:
                    r.font.color.rgb = color

    def warn(self, code: str, message: str, slide_dict: dict | None = None, el: dict | None = None, fallback: str | None = None) -> None:
        w = warning("pptx_generator", code, message, slide_dict.get("id") if slide_dict else None, el.get("id") if el else None, fallback)
        self.warnings.append(w)
        add_warning(self.p, w, slide_dict)

    @staticmethod
    def _set_office_metadata(prs, author: str) -> None:
        """docProps を整える。python-pptx の既定テンプレートは作成者・アプリ名が固定値のため、出力ごとに上書きする。"""
        prs.core_properties.author = author
        prs.core_properties.last_modified_by = author
        for part in prs.part.package.iter_parts():
            if str(part.partname) == "/docProps/app.xml":
                part._blob = re.sub(rb"<Application>[^<]*</Application>", b"<Application>Microsoft Office PowerPoint</Application>", part.blob)

    # --- 全体 ---
    @staticmethod
    def _blank_layout(prs: Any) -> Any:
        """土台 PPTX の中で最も部品の少ないレイアウト（白紙）を選ぶ。フッター類以外のプレースホルダが無いものを優先。"""
        from pptx.enum.shapes import PP_PLACEHOLDER as _PH

        trivial = {_PH.FOOTER, _PH.DATE, _PH.SLIDE_NUMBER}
        best, best_n = None, 10**6
        for lay in prs.slide_layouts:
            n = sum(1 for ph in lay.placeholders if ph.placeholder_format.type not in trivial)
            name = str(lay.name or "").lower()
            score = n * 10 + (0 if ("blank" in name or "白紙" in name) else 1)
            if score < best_n:
                best, best_n = lay, score
        return best if best is not None else prs.slide_layouts[len(prs.slide_layouts) - 1]

    @staticmethod
    def _layout_for(prs: Any, kind: str) -> Any | None:
        """土台 PPTX の中から、スライドの役割に合うレイアウトを選ぶ（表紙 → 題名、中身 → 題名+本文）。

        題名をレイアウトのプレースホルダへ入れると、PowerPoint 側でテーマのフォント・色・位置が
        そのまま効く（＝「デザインの適用」と同じ結果になる）。
        """
        from pptx.enum.shapes import PP_PLACEHOLDER as _PH

        titles = {_PH.TITLE, _PH.CENTER_TITLE}
        bodies = {_PH.BODY, _PH.OBJECT}
        best, best_score = None, -(10**6)
        for lay in prs.slide_layouts:
            kinds = [ph.placeholder_format.type for ph in lay.placeholders]
            has_title = any(k in titles for k in kinds)
            has_sub = any(k == _PH.SUBTITLE for k in kinds)
            has_body = any(k in bodies for k in kinds)
            extra = sum(1 for k in kinds if k not in titles | bodies | {_PH.SUBTITLE, _PH.FOOTER, _PH.DATE, _PH.SLIDE_NUMBER})
            if not has_title:
                continue
            if kind == "cover":
                score = 10 + (5 if has_sub else 0) - (3 if has_body else 0)
            else:
                score = 10 + (5 if has_body else 0) - (3 if has_sub else 0)
            score -= extra * 2
            name = str(lay.name or "").lower()
            if kind == "cover" and ("title slide" in name or "表紙" in name or "タイトル スライド" in name):
                score += 3
            if score > best_score:
                best, best_score = lay, score
        return best

    def _fill_title_placeholder(self, slide: Any, sd: dict) -> dict | None:
        """レイアウトの題名プレースホルダへ題名を入れ、その要素を返す（入れられなければ None）。"""
        try:
            ph = slide.shapes.title
        except Exception:  # noqa: BLE001
            ph = None
        if ph is None:
            return None
        el = next((e for e in sd.get("elements", []) if e.get("type") == "text" and e.get("role") == "title" and e.get("paragraphs")), None)
        if el is None:
            return None
        try:
            self.fill_paragraphs(ph.text_frame, el.get("paragraphs", []), "title", el=el)
            ph.name = f"title:{el.get('id', 'title')}"
            if el.get("user_bbox") and el.get("bbox"):  # 手で動かした枠だけ位置を上書きする
                b = el["bbox"]
                ph.left, ph.top, ph.width, ph.height = (Emu(pt_to_emu(float(b[k]))) for k in ("x", "y", "w", "h"))
        except Exception as e:  # noqa: BLE001
            log.warning("題名プレースホルダへの流し込みに失敗: %s", e)
            return None
        return el

    @staticmethod
    def _drop_empty_placeholders(slide: Any) -> None:
        """使わなかった空のプレースホルダを外す（PowerPoint で「テキストを入力」の枠が残らないように）。"""
        from pptx.enum.shapes import PP_PLACEHOLDER as _PH

        keep = {_PH.FOOTER, _PH.DATE, _PH.SLIDE_NUMBER}
        for ph in list(slide.placeholders):
            try:
                if ph.placeholder_format.type in keep:
                    continue
                if ph.has_text_frame and ph.text_frame.text.strip():
                    continue
                ph._element.getparent().remove(ph._element)
            except Exception:  # noqa: BLE001
                continue

    def _open_base(self) -> Any:
        """土台 PPTX を開き、既存スライドを全て外す（マスター・レイアウト・テーマだけを残す）。"""
        prs = Presentation(str(self.base_path))
        sld_id_lst = prs.slides._sldIdLst
        for sld_id in list(sld_id_lst):
            prs.part.drop_rel(sld_id.rId)
            sld_id_lst.remove(sld_id)
        bw, bh = float(prs.slide_width) / 12700.0, float(prs.slide_height) / 12700.0
        cw, ch = float(self.p["canvas"]["width_pt"]), float(self.p["canvas"]["height_pt"])
        if abs(bw - cw) > 1 or abs(bh - ch) > 1:
            self.warn("BASE_PPTX_SIZE_MISMATCH", f"土台 PPTX の寸法（{bw:.0f}×{bh:.0f}pt）が資料（{cw:.0f}×{ch:.0f}pt）と違うため、マスターの部品がずれることがあります。", fallback="資料の寸法で出力")
        return prs

    def build(self) -> bytes:
        if self.base_path is not None:
            try:
                prs = self._open_base()
            except Exception as e:  # noqa: BLE001
                self.warn("BASE_PPTX_UNAVAILABLE", f"土台 PPTX を開けないため通常の方法で出力します: {e}", fallback="新規 PPTX")
                self.skip_inherited_chrome = False
                prs = Presentation()
        else:
            prs = Presentation()
        prs.slide_width = Emu(pt_to_emu(float(self.p["canvas"]["width_pt"])))
        prs.slide_height = Emu(pt_to_emu(float(self.p["canvas"]["height_pt"])))
        use_master_layouts = self.base_path is not None and self.skip_inherited_chrome
        blank = self._blank_layout(prs) if use_master_layouts else prs.slide_layouts[6]
        layout_for_kind = {k: (self._layout_for(prs, k) if use_master_layouts else None) for k in ("cover", "content", "closing")}
        prs.core_properties.title = self.p.get("meta", {}).get("title", "")
        self._set_office_metadata(prs, self.p.get("meta", {}).get("author", "") or "")

        needs_raster = self.mode == "visual" or (self.mode == "hybrid" and any(e.get("type") == "unsupported" for s in self.p.get("slides", []) for e in s.get("elements", [])))
        if needs_raster:
            try:
                from . import rasterize

                self.images = rasterize.render_slide_images(self.p)
            except Exception as e:  # noqa: BLE001
                self.images = []
                if self.mode == "visual":
                    self.warn("VISUAL_MODE_UNAVAILABLE", f"画像化環境（Playwright/Chromium）が使えないため編集性優先モードで出力します: {e}", fallback="editable へフォールバック")
                    self.mode = "editable"
                else:
                    self.warn("HYBRID_RASTER_UNAVAILABLE", f"画像化環境が使えないため未対応要素は枠のみ出力します: {e}", fallback="枠のみ")

        for si, sd in enumerate(self.p.get("slides", [])):
            kind = template_kit.slide_kind(sd, self.p)
            layout = layout_for_kind.get(kind) if kind != "closing" else None
            slide = prs.slides.add_slide(layout or blank)
            self.placeholder_title = self._fill_title_placeholder(slide, sd) if layout is not None and self.mode != "visual" else None
            self.slide_text_color = (sd.get("background") or {}).get("text_color")
            # スライド種別（表紙 / 最終ページ）を cSld の name 属性へ残す。図形ではないので本文要素に混ざらない（Issue #8）
            if kind in ("cover", "closing"):
                slide._element.cSld.set("name", f"kind:{kind}")
            if self.mode != "visual":
                self.add_template_background(slide, sd, prs)
                self.add_template_chrome(slide, sd)  # 本文より先に置き、背面にする
            if self.mode == "visual":
                img = self.images[si] if si < len(self.images) else None
                if img:
                    pic = slide.shapes.add_picture(io.BytesIO(img), 0, 0, prs.slide_width, prs.slide_height)
                    pic.name = f"{sd['id']}_visual"
                    # 検索・アクセシビリティ用に文字列を代替テキストへ残す
                    self._set_alt_text(pic, sd)
                else:
                    self.warn("VISUAL_SLIDE_FAILED", "このスライドの画像化に失敗したため編集性優先で出力します。", sd, fallback="editable")
                    self._add_native_elements(slide, sd, si)
            else:
                self._add_native_elements(slide, sd, si)
            if layout is not None:
                self._drop_empty_placeholders(slide)
            notes = sd.get("notes")
            if notes:
                slide.notes_slide.notes_text_frame.text = str(notes)

        buf = io.BytesIO()
        prs.save(buf)
        return buf.getvalue()

    @staticmethod
    def _set_alt_text(pic: Any, sd: dict) -> None:
        from .model import element_plain_text

        text = "\n".join(t for t in (element_plain_text(e) for e in sd.get("elements", [])) if t)
        try:
            pic._element.nvPicPr.cNvPr.set("descr", text[:2000])
        except Exception:  # noqa: BLE001
            pass

    def add_diagram(self, slide: Any, el: dict) -> None:
        """図解要素を子（図形・文字・線）に展開して出力する。図形名で往復時に元へ戻せるようにする。"""
        children = diagrams.expand_diagram(el, self.template)
        if not children:
            return
        dtype = diagrams.normalize_type((el.get("diagram") or {}).get("type"))
        for i, child in enumerate(children):
            before = len(slide.shapes._spTree)
            t = child.get("type")
            if t == "shape":
                self.add_shape(slide, child)
            elif t == "text":
                self.add_text(slide, child)
            elif t == "line":
                self.add_line(slide, child)
            else:
                continue
            if len(slide.shapes._spTree) > before:
                slide.shapes[-1].name = f"diagram:{dtype}:{el.get('id', 'd')}:{i}"

    def _add_native_elements(self, slide: Any, sd: dict, si: int) -> None:
        elements = sorted(sd.get("elements", []), key=lambda e: int(e.get("z", 0) or 0))
        for el in elements:
            if el is getattr(self, "placeholder_title", None):
                continue  # レイアウトの題名プレースホルダへ入れたので二重に出さない
            if not el.get("bbox"):
                self.warn("ELEMENT_NO_BBOX", "座標の無い要素をスキップしました（事前にレイアウトを実行してください）。", sd, el, "スキップ")
                continue
            t = el.get("type")
            try:
                if t == "text":
                    self.add_text(slide, el)
                elif t == "shape":
                    self.add_shape(slide, el)
                elif t == "image":
                    self.add_image(slide, el, sd)
                elif t == "line":
                    self.add_line(slide, el)
                elif t == "table":
                    self.add_table(slide, el, sd)
                elif t == "diagram":
                    self.add_diagram(slide, el)
                else:
                    if self.mode == "hybrid" and self.add_cropped_image(slide, el, si):
                        self.warn("UNSUPPORTED_AS_IMAGE", "未対応要素を画像として貼り付けました。", sd, el, "画像化")
                    else:
                        self.warn("UNSUPPORTED_AS_BOX", f"未対応要素（{el.get('original_type') or t}）は枠のみ出力します。", sd, el, "枠のみ")
                        self.add_placeholder_box(slide, el["bbox"], str(el.get("alt") or "未対応要素"))
            except Exception as e:  # noqa: BLE001 - 1要素の失敗で全体を止めない
                log.exception("要素の出力に失敗: %s", el.get("id"))
                self.warn("ELEMENT_WRITE_FAILED", f"要素 {el.get('id')} の出力に失敗: {e}", sd, el, "枠のみ")
                try:
                    self.add_placeholder_box(slide, el["bbox"], "出力エラー")
                except Exception:  # noqa: BLE001
                    pass


def generate_pptx(presentation: dict, mode: str | None = None, use_base_pptx: bool = False) -> tuple[bytes, list[dict]]:
    """PPTX バイト列と、生成中に出た警告を返す。presentation には警告が追記される。

    use_base_pptx: テンプレートが `base_pptx`（PPTX から作ったテンプレートの元ファイル）を持つとき、それを土台にして出力する。
    """
    cfg = get_config()
    mode = mode or str(cfg.get("pptx_export.default_mode", "editable"))
    unknown = mode not in cfg.get("pptx_export.modes", ["editable", "visual", "hybrid"])
    if unknown:
        requested, mode = mode, "editable"
    gen = _Gen(presentation, mode, use_base_pptx=use_base_pptx)
    if unknown:
        gen.warn("MODE_UNKNOWN", f"出力モード '{requested}' は不明のため編集性優先で出力します。", fallback="editable")
    data = gen.build()
    log.info("PPTX 生成完了: mode=%s slides=%d warnings=%d bytes=%d", gen.mode, len(presentation.get("slides", [])), len(gen.warnings), len(data))
    return data, gen.warnings
