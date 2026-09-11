"""PowerPoint（表紙・中身・最終ページの小さなデッキ）を読み、テンプレート部品を推定する。

利用者の「自分のテンプレートを読み込ませたら、この解釈で入りますよと示してほしい」に応える解析器。
出力は config/templates.json と同じ形（cover / content / closing の部品定義、座標は 960×540 基準）で、
UI で位置を直してから `template_store.save_template` で保存する。

推定規則（docs/03 の「テンプレート推定」表と同じ）:
- スライドの役割: 先頭 = 表紙、末尾 = 最終ページ（3 枚以上、または 2 枚で末尾の文字が 40 字未満）、残り = 中身。UI から上書き可。
- 図形はスライド本体に加えてレイアウト・マスターの図形も見る（showMasterSp=0 のときはマスターを見ない）。
- 画像: 面積 80% 以上 → 背景画像。幅 25% 以下で端から 15% 以内 → ロゴ。それ以外 → 装飾画像。
- 文字の無い塗り図形: **細長く（縦横比 4 以上）、幅か高さが 30% 以上で、端から 15% 以内**なら帯（平行四辺形は adj から傾きを取る）。
  それ以外は「装飾」として提案に載せるが既定では使わない。同じ大きさの塗り図形が 3 個以上並ぶものは色見本と見なして部品にせず、色をテーマ候補に回す。
  帯は大きい順に 3 個まで。面積 80% 以上なら背景色。
- 中身スライドは、候補が複数あるとき「題名・本文プレースホルダを持ち、文字図形が多い」ものを選ぶ（色見本や装飾のスライドを避ける）。
- 部品には信頼度（high / low）を付け、判定に使わなかった文字は「文字候補」として理由付きで返す（画面で役割を選べる）。
- 題名 / 副題 / 本文 プレースホルダ → title / subtitle / body（本文領域）。
- 下部 15% の小さな文字: ページ番号フィールドや数字だけなら page_number、それ以外は footer（日付は {date} に置換）。
- 最終ページの文字 → message。
"""
from __future__ import annotations

import copy
import hashlib
import io
import re
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

from . import template_store
from .logging_setup import get_logger
from .model import new_presentation, new_slide, paragraph, run, text_element
from .pptx_parser import NS, _align_of, _fill_is_solid, _placeholder_type, _rgb_of, _transform, _xfrm_of
from .pptx_styles import TextStyleResolver, ThemeInfo, _xml_color, apply_brightness, reset_current_theme, set_current_theme
from .units import emu_to_pt

log = get_logger("template_from_pptx")

BASE_W, BASE_H = 960.0, 540.0
_KINDS = ("cover", "content", "closing")
_TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}
_BODY_TYPES = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.VERTICAL_BODY}
_MONTHS = "January|February|March|April|May|June|July|August|September|October|November|December"
_DATE_PATTERNS: list[tuple[re.Pattern[str], str]] = [
    (re.compile(r"\d{4}年\s?\d{1,2}月\s?\d{1,2}日"), "%Y年%-m月%-d日"),
    (re.compile(r"\d{4}年\s?\d{1,2}月"), "%Y年%-m月"),
    (re.compile(r"\d{4}/\d{1,2}/\d{1,2}"), "%Y/%m/%d"),
    (re.compile(r"\d{4}-\d{1,2}-\d{1,2}"), "%Y-%m-%d"),
    (re.compile(r"\d{4}\.\d{1,2}\.\d{1,2}"), "%Y.%m.%d"),
    (re.compile(rf"(?:{_MONTHS})\s+\d{{1,2}},\s*\d{{4}}"), "%B %-d, %Y"),
    (re.compile(rf"\d{{1,2}}\s+(?:{_MONTHS})\s+\d{{4}}"), "%-d %B %Y"),
    (re.compile(rf"(?:{_MONTHS})\s+\d{{4}}"), "%B %Y"),
]
_PAGE_RE = re.compile(r"^\s*(?:[‹<]#[›>]|#|\d{1,3})\s*$")
_PAGE_TOTAL_RE = re.compile(r"^\s*(?:[‹<]#[›>]|\d{1,3})\s*/\s*\d{1,3}\s*$")
PART_LABELS = {"background_image": "背景画像", "logo": "ロゴ", "images": "画像", "bar": "帯", "bars": "帯", "decor": "装飾", "title": "題名", "subtitle": "副題", "body": "本文領域", "message": "一言", "footer": "フッター", "page_number": "ページ番号", "candidates": "文字候補"}
MAX_BARS = 3
SWATCH_MIN = 3
SOURCE_LABELS = {"slide": "スライド", "layout": "レイアウト", "master": "マスター"}


# ---------------------------------------------------------------- 図形の収集
class _Shape:
    """解析用に正規化した 1 図形（座標は pt、スライド座標系）。"""

    def __init__(self, shape: Any, box: dict, source: str, slide_no: int):
        self.shape = shape
        self.box = box
        self.source = source
        self.slide_no = slide_no
        self.ph = _placeholder_type(shape)
        self.is_picture = shape.shape_type == MSO_SHAPE_TYPE.PICTURE or (self.ph == PP_PLACEHOLDER.PICTURE and hasattr(shape, "image"))
        self.is_line = shape.shape_type == MSO_SHAPE_TYPE.LINE
        self.text = ""
        try:
            if shape.has_text_frame:
                self.text = shape.text_frame.text.strip()
        except Exception:  # noqa: BLE001
            self.text = ""
        self.fields = [f.get("type") or "" for f in shape._element.findall(".//a:fld", NS)]

    @property
    def area(self) -> float:
        return float(self.box["w"]) * float(self.box["h"])

    def describe(self) -> str:
        name = str(getattr(self.shape, "name", "") or "図形")
        where = SOURCE_LABELS.get(self.source, self.source)
        if self.source == "slide":
            where = f"スライド {self.slide_no}"
        return f"{where} / {name}"


def _is_hidden(shape: Any) -> bool:
    try:
        cnv = shape._element.find(".//p:cNvPr", NS)
        return cnv is not None and cnv.get("hidden") == "1"
    except Exception:  # noqa: BLE001
        return False


def _iter_shapes(shapes: Any, tf: dict | None = None):
    for sh in shapes:
        if _is_hidden(sh):
            continue
        try:
            st = sh.shape_type
        except Exception:  # noqa: BLE001
            st = None
        if st == MSO_SHAPE_TYPE.GROUP:
            child_tf = None
            grp = sh._element.find(".//a:xfrm", NS)
            if grp is not None:
                ch_off, ch_ext = grp.find("a:chOff", NS), grp.find("a:chExt", NS)
                if ch_off is not None and ch_ext is not None and int(ch_ext.get("cx", 0)) and int(ch_ext.get("cy", 0)):
                    gx, gy, gw, gh = _transform(emu_to_pt(sh.left), emu_to_pt(sh.top), emu_to_pt(sh.width), emu_to_pt(sh.height), tf)
                    child_tf = {"ox": gx, "oy": gy, "cox": emu_to_pt(int(ch_off.get("x", 0))), "coy": emu_to_pt(int(ch_off.get("y", 0))), "sx": gw / emu_to_pt(int(ch_ext.get("cx"))), "sy": gh / emu_to_pt(int(ch_ext.get("cy")))}
            yield from _iter_shapes(sh.shapes, child_tf)
            continue
        yield sh, tf


def _box_of(shape: Any, tf: dict | None) -> dict | None:
    try:
        if shape.left is None or shape.top is None or shape.width is None or shape.height is None:
            return None
        x, y, w, h = _transform(emu_to_pt(shape.left), emu_to_pt(shape.top), emu_to_pt(shape.width), emu_to_pt(shape.height), tf)
    except Exception:  # noqa: BLE001
        return None
    if w <= 0 or h <= 0:
        return None
    return {"x": x, "y": y, "w": w, "h": h}


def _shows_master(el: Any) -> bool:
    return el.get("showMasterSp") != "0"


def collect_shapes(slide: Any, slide_no: int) -> list[_Shape]:
    """スライド本体 → レイアウト → マスター の順に図形を集める（後の判定で本体が優先されるよう本体を先に置く）。"""
    layout = slide.slide_layout
    master = layout.slide_master
    out: list[_Shape] = []
    seen_ph: set[Any] = set()
    for src, owner in (("slide", slide), ("layout", layout), ("master", master)):
        if src == "master" and not (_shows_master(slide._element) and _shows_master(layout._element)):
            break
        for sh, tf in _iter_shapes(owner.shapes):
            box = _box_of(sh, tf)
            if box is None:
                continue
            s = _Shape(sh, box, src, slide_no)
            if s.ph is not None:
                # プレースホルダは本体のものを使う。レイアウトの題名・副題・本文は位置の手がかりとして本体に無いときだけ使う。
                # フッター・日付・ページ番号はスライド本体に置かれたときだけ表示されるため、レイアウト・マスターのものは見ない。
                key = s.ph if s.ph in _TITLE_TYPES | _BODY_TYPES | {PP_PLACEHOLDER.SUBTITLE} else None
                if src == "master" or (src == "layout" and key is None):
                    continue
                if key is not None:
                    if key in seen_ph:
                        continue
                    seen_ph.add(key)
            out.append(s)
    return out


# ---------------------------------------------------------------- 文字スタイル
def _text_style(s: _Shape, resolver: TextStyleResolver | None) -> dict:
    """先頭段落の run の書式（明示 → 継承 の順）。{'size_pt','color','bold','font','align'}"""
    style: dict[str, Any] = {"size_pt": None, "color": None, "bold": None, "font": None, "align": None}
    try:
        tf = s.shape.text_frame
        para = tf.paragraphs[0] if tf.paragraphs else None
    except Exception:  # noqa: BLE001
        return style
    if para is None:
        return style
    inh: dict[str, Any] = {}
    if resolver is not None:
        try:
            inh = resolver.resolve(s.shape, int(para.level or 0), para._p.find("a:pPr", NS))
        except Exception:  # noqa: BLE001
            inh = {}
    r = para.runs[0] if para.runs else None
    if r is None:
        # ページ番号・日付のフィールド（a:fld）は runs に含まれないため XML から書式を読む
        fld = para._p.find("a:fld", NS)
        rpr = fld.find("a:rPr", NS) if fld is not None else None
        if rpr is not None:
            if rpr.get("sz"):
                style["size_pt"] = int(rpr.get("sz")) / 100.0
            if rpr.get("b") is not None:
                style["bold"] = rpr.get("b") in ("1", "true")
            style["color"] = _xml_color(rpr, resolver.theme if resolver else None)
    if r is not None:
        try:
            style["size_pt"] = r.font.size.pt if r.font.size else None
        except Exception:  # noqa: BLE001
            pass
        style["bold"] = r.font.bold
        style["color"] = _rgb_of(r.font.color)
        try:
            style["font"] = r.font.name
        except Exception:  # noqa: BLE001
            pass
    style["align"] = _align_of(para)
    if style["size_pt"] is None:
        style["size_pt"] = inh.get("sz")
    if style["bold"] is None:
        style["bold"] = inh.get("bold")
    if not style["color"]:
        style["color"] = inh.get("color")
    if not style["font"]:
        style["font"] = inh.get("font")
    return style


def _tokenize_footer(text: str) -> tuple[str, str | None]:
    """日付を {date} に、ページ番号を {page} に置き換える。返り値: (置換後, 日付書式 or None)"""
    fmt = None
    for pat, f in _DATE_PATTERNS:
        if pat.search(text):
            text = pat.sub("{date}", text, count=1)
            fmt = f
            break
    text = re.sub(r"[‹<]#[›>]", "{page}", text)
    return text, fmt


def _is_page_number(s: _Shape) -> bool:
    if s.ph == PP_PLACEHOLDER.SLIDE_NUMBER or any(f == "slidenum" for f in s.fields):
        return True
    return bool(s.text) and bool(_PAGE_RE.match(s.text) or _PAGE_TOTAL_RE.match(s.text))


# ---------------------------------------------------------------- 背景
def _background(slide: Any, theme: ThemeInfo) -> dict | None:
    """スライド → レイアウト → マスター の順に背景（単色 or 画像）を探す。"""
    chain = (("slide", slide), ("layout", slide.slide_layout), ("master", slide.slide_layout.slide_master))
    for source, owner in chain:
        bg = owner._element.find("p:cSld/p:bg", NS)
        if bg is None:
            continue
        bgpr = bg.find("p:bgPr", NS)
        if bgpr is not None:
            blip = bgpr.find("a:blipFill/a:blip", NS)
            if blip is not None and blip.get(f"{{{NS['r']}}}embed"):
                try:
                    part = owner.part.related_part(blip.get(f"{{{NS['r']}}}embed"))
                    return {"source": source, "image": part.blob, "ext": (part.partname.ext or "png").lstrip("."), "color": None}
                except Exception:  # noqa: BLE001
                    pass
            color = _xml_color(bgpr, theme)
            if color:
                return {"source": source, "image": None, "ext": None, "color": color}
            if bgpr.find("a:gradFill", NS) is not None:
                first = bgpr.find(".//a:gs", NS)
                color = _xml_color(first, theme) if first is not None else None
                return {"source": source, "image": None, "ext": None, "color": color, "gradient": True}
        ref = bg.find("p:bgRef", NS)
        if ref is not None:
            sch = ref.find("a:schemeClr", NS)
            if sch is not None:
                color = theme.scheme_hex(sch.get("val", ""))
                if color:
                    return {"source": source, "image": None, "ext": None, "color": color}
    return None


# ---------------------------------------------------------------- 解析本体
class _Assets:
    """画像を <assets_dir> に書き、同じ内容は同じファイルにする。"""

    def __init__(self, directory: Path):
        self.dir = directory
        self.by_hash: dict[str, str] = {}

    def put(self, blob: bytes, ext: str, stem: str) -> str:
        digest = hashlib.sha1(blob).hexdigest()
        if digest in self.by_hash:
            return self.by_hash[digest]
        ext = (ext or "png").lower().lstrip(".")
        if ext == "jpeg":
            ext = "jpg"
        path = self.dir / f"{stem}.{ext}"
        n = 2
        while path.exists():
            path = self.dir / f"{stem}_{n}.{ext}"
            n += 1
        path.write_bytes(blob)
        rel = template_store.rel_path(path)
        self.by_hash[digest] = rel
        return rel


def _to_base(box: dict, sx: float, sy: float) -> dict:
    """スライド座標（pt）→ 960×540 基準（`template_kit.scale` の逆変換）。"""
    return {"x": round(box["x"] * sx, 1), "y": round(box["y"] * sy, 1), "w": round(box["w"] * sx, 1), "h": round(box["h"] * sy, 1)}


def classify_slides(n: int, texts: list[int], overrides: dict[int, str] | None = None, scores: list[float] | None = None) -> list[str]:
    """各スライドの役割（cover / content / closing / skip）。texts は各スライドの文字数。

    scores（中身らしさ: 題名・本文プレースホルダの有無と文字図形の数）があれば、中間のスライドのうち
    最も点の高い 1 枚を中身にし、残りは skip にする（色見本や装飾のスライドが中身に選ばれないように）。
    """
    roles: list[str] = []
    for i in range(n):
        if i == 0:
            roles.append("cover")
        elif i == n - 1 and (n >= 3 or (n == 2 and texts[i] < 40)):
            roles.append("closing")
        else:
            roles.append("content")
    if scores is not None and len(scores) == n:
        middle = [i for i, r in enumerate(roles) if r == "content"]
        if len(middle) > 1:
            best = max(middle, key=lambda i: (scores[i], -i))
            for i in middle:
                if i != best:
                    roles[i] = "skip"
    for i, r in (overrides or {}).items():
        if 0 <= int(i) < n and r in ("cover", "content", "closing", "skip"):
            roles[int(i)] = r
    return roles


def _text_len(slide: Any) -> int:
    n = 0
    for sh in slide.shapes:
        try:
            if sh.has_text_frame:
                n += len(sh.text_frame.text.strip())
        except Exception:  # noqa: BLE001
            pass
    return n


def _content_score(slide: Any, cw: float, ch: float) -> float:
    """中身スライドらしさ。題名・本文プレースホルダがあり、文字図形が多く、同じ大きさの塗り図形（色見本）が少ないほど高い。"""
    score = 0.0
    fills: list[tuple[float, float]] = []
    for sh in slide.shapes:
        ph = _placeholder_type(sh)
        if ph in _TITLE_TYPES:
            score += 2.0
        elif ph in _BODY_TYPES:
            score += 2.0
        try:
            has_text = sh.has_text_frame and bool(sh.text_frame.text.strip())
        except Exception:  # noqa: BLE001
            has_text = False
        if has_text:
            score += 0.5
        elif sh.shape_type not in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINE, MSO_SHAPE_TYPE.GROUP):
            try:
                if sh.width and sh.height:
                    fills.append((round(emu_to_pt(sh.width)), round(emu_to_pt(sh.height))))
            except Exception:  # noqa: BLE001
                pass
    score -= 0.5 * (len(fills) - len(set(fills)))  # 同じ大きさの塗りが並ぶほど減点
    return score


def _classify_fill(b: dict, cw: float, ch: float) -> str:
    """文字の無い塗り図形を 帯（bar）か 装飾（decor）に分ける。

    帯 = 細長く（縦横比 4 以上）、幅か高さがスライドの 30% 以上で、上下左右いずれかの端から 15% 以内。
    """
    w, h = float(b["w"]), float(b["h"])
    if w <= 0 or h <= 0:
        return "decor"
    aspect = max(w, h) / min(w, h)
    horizontal = w >= h
    span = (w >= cw * 0.3) if horizontal else (h >= ch * 0.3)
    # 横帯は上下の端、縦帯は左右の端に寄っているか（横長の図形は大抵左端から始まるので x では判定しない）
    near_edge = (b["y"] <= ch * 0.15 or b["y"] + h >= ch * 0.85) if horizontal else (b["x"] <= cw * 0.15 or b["x"] + w >= cw * 0.85)
    return "bar" if aspect >= 4 and span and near_edge else "decor"


def _find_swatches(fills: list[dict]) -> set[int]:
    """同じ大きさ（±10%）の塗り図形が SWATCH_MIN 個以上あれば色見本と見なし、その index を返す。"""
    groups: list[list[int]] = []
    for i, f in enumerate(fills):
        placed = False
        for g in groups:
            r = fills[g[0]]
            if abs(f["_w"] - r["_w"]) <= max(2.0, r["_w"] * 0.1) and abs(f["_h"] - r["_h"]) <= max(2.0, r["_h"] * 0.1):
                g.append(i)
                placed = True
                break
        if not placed:
            groups.append([i])
    out: set[int] = set()
    for g in groups:
        if len(g) >= SWATCH_MIN and len({fills[i]["color"] for i in g}) >= 2:
            out.update(g)
    return out


def _hsl(hex_color: str) -> tuple[float, float, float]:
    try:
        r, g, b = (int(hex_color[i : i + 2], 16) / 255.0 for i in (1, 3, 5))
    except (ValueError, TypeError, IndexError):
        return (0.0, 0.0, 0.5)
    mx, mn = max(r, g, b), min(r, g, b)
    light = (mx + mn) / 2
    if mx == mn:
        return (0.0, 0.0, light)
    d = mx - mn
    sat = d / (1 - abs(2 * light - 1)) if (1 - abs(2 * light - 1)) else 0.0
    if mx == r:
        hue = ((g - b) / d) % 6
    elif mx == g:
        hue = (b - r) / d + 2
    else:
        hue = (r - g) / d + 4
    return (hue * 60.0, sat, light)


def _derive_colors(base: dict, candidates: list[str]) -> dict:
    """色見本などの候補色から primary / accent / text / muted / surface / line を決め、テーマ色に上書きする。

    候補が 3 色未満なら何もしない。判定は明度と彩度だけ（暗く鮮やか → primary、鮮やかで primary と色相が離れる → accent、
    暗い無彩色 → text、中間の無彩色 → muted、明るい色 → surface / line）。
    """
    uniq: list[str] = []
    for c in candidates:
        c = str(c or "").upper()
        if re.fullmatch(r"#[0-9A-F]{6}", c) and c not in uniq:
            uniq.append(c)
    if len(uniq) < 3:
        return base
    out = dict(base)
    info = {c: _hsl(c) for c in uniq}
    chroma = [c for c in uniq if info[c][1] >= 0.25 and 0.15 <= info[c][2] <= 0.7]
    grays = [c for c in uniq if info[c][1] < 0.25]
    if chroma:
        primary = min(chroma, key=lambda c: info[c][2])
        out["primary"] = primary
        others = [c for c in chroma if c != primary and abs(info[c][0] - info[primary][0]) % 360 > 25]
        if others:
            out["accent"] = max(others, key=lambda c: info[c][1] + (0.3 if info[c][2] > 0.35 else 0))
    dark = [c for c in grays if info[c][2] < 0.35]
    if dark:
        out["text"] = min(dark, key=lambda c: info[c][2])
    mids = [c for c in grays if 0.35 <= info[c][2] < 0.75]
    if mids:
        out["muted"] = mids[0]
    light = sorted([c for c in uniq if info[c][2] >= 0.75], key=lambda c: -info[c][2])
    if light:
        out["surface"] = light[0]
        if len(light) > 1:
            out["line"] = light[1]
    return out


def _slide_title(slide: Any) -> str:
    first = ""
    for sh in slide.shapes:
        try:
            if not sh.has_text_frame or not sh.text_frame.text.strip():
                continue
            t = sh.text_frame.text.strip().splitlines()[0][:40]
            if _placeholder_type(sh) in _TITLE_TYPES:
                return t
            first = first or t
        except Exception:  # noqa: BLE001
            continue
    return first


def _theme_palette(theme: ThemeInfo) -> tuple[dict, dict]:
    c = theme.colors
    dk1, lt1 = c.get("dk1", "#222222"), c.get("lt1", "#FFFFFF")
    dk2, lt2 = c.get("dk2", "#44546A"), c.get("lt2", "#E7E6E6")
    colors = {
        "primary": c.get("accent1", dk2),
        "secondary": dk2,
        "accent": c.get("accent2", c.get("accent1", "#E07A1F")),
        "background": lt1,
        "surface": lt2,
        "text": dk1,
        "muted": apply_brightness(dk1, 0.35),
        "line": apply_brightness(lt2, -0.15),
    }
    fonts = {
        "heading": theme.fonts.get("major_ea") or theme.fonts.get("major_latin") or "Meiryo",
        "body": theme.fonts.get("minor_ea") or theme.fonts.get("minor_latin") or "Meiryo",
    }
    return colors, fonts


def _analyze_part(kind: str, slide: Any, slide_no: int, theme: ThemeInfo, cw: float, ch: float, assets: _Assets, warnings: list[str]) -> dict:
    """1 枚のスライドから部品定義（960×540 基準）を作る。"""
    sx, sy = BASE_W / cw, BASE_H / ch
    part: dict[str, Any] = {}
    slide_theme = theme.with_override(TextStyleResolver.clr_map_override(slide))
    token = set_current_theme(slide_theme)
    try:
        resolver = TextStyleResolver(slide, slide_theme)
        bg = _background(slide, slide_theme)
        if bg:
            if bg.get("image"):
                part["background_image"] = assets.put(bg["image"], bg["ext"], f"{kind}_background")
                part["background_source"] = bg["source"]
            elif bg.get("color"):
                part["background_color"] = bg["color"]
                if bg.get("gradient"):
                    warnings.append(f"{kind}: グラデーション背景は先頭色の単色で近似しました。")
        shapes = collect_shapes(slide, slide_no)
        images: list[dict] = []
        bars: list[dict] = []
        fills: list[dict] = []
        texts: list[tuple[_Shape, dict]] = []
        for s in shapes:
            b = s.box
            ratio = s.area / (cw * ch)
            if s.is_picture:
                try:
                    img = s.shape.image
                    blob, ext = img.blob, img.ext
                except Exception as e:  # noqa: BLE001
                    warnings.append(f"{kind}: 画像 '{s.describe()}' を読めません: {e}")
                    continue
                near_edge = b["x"] <= cw * 0.15 or b["x"] + b["w"] >= cw * 0.85 or b["y"] <= ch * 0.15 or b["y"] + b["h"] >= ch * 0.85
                if ratio >= 0.8:
                    if "background_image" not in part:
                        part["background_image"] = assets.put(blob, ext, f"{kind}_background")
                        part["background_source"] = s.source
                    continue
                entry = {**_to_base(b, sx, sy), "source": s.source, "_desc": s.describe(), "_blob": blob, "_ext": ext}
                entry["_logo"] = b["w"] <= cw * 0.25 and near_edge
                if kind == "closing" and not entry["_logo"] and b["w"] <= cw * 0.5:
                    entry["_logo"] = True  # 最終ページ中央の大きめロゴ
                images.append(entry)
                continue
            if s.text:
                texts.append((s, b))
                continue
            if s.is_line:
                color = None
                try:
                    color = _rgb_of(s.shape.line.color)
                    width = float(s.shape.line.width.pt) if s.shape.line.width else 1.0
                except Exception:  # noqa: BLE001
                    width = 1.0
                if b["w"] >= b["h"] and b["w"] >= cw * 0.3:
                    bars.append({**_to_base({"x": b["x"], "y": b["y"] - width / 2, "w": b["w"], "h": max(width, 1.0)}, sx, sy), "color": color or "#666666", "slant_pt": 0, "source": s.source, "_desc": s.describe(), "confidence": "high" if s.source == "slide" else "low"})
                continue
            # 文字の無い図形: 塗りがあれば帯、面積 80% 以上なら背景色
            fill = _rgb_of(s.shape.fill.fore_color) if _fill_is_solid(s.shape) else None
            if not fill:
                continue
            if s.ph is not None and s.ph in _TITLE_TYPES | _BODY_TYPES | {PP_PLACEHOLDER.SUBTITLE}:
                texts.append((s, b))  # 塗り付きの空プレースホルダは位置として使う
                continue
            if ratio >= 0.8:
                part.setdefault("background_color", fill)
                continue
            if ratio < 0.0005:
                continue
            slant = 0.0
            try:
                if "PARALLELOGRAM" in str(s.shape.auto_shape_type):
                    slant = float(s.shape.adjustments[0]) * min(b["w"], b["h"])
            except Exception:  # noqa: BLE001
                slant = 0.0
            fills.append({**_to_base(b, sx, sy), "color": fill, "slant_pt": round(slant * sx, 1), "source": s.source, "_desc": s.describe(), "_kind": _classify_fill(b, cw, ch), "_w": b["w"], "_h": b["h"], "_ratio": ratio})

        # --- 塗り図形: 色見本を除き、帯と装飾に分ける
        swatch_idx = _find_swatches(fills)
        if swatch_idx:
            palette = []
            for i in sorted(swatch_idx):
                if fills[i]["color"] not in palette:
                    palette.append(fills[i]["color"])
            part["_palette"] = palette
            warnings.append(f"{kind}: 同じ大きさの塗り図形 {len(swatch_idx)} 個は色見本と判断し、部品にはせず色をテーマ候補に入れました。")
        decor: list[dict] = []
        for i, f in enumerate(fills):
            if i in swatch_idx:
                continue
            entry = {k: v for k, v in f.items() if k not in ("_kind", "_w", "_h", "_ratio")}
            if f["_kind"] == "bar":
                entry["confidence"] = "high" if f["source"] == "slide" else "low"
                bars.append(entry)
            else:
                entry["confidence"] = "low"
                decor.append(entry)

        # --- 画像: ロゴ（最初の 1 つ）と装飾画像
        logo_done = False
        extra: list[dict] = []
        for img in sorted(images, key=lambda i: (not i["_logo"], i["y"], i["x"])):
            if img["_logo"] and not logo_done:
                rel = assets.put(img["_blob"], img["_ext"], f"{kind}_logo")
                part["logo"] = {"image": rel, "x": img["x"], "y": img["y"], "w": img["w"], "source": img["source"], "confidence": "high" if img["source"] == "slide" else "low"}
                logo_done = True
            else:
                rel = assets.put(img["_blob"], img["_ext"], f"{kind}_image{len(extra) + 1}")
                extra.append({"image": rel, "x": img["x"], "y": img["y"], "w": img["w"], "h": img["h"], "source": img["source"]})
        if extra:
            part["images"] = extra[:6]
            if len(extra) > 6:
                warnings.append(f"{kind}: 画像が多いため 6 つまでを部品にしました。")
        if bars:
            bars.sort(key=lambda b: -(b["w"] * b["h"]))
            if len(bars) > MAX_BARS:
                warnings.append(f"{kind}: 帯らしい図形が {len(bars)} 個あるため、大きい {MAX_BARS} 個を帯にし、残りは装飾（既定では使わない）にしました。")
                decor = bars[MAX_BARS:] + decor
                bars = bars[:MAX_BARS]
            for b in bars:
                b.pop("_desc", None)
            part["bar"] = bars[0]
            if len(bars) > 1:
                part["bars"] = bars[1:]
        if decor:
            for d in decor:
                d.pop("_desc", None)
                d["enabled"] = False
            part["decor"] = decor[:12]

        # --- 文字: 役割の判定
        bottom_zone = ch * 0.85
        title_s = subtitle_s = body_s = None
        footers: list[tuple[_Shape, dict]] = []
        page_s = None
        others: list[tuple[_Shape, dict]] = []
        for s, b in texts:
            if s.ph in _TITLE_TYPES:
                title_s = title_s or (s, b)
            elif s.ph == PP_PLACEHOLDER.SUBTITLE:
                subtitle_s = subtitle_s or (s, b)
            elif s.ph in _BODY_TYPES:
                body_s = body_s or (s, b)
            elif _is_page_number(s) and (b["y"] >= ch * 0.7 or b["y"] + b["h"] <= ch * 0.15):
                page_s = page_s or (s, b)
            elif s.ph in (PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE) or (b["y"] >= bottom_zone - b["h"] and b["h"] <= ch * 0.12 and len(s.text) <= 120):
                footers.append((s, b))
            else:
                others.append((s, b))
        inferred: set[int] = set()  # プレースホルダ以外から推定した（信頼度 low）
        if kind == "cover" and (title_s is None or subtitle_s is None) and others:
            others.sort(key=lambda t: -(_text_style(t[0], resolver).get("size_pt") or 0))
            if title_s is None:
                title_s = others.pop(0)
                inferred.add(id(title_s[0]))
            if subtitle_s is None and others:
                subtitle_s = others.pop(0)
                inferred.add(id(subtitle_s[0]))
        if kind == "content" and title_s is None:
            top = [t for t in others if t[1]["y"] <= ch * 0.25]
            if top:
                title_s = sorted(top, key=lambda t: t[1]["y"])[0]
                others.remove(title_s)
                inferred.add(id(title_s[0]))
        if kind == "content" and body_s is None:
            big = [t for t in others if t[1]["h"] >= ch * 0.25]
            if big:
                body_s = max(big, key=lambda t: t[1]["w"] * t[1]["h"])
                others.remove(body_s)
                inferred.add(id(body_s[0]))

        def text_spec(s: _Shape, b: dict, default_size: float) -> dict:
            st = _text_style(s, resolver)
            spec = {**_to_base(b, sx, sy), "size_pt": float(st["size_pt"] or default_size), "color": st["color"] or "#222222", "bold": bool(st["bold"]), "align": st["align"] or "left", "source": s.source, "confidence": "low" if id(s) in inferred else "high"}
            if st.get("font"):
                spec["font"] = st["font"]
            return spec

        if title_s:
            part["title"] = text_spec(title_s[0], title_s[1], 32 if kind == "cover" else 28)
        if subtitle_s and kind == "cover":
            part["subtitle"] = text_spec(subtitle_s[0], subtitle_s[1], 20)
        if kind == "content":
            if body_s:
                part["body"] = {**_to_base(body_s[1], sx, sy), "source": body_s[0].source, "confidence": "low" if id(body_s[0]) in inferred else "high"}
            else:
                # 題名の下から下部の部品（帯・フッター・ロゴ）の上までを本文領域にする
                top = (part["title"]["y"] + part["title"]["h"] + 12) if part.get("title") else 36
                floor = BASE_H - 36
                for item in list(bars) + ([part["logo"]] if part.get("logo") else []) + [dict(_to_base(b, sx, sy)) for _s, b in footers]:
                    if item.get("y", 0) > BASE_H * 0.6:
                        floor = min(floor, item["y"] - 12)
                x = part["title"]["x"] if part.get("title") else 36
                w = part["title"]["w"] if part.get("title") else BASE_W - 72
                if floor - top >= 100:
                    part["body"] = {"x": x, "y": round(top, 1), "w": w, "h": round(floor - top, 1), "source": "auto", "confidence": "low"}
        if kind == "closing":
            cand = others or ([title_s] if title_s else []) or ([subtitle_s] if subtitle_s else [])
            if cand:
                s, b = max(cand, key=lambda t: t[1]["w"] * t[1]["h"])
                spec = text_spec(s, b, 16)
                spec["align"] = spec["align"] if spec["align"] != "left" else "center"
                part["message"] = spec
        if page_s:
            s, b = page_s
            spec = text_spec(s, b, 9)
            text = "{page}"
            if _PAGE_TOTAL_RE.match(s.text or ""):
                text = "{page} / {total}"
            part["page_number"] = {**spec, "text": text, "align": spec["align"] if spec["align"] != "left" else ("right" if b["x"] > cw * 0.5 else "left")}
        if footers:
            # 複数あるときは横に広いものを本体にし、残りは装飾として無視（警告）
            footers.sort(key=lambda t: -t[1]["w"])
            s, b = footers[0]
            text, fmt = _tokenize_footer(s.text)
            if any(f.startswith("datetime") for f in s.fields) or s.ph == PP_PLACEHOLDER.DATE:
                if "{date}" not in text:
                    text = "{date}"
            spec = text_spec(s, b, 9)
            spec["confidence"] = "high" if (s.ph in (PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE) or fmt or s.fields) else "low"
            part["footer"] = {**spec, "text": text}
            if fmt:
                part["date_format"] = fmt
            if len(footers) > 1:
                warnings.append(f"{kind}: 下部の文字が {len(footers)} 個あるため、一番幅の広いもの（'{s.text[:20]}'）をフッターにしました。")
        candidates: list[dict] = []
        for s, b in others:
            if kind == "closing":
                continue
            if b["y"] >= bottom_zone - b["h"]:
                reason = "下部にあるが幅が狭い、または文字が長い（フッターの候補）"
            elif b["y"] <= ch * 0.25:
                reason = "上部にあるが題名は別に決まった（副題の候補）"
            elif b["h"] < ch * 0.25:
                reason = "本文領域より小さい（説明文や注記の候補）"
            else:
                reason = "題名・本文・フッターのどれにも当たらない位置"
            spec = text_spec(s, b, 12)
            candidates.append({**spec, "text": s.text[:60], "reason": reason, "confidence": "low"})
            warnings.append(f"{kind}: 文字 '{s.text[:20]}'（{s.describe()}）は部品に含めませんでした（{reason}）。画面の「文字候補」で役割を選べます。")
        if candidates:
            part["candidates"] = candidates[:8]
    finally:
        reset_current_theme(token)
    return part


def analyze(data: bytes, filename: str, template_id: str | None = None, roles: dict[int, str] | None = None, name: str | None = None) -> dict:
    """PPTX を解析し、テンプレート定義の提案を返す。画像と土台 PPTX は assets_dir(id) に保存する。

    返り値: {"proposal": template, "parts": [...], "slides": [{index, role, title}], "warnings": [str]}
    """
    prs = Presentation(io.BytesIO(data))
    n = len(prs.slides)
    if n == 0:
        raise ValueError("スライドがありません。表紙・中身・最終ページの 1〜3 枚程度の PPTX を投入してください。")
    cw, ch = emu_to_pt(prs.slide_width), emu_to_pt(prs.slide_height)
    stem = Path(filename).stem or "template"
    tid = template_store.safe_id(template_id or stem)
    directory = template_store.assets_dir(tid)
    for old in directory.iterdir():  # 作り直し: 前回の画像を消す
        try:
            old.unlink()
        except OSError:
            pass
    assets = _Assets(directory)
    warnings: list[str] = []

    texts = [_text_len(s) for s in prs.slides]
    scores = [_content_score(s, cw, ch) for s in prs.slides]
    role_list = classify_slides(n, texts, roles, scores)
    slides_info = [{"index": i, "role": role_list[i], "title": _slide_title(s), "chars": texts[i]} for i, s in enumerate(prs.slides)]
    master = prs.slides[0].slide_layout.slide_master
    theme = ThemeInfo(master)
    colors, fonts = _theme_palette(theme)

    proposal: dict[str, Any] = {
        "id": tid,
        "name": name or f"{stem}（PPTX から作成）",
        "description": f"{filename} から推定したテンプレート。位置は「PPTX からテンプレート作成」画面で修正できます。",
        "source": "user",
        "fonts": fonts,
        "colors": colors,
        "footer": {"enabled": False, "text": "", "show_page_number": False},
        "confidential_mark": {"enabled": False, "text": ""},
        "date_format": "%Y-%m-%d",
    }
    for kind in _KINDS:
        idx = next((i for i, r in enumerate(role_list) if r == kind), None)
        if idx is None:
            if kind == "content" and n >= 1:
                # 中身スライドが無い場合は表紙のレイアウトのマスターから題名・本文の位置だけ取る
                warnings.append("中身のスライドが無いため、マスターの題名・本文プレースホルダから中身の部品を推定しました。")
                part = _analyze_part("content", prs.slides[0], 1, theme, cw, ch, assets, [])
                part = {k: v for k, v in part.items() if k in ("title", "body", "background_color", "logo", "bar", "bars", "footer", "page_number", "images", "background_image", "background_source")}
                if part:
                    proposal["content"] = part
            continue
        proposal[kind] = _analyze_part(kind, prs.slides[idx], idx + 1, theme, cw, ch, assets, warnings)
    palette: list[str] = []
    for kind in _KINDS:
        part = proposal.get(kind)
        if isinstance(part, dict) and part.get("_palette"):
            palette.extend(part.pop("_palette"))
    for i, r in enumerate(role_list):
        if r == "skip":  # 色見本のスライドは使わないが、色だけはテーマ候補に回す
            for c in _palette_from_slide(prs.slides[i], cw, ch):
                if c not in palette:
                    palette.append(c)
    if palette:
        warnings.append(f"色見本らしい塗り図形から {len(palette)} 色をテーマ候補にしました（primary / accent / text などに反映）。")
    if palette:
        proposal["colors"] = _derive_colors(proposal["colors"], palette)
        proposal["palette_candidates"] = palette
    proposal["guide"] = {kind: _guide_of(proposal[kind]) for kind in _KINDS if isinstance(proposal.get(kind), dict)}
    if proposal.get("content", {}).get("date_format"):
        proposal["date_format"] = proposal["content"]["date_format"]
    elif proposal.get("cover", {}).get("date_format"):
        proposal["date_format"] = proposal["cover"]["date_format"]
    # 土台として使うため元の PPTX も残す
    base = directory / "base.pptx"
    base.write_bytes(data)
    proposal["base_pptx"] = template_store.rel_path(base)
    proposal["base_canvas"] = {"width_pt": round(cw, 2), "height_pt": round(ch, 2)}
    if theme.unresolved:
        warnings.append("テーマ色の一部を解決できませんでした: " + ", ".join(sorted(theme.unresolved)))
    log.info("テンプレート推定: %s slides=%d parts=%s warnings=%d", filename, n, [k for k in _KINDS if k in proposal], len(warnings))
    return {"proposal": proposal, "parts": parts_of(proposal), "slides": slides_info, "warnings": warnings}


def _palette_from_slide(slide: Any, cw: float, ch: float) -> list[str]:
    """使わないスライド（色見本など）から、同じ大きさで並ぶ塗り図形の色だけを拾う。"""
    fills: list[dict] = []
    for sh, tf in _iter_shapes(slide.shapes):
        box = _box_of(sh, tf)
        if box is None:
            continue
        s = _Shape(sh, box, "slide", 0)
        if s.is_picture or s.is_line or s.text or s.ph is not None:
            continue
        fill = _rgb_of(sh.fill.fore_color) if _fill_is_solid(sh) else None
        if not fill:
            continue
        fills.append({"color": fill, "_w": box["w"], "_h": box["h"]})
    idx = _find_swatches(fills)
    out: list[str] = []
    for i in sorted(idx):
        if fills[i]["color"] not in out:
            out.append(fills[i]["color"])
    return out


def _guide_of(part: dict) -> str:
    """部品の一文要約（例: 「背景画像、ロゴ、題名、副題、フッター」）。"""
    names: list[str] = []
    for key in ("background_image", "logo", "images", "bar", "bars", "title", "subtitle", "body", "message", "footer", "page_number"):
        v = part.get(key)
        if not v:
            continue
        if key in ("images", "bars"):
            names.append(f"{PART_LABELS[key]} {len(v)}")
        elif key == "bar":
            names.append("帯")
        else:
            names.append(PART_LABELS[key])
    if part.get("decor"):
        names.append(f"装飾 {len(part['decor'])}（既定は使わない）")
    if part.get("candidates"):
        names.append(f"文字候補 {len(part['candidates'])}")
    return "、".join(names) if names else "部品なし"


# ---------------------------------------------------------------- UI 用: 部品の一覧と座標の書き戻し
def _logo_box(spec: dict) -> dict | None:
    from .template_kit import image_size

    size = image_size(spec.get("image"))
    if not size:
        return None
    w = float(spec.get("w", 120))
    return {"x": float(spec.get("x", 0)), "y": float(spec.get("y", 0)), "w": w, "h": round(w * size[1] / size[0], 1)}


def parts_of(template: dict) -> list[dict]:
    """表紙 / 中身 / 最終 の各部品を、UI が番号付きの枠として描ける一覧にする（座標は 960×540 基準）。"""
    out: list[dict] = []
    for kind in _KINDS:
        part = template.get(kind)
        if not isinstance(part, dict):
            continue
        no = 0

        def add(key: str, label: str, box: dict | None, ptype: str, source: str | None = None, extra: dict | None = None) -> None:
            nonlocal no
            no += 1
            out.append({"part": kind, "key": key, "no": no, "label": label, "box": box, "type": ptype, "source": source or "slide", **(extra or {})})

        if part.get("background_image"):
            add("background_image", PART_LABELS["background_image"], None, "background", part.get("background_source"))
        if isinstance(part.get("logo"), dict):
            add("logo", PART_LABELS["logo"], _logo_box(part["logo"]), "image", part["logo"].get("source"))
        for i, img in enumerate(part.get("images") or []):
            add(f"images.{i}", f"画像 {i + 1}", {k: float(img[k]) for k in ("x", "y", "w", "h")}, "image", img.get("source"))
        if isinstance(part.get("bar"), dict):
            add("bar", PART_LABELS["bar"], {k: float(part["bar"][k]) for k in ("x", "y", "w", "h")}, "bar", part["bar"].get("source"), {"color": part["bar"].get("color")})
        for i, b in enumerate(part.get("bars") or []):
            add(f"bars.{i}", f"帯 {i + 2}", {k: float(b[k]) for k in ("x", "y", "w", "h")}, "bar", b.get("source"), {"color": b.get("color"), "confidence": b.get("confidence", "high")})
        for key in ("title", "subtitle", "body", "message", "footer", "page_number"):
            spec = part.get(key)
            if isinstance(spec, dict) and all(k in spec for k in ("x", "y", "w", "h")):
                extra = {k: spec.get(k) for k in ("text", "size_pt", "color", "confidence") if k in spec}
                add(key, PART_LABELS[key], {k: float(spec[k]) for k in ("x", "y", "w", "h")}, "area" if key == "body" else "text", spec.get("source"), extra)
        for i, d in enumerate(part.get("decor") or []):
            add(f"decor.{i}", f"装飾 {i + 1}", {k: float(d[k]) for k in ("x", "y", "w", "h")}, "decor", d.get("source"), {"color": d.get("color"), "enabled": bool(d.get("enabled")), "confidence": "low"})
        for i, c in enumerate(part.get("candidates") or []):
            add(f"candidates.{i}", f"文字候補 {i + 1}", {k: float(c[k]) for k in ("x", "y", "w", "h")}, "candidate", c.get("source"), {"text": c.get("text"), "size_pt": c.get("size_pt"), "reason": c.get("reason"), "enabled": False, "confidence": "low"})
    return out


# 役割の変更で受け付ける先（UI のセレクトと同じ）
ROLE_TARGETS = ("bar", "decor", "logo", "images", "title", "subtitle", "body", "message", "footer", "page_number", "candidates", "skip")


def _pop_part(part: dict, key: str) -> dict | None:
    """key（"bar" / "bars.1" / "decor.0" …）の部品を提案から外して返す。"""
    if "." in key:
        base, idx = key.split(".", 1)
        items = part.get(base)
        if isinstance(items, list) and idx.isdigit() and int(idx) < len(items):
            spec = items.pop(int(idx))
            if not items:
                part.pop(base, None)
            return spec
        return None
    spec = part.pop(key, None)
    if key == "background_image":
        part.pop("background_source", None)
    if key == "bar" and part.get("bars"):
        part["bar"] = part["bars"].pop(0)
        if not part["bars"]:
            part.pop("bars")
    return spec if isinstance(spec, dict) else None


def set_part_role(template: dict, kind: str, key: str, role: str) -> dict:
    """画面で選んだ役割に部品を移す（帯 ⇄ 装飾、文字候補 → 題名/副題/フッター/一言/ページ番号、ロゴ ⇄ 画像、skip = 外す）。"""
    t = copy.deepcopy(template)
    part = t.get(kind)
    if not isinstance(part, dict) or role not in ROLE_TARGETS:
        return t
    spec = _pop_part(part, key)
    if spec is None or role == "skip":
        return t
    spec = dict(spec)
    spec.pop("enabled", None)
    spec.pop("reason", None)
    spec["confidence"] = "high"  # 人が決めた
    if role in ("bar", "decor", "images", "candidates"):
        if role == "bar":
            spec.setdefault("color", "#000000")
            spec.setdefault("slant_pt", 0)
            if not isinstance(part.get("bar"), dict):
                part["bar"] = spec
            else:
                part.setdefault("bars", []).append(spec)
        elif role == "decor":
            spec.setdefault("color", "#000000")
            spec["enabled"] = False
            part.setdefault("decor", []).append(spec)
        elif role == "images":
            if spec.get("image"):
                part.setdefault("images", []).append({k: spec[k] for k in ("image", "x", "y", "w", "h", "source") if k in spec})
        else:
            spec["enabled"] = False
            part.setdefault("candidates", []).append(spec)
        return t
    if role == "logo":
        if spec.get("image"):
            part["logo"] = {k: spec[k] for k in ("image", "x", "y", "w", "source", "confidence") if k in spec}
        return t
    # 文字系の役割
    text = str(spec.get("text") or "")
    if role == "page_number":
        spec["text"] = "{page}"
        spec.setdefault("size_pt", 9)
        spec["align"] = spec.get("align") or "right"
    elif role == "footer":
        tok, fmt = _tokenize_footer(text)
        spec["text"] = tok
        spec.setdefault("size_pt", 9)
        if fmt:
            t["date_format"] = fmt
    elif role == "body":
        spec = {k: spec[k] for k in ("x", "y", "w", "h", "source", "confidence")}
    else:
        spec.pop("text", None)
        if role == "message":
            spec["align"] = spec.get("align") if spec.get("align") not in (None, "left") else "center"
    part[role] = spec
    return t


def set_part_box(template: dict, kind: str, key: str, box: dict) -> dict:
    """UI で動かした枠の座標を提案へ書き戻す（ロゴは x/y/w のみ。高さは画像の縦横比で決まる）。"""
    t = copy.deepcopy(template)
    part = t.get(kind)
    if not isinstance(part, dict):
        return t
    target: Any = None
    if "." in key:
        base, idx = key.split(".", 1)
        items = part.get(base)
        if isinstance(items, list) and idx.isdigit() and int(idx) < len(items):
            target = items[int(idx)]
    else:
        target = part.get(key)
    if not isinstance(target, dict):
        return t
    for c in ("x", "y", "w", "h"):
        if c in box and box[c] is not None:
            if key == "logo" and c == "h":
                continue
            target[c] = round(float(box[c]), 1)
    return t


def sample_presentation(template: dict) -> dict:
    """解釈を確認するためのプレビュー資料（表紙 / 中身 / 最終ページ の 3 枚）。座標は 960×540。"""
    p = new_presentation("テンプレートの確認", "manual", "", None)
    p["canvas"] = {"width_pt": BASE_W, "height_pt": BASE_H, "aspect": "16:9"}
    p["theme"] = {"template_id": template.get("id", "user"), "fonts": dict(template.get("fonts", {})), "colors": dict(template.get("colors", {}))}
    cover = new_slide("tpl_cover", 0, layout="title", title="表紙の題名")
    cover["elements"] = [
        text_element("tpl_cover_title", [paragraph([run("表紙の題名（サンプル）", bold=True)])], role="title"),
        text_element("tpl_cover_sub", [paragraph([run("副題・発表者・日付などの副題")])], role="subtitle"),
    ]
    content = new_slide("tpl_content", 1, layout="title_body", title="中身の題名")
    content["elements"] = [
        text_element("tpl_content_title", [paragraph([run("中身の題名（サンプル）", bold=True)])], role="title"),
        text_element(
            "tpl_content_body",
            [
                paragraph([run("本文はこの領域（本文領域）に流し込まれます")], bullet="bullet"),
                paragraph([run("帯・ロゴ・フッターと重ならない位置に置かれます")], bullet="bullet"),
                paragraph([run("長い本文は自動で縮小・分割されます")], bullet="bullet"),
                paragraph([run("枠をドラッグして位置と大きさを直せます")], bullet="bullet"),
            ],
            role="body",
        ),
    ]
    closing = new_slide("tpl_closing", 2, layout="closing", title="最終ページ")
    closing["elements"] = [text_element("tpl_closing_msg", [paragraph([run("ご清聴ありがとうございました")], align="center")], role="body")]
    p["slides"] = [cover, content, closing]
    return p
