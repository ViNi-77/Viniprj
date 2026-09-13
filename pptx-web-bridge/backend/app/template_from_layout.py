"""スライドマスター / レイアウトを直接読むテンプレート方式（Phase N）。

`template_from_pptx.py` はスライドの見た目から部品を推測するが、推測は必ず外れる。
PowerPoint のテンプレートは、ロゴ・帯・フッターなどの設計が**スライドマスターとレイアウトに
明示的に入っている**ので、そこを読めば推測は要らない。

ただし日本の会社で出回る「テンプレート」は 1 ページ目に帯とロゴを貼ったものを複製する運用も多く、
その場合レイアウトは空なので、この方式は使えない。**レイアウト優先 + 明示的フォールバック**にして、
どちらを使うかと理由を必ず画面に出す（黙って 0 個にしない）。

描画は既存の `pptx_parser` をそのまま使う。`parse_pptx` はスライド本体の図形しか描かないので、
レイアウトを通せばそのレイアウトの装飾だけが正確に出る。レイアウトはスライドではないため、
`_LayoutAsSlide` で必要な属性だけを見せかける。
"""
from __future__ import annotations

import base64
import hashlib
import io
from pathlib import Path
from typing import Any

from pptx import Presentation

from . import pptx_parser, template_store
from .logging_setup import get_logger
from .model import new_presentation, new_slide, warning
from .pptx_styles import ThemeInfo, reset_current_theme, set_current_theme
from .units import emu_to_pt

log = get_logger("template_from_layout")

_KINDS = ("cover", "content", "closing")

#: 面積がこれ未満（キャンバス比）の図形は装飾として数えない
_MIN_AREA_RATIO = 0.00005


class _LayoutAsSlide:
    """レイアウト（または マスター）を `_parse_slide_shapes` に渡すための見せかけ。

    `pptx_parser` がスライドに求めるのは shapes / slide_layout / part / _element / background /
    has_notes_slide だけ。レイアウト自身をレイアウトとして返すと、文字スタイルの継承も正しく働く。
    """

    def __init__(self, layout: Any):
        self._layout = layout
        self.slide_layout = layout
        self.part = layout.part
        self._element = layout._element
        self.background = layout.background
        self.has_notes_slide = False
        # レイアウトはマスターの装飾を継承する（showMasterSp="0" のときだけ継承しない）。
        # マスターの装飾を先に置いて背面にする。マスターのプレースホルダはレイアウト側と重なるので入れない。
        master_chrome: list[Any] = []
        if layout._element.get("showMasterSp") != "0":
            try:
                master_chrome = [sh for sh in layout.slide_master.shapes if not sh.is_placeholder]
            except Exception:  # noqa: BLE001
                master_chrome = []
        self.shapes = [*master_chrome, *layout.shapes]

    @property
    def slide_master(self) -> Any:
        return self._layout.slide_master


def _all_layouts(prs: Presentation) -> list[tuple[int, int, Any]]:
    """(マスター番号, レイアウト番号, レイアウト) を全マスター分。

    `prs.slide_layouts` はマスター 0 の分しか返さないので使わない（部署がテンプレートを
    統合すると 2 マスターになることがある）。
    """
    out: list[tuple[int, int, Any]] = []
    for mi, master in enumerate(prs.slide_masters):
        for li, layout in enumerate(master.slide_layouts):
            out.append((mi, li, layout))
    return out


def _chrome_count(container: Any, cw: float, ch: float) -> int:
    """装飾（ロゴ・帯など）として数えられる非プレースホルダ図形の数。"""
    n = 0
    for sh in container.shapes:
        try:
            if sh.is_placeholder:
                continue
            if None in (sh.width, sh.height):
                continue
            if emu_to_pt(sh.width) * emu_to_pt(sh.height) < cw * ch * _MIN_AREA_RATIO:
                continue
        except Exception:  # noqa: BLE001 - 読めない図形は数えない
            continue
        n += 1
    return n


def _ph_types(layout: Any) -> list[str]:
    out: list[str] = []
    for ph in layout.placeholders:
        try:
            out.append(str(ph.placeholder_format.type).split(" ")[0].split(".")[-1])
        except Exception:  # noqa: BLE001
            continue
    return out


def enumerate_layouts(data: bytes, filename: str = "template.pptx") -> dict:
    """レイアウトの一覧と、それぞれの装飾の数を返す（解析も描画もしない。数えるだけ）。

    返り値: {"available": bool, "masters": [...], "layouts": [...], "warnings": [...]}
    """
    prs = Presentation(io.BytesIO(data))
    cw, ch = emu_to_pt(prs.slide_width), emu_to_pt(prs.slide_height)
    masters = []
    for mi, master in enumerate(prs.slide_masters):
        masters.append({"index": mi, "name": str(master.name or f"マスター {mi + 1}"), "chrome_count": _chrome_count(master, cw, ch)})
    layouts = []
    for mi, li, layout in _all_layouts(prs):
        layouts.append({
            "master": mi,
            "index": li,
            "name": str(layout.name or f"レイアウト {li + 1}"),
            "ph_types": _ph_types(layout),
            "chrome_count": _chrome_count(layout, cw, ch),
            "chrome_from_master": masters[mi]["chrome_count"],
        })
    available = any(la["chrome_count"] + la["chrome_from_master"] >= 1 for la in layouts)
    warnings: list[dict] = []
    if not available:
        warnings.append(warning(
            "template_from_layout", "LAYOUT_HAS_NO_CHROME",
            f"このファイルのスライドマスター / レイアウトには、ロゴや帯などの図形がありませんでした"
            f"（レイアウト {len(layouts)} 件を確認）。スライド上の図形から推測する方法に切り替えます。",
            fallback="ロゴや帯をレイアウトへ移した PowerPoint を使うと、推測せずに取り込めます。"))
    log.info("レイアウト一覧: %s masters=%d layouts=%d 装飾あり=%s", filename, len(masters), len(layouts), available)
    return {"available": available, "masters": masters, "layouts": layouts, "canvas": {"width_pt": round(cw, 2), "height_pt": round(ch, 2)}, "warnings": warnings}


def parse_layout(prs: Presentation, master: int, index: int, filename: str = "layout.pptx") -> dict:
    """1 つのレイアウトを Presentation JSON（1 枚）にする。描画は既存のパーサに任せる。"""
    masters = list(prs.slide_masters)
    if not (0 <= master < len(masters)):
        raise ValueError(f"マスター {master} がありません（{len(masters)} 件）")
    layouts = list(masters[master].slide_layouts)
    if not (0 <= index < len(layouts)):
        raise ValueError(f"レイアウト {index} がありません（マスター {master} には {len(layouts)} 件）")
    layout = layouts[index]

    cw, ch = emu_to_pt(prs.slide_width), emu_to_pt(prs.slide_height)
    pres = new_presentation(title=str(layout.name or "レイアウト"), source_type="pptx", filename=filename)
    pres["canvas"] = {**pres["canvas"], "width_pt": round(cw, 2), "height_pt": round(ch, 2)}
    ctx = pptx_parser._Ctx(pres, filename)
    sd = new_slide(ctx.ids.next("s"), 0)
    # テーマは「そのレイアウトが属するマスター」から取る（マスター 0 から取るとロゴが黒くなる）
    theme = ThemeInfo(layout.slide_master)
    token = set_current_theme(theme)
    shim = _LayoutAsSlide(layout)
    try:
        ctx.resolver = pptx_parser.TextStyleResolver(shim, theme)
        with pptx_parser.layout_mode():
            pptx_parser._parse_slide_shapes(shim, sd, ctx, pres, 0)
    finally:
        reset_current_theme(token)
    sd.pop("_has_title_ph", None)
    sd.pop("_has_subtitle", None)
    sd["layout"] = "title_body"
    pres["slides"] = [sd]
    return pres


def _scale_box(box: dict | None, sx: float, sy: float) -> dict | None:
    """スライド座標（pt）→ 960×540 基準。"""
    if not box:
        return None
    return {"x": round(float(box["x"]) * sx, 1), "y": round(float(box["y"]) * sy, 1),
            "w": round(float(box["w"]) * sx, 1), "h": round(float(box["h"]) * sy, 1)}


def _externalize(pres: dict, elements: list[dict], assets: "_AssetWriter") -> dict[str, dict]:
    """要素が指す画像を資産フォルダへ出し、テンプレート相対パスの表に置き換える。

    解析直後の要素は `presentation["assets"]` の id を指すだけなので、そのまま保存すると
    テンプレートを使うときに画像が見つからず、装飾が画面から消える（＝推測方式と同じ失敗）。
    """
    out: dict[str, dict] = {}
    for el in elements:
        aid = el.get("asset_id")
        if not aid:
            continue
        asset = (pres.get("assets") or {}).get(aid)
        if not asset:
            continue
        try:
            blob = base64.b64decode(asset.get("data_base64", ""))
        except Exception:  # noqa: BLE001
            continue
        ext = Path(str(asset.get("filename") or "x.png")).suffix.lstrip(".") or "png"
        rel = assets.put(blob, ext, f"layout_{aid}")
        out[aid] = {"path": rel, "mime": asset.get("mime") or "image/png",
                    "width_px": asset.get("width_px"), "height_px": asset.get("height_px")}
    return out


def build_part(prs: Presentation, master: int, index: int, cw: float, ch: float, assets: "_AssetWriter", filename: str) -> dict:
    """1 レイアウト → テンプレートの 1 部品（960×540 基準の装飾要素 + 領域の枠）。"""
    pres = parse_layout(prs, master, index, filename)
    sx, sy = 960.0 / cw if cw else 1.0, 540.0 / ch if ch else 1.0
    chrome: list[dict] = []
    for el in pres["slides"][0]["elements"]:
        if el["type"] not in ("shape", "image", "line", "table"):
            continue  # プレースホルダ由来の文字は「枠」として別に持つ（中身は資料側が入れる）
        e = dict(el)
        e["bbox"] = _scale_box(el.get("bbox"), sx, sy)
        chrome.append(e)
    part: dict[str, Any] = {"chrome_elements": chrome, "chrome_assets": _externalize(pres, chrome, assets)}
    part.update(region_boxes(prs, master, index))
    bg = (pres["slides"][0].get("background") or {}).get("color")
    if bg:
        part["background_color"] = bg
    return part


def region_boxes(prs: Presentation, master: int, index: int) -> dict[str, dict]:
    """レイアウトの題名・本文プレースホルダの枠（960×540 基準）。本文の流し込み先に使う。"""
    from pptx.enum.shapes import PP_PLACEHOLDER

    layout = list(list(prs.slide_masters)[master].slide_layouts)[index]
    cw, ch = emu_to_pt(prs.slide_width), emu_to_pt(prs.slide_height)
    sx, sy = 960.0 / cw if cw else 1.0, 540.0 / ch if ch else 1.0
    titles = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}
    bodies = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.VERTICAL_BODY, PP_PLACEHOLDER.VERTICAL_OBJECT}
    out: dict[str, dict] = {}
    for ph in layout.placeholders:
        try:
            kind = ph.placeholder_format.type
            if None in (ph.left, ph.top, ph.width, ph.height):
                continue
            box = {"x": round(emu_to_pt(ph.left) * sx, 1), "y": round(emu_to_pt(ph.top) * sy, 1),
                   "w": round(emu_to_pt(ph.width) * sx, 1), "h": round(emu_to_pt(ph.height) * sy, 1)}
        except Exception:  # noqa: BLE001
            continue
        if kind in titles and "title" not in out:
            out["title"] = box
        elif kind == PP_PLACEHOLDER.SUBTITLE and "subtitle" not in out:
            out["subtitle"] = box
        elif kind in bodies and "body" not in out:
            out["body"] = box
    return out


class _AssetWriter:
    """装飾の画像を <assets_dir> に書く（同じ内容は 1 ファイル）。`template_from_pptx._Assets` と同じ作法。"""

    def __init__(self, directory: Path):
        self.dir = directory
        self.by_hash: dict[str, str] = {}

    def put(self, blob: bytes, ext: str, stem: str) -> str:
        digest = hashlib.sha1(blob).hexdigest()
        if digest in self.by_hash:
            return self.by_hash[digest]
        ext = (ext or "png").lower().lstrip(".")
        if ext == "jpeg":
            ext = "jpg"
        path = self.dir / f"{stem}.{ext}"
        n = 2
        while path.exists():
            path = self.dir / f"{stem}_{n}.{ext}"
            n += 1
        path.write_bytes(blob)
        rel = template_store.rel_path(path)
        self.by_hash[digest] = rel
        return rel


def build_template(data: bytes, filename: str, layout_map: dict[str, dict], template_id: str | None = None, name: str | None = None) -> dict:
    """選んだレイアウトからテンプレート定義を作る。

    layout_map: {"cover": {"master": 0, "index": 0}, "content": {...}, "closing": {...}}
    役割が 1 つも指定されていなければ ValueError。
    """
    from .template_from_pptx import _theme_palette

    picked = {k: v for k, v in (layout_map or {}).items() if k in _KINDS and isinstance(v, dict)}
    if not picked:
        raise ValueError("表紙・中身・最終ページのうち、少なくとも 1 つにレイアウトを割り当ててください。")
    prs = Presentation(io.BytesIO(data))
    cw, ch = emu_to_pt(prs.slide_width), emu_to_pt(prs.slide_height)
    masters = list(prs.slide_masters)
    stem = Path(filename).stem or "template"
    tid = template_store.safe_id(template_id or stem)
    directory = template_store.assets_dir(tid)
    for old in directory.iterdir():  # 作り直し: 前回の画像を消す
        try:
            old.unlink()
        except OSError:
            pass
    assets = _AssetWriter(directory)

    first = next(iter(picked.values()))
    theme = ThemeInfo(masters[int(first.get("master", 0))])
    colors, fonts = _theme_palette(theme)
    proposal: dict[str, Any] = {
        "id": tid,
        "name": name or f"{stem}（レイアウトから作成）",
        "description": f"{filename} のスライドマスター / レイアウトから作ったテンプレート。推測はしていません。",
        "source": "user",
        "mode": "layout",
        "fonts": fonts,
        "colors": colors,
        "footer": {"enabled": False, "text": "", "show_page_number": False},
        "confidential_mark": {"enabled": False, "text": ""},
        "date_format": "%Y-%m-%d",
        "layout_map": {},
    }
    warnings: list[dict] = []
    for kind, ref in picked.items():
        mi, li = int(ref.get("master", 0)), int(ref.get("index", 0))
        try:
            proposal[kind] = build_part(prs, mi, li, cw, ch, assets, filename)
        except ValueError as e:
            warnings.append(warning("template_from_layout", "LAYOUT_NOT_FOUND", f"{kind}: {e}"))
            continue
        layout = list(masters[mi].slide_layouts)[li]
        proposal["layout_map"][kind] = {"master": mi, "index": li, "name": str(layout.name or "")}
    if not proposal["layout_map"]:
        raise ValueError("指定されたレイアウトが見つかりませんでした。")

    base = directory / "base.pptx"
    base.write_bytes(data)
    proposal["base_pptx"] = template_store.rel_path(base)
    proposal["base_canvas"] = {"width_pt": round(cw, 2), "height_pt": round(ch, 2)}
    log.info("レイアウトからテンプレート作成: %s → %s", filename, proposal["layout_map"])
    return {"proposal": proposal, "warnings": warnings}
