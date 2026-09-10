"""PPTX → Presentation JSON 変換器。

python-pptx で取れる情報を主に使い、取れないもの（SmartArt のテキスト、
グループ内座標、箇条書き種別）は OOXML（lxml）を直接読んで補完する。
未対応要素はエラーにせず、位置を残した unsupported 要素 + 警告に変換する。
"""
from __future__ import annotations

import base64
import hashlib
import io
from typing import Any

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.util import Emu

from .ids import IdFactory
from .logging_setup import get_logger
from .pptx_styles import TextStyleResolver, ThemeInfo, color_to_hex, current_theme, reset_current_theme, set_current_theme
from .model import (
    add_warning,
    bbox,
    cell,
    image_element,
    line_element,
    new_presentation,
    new_slide,
    paragraph,
    run,
    shape_element,
    table_element,
    text_element,
    unsupported_element,
    warning,
)
from .units import emu_to_pt

log = get_logger("pptx_parser")

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "dgm": "http://schemas.openxmlformats.org/drawingml/2006/diagram",
}

# 図形種別 → Presentation JSON の shape。ここに無いものは rect へ近似し警告する。
_SHAPE_MAP = {
    "RECTANGLE": "rect",
    "ROUNDED_RECTANGLE": "rounded_rect",
    "OVAL": "ellipse",
    "FLOWCHART_PROCESS": "rect",
    "FLOWCHART_ALTERNATE_PROCESS": "rounded_rect",
    "SNIP_1_RECTANGLE": "rect",
    "ISOSCELES_TRIANGLE": "triangle",
    "DIAMOND": "diamond",
    "RIGHT_ARROW": "arrow_right",
    "PARALLELOGRAM": "parallelogram",
}
_SKIP_PLACEHOLDERS = {PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.SLIDE_NUMBER}


def _rgb_of(color_obj: Any) -> str | None:
    """python-pptx の色オブジェクトから '#RRGGBB' を取る。テーマ色は現在のテーマで解決する（Issue #4）。"""
    return color_to_hex(color_obj)


def _placeholder_type(shape: Any) -> Any:
    try:
        if shape.is_placeholder:
            return shape.placeholder_format.type
    except (AttributeError, ValueError):
        return None
    return None


_NAME_ROLE_PREFIXES = ("title", "subtitle", "body", "caption", "card", "footer", "header")
from .template_kit import TEMPLATE_CHROME_NAMES as _TEMPLATE_CHROME_NAMES


def _role_from_name(shape: Any) -> str | None:
    """本アプリが生成した PPTX は図形名に 'role:id' を持つ。再読込時に役割を復元する。"""
    name = str(getattr(shape, "name", "") or "")
    prefix = name.split(":", 1)[0].strip().lower()
    if prefix in _NAME_ROLE_PREFIXES:
        return prefix  # 'title:el001' 形式、または 'title' 'subtitle' などの素の名前
    return None


def _role_for(shape: Any) -> str | None:
    named = _role_from_name(shape)
    if named:
        return named
    pt = _placeholder_type(shape)
    if pt in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE):
        return "title"
    if pt == PP_PLACEHOLDER.SUBTITLE:
        return "subtitle"
    if pt in (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.VERTICAL_BODY):
        return "body"
    return "body"


def _bullet_of(p_elm: etree._Element, is_body_placeholder: bool, inherited: str | None = None) -> str | None:
    """段落の箇条書き種別。段落の pPr → 継承チェーン → （本文プレースホルダなら bullet）の順で決める。"""
    ppr = p_elm.find("a:pPr", NS)
    if ppr is not None:
        if ppr.find("a:buNone", NS) is not None:
            return None
        if ppr.find("a:buAutoNum", NS) is not None:
            return "number"
        if ppr.find("a:buChar", NS) is not None:
            return "bullet"
    if inherited == "none":
        return None
    if inherited in ("bullet", "number"):
        return inherited
    return "bullet" if is_body_placeholder else None


def _align_of(paragraph_obj: Any) -> str | None:
    try:
        al = paragraph_obj.alignment
    except AttributeError:
        return None
    if al is None:
        return None
    name = str(al).split(".")[-1].split(" ")[0].lower()
    return {"left": "left", "center": "center", "right": "right", "justify": "justify"}.get(name)


def _run_font_name(r: Any) -> str | None:
    """run の明示フォント。日本語資料では ea（東アジア）を優先し、無ければ latin。継承層（_level_props）と同じ優先順。"""
    ea = r._r.find("a:rPr/a:ea", NS)
    if ea is not None and ea.get("typeface"):
        return ea.get("typeface")
    return r.font.name


def _paragraphs_of(text_frame: Any, is_body_placeholder: bool, shape: Any = None, resolver: TextStyleResolver | None = None) -> list[dict]:
    """段落列を JSON へ。run に無い書式は継承チェーン（図形 → レイアウト → マスター）から補う（Issue #5）。

    継承で補った属性名は run["inherited"] に列挙する。テンプレート適用時はそれらを上書きしてよく、
    明示指定（作成者が意図した書式）は保持する。
    """
    paras: list[dict] = []
    scale = TextStyleResolver.font_scale(shape) if (shape is not None and resolver is not None) else 1.0
    theme = resolver.theme if resolver is not None else current_theme()
    empty = {"sz": None, "bullet": None, "bold": None, "color": None, "font": None}
    for p in text_frame.paragraphs:
        level = int(p.level or 0)
        inherited = resolver.resolve(shape, level, p._p.find("a:pPr", NS)) if (resolver is not None and shape is not None) else dict(empty)
        runs: list[dict] = []
        for r in p.runs:
            f = r.font
            from_inherit: list[str] = []
            size = f.size.pt if f.size is not None else None
            if size is None and inherited["sz"] is not None:
                size, from_inherit = inherited["sz"], from_inherit + ["size_pt"]
            if size is not None and scale != 1.0:
                size = round(size * scale, 1)
            bold = f.bold
            if bold is None and inherited["bold"] is not None:
                bold, from_inherit = inherited["bold"], from_inherit + ["bold"]
            color = _rgb_of(f.color)
            if color is None and inherited["color"] is not None:
                color, from_inherit = inherited["color"], from_inherit + ["color"]
            font = _run_font_name(r)
            font = theme.resolve_font(font) if theme is not None else font
            if font is None and inherited["font"] is not None:
                font, from_inherit = inherited["font"], from_inherit + ["font"]
            href = None
            try:
                href = r.hyperlink.address
            except (AttributeError, KeyError):
                href = None
            rd = run(r.text, bold=bold or None, italic=f.italic or None, underline=f.underline or None, size_pt=size, color=color, font=font, href=href)
            if from_inherit:
                rd["inherited"] = from_inherit
            runs.append(rd)
        if not runs:
            runs.append(run(p.text or ""))
        bullet = _bullet_of(p._p, is_body_placeholder, inherited["bullet"])
        paras.append(paragraph(runs, level=level, bullet=bullet, align=_align_of(p)))
    return paras


def _xfrm_of(shape_elm: etree._Element) -> tuple[int, int, int, int, bool, bool] | None:
    """OOXML の xfrm（off/ext/flip）を EMU で返す。"""
    xfrm = shape_elm.find(".//a:xfrm", NS)
    if xfrm is None:
        xfrm = shape_elm.find(".//p:xfrm", NS)
    if xfrm is None:
        return None
    off, ext = xfrm.find("a:off", NS), xfrm.find("a:ext", NS)
    if off is None or ext is None:
        return None
    return (int(off.get("x", 0)), int(off.get("y", 0)), int(ext.get("cx", 0)), int(ext.get("cy", 0)), xfrm.get("flipH") == "1", xfrm.get("flipV") == "1")


class _Ctx:
    """変換中の状態（資料、ID、資産の重複排除）をまとめる。"""

    def __init__(self, presentation: dict, filename: str):
        self.presentation = presentation
        self.ids = IdFactory()
        self.asset_by_hash: dict[str, str] = {}
        self.filename = filename
        self.resolver: TextStyleResolver | None = None  # スライドごとに差し替える

    def add_asset(self, blob: bytes, mime: str, ext: str) -> str:
        digest = hashlib.sha1(blob).hexdigest()
        if digest in self.asset_by_hash:
            return self.asset_by_hash[digest]
        asset_id = self.ids.next("img")
        width = height = None
        try:
            from PIL import Image

            with Image.open(io.BytesIO(blob)) as im:
                width, height = im.size
        except Exception:  # noqa: BLE001 - 画像が読めない場合も資産としては保持する
            pass
        self.presentation["assets"][asset_id] = {"mime": mime, "filename": f"{asset_id}.{ext}", "data_base64": base64.b64encode(blob).decode("ascii"), "width_px": width, "height_px": height}
        self.asset_by_hash[digest] = asset_id
        return asset_id


def _transform(x: float, y: float, w: float, h: float, tf: dict | None) -> tuple[float, float, float, float]:
    """グループ内座標を親座標へ写像する（chOff/chExt → off/ext のスケール）。"""
    if not tf:
        return x, y, w, h
    sx, sy = tf["sx"], tf["sy"]
    return tf["ox"] + (x - tf["cox"]) * sx, tf["oy"] + (y - tf["coy"]) * sy, w * sx, h * sy


def _shape_bbox(shape: Any, tf: dict | None) -> dict | None:
    """図形の bbox。位置情報（xfrm）が無い図形は None を返し、自動レイアウトへ回す。"""
    if shape.left is None or shape.top is None or shape.width is None or shape.height is None:
        return None
    x, y, w, h = emu_to_pt(shape.left), emu_to_pt(shape.top), emu_to_pt(shape.width), emu_to_pt(shape.height)
    x, y, w, h = _transform(x, y, w, h, tf)
    return bbox(x, y, w, h)


def _crop_blob(blob: bytes, crop: dict[str, float]) -> bytes | None:
    """PowerPoint のトリミング（各辺の比率）を画像に焼き込む。拡張トリミング（負値）は対象外。"""
    if any(v < 0 for v in crop.values()) or crop["left"] + crop["right"] >= 1 or crop["top"] + crop["bottom"] >= 1:
        return None
    try:
        from PIL import Image

        with Image.open(io.BytesIO(blob)) as im:
            w, h = im.size
            box = (int(round(w * crop["left"])), int(round(h * crop["top"])), int(round(w * (1 - crop["right"]))), int(round(h * (1 - crop["bottom"]))))
            if box[2] - box[0] < 1 or box[3] - box[1] < 1:
                return None
            out = io.BytesIO()
            im.crop(box).save(out, format="PNG")
            return out.getvalue()
    except Exception:  # noqa: BLE001
        return None


def _rotation_of(shape: Any) -> float | None:
    try:
        rot = float(getattr(shape, "rotation", 0) or 0)
    except (TypeError, ValueError):
        return None
    return round(rot, 2) if abs(rot) > 0.01 else None


def _convert_shape(shape: Any, ctx: _Ctx, slide: dict, tf: dict | None, z: int) -> list[dict]:
    """1 図形を 0 個以上の要素へ変換する。回転は rotation_deg として要素に残す。"""
    els = _convert_shape_inner(shape, ctx, slide, tf, z)
    rot = _rotation_of(shape) if getattr(shape, "shape_type", None) != MSO_SHAPE_TYPE.GROUP else None
    if rot:
        for el in els:
            el["rotation_deg"] = rot
    return els


def _convert_shape_inner(shape: Any, ctx: _Ctx, slide: dict, tf: dict | None, z: int) -> list[dict]:
    pres = ctx.presentation
    st = shape.shape_type
    pt = _placeholder_type(shape)
    if pt in _SKIP_PLACEHOLDERS or str(getattr(shape, "name", "")) in _TEMPLATE_CHROME_NAMES:
        return []  # フッター・日付・ページ番号・機密表示はテンプレート側で再生成する

    # グループ: 子図形を親座標へ写像して平坦化
    if st == MSO_SHAPE_TYPE.GROUP:
        xf = _xfrm_of(shape._element)
        grp = shape._element.find(".//a:xfrm", NS)
        child_tf = None
        if grp is not None:
            ch_off, ch_ext = grp.find("a:chOff", NS), grp.find("a:chExt", NS)
            if ch_off is not None and ch_ext is not None and int(ch_ext.get("cx", 0)) and int(ch_ext.get("cy", 0)):
                gx, gy, gw, gh = _transform(emu_to_pt(shape.left), emu_to_pt(shape.top), emu_to_pt(shape.width), emu_to_pt(shape.height), tf)
                child_tf = {
                    "ox": gx,
                    "oy": gy,
                    "cox": emu_to_pt(int(ch_off.get("x", 0))),
                    "coy": emu_to_pt(int(ch_off.get("y", 0))),
                    "sx": gw / emu_to_pt(int(ch_ext.get("cx"))),
                    "sy": gh / emu_to_pt(int(ch_ext.get("cy"))),
                }
        out: list[dict] = []
        for i, child in enumerate(shape.shapes):
            out.extend(_convert_shape(child, ctx, slide, child_tf, z + i))
        return out

    box = _shape_bbox(shape, tf)

    if st == MSO_SHAPE_TYPE.PICTURE or (pt == PP_PLACEHOLDER.PICTURE and hasattr(shape, "image")):
        try:
            img = shape.image
            blob, content_type, ext = img.blob, img.content_type, img.ext
            crop = {side: float(getattr(shape, f"crop_{side}", 0) or 0) for side in ("left", "right", "top", "bottom")}
            cropped = any(abs(c) > 0.001 for c in crop.values())
            baked = _crop_blob(blob, crop) if cropped else None
            if baked:
                blob, content_type, ext = baked, "image/png", "png"
            asset_id = ctx.add_asset(blob, content_type, ext)
            el = image_element(ctx.ids.next("el"), asset_id, box, alt=shape.name, fit="stretch", z=z)
            if baked:
                el["crop"] = crop
                add_warning(pres, warning("pptx_parser", "IMAGE_CROP_BAKED", "画像のトリミングを画像そのものに反映しました。", slide["id"], el["id"], "トリミング済み画像を配置"), slide)
            elif cropped:
                add_warning(pres, warning("pptx_parser", "IMAGE_CROP_IGNORED", "画像のトリミングは再現されません（元画像を枠に合わせます）。", slide["id"], el["id"], "元画像をそのまま配置"), slide)
            return [el]
        except Exception as e:  # noqa: BLE001
            add_warning(pres, warning("pptx_parser", "IMAGE_READ_FAILED", f"画像を読み取れません: {e}", slide["id"], None, "位置のみ保持"), slide)
            return [unsupported_element(ctx.ids.next("el"), "picture", box, "画像を読み取れませんでした")]

    if shape.has_text_frame and st not in (MSO_SHAPE_TYPE.AUTO_SHAPE,) and not (st == MSO_SHAPE_TYPE.PLACEHOLDER and _is_geometric_placeholder(shape)):
        role = _role_for(shape)
        paras = _paragraphs_of(shape.text_frame, role == "body" and pt is not None, shape, ctx.resolver)
        fill = _rgb_of(shape.fill.fore_color) if _fill_is_solid(shape) else None
        has_text = any(r.get("text", "").strip() for p_ in paras for r in p_["runs"])
        if not has_text and not fill:
            return []  # 文字も塗りも無いテキストボックス（未入力プレースホルダ等）は出力しない
        el = text_element(ctx.ids.next("el"), paras, role=role, box=box, z=z, vertical_align=_anchor_of(shape))
        if fill:
            el["fill"] = fill
        return [el]

    if st in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.PLACEHOLDER):
        return [_convert_autoshape(shape, ctx, slide, box, z)]

    if st == MSO_SHAPE_TYPE.LINE:
        xf = _xfrm_of(shape._element)
        flip_h = xf[4] if xf else False
        flip_v = xf[5] if xf else False
        x1, y1 = box["x"], box["y"]
        x2, y2 = box["x"] + box["w"], box["y"] + box["h"]
        if flip_h:
            x1, x2 = x2, x1
        if flip_v:
            y1, y2 = y2, y1
        width = shape.line.width.pt if shape.line.width else 1.0
        return [line_element(ctx.ids.next("el"), x1, y1, x2, y2, stroke=_rgb_of(shape.line.color), stroke_width_pt=width)]

    if st == MSO_SHAPE_TYPE.TABLE or getattr(shape, "has_table", False):
        return [_convert_table(shape, ctx, box, z)]

    # SmartArt / グラフ / OLE / メディアなど
    return [_convert_graphic_frame(shape, ctx, slide, box, z)]


def _is_geometric_placeholder(shape: Any) -> bool:
    return False


def _anchor_of(shape: Any) -> str | None:
    try:
        bodypr = shape.text_frame._txBody.find("a:bodyPr", NS)
        anchor = bodypr.get("anchor") if bodypr is not None else None
    except AttributeError:
        return None
    return {"t": "top", "ctr": "middle", "b": "bottom"}.get(anchor or "", None)


def _fill_is_solid(shape: Any) -> bool:
    try:
        from pptx.enum.dml import MSO_FILL

        return shape.fill.type == MSO_FILL.SOLID
    except (AttributeError, TypeError, ValueError):
        return False


def _convert_autoshape(shape: Any, ctx: _Ctx, slide: dict, box: dict, z: int) -> dict:
    pres = ctx.presentation
    try:
        auto_name = str(shape.auto_shape_type).split(".")[-1].split(" ")[0]
    except (AttributeError, ValueError):
        auto_name = "RECTANGLE"
    kind = _SHAPE_MAP.get(auto_name)
    if kind is None:
        kind = "rect"
        add_warning(pres, warning("pptx_parser", "SHAPE_APPROXIMATED", f"図形 '{auto_name}' は矩形に近似しました。", slide["id"], None, "矩形へ近似"), slide)
    fill = _rgb_of(shape.fill.fore_color) if _fill_is_solid(shape) else None
    stroke = None
    stroke_w = None
    try:
        stroke = _rgb_of(shape.line.color)
        stroke_w = shape.line.width.pt if shape.line.width else None
    except (AttributeError, TypeError, ValueError):
        pass
    paras = _paragraphs_of(shape.text_frame, False, shape, ctx.resolver) if shape.has_text_frame else []
    # 文字だけで塗り・枠線が無いものはテキスト要素として扱う（編集性優先）
    if not fill and not stroke and paras and any(r.get("text") for p in paras for r in p["runs"]):
        return text_element(ctx.ids.next("el"), paras, role=_role_for(shape), box=box, z=z, vertical_align=_anchor_of(shape))
    return shape_element(ctx.ids.next("el"), kind, box, fill=fill, stroke=stroke, stroke_width_pt=stroke_w, paragraphs=paras, z=z, vertical_align=_anchor_of(shape))


def _convert_table(shape: Any, ctx: _Ctx, box: dict, z: int) -> dict:
    tbl = shape.table
    rows: list[list[dict]] = []
    for r in tbl.rows:
        row: list[dict] = []
        for c in r.cells:
            if c.is_spanned:
                continue  # 結合先セルは出力しない（colspan/rowspan で表現）
            bold = any(run_.font.bold for p in c.text_frame.paragraphs for run_ in p.runs if run_.font.bold)
            fill = _rgb_of(c.fill.fore_color) if _cell_fill_solid(c) else None
            cd = cell(c.text, bold=bold, fill=fill, colspan=int(c.span_width or 1), rowspan=int(c.span_height or 1))
            # ラン単位の書式（色・サイズ）を持つセルは paragraphs も保持し、Web/PPTX で再現する
            paras = _paragraphs_of(c.text_frame, False)  # 表セルはマスター継承を持たない（テーマフォント名の解決だけ current_theme で行う）
            if any(r.get("color") or r.get("size_pt") for p_ in paras for r in p_["runs"]):
                cd["paragraphs"] = paras
            row.append(cd)
        rows.append(row)
    col_widths = [emu_to_pt(col.width) for col in tbl.columns]
    header_rows = 1 if rows else 0
    return table_element(ctx.ids.next("el"), rows, box, header_rows=header_rows, col_widths_pt=col_widths, z=z)


def _cell_fill_solid(c: Any) -> bool:
    try:
        from pptx.enum.dml import MSO_FILL

        return c.fill.type == MSO_FILL.SOLID
    except (AttributeError, TypeError, ValueError):
        return False


def _convert_graphic_frame(shape: Any, ctx: _Ctx, slide: dict, box: dict, z: int) -> dict:
    """SmartArt・グラフ等。SmartArt はデータ部の文字列を箇条書きへ退避する。"""
    pres = ctx.presentation
    elm = shape._element
    rel_ids = elm.find(".//dgm:relIds", NS)
    if rel_ids is not None:
        texts = _smartart_texts(shape, rel_ids)
        if texts:
            paras = [paragraph([run(t)], bullet="bullet") for t in texts]
            add_warning(pres, warning("pptx_parser", "SMARTART_TEXT_ONLY", "SmartArt は図形を再現せず、文字列を箇条書きへ変換しました。", slide["id"], None, "箇条書きへ退避"), slide)
            return text_element(ctx.ids.next("el"), paras, role="body", box=box, z=z, original_type="smartart")
        add_warning(pres, warning("pptx_parser", "SMARTART_UNSUPPORTED", "SmartArt を読み取れません。", slide["id"], None, "位置のみ保持"), slide)
        return unsupported_element(ctx.ids.next("el"), "smartart", box, "SmartArt（未対応）")
    if getattr(shape, "has_chart", False):
        title = ""
        try:
            if shape.chart.has_title:
                title = shape.chart.chart_title.text_frame.text
        except Exception:  # noqa: BLE001
            title = ""
        add_warning(pres, warning("pptx_parser", "CHART_UNSUPPORTED", "グラフは未対応です。枠と題名のみ残します。", slide["id"], None, "枠のみ保持"), slide)
        return unsupported_element(ctx.ids.next("el"), "chart", box, f"グラフ（未対応）{(' ' + title) if title else ''}")
    kind = str(shape.shape_type).split(".")[-1].split(" ")[0].lower() if shape.shape_type is not None else "unknown"
    add_warning(pres, warning("pptx_parser", "ELEMENT_UNSUPPORTED", f"未対応要素 '{kind}' は位置のみ保持します。", slide["id"], None, "枠のみ保持"), slide)
    return unsupported_element(ctx.ids.next("el"), kind, box, f"{kind}（未対応）")


def _smartart_texts(shape: Any, rel_ids: etree._Element) -> list[str]:
    """SmartArt のデータ部（dgm:dataModel）からテキストを順に集める。"""
    dm_rid = rel_ids.get(f"{{{NS['r']}}}dm")
    if not dm_rid:
        return []
    try:
        part = shape.part.related_part(dm_rid)
        root = etree.fromstring(part.blob)
    except Exception:  # noqa: BLE001
        return []
    texts: list[str] = []
    for pt in root.findall(".//dgm:pt", NS):
        if pt.get("type") in ("doc", "pres"):
            continue
        t = "".join(x.text or "" for x in pt.findall(".//a:t", NS)).strip()
        if t:
            texts.append(t)
    return texts


def _infer_title_role(sd: dict, canvas: dict) -> None:
    """プレースホルダ由来のタイトルが無い場合、上部 25% にある最大フォントの短いテキストをタイトルとみなす。"""
    if any(el.get("role") == "title" for el in sd["elements"]):
        return
    limit_y = float(canvas["height_pt"]) * 0.25
    best, best_size = None, 0.0
    for el in sd["elements"]:
        if el.get("type") != "text" or not el.get("bbox") or el["bbox"]["y"] > limit_y:
            continue
        paras = el.get("paragraphs", [])
        text = "".join(r.get("text", "") for p in paras for r in p.get("runs", []))
        if not text.strip() or len(text) > 80 or len(paras) > 2:
            continue
        size = max((float(r.get("size_pt") or 0) for p in paras for r in p.get("runs", [])), default=0.0)
        if size > best_size and size >= 20:
            best, best_size = el, size
    if best is not None:
        best["role"] = "title"


def _classify_layout(slide_dict: dict, has_title_placeholder: bool, has_subtitle: bool) -> str:
    types = [el["type"] for el in slide_dict["elements"]]
    roles = [el.get("role") for el in slide_dict["elements"]]
    if has_subtitle and "title" in roles:
        return "title"
    # 先頭スライドが「題名 + 少数の文字/図形のみ」なら表紙とみなす（テンプレートの表紙部品を適用するため）
    if int(slide_dict.get("index", 0)) == 0 and "title" in roles and types.count("text") <= 3 and all(t in ("text", "shape", "line") for t in types):
        return "title"
    if "table" in types:
        return "table"
    if types.count("image") >= 1 and all(t in ("image", "text") for t in types) and roles.count("body") == 0:
        return "image"
    if roles.count("body") >= 2:
        return "two_column"
    if not slide_dict["elements"]:
        return "blank"
    return "title_body"


def _parse_slide_shapes(slide: Any, sd: dict, ctx: "_Ctx", presentation: dict, idx: int) -> None:
    """1 スライド分の図形・ノート・背景を sd へ詰める（テーマは呼び出し側が contextvar に設定済み）。"""
    has_title_ph = has_subtitle = False
    z = 0
    for shape in slide.shapes:
        pt = _placeholder_type(shape)
        if pt in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            has_title_ph = True
        if pt == PP_PLACEHOLDER.SUBTITLE:
            has_subtitle = True
        try:
            els = _convert_shape(shape, ctx, sd, None, z)
        except Exception as e:  # noqa: BLE001 - 1図形の失敗で全体を止めない
            log.exception("図形の変換に失敗: slide=%s shape=%s", idx + 1, getattr(shape, "name", "?"))
            add_warning(presentation, warning("pptx_parser", "SHAPE_CONVERT_FAILED", f"図形 '{getattr(shape, 'name', '?')}' の変換に失敗: {e}", sd["id"], None, "位置のみ保持"), sd)
            try:
                els = [unsupported_element(ctx.ids.next("el"), "error", _shape_bbox(shape, None), "変換失敗")]
            except Exception:  # noqa: BLE001
                els = []
        for el in els:
            el["z"] = z
            z += 1
            sd["elements"].append(el)
    _infer_title_role(sd, presentation["canvas"])
    # スライド題名
    for el in sd["elements"]:
        if el.get("role") == "title":
            t = "".join(r.get("text", "") for p in el.get("paragraphs", []) for r in p.get("runs", [])).strip()
            if t:
                sd["title"] = t
                break
    # ノート
    try:
        if slide.has_notes_slide:
            sd["notes"] = slide.notes_slide.notes_text_frame.text or None
    except Exception:  # noqa: BLE001
        sd["notes"] = None
    # 背景（単色のみ）
    try:
        bg = slide.background.fill
        from pptx.enum.dml import MSO_FILL

        if bg.type == MSO_FILL.SOLID:
            c = _rgb_of(bg.fore_color)
            if c:
                sd["background"] = {"color": c}
    except Exception:  # noqa: BLE001
        pass
    sd["_has_title_ph"], sd["_has_subtitle"] = has_title_ph, has_subtitle


def parse_pptx(data: bytes, filename: str = "input.pptx", template_id: str | None = None) -> dict:
    """PPTX バイト列を Presentation JSON へ変換する。"""
    prs = Presentation(io.BytesIO(data))
    title = ""
    try:
        title = prs.core_properties.title or ""
    except Exception:  # noqa: BLE001
        title = ""
    presentation = new_presentation(title=title, source_type="pptx", filename=filename, template_id=template_id)
    presentation["canvas"]["width_pt"] = round(emu_to_pt(prs.slide_width), 2)
    presentation["canvas"]["height_pt"] = round(emu_to_pt(prs.slide_height), 2)
    ratio = presentation["canvas"]["width_pt"] / presentation["canvas"]["height_pt"]
    presentation["canvas"]["aspect"] = "16:9" if abs(ratio - 16 / 9) < 0.02 else ("4:3" if abs(ratio - 4 / 3) < 0.02 else f"{ratio:.3f}")
    if presentation["canvas"]["aspect"] != "16:9":
        add_warning(presentation, warning("pptx_parser", "ASPECT_NOT_16_9", f"スライド比率が 16:9 ではありません（{presentation['canvas']['aspect']}）。", None, None, "元の比率を保持"))
    ctx = _Ctx(presentation, filename)

    themes: dict[str, ThemeInfo] = {}
    for idx, slide in enumerate(prs.slides):
        sd = new_slide(ctx.ids.next("s"), idx)
        has_title_ph = has_subtitle = False
        z = 0
        master = slide.slide_layout.slide_master
        master_key = str(master.part.partname)
        theme = themes.get(master_key)
        if theme is None:
            theme = ThemeInfo(master)
            themes[master_key] = theme
        slide_theme = theme.with_override(TextStyleResolver.clr_map_override(slide))  # clrMapOvr を反映（Issue #14）
        token = set_current_theme(slide_theme)
        try:
            ctx.resolver = TextStyleResolver(slide, slide_theme)
            _parse_slide_shapes(slide, sd, ctx, presentation, idx)
        finally:
            reset_current_theme(token)  # 例外時もテーマを残さない（CLI からの連続呼び出し対策）
        sd["layout"] = _classify_layout(sd, sd.pop("_has_title_ph"), sd.pop("_has_subtitle"))
        # 本アプリが出力した PPTX は cSld name に種別を持つ（Issue #8）。表紙・最終ページを確実に復元する
        kind_name = str(slide._element.cSld.get("name") or "")
        if kind_name == "kind:closing":
            sd["layout"] = "closing"
            sd["title"] = sd.get("title") or "最終ページ"
        elif kind_name == "kind:cover":
            sd["layout"] = "title"
        presentation["slides"].append(sd)

    unresolved = sorted({k for t in themes.values() for k in t.unresolved})
    if unresolved:
        add_warning(presentation, warning("pptx_parser", "THEME_COLOR_UNRESOLVED", f"テーマ色を解決できませんでした: {', '.join(unresolved)}", None, None, "役割既定色で描画"))
    if not presentation["meta"]["title"] and presentation["slides"]:
        presentation["meta"]["title"] = presentation["slides"][0].get("title") or filename.rsplit(".", 1)[0]
    log.info("PPTX 解析完了: %s slides=%d assets=%d warnings=%d", filename, len(presentation["slides"]), len(presentation["assets"]), len(presentation["warnings"]))
    return presentation
