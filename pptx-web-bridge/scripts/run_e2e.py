"""E2E 試験: サンプル PPTX / HTML で往復変換を実行し、結果を docs/試験結果.md に書く。

  python scripts/run_e2e.py            … 実行して docs/試験結果.md と output/e2e/ を更新
  python scripts/run_e2e.py --no-raster … 画像化（Playwright）を使う工程を省く
"""
from __future__ import annotations

import copy
import io
import json
import sys
import time
import zipfile
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

from pptx import Presentation  # noqa: E402

from app import pipeline  # noqa: E402
from app.layout import layout_presentation  # noqa: E402
from app.pptx_parser import parse_pptx  # noqa: E402
from app.quality_check import check_presentation, summarize  # noqa: E402
from app.rasterize import is_available  # noqa: E402
from app.validate import validate  # noqa: E402
from app.web_renderer import write_bundle  # noqa: E402

OUT = ROOT / "output" / "e2e"
REPORT = ROOT / "docs" / "試験結果.md"


def _reread(data: bytes) -> tuple[int, list[str]]:
    prs = Presentation(io.BytesIO(data))
    p = parse_pptx(data, "reread.pptx")
    return len(prs.slides), [s["title"] for s in p["slides"]]


def main() -> int:
    use_raster = "--no-raster" not in sys.argv and is_available()
    OUT.mkdir(parents=True, exist_ok=True)
    rows: list[tuple[str, str, str]] = []
    ok_all = True

    def record(name: str, ok: bool, detail: str) -> None:
        nonlocal ok_all
        ok_all = ok_all and ok
        rows.append((name, "合格" if ok else "不合格", detail))

    t0 = time.time()
    # --- PPTX → JSON → Web ---
    deck = ROOT / "samples" / "sample_deck.pptx"
    if not deck.exists():
        sys.path.insert(0, str(ROOT / "samples"))
        from make_sample_pptx import build  # type: ignore

        build(deck)
    res = pipeline.import_pptx(deck.read_bytes(), deck.name)
    pres = res["presentation"]
    record("PPTX→JSON 解析", not res["schema_errors"] and len(pres["slides"]) == 6, f"slides={len(pres['slides'])} assets={len(pres['assets'])} warnings={len(res['warnings'])}")
    (OUT / "deck.json").write_text(json.dumps(pres, ensure_ascii=False, indent=2), encoding="utf-8")
    web_dir = write_bundle(pres, OUT / "deck_web")
    index_html = (web_dir / "index.html").read_text(encoding="utf-8")
    record("JSON→Web 一式生成", all(f'id="{s["id"]}"' in index_html for s in pres["slides"]) and (web_dir / "assets").exists(), f"{web_dir}")
    q = check_presentation(pres)
    record("PPTX 由来の品質検査", summarize(q)["ok"], f"issues={len(q)}")
    # JSON 保存→再読込
    again = json.loads((OUT / "deck.json").read_text(encoding="utf-8"))
    record("JSON 保存・再読込", validate(again) == [] and again == pres, "同一内容で復元")
    # PPTX 往復
    data, warns = pipeline.export_pptx(pres, "editable")[0:2]
    (OUT / "deck_roundtrip.pptx").write_bytes(data)
    n, titles = _reread(data)
    record("JSON→PPTX（編集性優先）→再読込", n == 6 and titles == [s["title"] for s in pres["slides"]], f"slides={n} titles一致")

    # --- HTML → JSON → PPTX ---
    base = ROOT / "samples" / "sample_html"
    files = {str(p.relative_to(base)): p.read_bytes() for p in base.rglob("*") if p.is_file()}
    for name in ("long.html", "slides.html", "cards.html", "mixed.html"):
        res = pipeline.import_html(files[name], name, extra_files=files)
        pres_h = res["presentation"]
        record(f"HTML→JSON（{name}）", not res["schema_errors"] and len(pres_h["slides"]) >= 3, f"slides={len(pres_h['slides'])} warnings={[w['code'] for w in res['warnings']]}")
        (OUT / f"{name}.json").write_text(json.dumps(pres_h, ensure_ascii=False, indent=2), encoding="utf-8")
        for mode in ("editable", "hybrid") + (("visual",) if use_raster else ()):
            data, _p, warns = pipeline.export_pptx(pres_h, mode)
            (OUT / f"{name}.{mode}.pptx").write_bytes(data)
            n, _t = _reread(data)
            bad = [w["code"] for w in warns if w["code"] in ("ELEMENT_WRITE_FAILED", "VISUAL_MODE_UNAVAILABLE")]
            record(f"JSON→PPTX（{name}, {mode}）", n == len(pres_h["slides"]) and not bad, f"slides={n} warnings={[w['code'] for w in warns]}")
        write_bundle(pres_h, OUT / f"{name}_web")
        q = check_presentation(pres_h)
        record(f"HTML 由来の品質検査（{name}）", summarize(q)["ok"], f"issues={[i['code'] for i in q]}")

    # --- 模擬 12 枚 PPTX（mock-pptx）と コーポレートテンプレート ---
    mock = ROOT / "mock-pptx" / "output" / "mock_bidirectional_conversion_12slides.pptx"
    if not mock.exists():
        import subprocess

        subprocess.run([sys.executable, str(ROOT / "mock-pptx" / "src" / "generate_mock_pptx.py")], check=True)
    res = pipeline.import_pptx(mock.read_bytes(), mock.name)
    pres_m = res["presentation"]
    record("模擬12枚 PPTX→JSON", len(pres_m["slides"]) == 12 and not res["warnings"], f"slides={len(pres_m['slides'])} warnings={[w['code'] for w in res['warnings']]}")
    write_bundle(pres_m, OUT / "mock_web")
    q = check_presentation(pres_m)
    record("模擬12枚 品質検査", summarize(q)["ok"] and not q, f"issues={[i['code'] for i in q]}")
    data, _p, warns = pipeline.export_pptx(pres_m, "editable")
    (OUT / "mock_roundtrip.pptx").write_bytes(data)
    n, titles = _reread(data)
    record("模擬12枚 JSON→PPTX→再読込", n == 12 and titles == [s["title"] for s in pres_m["slides"]], f"slides={n} warnings={[w['code'] for w in warns]}")
    from app.report import build_report

    rep = build_report(pres_m)["summary"]
    record("要素判別レポート（文字はテキストシェイプ）", rep["text_as_shapes"] == rep["text_with_bbox"] and rep["editable_ratio"] == 1.0, f"{rep}")
    pres_t = pipeline.import_pptx(mock.read_bytes(), mock.name, "corporate_standard")["presentation"]
    from app.template_kit import make_closing_slide

    pres_t["slides"].append(make_closing_slide("s_end", len(pres_t["slides"]), ""))
    pres_t, _e, _f = pipeline.prepare(pres_t)
    data, _p, warns = pipeline.export_pptx(pres_t, "editable")
    (OUT / "mock_corporate_template.pptx").write_bytes(data)
    write_bundle(pres_t, OUT / "mock_corporate_web")
    prs_t = Presentation(io.BytesIO(data))
    names0 = {sh.name for sh in prs_t.slides[0].shapes}
    record("コーポレートテンプレート適用（表紙背景・ロゴ・帯・最終ページ）", "cover_background" in names0 and "logo" in names0 and len(prs_t.slides) == 13 and not warns, f"slides={len(prs_t.slides)} cover={sorted(names0)[:4]}")

    # --- PPTX からテンプレート作成（Phase C）: 部品の推定 → 保存 → 適用 → 土台 PPTX で出力 → 再読込 ---
    brand = ROOT / "samples" / "brand_template.pptx"
    if not brand.exists():
        sys.path.insert(0, str(ROOT / "samples"))
        from make_brand_template_pptx import build as build_brand  # type: ignore

        build_brand(brand)
    import tempfile

    from app import config as app_config
    from app import template_store
    from app.template_from_pptx import analyze

    tmp_store = Path(tempfile.mkdtemp(prefix="pwb_tpl_"))
    base_cfg = json.loads((ROOT / "config" / "app_config.json").read_text(encoding="utf-8"))
    base_cfg["paths"]["user_templates_file"] = str(tmp_store / "user_templates.json")
    base_cfg["paths"]["user_template_assets_dir"] = str(tmp_store / "assets")
    (tmp_store / "cfg.json").write_text(json.dumps(base_cfg, ensure_ascii=False), encoding="utf-8")
    saved_cfg = app_config._config_singleton
    app_config._config_singleton = app_config.AppConfig(tmp_store / "cfg.json")
    try:
        an = analyze(brand.read_bytes(), brand.name, template_id="e2e_brand")
        prop = an["proposal"]
        got = {k: sorted(prop.get(k, {}).keys()) for k in ("cover", "content", "closing")}
        ok_parts = "background_image" in got["cover"] and "logo" in got["cover"] and "bar" in got["content"] and "body" in got["content"] and "page_number" in got["content"] and "message" in got["closing"]
        record("PPTX からテンプレート推定（背景・ロゴ・帯・本文領域・ページ番号・一言）", ok_parts, f"cover={got['cover']} content={got['content']} closing={got['closing']} warnings={len(an['warnings'])}")
        template_store.save_template(prop)
        ids = [t["id"] for t in app_config.get_config().templates()]
        record("ユーザーテンプレートの保存・一覧", "e2e_brand" in ids and any(t.get("source") == "user" for t in app_config.get_config().templates()), f"templates={ids}")
        pres_u = pipeline.import_pptx(deck.read_bytes(), deck.name, "e2e_brand")["presentation"]
        html_u = pipeline.import_html((ROOT / "samples" / "sample_html" / "long.html").read_bytes(), "long.html", "e2e_brand")["presentation"]
        html_u, _e, _f = pipeline.prepare(html_u)
        from app.template_kit import content_area

        area = content_area(prop, html_u["canvas"])
        bodies = [el for s in html_u["slides"] if s.get("layout") == "title_body" for el in s["elements"] if el.get("role") != "title"]
        inside = all(el["bbox"]["x"] >= area["x"] - 0.5 and el["bbox"]["x"] + el["bbox"]["w"] <= area["x"] + area["w"] + 0.5 and el["bbox"]["y"] >= area["y"] - 0.5 for el in bodies)
        record("本文がテンプレートの本文領域に収まる（HTML 取込）", bool(bodies) and inside, f"area={area} bodies={len(bodies)}")
        data, _p, warns = pipeline.export_pptx(pres_u, "editable", use_base_pptx=True)
        (OUT / "user_template_base.pptx").write_bytes(data)
        prs_u = Presentation(io.BytesIO(data))
        names1 = {sh.name for sh in prs_u.slides[1].shapes}
        n_u, titles_u = _reread(data)
        record("土台 PPTX で出力 → 再読込（マスター維持・部品名・題名一致）", n_u == len(pres_u["slides"]) and "bar" in names1 and "logo" in names1 and titles_u == [s["title"] for s in pres_u["slides"]] and not warns, f"slides={n_u} names={sorted(names1)[:5]} warnings={[w['code'] for w in warns]}")
        write_bundle(pipeline.prepare(pres_u)[0], OUT / "user_template_web")
        template_store.delete_template("e2e_brand")
    finally:
        app_config._config_singleton = saved_cfg

    # --- Copilot 連携（Phase D、API 不使用）: 渡す形（Markdown / Word / 一式）と回答の取込 ---
    from app import copilot_handoff as ch

    md = ch.to_markdown(pres)
    built = ch.build_prompt(pres, "brand_deck")
    bundle = ch.bundle_zip(pres, "brand_deck")
    (OUT / "copilot_handoff.zip").write_bytes(bundle)
    with zipfile.ZipFile(io.BytesIO(bundle)) as zb:
        names_b = set(zb.namelist())
        docx_ok = zb.read("outline.docx")[:2] == b"PK"
    record("Copilot へ渡す一式（prompt / Markdown / JSON / Word / 画像）", "## 2. 背景と課題" in md and "ブランドキット" in built["instruction"] and {"prompt.txt", "outline.md", "outline.json", "outline.docx"} <= names_b and docx_ok, f"chars={built['chars']} files={sorted(names_b)[:5]}")
    back = ch.import_markdown(md)
    back = layout_presentation(back)
    record("Copilot の回答（Markdown）を資料に戻す（往復で枚数一致）", len(back["slides"]) == len(pres["slides"]) and all(el.get("bbox") for s_ in back["slides"] for el in s_["elements"]), f"slides={len(back['slides'])}")
    reply = "# 図解版\n\n## 1. 背景\n型: カード\n- 課題: 二重作業\n- 原因: 形式が違う\n- 対策: 共通形式\nノート: 2 分で\n\n## 2. 効果\n| 項目 | 前 | 後 |\n| --- | --- | --- |\n| 時間 | 10h | 5h |\n"
    fig = layout_presentation(ch.import_markdown(reply))
    cards = [e for e in fig["slides"][1]["elements"] if e["type"] == "diagram"] if len(fig["slides"]) > 1 else []
    items = cards[0]["diagram"]["items"] if cards else []
    record("「図解風に仕上げる」の回答をカード図解・表として取り込む", len(items) == 3 and cards[0]["diagram"]["type"] == "cards" and fig["slides"][1]["notes"] == "2 分で" and fig["slides"][2]["layout"] == "table", f"slides={len(fig['slides'])} items={len(items)}")

    # --- 図解部品（Phase F）: Copilot の回答 → 図解 → PPTX → 再読込 ---
    from app import diagrams as dg

    fig_md = (
        "# 図解の資料\n\n"
        "## 1. 進め方\n型: フロー\n- 受付: 窓口で受け取る\n- 審査: 担当が確認\n- 完了: 通知する\nノート: 3 分で説明\n\n"
        "## 2. 効果\n型: 数値\n- 40%｜作業時間の削減\n- 12h｜月に浮く時間\n\n"
        "## 3. 現状と改善後\n型: 比較\n- Before: 二重作業\n- After: 一度で済む\n"
    )
    fig_pres = layout_presentation(ch.import_markdown(fig_md))
    fig_els = [e for s_ in fig_pres["slides"] for e in s_["elements"] if e["type"] == "diagram"]
    types_in = [e["diagram"]["type"] for e in fig_els]
    record("Copilot の回答から図解（フロー・数値・比較）を作る", types_in == ["flow", "kpi", "compare"] and all(e.get("bbox") for e in fig_els), f"types={types_in} slides={len(fig_pres['slides'])}")

    fig_bytes, _fp, fig_warns = pipeline.export_pptx(fig_pres)
    (OUT / "diagrams.pptx").write_bytes(fig_bytes)
    fig_back = parse_pptx(fig_bytes, "diagrams.pptx")
    back_els = [e for s_ in fig_back["slides"] for e in s_["elements"] if e["type"] == "diagram"]
    titles_in = [i["title"] for e in fig_els for i in e["diagram"]["items"]]
    titles_back = [i["title"] for e in back_els for i in e["diagram"]["items"]]
    record("図解を PPTX に出して読み戻す（型と項目が保たれる）", [e["diagram"]["type"] for e in back_els] == types_in and titles_back == titles_in and not fig_warns, f"back={[e['diagram']['type'] for e in back_els]} items={len(titles_back)}")
    write_bundle(pipeline.prepare(fig_pres)[0], OUT / "diagrams_web")
    md_again = ch.to_markdown(fig_pres)
    record("図解を Markdown に戻して Copilot へ渡せる", f"型: {dg.word_of('flow')}" in md_again and "40%｜作業時間の削減" in md_again, f"chars={len(md_again)}")

    # --- Copilot エージェント一式（Phase E、API 不使用） ---
    from app import copilot_agent_kit as kit

    kit_bytes = kit.build_kit_zip()
    (OUT / "copilot_agent.zip").write_bytes(kit_bytes)
    with zipfile.ZipFile(io.BytesIO(kit_bytes)) as zk:
        names_k = set(zk.namelist())
        manifest = json.loads(zk.read("declarativeAgent.json").decode("utf-8"))
        bodies_k = [n for n in names_k if n.startswith("knowledge/") and "INDEX" not in n]
        index_k = zk.read(f"knowledge/{kit.KNOWLEDGE_PREFIX}_INDEX.txt").decode("utf-8")
        within = all(len(zk.read(n).decode("utf-8")) <= kit.MAX_CHARS for n in bodies_k)
    ok_kit = (
        {"declarativeAgent.json", "instructions.md", "kit.json", "READ_ME.txt"} <= names_k
        and manifest["instructions"] == "$[file('instructions.md')]"
        and all(n.split("/")[-1] in index_k for n in bodies_k)
        and len(bodies_k) <= kit.MAX_FILES
        and within
    )
    record("Copilot エージェント一式（定義・指示文・ナレッジ・目次）", ok_kit, f"files={len(names_k)} knowledge={len(bodies_k)} starters={len(manifest['conversation_starters'])}")

    # --- テンプレート推定の精度（Phase I）: 色見本・装飾入りの 4 枚デッキ ---
    from app import template_store
    from app.template_from_pptx import analyze as tpl_analyze, set_part_role

    sys.path.insert(0, str(ROOT / "samples"))
    from make_brand_template_pptx import build_extended  # noqa: E402

    ext_path = build_extended(ROOT / "samples" / "brand_template_ext.pptx")
    ext = tpl_analyze(ext_path.read_bytes(), "brand_template_ext.pptx", template_id="e2e_ext")
    try:
        roles = [s_["role"] for s_ in ext["slides"]]
        c = ext["proposal"]["content"]
        record("色見本のスライドを避けて中身を選び、帯は 1 本・飾りは装飾にする", roles == ["cover", "content", "skip", "closing"] and "bars" not in c and len(c.get("decor", [])) == 2 and ext["proposal"]["colors"]["primary"] == "#001A72", f"roles={roles} decor={len(c.get('decor', []))} primary={ext['proposal']['colors']['primary']}")
        moved = set_part_role(ext["proposal"], "content", "decor.0", "bar")
        record("画面の表で装飾を帯に変えられる（提案が組み替わる）", len(moved["content"].get("bars", [])) == 1 and len(moved["content"]["decor"]) == 1 and "帯" in ext["proposal"]["guide"]["content"], f"guide={ext['proposal']['guide']['content']}")
    finally:
        template_store.delete_template("e2e_ext")

    # --- 差分マージ再取込（Phase G） ---
    from app import merge as mg

    html_v1 = "<html><body><section><h2>背景</h2><p>課題は二重作業</p></section><section><h2>効果</h2><p>時間が半分</p></section></body></html>".encode("utf-8")
    html_v2 = "<html><body><section><h2>背景</h2><p>課題は二重作業と転記</p></section><section><h2>効果</h2><p>時間が半分</p></section><section><h2>今後</h2><p>全社展開</p></section></body></html>".encode("utf-8")
    imported_v1 = pipeline.import_html(html_v1, "plan.html")["presentation"]
    edited = copy.deepcopy(imported_v1)
    body_el = [e for e in edited["slides"][0]["elements"] if e.get("role") != "title"][-1]
    body_el["bbox"] = {"x": 50.0, "y": 300.0, "w": 200.0, "h": 60.0}
    body_el["user_bbox"] = True
    edited["slides"][1]["notes"] = "自分のノート"
    incoming = pipeline.import_html(html_v2, "plan.html")["presentation"]
    merged = pipeline.merge_import(edited, incoming)
    m_pres, m_rep = merged["presentation"], merged["report"]
    kept = [e for e in m_pres["slides"][0]["elements"] if e.get("user_bbox")]
    texts = [mg.element_text(e) for e in m_pres["slides"][0]["elements"]]
    ok_merge = bool(kept) and kept[0]["bbox"]["x"] == 50.0 and any("転記" in t for t in texts) and m_pres["slides"][1]["notes"] == "自分のノート" and len(m_pres["slides"]) == 3
    record("差分マージ再取込（座標・ノートを残して本文だけ更新、ページ追加）", ok_merge and m_rep["updated"] >= 1 and not m_rep["conflicts"], f"{merged['summary']} slides={len(m_pres['slides'])}")

    conflicted = copy.deepcopy(imported_v1)
    c_el = [e for e in conflicted["slides"][0]["elements"] if e.get("role") != "title"][-1]
    c_el["paragraphs"][0]["runs"][0]["text"] = "課題は私が直した"
    merged2 = pipeline.merge_import(conflicted, pipeline.import_html(html_v2, "plan.html")["presentation"])
    conflicts = merged2["report"]["conflicts"]
    record("両方で変わった箇所を競合として報告する", len(conflicts) == 1 and conflicts[0]["ours"] == "課題は私が直した" and "転記" in conflicts[0]["theirs"], f"conflicts={len(conflicts)} applied={conflicts[0]['applied'] if conflicts else '-'}")

    # --- 版の表示 ---
    from app import version as version_mod

    vinfo = version_mod.version_info()
    feature_ids = {f["id"] for f in vinfo["features"]}
    record("版と機能一覧を返す（/api/version の中身）", vinfo["version"] == version_mod.APP_VERSION and {"diagrams", "merge_import", "copilot_agent_kit"} <= feature_ids and bool(vinfo["commit"]), f"v{vinfo['version']} commit={vinfo['commit']} features={len(feature_ids)}")

    # --- 例外系 ---
    try:
        pipeline.import_pptx(b"broken", "broken.pptx")
        record("壊れた PPTX の拒否", False, "例外が出なかった")
    except Exception as e:  # noqa: BLE001
        record("壊れた PPTX の拒否", True, type(e).__name__)
    broken_json = {"slides": [{"elements": [{"type": "text", "paragraphs": [{"runs": [{"text": "修復"}]}]}]}]}
    from app.validate import validate_and_repair

    fixed, errors, fixes = validate_and_repair(broken_json)
    record("壊れた JSON の自動修復", not errors and fixes, f"repairs={[f['code'] for f in fixes]}")
    html_ext = b"<html><body><section><h2>t</h2><img src='https://example.invalid/x.png'></section></body></html>"
    res = pipeline.import_html(html_ext, "ext.html")
    record("外部画像の非取得と警告", any(w["code"] == "REMOTE_IMAGE_SKIPPED" for w in res["warnings"]), "REMOTE_IMAGE_SKIPPED")

    elapsed = time.time() - t0
    lines = [
        "# 試験結果（E2E）",
        "",
        f"- 実行日時: {datetime.now().isoformat(timespec='seconds')}",
        f"- 実行環境: Python {sys.version.split()[0]} / 画像化（Playwright）: {'利用可' if use_raster else '利用不可・省略'}",
        f"- 所要時間: {elapsed:.1f} 秒",
        f"- 総合: {'合格' if ok_all else '不合格あり'}（{sum(1 for r in rows if r[1] == '合格')} / {len(rows)}）",
        f"- 成果物: `{OUT.relative_to(ROOT)}/`（リポジトリには含めない）",
        "",
        "| 試験項目 | 結果 | 詳細 |",
        "|---|---|---|",
    ]
    lines += [f"| {n} | {r} | {d.replace('|', '／')} |" for n, r, d in rows]
    lines += ["", "再実行: `python scripts/run_e2e.py`（単体試験は `python -m pytest`）", ""]
    REPORT.write_text("\n".join(lines), encoding="utf-8")
    print("\n".join(lines))
    return 0 if ok_all else 1


if __name__ == "__main__":
    raise SystemExit(main())
