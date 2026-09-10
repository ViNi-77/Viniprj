"""テンプレート推定の試験用: 架空ブランドの 3 枚デッキ（表紙・中身・最終ページ）を生成する。

- 表紙: 全面の背景画像、左上の小さなロゴ、題名・副題プレースホルダ、右下に日付の文字
- 中身: 題名・本文プレースホルダ、下部の紺の平行四辺形（帯）、左下ロゴ、右下にページ番号フィールド
- 最終ページ: 中央の大きめロゴと一言
実案件の素材は含まない（PIL で描いた図形のみ）。
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml import parse_xml
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parent / "brand_template.pptx"
NAVY = RGBColor(0x00, 0x1A, 0x72)


def _logo(w: int = 300, h: int = 120, color=(0, 26, 114), bg=(255, 255, 255, 0)) -> bytes:
    im = Image.new("RGBA", (w, h), bg)
    d = ImageDraw.Draw(im)
    d.ellipse([10, 10, h - 10, h - 10], fill=color)
    d.rectangle([h + 10, h * 0.35, w - 10, h * 0.65], fill=color)
    buf = io.BytesIO()
    im.save(buf, format="PNG")
    return buf.getvalue()


def _background(w: int = 1280, h: int = 720) -> bytes:
    im = Image.new("RGB", (w, h), (0, 26, 114))
    d = ImageDraw.Draw(im)
    for i in range(0, w, 80):
        d.line([(i, 0), (i + 200, h)], fill=(20, 50, 140), width=3)
    buf = io.BytesIO()
    im.save(buf, format="JPEG", quality=80)
    return buf.getvalue()


def build(out: Path = OUT) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1. 表紙
    s = prs.slides.add_slide(prs.slide_layouts[0])
    bg = s.shapes.add_picture(io.BytesIO(_background()), 0, 0, prs.slide_width, prs.slide_height)
    bg.name = "background_photo"
    # 背景を最背面へ
    s.shapes._spTree.remove(bg._element)
    s.shapes._spTree.insert(2, bg._element)
    logo = s.shapes.add_picture(io.BytesIO(_logo(color=(255, 255, 255))), Inches(0.5), Inches(0.3), width=Inches(1.8))
    logo.name = "brand_logo_white"
    s.shapes.title.text = "題名がここに入ります"
    for r in s.shapes.title.text_frame.paragraphs[0].runs:
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(36)
        r.font.bold = True
    s.placeholders[1].text = "副題がここに入ります"
    for r in s.placeholders[1].text_frame.paragraphs[0].runs:
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(20)
    date = s.shapes.add_textbox(Inches(8.3), Inches(6.9), Inches(4.7), Inches(0.35))
    date.name = "date_text"
    date.text_frame.text = "2026年9月10日 / © Sample Brand"
    date.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
    date.text_frame.paragraphs[0].runs[0].font.size = Pt(9)
    date.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # 2. 中身
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "中身の題名"
    for r in s.shapes.title.text_frame.paragraphs[0].runs:
        r.font.color.rgb = NAVY
        r.font.size = Pt(28)
        r.font.bold = True
    s.placeholders[1].text_frame.text = "本文の箇条書き"
    bar = s.shapes.add_shape(MSO_SHAPE.PARALLELOGRAM, Inches(2.8), Inches(7.15), Inches(10.53), Inches(0.25))
    bar.name = "bottom_band"
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    bar.adjustments[0] = 0.6
    logo2 = s.shapes.add_picture(io.BytesIO(_logo()), Inches(0.35), Inches(7.1), width=Inches(2.1))
    logo2.name = "brand_logo_blue"
    num = s.shapes.add_textbox(Inches(12.2), Inches(6.85), Inches(0.8), Inches(0.25))
    num.name = "page_no"
    num.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
    num.text_frame.paragraphs[0]._p.append(parse_xml('<a:fld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" id="{B6F5F7C2-5C4C-4C6E-9D0E-1E2F3A4B5C6D}" type="slidenum"><a:rPr lang="ja-JP" sz="900"/><a:t>2</a:t></a:fld>'))

    # 3. 最終ページ
    s = prs.slides.add_slide(prs.slide_layouts[6])
    logo3 = s.shapes.add_picture(io.BytesIO(_logo()), Inches(4.6), Inches(2.4), width=Inches(4.1))
    logo3.name = "brand_logo_center"
    msg = s.shapes.add_textbox(Inches(1.7), Inches(4.6), Inches(10.0), Inches(0.6))
    msg.text_frame.text = "ご清聴ありがとうございました"
    msg.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    msg.text_frame.paragraphs[0].runs[0].font.size = Pt(18)

    prs.core_properties.title = "架空ブランドのテンプレート"
    prs.core_properties.author = "サンプル株式会社（架空）"
    prs.core_properties.last_modified_by = "サンプル株式会社（架空）"
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out


if __name__ == "__main__":
    print(build(Path(sys.argv[1]) if len(sys.argv) > 1 else OUT))
