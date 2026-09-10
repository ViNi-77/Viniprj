"""匿名化サンプル PPTX を生成する（必須6: サンプルPPTX）。

実案件名・実データは含めない。架空の「サンプル株式会社」の業務改善説明資料を模す。
表紙、本文（箇条書き）、2列、図形、表、画像、グループ、線を含む。
"""
from __future__ import annotations

import io
import re
import sys
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parent / "sample_deck.pptx"


def _sample_image(w: int = 800, h: int = 450) -> bytes:
    """架空のグラフ風画像。外部画像へ依存しない。"""
    im = Image.new("RGB", (w, h), (244, 246, 249))
    d = ImageDraw.Draw(im)
    bars = [120, 200, 260, 310, 380]
    for i, v in enumerate(bars):
        x0 = 80 + i * 140
        d.rectangle([x0, h - 40 - v, x0 + 90, h - 40], fill=(31, 58, 95))
    d.line([(60, h - 40), (w - 40, h - 40)], fill=(90, 90, 90), width=3)
    buf = io.BytesIO()
    im.save(buf, format="PNG")
    return buf.getvalue()


def build(out: Path = OUT) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    primary = RGBColor(0x1F, 0x3A, 0x5F)
    accent = RGBColor(0xE0, 0x7A, 0x1F)

    # 1. 表紙
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "業務改善提案（サンプル）"
    s.placeholders[1].text = "サンプル株式会社 業務推進部\n2026年 架空データ"

    # 2. 箇条書き
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "背景と課題"
    tf = s.placeholders[1].text_frame
    tf.text = "資料作成が PowerPoint と Web の二重作業になっている"
    for txt, lvl in [("同じ内容を2回作り直している", 1), ("修正が片方にしか反映されない", 1), ("共通の中間形式で双方向に変換したい", 0), ("編集性と見た目のバランスを選べるようにする", 1)]:
        p = tf.add_paragraph()
        p.text = txt
        p.level = lvl
    s.notes_slide.notes_text_frame.text = "発表者ノート: 背景説明は2分で終える。"

    # 3. 2列（図形 + テキスト）
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "変換の流れ"
    labels = ["PPTX 解析", "Presentation JSON", "Web レンダリング"]
    for i, label in enumerate(labels):
        shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i * 4.2), Inches(2.5), Inches(3.4), Inches(1.4))
        shp.fill.solid()
        shp.fill.fore_color.rgb = primary if i != 1 else accent
        shp.line.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shp.text_frame.text = label
        shp.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
        shp.text_frame.paragraphs[0].runs[0].font.bold = True
        shp.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        if i < 2:
            con = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(4.2 + i * 4.2), Inches(3.2), Inches(5.0 + i * 4.2), Inches(3.2))
            con.line.color.rgb = RGBColor(0x66, 0x66, 0x66)
            con.line.width = Pt(2)
    tb = s.shapes.add_textbox(Inches(0.8), Inches(4.4), Inches(11.5), Inches(1.2))
    tb.text_frame.word_wrap = True
    tb.text_frame.text = "AI は構造理解と分類を担当し、座標計算とファイル生成は決定的なプログラムが行う。"
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(16)

    # 4. 表
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "MVP 対応範囲"
    rows = [["区分", "MVP で対応", "初期対象外"], ["PPTX 要素", "タイトル、本文、画像、矩形、線、表", "SmartArt、動画、マクロ"], ["HTML 要素", "section、h1-h3、p、リスト、画像、表", "任意 SPA、ログイン後ページ"], ["出力", "編集性 / 見た目 / ハイブリッド", "完全ピクセル一致"]]
    tbl = s.shapes.add_table(len(rows), 3, Inches(0.8), Inches(1.8), Inches(11.7), Inches(3.0)).table
    for r, row in enumerate(rows):
        for c, txt in enumerate(row):
            tbl.cell(r, c).text = txt
            for p in tbl.cell(r, c).text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(14)
                    if r == 0:
                        run.font.bold = True

    # 5. 画像 + キャプション + グループ
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "効果の見込み（架空データ）"
    s.shapes.add_picture(io.BytesIO(_sample_image()), Inches(0.8), Inches(1.7), width=Inches(7.0))
    grp = s.shapes.add_group_shape()
    box = grp.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.3), Inches(1.9), Inches(4.2), Inches(2.2))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0xF4, 0xF6, 0xF9)
    box.line.color.rgb = RGBColor(0xC9, 0xD1, 0xDB)
    t = grp.shapes.add_textbox(Inches(8.5), Inches(2.0), Inches(3.9), Inches(2.0))
    t.text_frame.word_wrap = True
    t.text_frame.text = "作成時間を 5 割削減（試算）"
    p = t.text_frame.add_paragraph()
    p.text = "※ 数値は架空のサンプル"
    p.runs[0].font.size = Pt(11)
    cap = s.shapes.add_textbox(Inches(0.8), Inches(5.8), Inches(7.0), Inches(0.5))
    cap.text_frame.text = "図: 月別の作成件数（架空）"
    cap.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
    # テーマ色（Issue #4 の試験用）: 文字は accent1、明るさ補正付きの塗りは accent2
    cap.text_frame.paragraphs[0].runs[0].font.color.theme_color = MSO_THEME_COLOR.ACCENT_1
    badge = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.3), Inches(4.4), Inches(4.2), Inches(0.7))
    badge.name = "theme_badge"
    badge.fill.solid()
    badge.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_2
    badge.fill.fore_color.brightness = 0.4
    badge.line.fill.background()
    badge.text_frame.text = "テーマ色の塗り（accent2 +40%）"
    badge.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    badge.text_frame.paragraphs[0].runs[0].font.color.theme_color = MSO_THEME_COLOR.TEXT_2

    # 6. 区切り + まとめ
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "まとめ"
    tf = s.placeholders[1].text_frame
    tf.text = "一度作れば PowerPoint と Web の両方へ展開できる"
    for txt in ["未対応要素は警告と代替処理で止めない", "Presentation JSON を保存すれば作業を再開できる"]:
        p = tf.add_paragraph()
        p.text = txt

    prs.core_properties.title = "業務改善提案（サンプル）"
    prs.core_properties.author = "サンプル株式会社（架空）"
    prs.core_properties.last_modified_by = "サンプル株式会社（架空）"
    for part in prs.part.package.iter_parts():
        if str(part.partname) == "/docProps/app.xml":
            part._blob = re.sub(rb"<Application>[^<]*</Application>", b"<Application>Microsoft Office PowerPoint</Application>", part.blob)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else OUT
    print(build(target))
