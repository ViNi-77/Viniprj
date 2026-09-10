"""匿名化 12 枚模擬 PPTX の生成器（模擬PPTX生成指示書 / 匿名化スライド仕様書 準拠）。

- 実在情報は一切含めない。数値は架空、画像はコード内で生成した抽象図形のみ。
- 採用技術: Python 3 + python-pptx + Pillow（指示書 4 章の代替候補。理由は README 参照）。
- 実行: python mock-pptx/src/generate_mock_pptx.py  → mock-pptx/output/ に PPTX とマニフェストを出力。
"""
from __future__ import annotations

import io
import re
import json
import math
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from theme import THEME  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "output"
ASSET_DIR = ROOT / "assets" / "generated_shapes"
PPTX_NAME = "mock_bidirectional_conversion_12slides.pptx"
MANIFEST_NAME = "mock_bidirectional_conversion_manifest.json"

C = {k: RGBColor.from_string(v) for k, v in THEME["colors"].items()}
FONT = THEME["font"]
SZ = THEME["sizes"]
SW, SH = THEME["slide_width_in"], THEME["slide_height_in"]
M = THEME["margin_in"]


# ---------------------------------------------------------------- 画像生成（外部素材禁止）
def _abstract_gradient(w: int, h: int, seed: int) -> bytes:
    """幾何学グラデーション画像。乱数を使わず seed から決定的に描く。"""
    im = Image.new("RGB", (w, h))
    px = im.load()
    for y in range(h):
        for x in range(w):
            t = (x / w) * 0.7 + (y / h) * 0.3
            px[x, y] = (int(8 + 30 * t), int(43 + 90 * t + 20 * math.sin(seed + t * 6)), int(92 + 120 * t))
    d = ImageDraw.Draw(im, "RGBA")
    for i in range(6):
        r = int(min(w, h) * (0.15 + 0.08 * i))
        cx, cy = int(w * (0.2 + 0.12 * i + 0.05 * seed) % w), int(h * (0.3 + 0.1 * ((i * seed) % 5)))
        d.ellipse((cx - r, cy - r, cx + r, cy + r), outline=(255, 255, 255, 90), width=3)
    im = im.filter(ImageFilter.SMOOTH)
    buf = io.BytesIO()
    im.save(buf, format="PNG")
    return buf.getvalue()


def _geometric_pattern(w: int, h: int) -> bytes:
    im = Image.new("RGB", (w, h), (238, 244, 248))
    d = ImageDraw.Draw(im)
    step = 48
    for y in range(0, h, step):
        for x in range(0, w, step):
            if ((x // step) + (y // step)) % 2 == 0:
                d.polygon([(x, y + step), (x + step // 2, y), (x + step, y + step)], fill=(18, 97, 160))
            else:
                d.ellipse((x + 8, y + 8, x + step - 8, y + step - 8), outline=(0, 166, 199), width=3)
    buf = io.BytesIO()
    im.save(buf, format="PNG")
    return buf.getvalue()


def save_assets() -> dict[str, bytes]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    assets = {"wide_gradient.png": _abstract_gradient(1200, 675, 2), "tall_pattern.png": _geometric_pattern(600, 900)}
    for name, blob in assets.items():
        (ASSET_DIR / name).write_bytes(blob)
    return assets


# ---------------------------------------------------------------- 描画ヘルパー
def _text(tf, text: str, size: int, bold: bool = False, color=None, align=None, font: str = FONT):
    tf.word_wrap = True
    p = tf.paragraphs[0] if not tf.paragraphs[0].runs and len(tf.paragraphs) == 1 else tf.add_paragraph()
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = font
    r.font.color.rgb = color or C["ink"]
    if align:
        p.alignment = align
    return p


def _bullets(tf, items: list[str], size: int = SZ["body"], color=None):
    tf.word_wrap = True
    first = True
    for it in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        r = p.add_run()
        r.text = "• " + it
        r.font.size = Pt(size)
        r.font.name = FONT
        r.font.color.rgb = color or C["ink"]
        p.space_after = Pt(6)


def add_title(slide, text: str):
    tb = slide.shapes.add_textbox(Inches(M), Inches(THEME["title_top_in"]), Inches(SW - 2 * M), Inches(THEME["title_height_in"]))
    tb.name = "title"
    _text(tb.text_frame, text, SZ["title"], bold=True, color=C["navy"])
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(M), Inches(THEME["title_top_in"] + THEME["title_height_in"]), Inches(M + 1.2), Inches(THEME["title_top_in"] + THEME["title_height_in"]))
    ln.line.color.rgb = C["cyan"]
    ln.line.width = Pt(3)
    ln.name = "title_rule"
    return tb


def add_footer(slide, page: int):
    ft = slide.shapes.add_textbox(Inches(M), Inches(SH - 0.45), Inches(6), Inches(0.3))
    ft.name = "footer"
    _text(ft.text_frame, THEME["footer_text"], THEME["footer_size"], color=C["muted"])
    pn = slide.shapes.add_textbox(Inches(SW - M - 1.0), Inches(SH - 0.45), Inches(1.0), Inches(0.3))
    pn.name = "page_number"
    _text(pn.text_frame, f"{page:02d}", THEME["footer_size"], color=C["muted"], align=PP_ALIGN.RIGHT)


def add_card(slide, x: float, y: float, w: float, h: float, title: str, body: list[str] | str, fill=None, icon: str | None = None, name: str = "card"):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    box.name = name
    box.fill.solid()
    box.fill.fore_color.rgb = fill or C["light"]
    box.line.color.rgb = C["line"]
    box.shadow.inherit = False
    box.text_frame.text = ""
    top = y + 0.2
    if icon:
        shp = {"circle": MSO_SHAPE.OVAL, "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE, "diamond": MSO_SHAPE.DIAMOND}[icon]
        ic = slide.shapes.add_shape(shp, Inches(x + 0.25), Inches(top), Inches(0.5), Inches(0.5))
        ic.name = f"{name}_icon"
        ic.fill.solid()
        ic.fill.fore_color.rgb = C["cyan"]
        ic.line.fill.background()
        top += 0.7
    tt = slide.shapes.add_textbox(Inches(x + 0.2), Inches(top), Inches(w - 0.4), Inches(0.5))
    tt.name = f"{name}_title"
    _text(tt.text_frame, title, SZ["card_title"], bold=True, color=C["blue"])
    bt = slide.shapes.add_textbox(Inches(x + 0.2), Inches(top + 0.55), Inches(w - 0.4), Inches(h - (top - y) - 0.7))
    bt.name = f"{name}_body"
    if isinstance(body, list):
        _bullets(bt.text_frame, body, SZ["small"] + 2)
    else:
        _text(bt.text_frame, body, SZ["small"] + 2)


def add_shape_text(slide, shape, x, y, w, h, text, fill, size=SZ["body"], color=None, name="shape"):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.name = name
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    tf = s.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    _text(tf, text, size, bold=True, color=color or C["white"], align=PP_ALIGN.CENTER)
    return s


# ---------------------------------------------------------------- スライド定義
LONG_A = (
    "この段落は文字切れとオーバーフロー検査のための架空の説明文です。共通の中間形式を用いることで、"
    "スライド資料と Web 図解の間で見出し、本文、箇条書き、表、画像といった要素を対応付け、"
    "どちらの形式からでも同じ内容を再構成できるようにします。変換の各段階では要素の種別と座標、"
    "文字サイズを数値として保持し、画像化に頼らずに編集可能な状態を維持します。数値はすべて例示であり、"
    "実在の案件や成果を示すものではありません。長文の折返しと段組の扱いを確認するための文章です。"
    "見出しと本文の間隔、行間、段落間の余白も同時に確認します。"
)
LONG_B = (
    "二段目の本文も同じく架空の説明です。未対応の要素が含まれる場合でも変換を停止せず、警告として記録し、"
    "代替の表現へ置き換えることで、閲覧者が欠落に気づける状態を保ちます。文字が領域を超える場合は"
    "縮小または分割の規則を適用し、結果を検査項目として記録します。ここで述べる仕組みや効果はサンプル用の"
    "記述であり、特定の組織や製品を指すものではありません。段組の両側が同程度の分量になるよう調整しています。"
    "検査では推定行数と枠の高さを比較し、超過があれば記録します。フッターとページ番号が本文と重ならないことも併せて確認します。"
)


def build(out_dir: Path = OUT_DIR) -> tuple[Path, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    assets = save_assets()
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
    blank = prs.slide_layouts[6]
    manifest_slides: list[dict] = []

    def new_slide(idx: int, sid: str, stype: str, title: str, req: list[str]) -> object:
        s = prs.slides.add_slide(blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = C["white"] if idx != 1 else C["light"]
        manifest_slides.append({"index": idx, "id": sid, "type": stype, "title": title, "requiredElementTypes": req})
        return s

    # 01 表紙
    s = new_slide(1, "slide-01", "cover", "PowerPoint・Web図解 双方向変換デモ", ["text", "shape", "footer", "pageNumber"])
    for i, (cx, r, col) in enumerate([(9.6, 1.6, C["blue"]), (11.0, 1.1, C["cyan"]), (10.2, 0.6, C["navy"])]):
        o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx), Inches(1.4 + i * 1.5), Inches(r * 2), Inches(r * 2))
        o.name = f"cover_circle_{i + 1}"
        o.fill.solid()
        o.fill.fore_color.rgb = col
        o.line.fill.background()
        o.shadow.inherit = False
    tb = s.shapes.add_textbox(Inches(M), Inches(2.2), Inches(8.4), Inches(1.9))
    tb.name = "title"
    _text(tb.text_frame, "PowerPoint・Web図解 双方向変換デモ", 40, bold=True, color=C["navy"])
    sb = s.shapes.add_textbox(Inches(M), Inches(4.1), Inches(8.2), Inches(0.6))
    sb.name = "subtitle"
    _text(sb.text_frame, "匿名化サンプル資料", 24, color=C["blue"])
    ds = s.shapes.add_textbox(Inches(M), Inches(4.8), Inches(8.2), Inches(1.0))
    ds.name = "description"
    _text(ds.text_frame, "変換アプリの受入試験に使う模擬資料です。すべての内容・数値は架空（サンプル）です。", SZ["body"], color=C["ink"])
    add_footer(s, 1)

    # 02 目的と対象
    s = new_slide(2, "slide-02", "two_column_cards", "目的と対象", ["text", "shape", "bullets", "footer", "pageNumber"])
    add_title(s, "目的と対象")
    bt = s.shapes.add_textbox(Inches(M), Inches(1.6), Inches(5.6), Inches(4.5))
    bt.name = "body"
    _bullets(bt.text_frame, ["一度作った資料を説明用・閲覧用・再利用用へ展開する", "スライドと Web 図解の両方を同じ構造から生成する", "非エンジニアがブラウザ操作だけで変換できるようにする", "未対応要素は警告と代替処理で扱い、停止させない"])
    for i, (t, b) in enumerate([("説明用", "会議での投影・配布に使う 16:9 のスライド"), ("閲覧用", "ブラウザで読む図解。端末幅に応じて並べ替える"), ("再利用用", "共通 JSON から別テンプレートへ再展開する")]):
        add_card(s, 6.9, 1.5 + i * 1.6, 5.8, 1.4, t, b, name=f"card_{i + 1}")
    add_footer(s, 2)

    # 03 現状課題
    s = new_slide(3, "slide-03", "three_column_cards", "現状課題", ["text", "shape", "footer", "pageNumber"])
    add_title(s, "現状課題")
    for i, (t, b, ic) in enumerate([("二重作業", ["同じ内容を二つの形式で作り直す", "作成時間が倍になる（例示）"], "circle"), ("更新ずれ", ["修正が片方にしか反映されない", "どちらが最新か分からない"], "triangle"), ("属人化", ["変換手順が担当者の手作業に依存", "引き継ぎ資料が残らない"], "diamond")]):
        add_card(s, M + i * 4.1, 1.6, 3.8, 4.2, t, b, icon=ic, name=f"card_{i + 1}")
    add_footer(s, 3)

    # 04 変換フロー
    s = new_slide(4, "slide-04", "process_flow", "変換フロー", ["text", "shape", "connector", "footer", "pageNumber"])
    add_title(s, "変換フロー")
    steps = ["PPTX入力", "抽出", "Presentation JSON", "AI再構成", "Web描画", "検証"]
    bw, gap = 1.75, 0.3
    x0 = (SW - (bw * 6 + gap * 5)) / 2
    for i, st in enumerate(steps):
        x = x0 + i * (bw + gap)
        add_shape_text(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 3.0, bw, 1.2, st, C["blue"] if i != 2 else C["cyan"], size=14, name=f"step_{i + 1}")
        if i < 5:
            con = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x + bw), Inches(3.6), Inches(x + bw + gap), Inches(3.6))
            con.name = f"arrow_{i + 1}"
            con.line.color.rgb = C["navy"]
            con.line.width = Pt(2.5)
            # python-pptx は矢頭 API を持たないため XML（a:tailEnd）を直接追加する
            ln = con.line._get_or_add_ln()
            from lxml import etree
            from pptx.oxml.ns import qn

            tail = etree.SubElement(ln, qn("a:tailEnd"))
            tail.set("type", "triangle")
    nt = s.shapes.add_textbox(Inches(M), Inches(4.6), Inches(SW - 2 * M), Inches(0.8))
    nt.name = "note"
    _text(nt.text_frame, "AI は構造の分類と候補提示のみを担当し、座標計算とファイル生成は決定的な処理で行う（例示）。", SZ["small"] + 2, color=C["muted"])
    add_footer(s, 4)

    # 05 Before / After
    s = new_slide(5, "slide-05", "before_after", "Before / After", ["text", "shape", "bullets", "footer", "pageNumber"])
    add_title(s, "Before / After")
    for i, (t, items, fill, col) in enumerate([("Before", ["資料ごとに手作業で作り直す", "修正漏れが月に数件（例示）", "Web 公開まで数日（例示）"], C["orange_light"], C["warning"]), ("After", ["共通 JSON から両形式を生成", "修正は 1 か所で反映", "公開までの日数を短縮（例示）"], C["green_light"], C["success"])]):
        x = M + i * 6.9
        panel = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(1.6), Inches(5.2), Inches(4.4))
        panel.name = f"panel_{t.lower()}"
        panel.fill.solid()
        panel.fill.fore_color.rgb = fill
        panel.line.fill.background()
        panel.shadow.inherit = False
        hd = s.shapes.add_textbox(Inches(x + 0.3), Inches(1.8), Inches(4.6), Inches(0.6))
        hd.name = f"heading_{t.lower()}"
        _text(hd.text_frame, t, 24, bold=True, color=col)
        bb = s.shapes.add_textbox(Inches(x + 0.3), Inches(2.5), Inches(4.6), Inches(3.2))
        bb.name = f"bullets_{t.lower()}"
        _bullets(bb.text_frame, items)
    arrow = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(6.05), Inches(3.3), Inches(1.2), Inches(1.0))
    arrow.name = "big_arrow"
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = C["navy"]
    arrow.line.fill.background()
    arrow.shadow.inherit = False
    add_footer(s, 5)

    # 06 数値カード
    s = new_slide(6, "slide-06", "kpi_cards", "数値カード", ["text", "shape", "footer", "pageNumber"])
    add_title(s, "数値カード")
    for i, (v, lab) in enumerate([("12ページ（仮）", "模擬資料の枚数"), ("3形式（仮）", "PPTX / JSON / Web"), ("100%架空データ", "実在情報を含まない")]):
        x = M + i * 4.1
        box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.8), Inches(3.8), Inches(3.0))
        box.name = f"kpi_card_{i + 1}"
        box.fill.solid()
        box.fill.fore_color.rgb = C["light"]
        box.line.color.rgb = C["line"]
        box.shadow.inherit = False
        vt = s.shapes.add_textbox(Inches(x + 0.2), Inches(2.1), Inches(3.4), Inches(1.4))
        vt.name = f"kpi_value_{i + 1}"
        _text(vt.text_frame, v, SZ["kpi"] if i < 2 else 26, bold=True, color=C["navy"], align=PP_ALIGN.CENTER)
        lt = s.shapes.add_textbox(Inches(x + 0.2), Inches(3.7), Inches(3.4), Inches(0.6))
        lt.name = f"kpi_label_{i + 1}"
        _text(lt.text_frame, lab, SZ["body"], color=C["muted"], align=PP_ALIGN.CENTER)
    nt = s.shapes.add_textbox(Inches(M), Inches(5.2), Inches(SW - 2 * M), Inches(0.6))
    nt.name = "note"
    _text(nt.text_frame, "効果数値ではなく変換試験用の表示サンプル", SZ["note"], color=C["muted"])
    add_footer(s, 6)

    # 07 表
    s = new_slide(7, "slide-07", "table", "要素別の変換方式", ["text", "table", "footer", "pageNumber"])
    add_title(s, "要素別の変換方式")
    rows = [["要素", "Web化", "PPTX化", "方式"], ["タイトル", "見出しへ変換", "テキストボックス", "編集可能な文字"], ["本文", "段落・箇条書き", "テキストボックス", "編集可能な文字"], ["画像", "img 要素", "画像として配置", "縦横比を保持"], ["表", "table 要素", "表オブジェクト", "セル単位で編集可能"], ["複雑図解", "画像化フォールバック", "画像 + 警告", "条件付き"]]
    gf = s.shapes.add_table(len(rows), 4, Inches(M), Inches(1.6), Inches(SW - 2 * M), Inches(4.2))
    gf.name = "table"
    tbl = gf.table
    for r, row in enumerate(rows):
        for c, txt in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = txt
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(SZ["small"] + 2)
                    run.font.name = FONT
                    run.font.bold = r == 0
                    run.font.color.rgb = C["white"] if r == 0 else C["ink"]
            cell.fill.solid()
            cell.fill.fore_color.rgb = C["navy"] if r == 0 else (C["light"] if r % 2 == 0 else C["white"])
    add_footer(s, 7)

    # 08 画像配置
    s = new_slide(8, "slide-08", "images", "画像配置", ["text", "image", "footer", "pageNumber"])
    add_title(s, "画像配置")
    pic = s.shapes.add_picture(io.BytesIO(assets["wide_gradient.png"]), Inches(M), Inches(1.7), width=Inches(7.2))
    pic.name = "image_wide"
    cap = s.shapes.add_textbox(Inches(M), Inches(1.7 + 7.2 * 675 / 1200 + 0.1), Inches(7.2), Inches(0.4))
    cap.name = "caption_wide"
    _text(cap.text_frame, "図1: 横長の抽象グラデーション（生成画像・サンプル）", SZ["note"], color=C["muted"])
    pic2 = s.shapes.add_picture(io.BytesIO(assets["tall_pattern.png"]), Inches(8.3), Inches(1.7), height=Inches(4.2))
    pic2.name = "image_tall"
    ds = s.shapes.add_textbox(Inches(11.3), Inches(1.7), Inches(1.6), Inches(4.2))
    ds.name = "image_description"
    _text(ds.text_frame, "縦長の幾何学模様。縦横比を保持したまま配置されることを確認する。", SZ["small"], color=C["ink"])
    add_footer(s, 8)

    # 09 複合レイアウト
    s = new_slide(9, "slide-09", "composite", "複合レイアウト", ["text", "shape", "connector", "footer", "pageNumber"])
    add_title(s, "複合レイアウト")
    add_card(s, M, 1.5, 5.9, 2.0, "上段カード A", "複数領域の整列と余白を確認するためのカード（例示）。", name="card_a")
    add_card(s, M + 6.2, 1.5, 5.9, 2.0, "上段カード B", ["箇条書きを含むカード", "折返しと行間の確認"], name="card_b")
    line = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(M + 0.4), Inches(5.0), Inches(SW - M - 0.4), Inches(5.0))
    line.name = "timeline_axis"
    line.line.color.rgb = C["line"]
    line.line.width = Pt(3)
    for i, lab in enumerate(["準備", "変換", "検査", "修正", "公開"]):
        x = M + 0.4 + i * (SW - 2 * M - 0.8) / 4
        dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - 0.18), Inches(4.82), Inches(0.36), Inches(0.36))
        dot.name = f"milestone_{i + 1}"
        dot.fill.solid()
        dot.fill.fore_color.rgb = C["cyan"] if i % 2 == 0 else C["blue"]
        dot.line.fill.background()
        dot.shadow.inherit = False
        lt = s.shapes.add_textbox(Inches(x - 0.8), Inches(5.3), Inches(1.6), Inches(0.5))
        lt.name = f"milestone_label_{i + 1}"
        _text(lt.text_frame, f"{i + 1}. {lab}", SZ["small"] + 2, color=C["ink"], align=PP_ALIGN.CENTER)
    add_footer(s, 9)

    # 10 長文・オーバーフロー試験
    s = new_slide(10, "slide-10", "long_text", "長文・オーバーフロー試験", ["text", "shape", "footer", "pageNumber"])
    add_title(s, "長文・オーバーフロー試験")
    for i, txt in enumerate([LONG_A, LONG_B]):
        col = s.shapes.add_textbox(Inches(M + i * 6.2), Inches(1.5), Inches(5.9), Inches(3.9))
        col.name = f"column_{i + 1}"
        _text(col.text_frame, txt, SZ["small"] + 2, color=C["ink"])
    wb = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(M), Inches(5.55), Inches(SW - 2 * M), Inches(0.75))
    wb.name = "warning_box"
    wb.fill.solid()
    wb.fill.fore_color.rgb = C["orange_light"]
    wb.line.color.rgb = C["warning"]
    wb.line.width = Pt(2)
    wb.shadow.inherit = False
    wb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    _text(wb.text_frame, "警告: 本文が領域を超える場合は縮小または分割の規則を適用する（サンプル表示）。", SZ["small"] + 2, bold=True, color=C["ink"])
    add_footer(s, 10)

    # 11 対象外要素の説明
    s = new_slide(11, "slide-11", "scope", "対象外要素の説明", ["text", "shape", "footer", "pageNumber"])
    add_title(s, "対象外要素の説明")
    for i, (t, items, col) in enumerate([("対応", ["タイトル・本文・箇条書き", "画像・図形・線", "表"], C["success"]), ("条件付き", ["グループ図形（2 階層まで）", "複雑図解は画像化", "SmartArt は文字列のみ"], C["warning"]), ("対象外", ["動画", "マクロ", "3D", "複雑アニメーション"], C["muted"])]):
        x = M + i * 4.1
        hd = add_shape_text(s, MSO_SHAPE.RECTANGLE, x, 1.6, 3.8, 0.6, t, col, size=SZ["body"], name=f"scope_header_{i + 1}")
        bb = s.shapes.add_textbox(Inches(x), Inches(2.3), Inches(3.8), Inches(3.4))
        bb.name = f"scope_items_{i + 1}"
        _bullets(bb.text_frame, items)
    nt = s.shapes.add_textbox(Inches(M), Inches(5.8), Inches(SW - 2 * M), Inches(0.5))
    nt.name = "note"
    _text(nt.text_frame, "対象外要素は本資料に含めず、変換時は警告メタデータとして扱う。", SZ["note"], color=C["muted"])
    add_footer(s, 11)

    # 12 まとめ
    s = new_slide(12, "slide-12", "summary", "まとめ", ["text", "shape", "footer", "pageNumber"])
    add_title(s, "まとめ")
    for i, (t, b) in enumerate([("共通構造", "1 つの JSON から両形式を生成する"), ("双方向変換", "PPTX と Web のどちらからでも往復できる"), ("人による確認", "警告と検査結果を人が最終確認する")]):
        add_card(s, M + i * 4.1, 1.6, 3.8, 2.6, t, b, name=f"message_{i + 1}")
    nx = s.shapes.add_textbox(Inches(M), Inches(4.5), Inches(6), Inches(0.5))
    nx.name = "next_action_label"
    _text(nx.text_frame, "次のアクション", SZ["body"], bold=True, color=C["navy"])
    for i, lab in enumerate(["サンプルで往復変換", "検査結果を確認", "テンプレートを適用"]):
        add_shape_text(s, MSO_SHAPE.ROUNDED_RECTANGLE, M + i * 4.1, 5.1, 3.8, 0.7, lab, C["blue"], size=SZ["small"] + 2, name=f"action_button_{i + 1}")
    add_footer(s, 12)

    # プロパティは匿名化
    cp = prs.core_properties
    cp.author = THEME["author"]
    cp.last_modified_by = THEME["author"]
    cp.title = "PowerPoint・Web図解 双方向変換デモ"
    cp.subject = "Sample Only"
    cp.comments = ""
    cp.keywords = ""
    cp.category = ""
    for part in prs.part.package.iter_parts():
        if str(part.partname) == "/docProps/app.xml":
            part._blob = re.sub(rb"<Application>[^<]*</Application>", b"<Application>Microsoft Office PowerPoint</Application>", part.blob)

    pptx_path = out_dir / PPTX_NAME
    prs.save(str(pptx_path))
    manifest = {
        "schemaVersion": "1.0",
        "documentId": "mock-bidirectional-conversion",
        "title": "PowerPoint・Web図解 双方向変換デモ",
        "classification": "Sample Only",
        "aspectRatio": "16:9",
        "theme": THEME["name"],
        "slideCount": len(manifest_slides),
        "generator": "python-pptx + Pillow (mock-pptx/src/generate_mock_pptx.py)",
        "slides": manifest_slides,
    }
    manifest_path = out_dir / MANIFEST_NAME
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return pptx_path, manifest_path


if __name__ == "__main__":
    p, m = build()
    print(p)
    print(m)
