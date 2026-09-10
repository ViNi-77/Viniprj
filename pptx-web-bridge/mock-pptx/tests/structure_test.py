"""構造検査（模擬PPTX生成指示書 8 章）: 枚数、題名、ページ番号、比率、画像縦横比、範囲外要素、再読込。"""
from __future__ import annotations

import io
import re

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def _shapes(slide):
    out = []
    for sh in slide.shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            out.extend(sh.shapes)
        else:
            out.append(sh)
    return out


def test_reload_and_slide_count(pptx_path, manifest):
    prs = Presentation(str(pptx_path))
    assert len(prs.slides) == 12 == manifest["slideCount"]
    assert abs(prs.slide_width / prs.slide_height - 16 / 9) < 0.01


def test_every_slide_has_title_matching_manifest(pptx_path, manifest):
    prs = Presentation(str(pptx_path))
    for slide, entry in zip(prs.slides, manifest["slides"]):
        titles = [sh for sh in slide.shapes if sh.name == "title"]
        assert titles, f"{entry['id']}: title 図形がありません"
        assert titles[0].text_frame.text.strip() == entry["title"]


def test_page_numbers_sequential(pptx_path):
    prs = Presentation(str(pptx_path))
    nums = []
    for slide in prs.slides:
        pn = [sh for sh in slide.shapes if sh.name == "page_number"]
        assert pn, "ページ番号がありません"
        nums.append(pn[0].text_frame.text.strip())
    assert nums == [f"{i:02d}" for i in range(1, 13)]


def test_footer_on_every_slide(pptx_path):
    prs = Presentation(str(pptx_path))
    for slide in prs.slides:
        ft = [sh for sh in slide.shapes if sh.name == "footer"]
        assert ft and ft[0].text_frame.text == "Anonymous Conversion Test / Sample Only"


def test_required_element_types(pptx_path, manifest):
    prs = Presentation(str(pptx_path))
    for slide, entry in zip(prs.slides, manifest["slides"]):
        kinds = set()
        for sh in _shapes(slide):
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                kinds.add("image")
            elif getattr(sh, "has_table", False) and sh.has_table:
                kinds.add("table")
            elif sh.shape_type == MSO_SHAPE_TYPE.LINE:
                kinds.add("connector")
            elif sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                kinds.add("shape")
            if sh.has_text_frame and sh.text_frame.text.strip():
                kinds.add("text")
                if "• " in sh.text_frame.text:
                    kinds.add("bullets")
            if sh.name == "footer":
                kinds.add("footer")
            if sh.name == "page_number":
                kinds.add("pageNumber")
        missing = set(entry["requiredElementTypes"]) - kinds
        assert not missing, f"{entry['id']}: 不足 {missing}"


def test_elements_inside_slide(pptx_path):
    prs = Presentation(str(pptx_path))
    W, H = prs.slide_width, prs.slide_height
    for i, slide in enumerate(prs.slides, 1):
        for sh in _shapes(slide):
            assert sh.left >= 0 and sh.top >= 0, f"slide {i} {sh.name}: 負の座標"
            assert sh.left + sh.width <= W + 1 and sh.top + sh.height <= H + 1, f"slide {i} {sh.name}: スライド外"


def test_images_keep_aspect_ratio(pptx_path):
    from PIL import Image

    prs = Presentation(str(pptx_path))
    found = 0
    for slide in prs.slides:
        for sh in _shapes(slide):
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                with Image.open(io.BytesIO(sh.image.blob)) as im:
                    ratio_img = im.width / im.height
                ratio_box = sh.width / sh.height
                assert abs(ratio_img - ratio_box) / ratio_img < 0.02, f"{sh.name}: 縦横比が崩れています"
                found += 1
    assert found == 2


def test_long_text_within_box(pptx_path):
    """長文ページ: 近似計測で本文が枠に収まる（全角 1 文字 = サイズ幅、行間 1.35）。"""
    prs = Presentation(str(pptx_path))
    slide = prs.slides[9]
    for sh in slide.shapes:
        if sh.name.startswith("column_"):
            text = sh.text_frame.text
            size = sh.text_frame.paragraphs[0].runs[0].font.size.pt
            width_pt = sh.width / 12700 - 14
            chars_per_line = max(1, int(width_pt / size))
            lines = -(-len(text) // chars_per_line)
            need_pt = lines * size * 1.35
            assert need_pt <= sh.height / 12700, f"{sh.name}: 文字あふれの可能性（{need_pt:.0f}pt > {sh.height / 12700:.0f}pt）"
            assert 250 <= len(text) <= 350


def test_text_is_editable_shapes_not_images(pptx_path):
    prs = Presentation(str(pptx_path))
    text_shapes = sum(1 for s in prs.slides for sh in _shapes(s) if sh.has_text_frame and sh.text_frame.text.strip())
    pictures = sum(1 for s in prs.slides for sh in _shapes(s) if sh.shape_type == MSO_SHAPE_TYPE.PICTURE)
    assert text_shapes > 80 and pictures == 2


def test_no_hyperlinks_macros_media(pptx_path):
    import zipfile

    with zipfile.ZipFile(pptx_path) as z:
        names = z.namelist()
        assert not any(n.endswith(".bin") and "vba" in n.lower() for n in names)
        assert not any("/media/" in n and not n.endswith(".png") for n in names)
        for n in names:
            if n.endswith(".rels"):
                assert "hyperlink" not in z.read(n).decode("utf-8", "ignore")
    prs = Presentation(str(pptx_path))
    for slide in prs.slides:
        for sh in _shapes(slide):
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        assert r.hyperlink.address is None
