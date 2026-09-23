"""合成資料を実アプリで操作し、オフライン操作マニュアル用の画像を撮影する。

利用者の原本・設定には触れない。本文やステータスをDOMに注入しない。
配布用画像はWindowsで撮影する。撮影環境は開発者向け証跡に記録する。
"""
from __future__ import annotations

import argparse
from contextlib import contextmanager
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile

from docx import Document
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from playwright.sync_api import expect, sync_playwright

from ui_smoke import ROOT, assert_laptop, browser_path, reserve_port, wait_jobs, wait_server


@contextmanager
def capture_workspace(path: Path | None):
    """画面に表示するサンプルパスを指定可能にする。既存フォルダは使わない。"""
    if path is None:
        with tempfile.TemporaryDirectory(prefix="contextgen-manual-") as directory:
            yield Path(directory)
        return
    path.mkdir(parents=True, exist_ok=False)
    try:
        yield path
    finally:
        shutil.rmtree(path)


def capture(output: Path, workspace: Path | None = None, record: Path | None = None):
    output.mkdir(parents=True, exist_ok=True)
    pictures = []
    with capture_workspace(workspace) as temp:
        sources = temp / "操作例 資料"
        sources.mkdir()
        note = sources / "01_点検手順.txt"
        original = (
            "設備点検の手順（操作説明用サンプル）\n\n"
            "1. 始業前に電源・カバー・周囲の安全を確認する。\n"
            "2. 点検日と担当者、確認結果を点検表に記録する。\n"
            "3. 異常があった場合は運転を停止し、責任者に報告する。\n\n"
            "確認事項：交換部品の在庫と、次回点検日を確認してください。\n"
        )
        note.write_text(original, encoding="utf-8")
        doc = Document()
        doc.add_heading("改善案の比較", 0)
        doc.add_paragraph("操作説明のために作成した架空の資料です。")
        table = doc.add_table(rows=1, cols=3)
        for cell, text in zip(table.rows[0].cells, ["案", "変更内容", "確認事項"]):
            cell.text = text
        for values in (["A案", "点検表の記入欄を整理", "記入時間を測定"], ["B案", "予備部品の置場を統一", "必要な棚数を確認"]):
            for cell, text in zip(table.add_row().cells, values):
                cell.text = text
        doc.save(sources / "02_改善案.docx")
        book = Workbook()
        sheet = book.active
        sheet.title = "点検一覧"
        for row in (["対象", "結果", "確認事項"], ["設備A", "異常なし", "次回は月曜日"], ["設備B", "要確認", "交換部品の在庫"]):
            sheet.append(row)
        equipment_table = Table(displayName="EquipmentChecks", ref="A1:C3")
        equipment_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(equipment_table)
        book.save(sources / "03_点検一覧.xlsx")
        upload = temp / "今回の確認事項.txt"
        upload.write_text("今回だけの資料（操作例）\n改善案AとBの共通点・相違点を整理する。", encoding="utf-8")
        port = reserve_port()
        base = f"http://127.0.0.1:{port}"
        log = temp / "server.log"
        with log.open("w", encoding="utf-8") as stream:
            proc = subprocess.Popen([sys.executable, "-m", "contextgen_kai", "--state-dir", str(temp / "state"), "--port", str(port), "--no-browser"], cwd=ROOT, stdout=stream, stderr=subprocess.STDOUT)
            try:
                wait_server(base, proc, log)
                with sync_playwright() as playwright:
                    browser = playwright.chromium.launch(executable_path=browser_path(), headless=True)
                    page = browser.new_page(base_url=base, viewport={"width": 1440, "height": 940}, locale="ja-JP")
                    errors = []
                    page.on("pageerror", lambda error: errors.append(str(error)))
                    page.on("dialog", lambda dialog: dialog.accept())
                    page.goto(base, wait_until="networkidle")
                    expect(page.locator("#connection")).to_contain_text("ローカル接続中")

                    def shot(name):
                        expect(page.locator(".toast")).to_have_count(0, timeout=15000)
                        assert_laptop(page, name)
                        page.screenshot(path=str(output / name), full_page=True)
                        pictures.append(name)

                    shot("01_home.png")
                    page.locator("#add-library").click()
                    page.locator("#library-name").fill("設備点検・改善資料（操作例）")
                    page.locator("#library-path").fill(str(sources))
                    shot("02_register.png")
                    page.locator("#library-form [type=submit]").click()
                    wait_jobs(page, 3)
                    page.locator("#refresh-status").click()
                    expect(page.locator("#stat-ok")).to_contain_text("3")
                    shot("03_read_complete.png")
                    page.locator('[data-view="documents"]').click()
                    expect(page.locator("#documents tr")).to_have_count(3)
                    shot("04_documents.png")
                    page.get_by_role("button", name="01_点検手順.txt", exact=True).click()
                    page.locator("#edited-text").fill(original.replace("次回点検日を確認してください。", "次回点検日は月曜日です。"))
                    shot("05_edit.png")
                    page.locator("#save-document").click()
                    expect(page.locator("#document-dialog")).not_to_be_visible()
                    page.locator("#documents input[type=checkbox]").first.check()
                    page.locator("#collection-from-selection").click()
                    page.locator("#collection-name").fill("点検前に確認したいこと")
                    page.locator("#collection-purpose").select_option("questions")
                    page.locator("#collection-target").select_option("builder")
                    page.locator("#collection-advanced > summary").click()
                    page.locator("#collection-audience").fill("設備の点検担当者")
                    page.locator("#collection-answer-scope").fill("始業前の確認事項と記録方法")
                    page.locator("#collection-out-of-scope").fill("設備の設計変更と故障原因の断定")
                    page.locator("#collection-instructions").fill("未確認の事項を箇条書きにし、根拠となる資料を示してください。")
                    shot("06_collection.png")
                    page.locator("#collection-form [type=submit]").click()
                    expect(page.locator("#collection-dialog")).not_to_be_visible()
                    page.get_by_role("button", name="ナレッジを生成", exact=True).click()
                    page.locator("#confirm-export").click()
                    wait_jobs(page)
                    page.locator("#refresh-exports").click()
                    expect(page.get_by_role("link", name="一式をダウンロード")).to_be_visible()
                    shot("07_exports.png")
                    page.get_by_role("button", name="指示文", exact=True).first.click()
                    expect(page.locator("#prompt-text")).not_to_have_value("")
                    shot("08_prompt.png")
                    page.locator("#prompt-dialog [data-close]").first.click()
                    page.locator('[data-view="schedules"]').click()
                    page.locator("#add-schedule").click()
                    page.locator("#schedule-name").fill("平日の点検資料を更新")
                    page.locator("#schedule-collection").select_option(index=1)
                    page.locator("#schedule-frequency").select_option("weekly")
                    page.locator("#schedule-time").fill("09:00")
                    shot("09_schedule.png")
                    # 撮影時に予約が偶発実行されないよう停止状態で保存する。
                    page.locator("#schedule-enabled").uncheck()
                    page.locator("#schedule-form [type=submit]").click()
                    expect(page.locator("#schedule-dialog")).not_to_be_visible()
                    shot("10_settings.png")
                    # 保存済みの手修正に対して原本が変わった実際の競合を作る。
                    note.write_text(original + "\n追記：点検表の様式を改訂しました。\n", encoding="utf-8")
                    page.locator('[data-view="home"]').click()
                    page.locator("#libraries").get_by_role("button", name="読み取る", exact=True).first.click()
                    wait_jobs(page, 3)
                    page.locator('[data-view="documents"]').click()
                    expect(page.locator("#documents")).to_contain_text("修正の確認待ち")
                    page.get_by_role("button", name="01_点検手順.txt", exact=True).click()
                    expect(page.locator("#document-conflict")).to_be_visible()
                    shot("11_conflict.png")
                    page.locator("#document-dialog [data-close]").first.click()
                    page.locator('[data-view="exports"]').click()
                    page.get_by_role("button", name="ナレッジを生成", exact=True).click()
                    page.locator("#confirm-export").click()
                    wait_jobs(page)
                    page.locator("#refresh-exports").click()
                    expect(page.locator("#exports")).to_contain_text("確認待ち")
                    shot("12_held.png")
                    page.locator('[data-view="home"]').click()
                    page.locator("#upload-files").set_input_files(upload)
                    wait_jobs(page, 4)
                    page.locator("#refresh-status").click()
                    expect(page.locator("#stat-total")).to_contain_text("4")
                    shot("13_upload.png")
                    # 既存Officeの抽出結果と別登録元の単発資料を、実画面で選び直す。
                    page.locator('[data-view="exports"]').click()
                    page.locator("#add-collection").click()
                    page.locator("#collection-name").fill("設備点検と改善案の案件セット")
                    for choice in page.locator("#collection-libraries input").all():
                        choice.check()
                    page.locator("#collection-mode").select_option("fixed")
                    page.locator("#collection-target").select_option("studio")
                    page.locator("#collection-cleanup").select_option("standard")
                    page.locator("#collection-purpose").select_option("compare")
                    page.locator("#collection-advanced > summary").click()
                    page.locator("#collection-description").fill("設備点検一覧と、今回の改善案確認事項を集めた資料です。")
                    page.locator("#collection-audience").fill("点検担当者と改善案の確認者")
                    page.locator("#collection-answer-scope").fill("点検結果、改善案の共通点・相違点、追加確認事項")
                    page.locator("#collection-out-of-scope").fill("資料に記載のない数値や故障原因の推測")
                    page.locator("#preview-collection").click()
                    expect(page.locator("#collection-preview .candidate-row")).to_have_count(4)
                    for name in ("03_点検一覧.xlsx", "今回の確認事項.txt"):
                        page.locator("#collection-preview .candidate-row").filter(has_text=name).locator("input").check()
                    page.locator("#collection-preview").scroll_into_view_if_needed()
                    shot("14_multifolder.png")
                    page.locator("#collection-form [type=submit]").click()
                    card = page.locator("#collections .collection-card").filter(has_text="設備点検と改善案の案件セット")
                    card.get_by_role("button", name="編集", exact=True).click()
                    page.locator("#preview-collection").click()
                    page.locator("#collection-preview .candidate-row").filter(has_text="03_点検一覧.xlsx").get_by_role("button", name="整理前後", exact=True).click()
                    expect(page.locator("#cleanup-before")).not_to_have_value("")
                    expect(page.locator("#cleanup-after")).not_to_have_value("")
                    assert page.locator("#cleanup-before").input_value() != page.locator("#cleanup-after").input_value()
                    shot("15_cleanup.png")
                    page.locator("#cleanup-dialog [data-close]").first.click()
                    page.locator("#collection-form [type=submit]").click()
                    card.get_by_role("button", name="ナレッジを生成", exact=True).click()
                    expect(page.locator("#preflight-summary")).to_contain_text("設備点検と改善案")
                    shot("16_preflight.png")
                    page.locator("#confirm-export").click()
                    wait_jobs(page)
                    page.locator("#refresh-exports").click()
                    expect(page.locator("#exports .export-row").first).to_contain_text("現在の出力")
                    shot("17_studio.png")
                    page.locator("#exports .export-row").first.get_by_role("button", name="登録状況", exact=True).click()
                    expect(page.locator("#handoff-summary")).to_contain_text("Copilot Studio")
                    page.locator("#handoff-select-all").click()
                    shot("18_handoff.png")
                    # 合成資料だけの記録例。外部Copilotへの接続・アップロードは行わない。
                    page.locator("#save-handoff").click()
                    expect(page.locator("#handoff-dialog")).not_to_be_visible()
                    assert not errors, errors
                    version = page.locator("#version").inner_text()
                    browser.close()
            finally:
                proc.terminate()
                try:
                    proc.wait(timeout=10)
                except subprocess.TimeoutExpired:
                    proc.kill()
                    proc.wait(timeout=5)
    result = {"application_version": version, "platform": sys.platform, "viewport": "1440x940", "fixtures": "synthetic only", "page_errors": errors, "screenshots": pictures, "note": "実アプリ・実APIの操作を撮影。Windows固有機能の受入証跡ではありません。"}
    record = record or ROOT / "build/manual-capture.json"
    record.parent.mkdir(parents=True, exist_ok=True)
    record.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=ROOT / "images")
    parser.add_argument("--workspace", type=Path, help="撮影専用の未作成フォルダ（終了時に削除）")
    parser.add_argument("--record", type=Path, default=ROOT / "build/manual-capture.json")
    args = parser.parse_args()
    capture(args.output, args.workspace, args.record)
