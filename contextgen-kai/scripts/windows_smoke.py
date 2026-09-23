"""実配布 EXE を開発 PATH なしで検証。実予約は一時 state 配下で作成し必ず解除する。"""
from __future__ import annotations
import argparse
from contextlib import nullcontext
from datetime import datetime, timedelta, timezone
import io
import hashlib
from html.parser import HTMLParser
import json
import os
from pathlib import Path
import shutil
import socket
import sqlite3
import subprocess
import sys
import tempfile
import time
import urllib.error
import urllib.request
from urllib.parse import quote
import zipfile
import xml.etree.ElementTree as ET

from contextgen_kai import __version__


class ManualImages(HTMLParser):
    def __init__(self):
        super().__init__()
        self.sources = []

    def handle_starttag(self, tag, attrs):
        if tag == 'img':
            attributes = dict(attrs)
            # 拡大用の空imgにはクリック後に選択した画像を設定する。
            if attributes.get('id') == 'image-large' and not attributes.get('src'):
                return
            self.sources.append(attributes.get('src', ''))


def verify_distribution_layout(root: Path):
    """利用者ZIPの入口は4点だけ。開発用文書・検証来歴が紛れたら不合格にする。"""
    from contextgen_kai.api import MANUAL_ASSETS
    assert {path.name for path in root.iterdir()} == {
        'ContextgenKai.exe', 'contextgen改_操作マニュアル.html', '最初にお読みください.txt', '_internal',
    }, 'Distribution root must contain exactly the four user-facing items'
    internal = root / '_internal'
    assert internal.is_dir()
    ocr = internal / 'ocr'
    assert {path.name for path in ocr.iterdir()} == {'tesseract.exe', 'tessdata'}
    assert {path.name for path in (ocr / 'tessdata').iterdir()} == {'eng.traineddata', 'jpn.traineddata', 'LICENSE'}
    licenses = internal / 'THIRD_PARTY_LICENSES'
    assert (licenses / 'tesseract-LICENSE.txt').is_file()
    assert (licenses / 'tessdata-LICENSE.txt').is_file()
    assert (licenses / 'Tcl-LICENSE.txt').is_file()
    assert (licenses / 'Tk-LICENSE.txt').is_file()
    assert (licenses / 'python/Python-LICENSE.txt').is_file()
    images = internal / 'manual/images'
    assert images.is_dir()
    assert {path.name for path in images.iterdir()} == MANUAL_ASSETS
    assert all(path.is_file() for path in images.iterdir())
    for name in ('BUILD-MANIFEST.json', 'BUNDLE-MANIFEST.json', 'capture-record.json', 'IMPLEMENTATION_CONTRACT.md',
                 '仕様書兼要件定義書.md', 'アプリ概要とバージョン履歴.md', 'アプリ基本設計基準書.md',
                 'start_windows.bat', 'update_windows.bat', '起動.bat', '更新.bat'):
        assert not list(root.rglob(name)), f'Developer evidence must not be included: {name}'
    # 展開された配布ファイルだけを確認。Python実行に必要なPYZ/base_library.zip内部は対象外。
    for path in root.rglob('*'):
        assert not any(part in {'.DS_Store', '__MACOSX', '__pycache__'} or part.startswith('._')
                       for part in path.relative_to(root).parts), f'Unwanted host metadata in distribution: {path}'
        if path.is_file():
            assert path.suffix.lower() not in {'.py', '.pyc', '.pdb'}, f'Developer source/cache/symbol file in distribution: {path}'


def verify_manual_assets(root: Path, fetch, *, packaged=False):
    """配布ファイルと実HTTPの一致を照合。OS固有処理から分けてローカルでも検証する。"""
    manual_bytes = fetch('/manual/')
    # read_textによるWindows改行変換を避け、配信実体をそのまま比べる。
    assert (root / 'contextgen改_操作マニュアル.html').read_bytes() == manual_bytes
    parser = ManualImages()
    parser.feed(manual_bytes.decode('utf-8'))
    assert len(parser.sources) >= 4, 'Manual must include actual app screenshots'
    prefix = '_internal/manual/images/' if packaged else 'images/'
    for source in parser.sources:
        assert source.startswith(prefix) and '/' not in source[len(prefix):] and '..' not in source, source
        image_bytes = fetch('/manual/' + quote(source))
        assert image_bytes == (root / source).read_bytes()
        assert image_bytes.startswith(b'\x89PNG\r\n\x1a\n'), source
    assert fetch('/manual/' + prefix + 'manual.js') == (root / prefix / 'manual.js').read_bytes()
    if not packaged:
        for document in ('README.md', '仕様書兼要件定義書.md', 'アプリ概要とバージョン履歴.md', 'アプリ基本設計基準書.md'):
            assert fetch('/manual/' + quote(document)) == (root / document).read_bytes()
    return len(parser.sources)


def fixtures(folder):
    from PIL import Image, ImageDraw, ImageFont
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from reportlab.pdfgen.canvas import Canvas
    folder.mkdir()
    (folder / '基本 資料.txt').write_text('安全確認の手順。設備点検を毎日行う。 contextgen acceptance evidence.', encoding='utf-8')
    image = Image.new('RGB', (1500, 240), 'white')
    draw = ImageDraw.Draw(image)
    fonts = Path(os.environ['WINDIR']) / 'Fonts'
    draw.text((40, 25), 'CONTEXTGEN SAFETY CHECK 123', font=ImageFont.truetype(str(fonts / 'arial.ttf'), 46), fill='black')
    japanese = next((fonts / f for f in ('msgothic.ttc', 'meiryo.ttc', 'YuGothM.ttc') if (fonts / f).exists()), None)
    if japanese is None:
        # CI の Server イメージに日本語フォントがない場合のみ、試験用画像の作成時に取得する。
        # 配布アプリの実行中にはダウンロードしない。
        japanese = folder.parent / 'NotoSansCJKjp-Regular.otf'
        url = 'https://raw.githubusercontent.com/notofonts/noto-cjk/523d033d6cb47f4a80c58a35753646f5c3608a78/Sans/OTF/Japanese/NotoSansCJKjp-Regular.otf'
        font_data = urllib.request.urlopen(url, timeout=60).read()
        assert hashlib.sha256(font_data).hexdigest() == '68a3fc98800b2a27b371f2fb79991daf3633bd89309d4ffaa6946fd587f375b5'
        japanese.write_bytes(font_data)
    draw.text((40, 120), '安全確認 毎日の設備点検', font=ImageFont.truetype(str(japanese), 46), fill='black')
    image.save(folder / '画像.png')
    doc = Document()
    doc.add_paragraph('WORD acceptance table and embedded image')
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = 'Item', 'Inspection'
    doc.add_picture(str(folder / '画像.png'))
    doc.save(folder / '手順.docx')
    workbook = Workbook()
    workbook.active.title = '点検'
    workbook.active.append(['安全確認', 'EXCEL evidence'])
    workbook.save(folder / '一覧.xlsx')
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = 'POWERPOINT evidence'
    slide.placeholders[1].text = 'Document traceability'
    deck.save(folder / '説明.pptx')
    canvas = Canvas(str(folder / '混在.pdf'))
    canvas.drawString(50, 760, 'PDF text page safety inspection evidence')
    canvas.showPage()
    canvas.drawImage(str(folder / '画像.png'), 25, 600, width=550, height=88)
    canvas.showPage()
    canvas.save()
    with zipfile.ZipFile(folder / '資料.zip', 'w') as archive:
        archive.writestr('内包.txt', 'Archive text evidence')


def request(base, route, *, method='GET', data=None, token=None):
    headers = {'Origin': base}
    if token:
        headers['X-Contextgen-Token'] = token
    raw = json.dumps(data).encode('utf-8') if data is not None else None
    if data is not None:
        headers['Content-Type'] = 'application/json'
    req = urllib.request.Request(base + route, data=raw, method=method, headers=headers)
    with urllib.request.urlopen(req, timeout=30) as response:
        body = response.read()
        return json.loads(body) if 'json' in response.headers.get('Content-Type', '') else body


def wait_for(callback, *, seconds=180, label='operation'):
    deadline = time.monotonic() + seconds
    while time.monotonic() < deadline:
        result = callback()
        if result:
            return result
        time.sleep(0.5)
    raise AssertionError(f'Timeout waiting for {label}')


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--exe', type=Path, required=True)
    exe = parser.parse_args().exe.resolve()
    evidence = {'version': __version__, 'checks': [], 'user_machine_acceptance': 'pending'}
    output = Path('build')
    output.mkdir(exist_ok=True)
    def checked(name):
        evidence['checks'].append(name)
        print(name, flush=True)
    process, task_name, scheduler, schedule_id, root, temporary_root = None, None, None, None, None, None
    try:
        assert os.name == 'nt', 'Windows only'
        assert exe.is_file()
        temporary_root = tempfile.TemporaryDirectory(prefix='contextgen-kai-smoke-', ignore_cleanup_errors=True)
        with nullcontext(temporary_root.name) as temp:
            root = Path(temp)
            # 配布フォルダ自体を日本語・空白入りに移し、OCRの言語データ探索も検証する。
            relocated = root / '配布 アプリ'
            shutil.copytree(exe.parent, relocated)
            exe = relocated / exe.name
            verify_distribution_layout(exe.parent)
            checked('distribution has exactly four user-facing root items and no developer evidence')
            state = root / '状態 日本語'
            source = root / '資料 日本語'
            fixtures(source)
            env = dict(os.environ, PATH=os.environ['SystemRoot'] + '\\System32;' + os.environ['SystemRoot'])
            env.pop('PYTHONPATH', None)
            env.pop('CONTEXTGEN_TESSERACT', None)
            env.pop('TESSDATA_PREFIX', None)
            ocr = exe.parent / '_internal' / 'ocr' / 'tesseract.exe'
            result = subprocess.run([str(ocr), str(source / '画像.png'), 'stdout', '-l', 'eng+jpn', '--tessdata-dir', str(ocr.parent / 'tessdata')],
                                    env=env, capture_output=True, timeout=60)
            text = result.stdout.decode('utf-8', errors='replace')
            assert result.returncode == 0, result.stderr.decode(errors='replace')
            assert 'SAFETY' in text and ('安全' in text or '確認' in text), text
            checked('bundled English/Japanese OCR under system-only PATH')
            with socket.socket() as sock:
                sock.bind(('127.0.0.1', 0))
                port = sock.getsockname()[1]
            base = f'http://127.0.0.1:{port}'
            with (output / 'windows-smoke.log').open('w', encoding='utf-8') as log:
                process = subprocess.Popen([str(exe), '--no-browser', '--state-dir', str(state), '--port', str(port)],
                                           cwd=root, env=env, stdout=log, stderr=log)
                def ready():
                    if process.poll() is not None:
                        raise AssertionError(f'EXE exited: {process.returncode}')
                    try:
                        return request(base, '/api/status')
                    except (urllib.error.URLError, TimeoutError):
                        return None
                status = wait_for(ready, seconds=60, label='packaged API')
                token = status['token']
                assert status['version'] == __version__
                assert b'contextgen' in request(base, '/').lower()
                checked('standalone EXE serves local UI/API without developer Python PATH')
                # 配布HTMLは直下、画像とJSは_internal。実HTTPの内容が同じ配布実体と一致すること。
                verify_manual_assets(exe.parent, lambda route: request(base, route), packaged=True)
                for path in ('README.md', '仕様書兼要件定義書.md', '_internal/manual/images/capture-record.json',
                             '_internal/ocr/tesseract.exe', '_internal/THIRD_PARTY_LICENSES/tesseract-LICENSE.txt'):
                    try:
                        request(base, '/manual/' + quote(path))
                    except urllib.error.HTTPError as error:
                        assert error.code == 404, (path, error.code)
                    else:
                        raise AssertionError(f'Unexpected distribution file exposed by manual API: {path}')
                checked('packaged illustrated manual serves internal images and blocks non-manual files')
                library = request(base, '/api/libraries', method='POST', data={'name': 'Windows 受入資料', 'path': str(source)}, token=token)
                job = request(base, '/api/jobs', method='POST', data={'library_id': library['id'], 'export_after': False}, token=token)
                def job_done(job_id):
                    rows = request(base, '/api/jobs')
                    row = next(j for j in rows if j['id'] == job_id)
                    return row if row['state'] in ('completed', 'held', 'failed', 'stopped') else None
                finished = wait_for(lambda: job_done(job['id']), label='Office/PDF/image scan')
                assert finished['state'] == 'completed', finished
                assert finished['errors'] == 0, finished
                docs = request(base, '/api/documents?limit=50')['items']
                assert len(docs) == 7, docs
                for doc in docs:
                    detail = request(base, '/api/documents/' + doc['id'])
                    assert detail['status'] == 'ok', detail
                    assert detail['effective_text'].strip(), detail
                    assert detail['units'], detail
                    if doc['relative_path'] == '混在.pdf':
                        assert 'SAFETY' in detail['effective_text'] and 'text page' in detail['effective_text'], detail
                checked('packaged Office/text/mixed-PDF/image/ZIP extraction and provenance')
                collection = request(base, '/api/collections', method='POST', data={'name': '受入セット', 'library_id': library['id'], 'purpose': 'overview'}, token=token)
                exported = request(base, '/api/exports', method='POST', data={'collection_id': collection['id']}, token=token)
                completed = wait_for(lambda: job_done(exported['id']), label='export')
                assert completed['state'] == 'completed', completed
                exports = request(base, '/api/exports')
                generation = next(e for e in exports if e['id'] == completed['export_id'])
                assert generation['state'] == 'published' and generation['is_active'], generation
                archive = request(base, '/api/exports/' + generation['id'] + '/download')
                with zipfile.ZipFile(io.BytesIO(archive)) as zip_out:
                    assert any(p.lower().endswith('.txt') for p in zip_out.namelist())
                    assert any(p.lower().endswith('.jsonl') for p in zip_out.namelist())
                checked('packaged export automatically published and ZIP downloaded')
                # 最初は翌日の予約にし、起動中 tick と競合しないようにする。
                schedule = request(base, '/api/schedules', method='POST', token=token, data={
                    'name': 'CI 一時予約', 'library_id': library['id'], 'collection_id': collection['id'],
                    'enabled': True, 'mode': 'background', 'frequency': 'daily', 'time': datetime.now().strftime('%H:%M')})
                schedule_id = schedule['id']
                from contextgen_kai.scheduling import WindowsTasks
                scheduler = WindowsTasks(state)
                task_name = scheduler.name(schedule_id)
                query = subprocess.run(['schtasks.exe', '/Query', '/TN', task_name, '/XML'], capture_output=True)
                raw_xml = query.stdout
                encoding = 'utf-16' if raw_xml.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-16-le' if b'\x00' in raw_xml[:200] else 'utf-8'
                xml = raw_xml.decode(encoding, errors='replace')
                assert query.returncode == 0, query.stderr.decode(errors='replace')
                task_xml = ET.fromstring(xml)
                ns = {'t': 'http://schemas.microsoft.com/windows/2004/02/mit/task'}
                assert task_xml.find('.//t:LogonType', ns).text == 'InteractiveToken'
                wake = task_xml.find('.//t:WakeToRun', ns)
                assert wake is None or wake.text == 'false'
                checked('real Windows InteractiveToken limited-privilege task registration')
                request(base, '/api/shutdown', method='POST', data={}, token=token)
                process.wait(timeout=30)
                # 実行漏れの状態を作り、アプリ終了後に実タスクから処理させる。
                db_path = state / 'schedules.sqlite3'
                with sqlite3.connect(db_path) as conn:
                    data = json.loads(conn.execute('SELECT data FROM schedules WHERE id=?', (schedule_id,)).fetchone()[0])
                    data['next_run'] = (datetime.now(timezone.utc) - timedelta(minutes=2)).isoformat()
                    conn.execute('UPDATE schedules SET data=? WHERE id=?', (json.dumps(data), schedule_id))
                run = subprocess.run(['schtasks.exe', '/Run', '/TN', task_name], capture_output=True, text=True, errors='replace')
                assert run.returncode == 0, run.stderr or run.stdout
                def background_finished():
                    with sqlite3.connect(db_path) as conn:
                        row = conn.execute('SELECT data FROM schedules WHERE id=?', (schedule_id,)).fetchone()
                    data = json.loads(row[0])
                    return data if data.get('last_result') in ('completed', 'held', 'failed', 'stopped') else None
                result = wait_for(background_finished, seconds=180, label='real headless task after EXE shutdown')
                assert result['last_result'] == 'completed', result
                checked('background task scans/exports with app closed and reports completion')
                scheduler.delete(schedule_id)
                task_name = None
                checked('temporary Windows task unregistered')
    finally:
        if process and process.poll() is None:
            process.terminate()
            process.wait(timeout=15)
        if task_name:
            diagnosis = subprocess.run(['schtasks.exe', '/Query', '/TN', task_name, '/V', '/FO', 'CSV'], capture_output=True)
            (output / 'windows-task-details.txt').write_bytes(diagnosis.stdout + diagnosis.stderr)
            subprocess.run(['schtasks.exe', '/End', '/TN', task_name], capture_output=True)
            subprocess.run(['schtasks.exe', '/Delete', '/TN', task_name, '/F'], capture_output=True)
        if root:
            app_log = root / '状態 日本語' / 'app.log'
            if app_log.is_file():
                shutil.copy2(app_log, output / 'windows-app.log')
        if temporary_root:
            temporary_root.cleanup()
        (output / 'windows-smoke.json').write_text(json.dumps(evidence, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(evidence, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
