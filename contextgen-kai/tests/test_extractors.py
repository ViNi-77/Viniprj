"""生成した資料で出典・欠落・復旧可能な失敗を確かめる。"""
import io
import shutil
import stat
import subprocess
from pathlib import Path
import zipfile

from PIL import Image, ImageDraw, ImageFont
import pytest

from contextgen_kai import extractors as ex


def add_zip_entries(path, entries):
    with zipfile.ZipFile(path, "a", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, data in entries.items():
            archive.writestr(name, data)


def make_text_pdf(path):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    page = writer.add_blank_page(width=600, height=800)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 18 Tf 50 700 Td (Production report ALPHA) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(content)
    writer.write(path)


def make_ocr_image(path):
    image = Image.new("RGB", (1300, 400), "white")
    draw = ImageDraw.Draw(image)
    candidates = ["/System/Library/Fonts/Supplemental/Arial.ttf", "C:/Windows/Fonts/arial.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"]
    font = next((ImageFont.truetype(p, 64) for p in candidates if Path(p).is_file()), ImageFont.load_default(size=64))
    draw.text((45, 140), "FACTORY REPORT 2026", fill="black", font=font)
    image.save(path)
    return image


def test_text_japanese_encoding_and_non_destructive(tmp_path):
    path = tmp_path / "日本語 の資料.csv"
    original = "品名,台数\n設備A,24\n".encode("cp932")
    path.write_bytes(original)
    result = ex.extract_document(path)
    assert result.status == "ok"
    assert "設備A,24" in result.text
    assert result.metadata["encoding"] == "cp932"
    assert result.units[0]["locator"] == "本文"
    assert path.read_bytes() == original
    assert result.as_dict()["warnings"] == []


def test_word_tables_headers_and_embedded_workbook(tmp_path):
    import docx
    import openpyxl
    workbook_path = tmp_path / "embedded.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "設備"
    workbook.active.append(["稼働時間", 72])
    workbook.save(workbook_path)
    path = tmp_path / "手順.docx"
    document = docx.Document()
    document.add_paragraph("保全の手順")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "対象"
    table.cell(0, 1).text = "設備A"
    document.sections[0].header.paragraphs[0].text = "検査記録"
    document.save(path)
    add_zip_entries(path, {"word/embeddings/設備.xlsx": workbook_path.read_bytes(), "word/embeddings/old.bin": ex.OLE_MAGIC})
    result = ex.extract_document(path)
    assert "保全の手順" in result.text and "対象\t設備A" in result.text
    assert "検査記録" in result.text
    assert "稼働時間" in result.text and "B1: 72" in result.text
    assert any("word/embeddings/設備.xlsx" in unit["locator"] for unit in result.units)
    assert any("旧OLE" in warning for warning in result.warnings)
    assert result.status == "partial"
    assert result.metadata["embedded_documents"] == 2


def test_excel_sheet_cells_and_formula_explanation(tmp_path):
    import openpyxl
    path = tmp_path / "report.xlsm"
    workbook = openpyxl.Workbook()
    workbook.active.title = "工程"
    workbook.active.append(["設備B", 9, "=B1*2"])
    workbook.create_sheet("次工程").append(["検査", 4])
    workbook.save(path)
    result = ex.extract_document(path)
    assert "シート「工程」 行 1" in result.text
    assert "A1: 設備B" in result.text and "C1: =B1*2" in result.text
    assert "次工程" in result.text
    assert any("数式" in warning for warning in result.warnings)
    assert result.status == "partial"


def test_powerpoint_table_and_notes(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "安全確認"
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(6), Inches(1)).table
    table.cell(0, 0).text = "箇所"
    table.cell(0, 1).text = "ラインA"
    slide.notes_slide.notes_text_frame.text = "講師の説明"
    path = tmp_path / "手順.pptx"
    deck.save(path)
    result = ex.extract_document(path)
    assert result.status == "ok"
    assert "安全確認" in result.text and "ラインA" in result.text
    assert "講師の説明" in result.text
    assert any(unit["locator"] == "スライド 1 ノート" for unit in result.units)


def test_pdf_page_text_and_limits(tmp_path):
    from pypdf import PdfReader, PdfWriter
    path = tmp_path / "text.pdf"
    make_text_pdf(path)
    writer = PdfWriter()
    page = PdfReader(path).pages[0]
    writer.add_page(page)
    writer.add_page(page)
    writer.write(tmp_path / "two.pdf")
    result = ex.extract_document(tmp_path / "two.pdf", max_pages=1)
    assert "Production report ALPHA" in result.text
    assert result.units[0]["locator"] == "ページ 1"
    assert result.status == "partial"
    assert any("2 ページ中 1" in warning for warning in result.warnings)


def test_protected_and_corrupt_documents(tmp_path):
    from pypdf import PdfWriter
    protected = tmp_path / "password.pdf"
    writer = PdfWriter()
    writer.add_blank_page(100, 100)
    writer.encrypt("secret")
    writer.write(protected)
    assert ex.extract_document(protected).status == "protected"
    office = tmp_path / "secret.docx"
    office.write_bytes(ex.OLE_MAGIC + b"encrypted")
    assert ex.extract_document(office).status == "protected"
    broken = tmp_path / "broken.pdf"
    broken.write_bytes(b"broken file")
    result = ex.extract_document(broken)
    assert result.status == "error" and result.warnings
    assert ex.extract_document(tmp_path / "missing.txt").status == "error"


def test_missing_ocr_and_office_image_visible(tmp_path, monkeypatch):
    import docx
    monkeypatch.setattr(ex, "tesseract_path", lambda: None)
    path = tmp_path / "picture.png"
    make_ocr_image(path)
    result = ex.extract_document(path)
    assert result.status == "needs_ocr"
    assert "Tesseract" in result.warnings[0]
    document = docx.Document()
    document.add_paragraph("画像のある資料")
    document.add_picture(str(path))
    office = tmp_path / "images.docx"
    document.save(office)
    result = ex.extract_document(office)
    assert result.status == "partial"
    assert result.metadata["embedded_images"] == 1
    assert any("word/media/" in warning for warning in result.warnings)


def test_ocr_timeout_reports_unread_part(tmp_path, monkeypatch):
    path = tmp_path / "picture.png"
    make_ocr_image(path)
    monkeypatch.setattr(ex, "tesseract_path", lambda: "/fake/tesseract")
    def timeout(*args, **kwargs):
        assert kwargs["timeout"] <= 30
        raise subprocess.TimeoutExpired(args[0], kwargs["timeout"])
    monkeypatch.setattr(ex.subprocess, "run", timeout)
    result = ex.extract_document(path)
    assert result.status == "needs_ocr"
    assert any("制限時間" in warning for warning in result.warnings)


def test_zip_traversal_symlinks_and_valid_member(tmp_path):
    path = tmp_path / "files.zip"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("正常/資料.txt", "普通に読める")
        archive.writestr("../../outside.txt", "outside")
        archive.writestr("C:\\outside.txt", "outside")
        link = zipfile.ZipInfo("link.txt")
        link.create_system = 3
        link.external_attr = (stat.S_IFLNK | 0o777) << 16
        archive.writestr(link, "../../target")
        archive.writestr("old.doc", "old doc")
    result = ex.extract_document(path)
    assert result.status == "partial"
    assert "普通に読める" in result.text
    assert "outside" not in result.text
    assert any("安全でない" in warning for warning in result.warnings)
    assert any("シンボリック" in warning for warning in result.warnings)
    assert any("新形式" in warning for warning in result.warnings)
    assert not (tmp_path.parent / "outside.txt").exists()


def test_zip_bomb_and_nesting_limits_are_reported(tmp_path, monkeypatch):
    monkeypatch.setattr(ex, "MAX_COMPRESSION_RATIO", 10)
    path = tmp_path / "large.zip"
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("large.txt", "A" * 100_000)
        archive.writestr("good.txt", "Small valid text")
    result = ex.extract_document(path)
    assert "Small valid text" in result.text
    assert any("展開率" in warning for warning in result.warnings)
    monkeypatch.setattr(ex, "MAX_COMPRESSION_RATIO", 2000)
    raw = b""
    for level in range(5):
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("inner.zip" if level else "file.txt", raw if level else "bottom")
        raw = buffer.getvalue()
    nested = tmp_path / "nested.zip"
    nested.write_bytes(raw)
    result = ex.extract_document(nested)
    assert any("入れ子" in warning for warning in result.warnings)
    assert not result.text


def test_text_budget_does_not_silently_truncate(tmp_path, monkeypatch):
    monkeypatch.setattr(ex, "MAX_TEXT_CHARS", 10)
    path = tmp_path / "text.txt"
    path.write_text("ABCDEF" * 10)
    result = ex.extract_document(path)
    assert result.status == "partial"
    assert result.units[0]["text"] == "ABCDEFABCD"
    assert any("上限 10 字" in warning for warning in result.warnings)


def test_deadline_and_validation(tmp_path):
    path = tmp_path / "source.txt"
    path.write_text("data")
    result = ex.extract_document(path, timeout=0.00000001)
    assert result.status == "error"
    assert any("処理時間" in warning for warning in result.warnings)
    with pytest.raises(ValueError):
        ex.extract_document(path, max_pages=0)


def test_legacy_formats_visible_to_scanner(tmp_path):
    for suffix in [".doc", ".xls", ".ppt"]:
        assert suffix in ex.SUPPORTED_EXTENSIONS
        path = tmp_path / ("legacy" + suffix)
        path.write_bytes(b"content")
        assert ex.extract_document(path).status == "needs_conversion"


def test_mixed_pdf_ocr_only_image_page(tmp_path, monkeypatch):
    from pypdf import PdfReader, PdfWriter
    text_pdf = tmp_path / "text.pdf"
    image_pdf = tmp_path / "scan.pdf"
    make_text_pdf(text_pdf)
    make_ocr_image(tmp_path / "scan.png").save(image_pdf)
    writer = PdfWriter()
    writer.add_page(PdfReader(text_pdf).pages[0])
    writer.add_page(PdfReader(image_pdf).pages[0])
    mixed = tmp_path / "mixed.pdf"
    writer.write(mixed)
    monkeypatch.setattr(ex, "tesseract_path", lambda: "/fake/tesseract")
    calls = []
    def ocr(image, result, budget, locator):
        assert image.width > 0 and image.height > 0
        calls.append(locator)
        budget.add(result, locator, "FACTORY REPORT 2026", "ocr")
    monkeypatch.setattr(ex, "_ocr_image", ocr)
    result = ex.extract_document(mixed)
    assert "Production report ALPHA" in result.text
    assert "FACTORY REPORT 2026" in result.text
    assert calls == ["ページ 2"]
    assert result.status == "ok"


@pytest.mark.skipif(not ex.tesseract_path(), reason="Tesseract binary unavailable in this environment")
def test_real_tesseract_image_and_mixed_pdf(tmp_path, monkeypatch):
    """OCR同梱Windows CIと、開発環境にTesseractがあるとき実際の文字を検証。"""
    from pypdf import PdfReader, PdfWriter
    monkeypatch.setenv("CONTEXTGEN_OCR_LANG", "eng")
    image_path = tmp_path / "actual.png"
    image = make_ocr_image(image_path)
    result = ex.extract_document(image_path)
    assert "FACTORY REPORT 2026" in result.text, result.as_dict()
    assert result.metadata["ocr_engine"] == "Tesseract"
    image_pdf = tmp_path / "scan.pdf"
    image.save(image_pdf)
    text_pdf = tmp_path / "text.pdf"
    make_text_pdf(text_pdf)
    writer = PdfWriter()
    writer.add_page(PdfReader(text_pdf).pages[0])
    writer.add_page(PdfReader(image_pdf).pages[0])
    mixed = tmp_path / "mixed.pdf"
    writer.write(mixed)
    result = ex.extract_document(mixed)
    assert "Production report ALPHA" in result.text
    assert "FACTORY REPORT 2026" in result.text
    assert result.metadata["ocr_units"] == 1
    assert result.status == "ok"


@pytest.mark.skipif(not ex.tesseract_path(), reason="Tesseract binary unavailable in this environment")
def test_real_japanese_ocr_image_and_office_media(tmp_path, monkeypatch):
    import docx
    import unicodedata
    candidates = [Path("C:/Windows/Fonts/msgothic.ttc"), Path("C:/Windows/Fonts/YuGothR.ttc")]
    candidates += [p for p in Path("/System/Library/Fonts").glob("*W3.ttc") if "ヒラギノ" in unicodedata.normalize("NFC", p.name)]
    font_path = next((p for p in candidates if p.is_file()), None)
    if font_path is None:
        pytest.skip("Japanese rendering font unavailable")
    languages = subprocess.run([ex.tesseract_path(), "--list-langs"], capture_output=True, text=True, timeout=10)
    if "jpn" not in languages.stdout:
        pytest.skip("Japanese OCR language data unavailable")
    monkeypatch.setenv("CONTEXTGEN_OCR_LANG", "jpn+eng")
    image = Image.new("RGB", (1600, 400), "white")
    ImageDraw.Draw(image).text((60, 140), "安全確認 生産設備 点検記録", fill="black", font=ImageFont.truetype(str(font_path), 72))
    path = tmp_path / "日本語 画像.png"
    image.save(path)
    result = ex.extract_document(path)
    compact = "".join(result.text.split())
    assert "生産設備" in compact and "点検記録" in compact, result.as_dict()
    document = docx.Document()
    document.add_paragraph("設備資料の画像")
    document.add_picture(str(path))
    office = tmp_path / "日本語埋め込み.docx"
    document.save(office)
    result = ex.extract_document(office)
    assert result.status == "ok", result.as_dict()
    assert "点検記録" in "".join(result.text.split())
    assert any("word/media/" in unit["locator"] and unit["status"] == "ocr" for unit in result.units)
