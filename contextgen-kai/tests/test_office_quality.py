"""Officeの原文を残し、整理に必要な構造・全出典だけを追加する回帰試験。"""
import hashlib
from pathlib import Path
from xml.etree import ElementTree as ET
import zipfile

from PIL import Image
import openpyxl
from openpyxl.worksheet.table import Table
from pptx import Presentation
from pptx.util import Inches

from contextgen_kai import extractors as ex


def rewrite_zip(path, updates):
    with zipfile.ZipFile(path) as source:
        entries = {name: source.read(name) for name in source.namelist()}
    entries.update(updates)
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as target:
        for name, data in entries.items():
            target.writestr(name, data)


def mock_ocr(monkeypatch, text="原文 100 / 条件A"):
    calls = []
    monkeypatch.setattr(ex, "tesseract_path", lambda: "/fake/tesseract")

    def read(image, result, budget, locator):
        calls.append(locator)
        result.metadata["ocr_units"] = 1
        budget.add(result, locator, text, "ocr")

    monkeypatch.setattr(ex, "_ocr_image", read)
    return calls


def test_excel_real_table_no_header_guess_no_identical_row_deletion(tmp_path):
    path = tmp_path / "案件別.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "測定結果"
    sheet.append(["単独の前置き", 0, False])
    sheet.append(["設備", "測定値", "条件"])
    sheet.append(["A", 100, "通常"])
    sheet.append(["A", 100, "通常"])
    sheet.add_table(Table(displayName="Measurements", ref="A2:C4"))
    sheet["A8"] = "別の表の見出しかどうかは未確定"
    sheet["B8"] = 9
    book.save(path)
    before = path.read_bytes()
    result = ex.extract_document(path)
    assert result.status == "ok", result.as_dict()
    rows = [u for u in result.units if u["table_id"]]
    assert len(rows) == 3 and rows[0]["table_header_row"]
    assert rows[1]["values"] == rows[2]["values"] == ["A", "100", "通常"]
    assert rows[1]["table_headers"] == ["設備", "測定値", "条件"]
    assert rows[1]["cells"] == ["A3", "B3", "C3"]
    assert rows[1]["clean_text"] == "A\t100\t通常"
    assert "A3: A\tB3: 100" in rows[1]["text"]
    others = [u for u in result.units if not u["table_id"]]
    assert all("clean_text" not in u and not u["table_headers"] for u in others)
    assert "B1: 0" in result.text and "C1: False" in result.text
    assert path.read_bytes() == before


def test_excel_hidden_columns_rows_and_sheets_remain_in_raw(tmp_path):
    path = tmp_path / "非表示含む.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["表示列", "非表示列"])
    sheet.append(["条件", 42])
    sheet.column_dimensions["B"].hidden = True
    sheet.row_dimensions[2].hidden = True
    sheet.add_table(Table(displayName="Conditions", ref="A1:B2"))
    hidden = book.create_sheet("補足")
    hidden.sheet_state = "veryHidden"
    hidden["A1"] = "隠れていても重要な上限値 900"
    book.save(path)
    result = ex.extract_document(path)
    by_cell = {cell: u for u in result.units if u["parent_locator"] == "シート「Sheet」" for cell in u["cells"]}
    assert by_cell["A1"]["hidden"] is False
    assert by_cell["B1"]["hidden"] and by_cell["B1"]["hidden_columns"] == ["B1"]
    assert by_cell["A2"]["hidden_row"] and by_cell["A2"]["hidden"]
    assert by_cell["B1"]["table_headers"] == ["非表示列"]
    supplement = next(u for u in result.units if "900" in u["text"])
    assert supplement["hidden_sheet"] and supplement["hidden"]
    assert "上限値 900" in result.text and "B2: 42" in result.text


def test_excel_cached_formula_is_explicitly_unrecalculated(tmp_path):
    path = tmp_path / "計算.xlsx"
    book = openpyxl.Workbook()
    book.active.append([5, "=A1*2", "=A1+7", "=IF(A1=5,\"\",\"x\")"])
    book.save(path)
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
    root.find(".//s:c[@r='B1']/s:v", ns).text = "10"
    rewrite_zip(path, {"xl/worksheets/sheet1.xml": ET.tostring(root)})
    result = ex.extract_document(path)
    formulas = result.units[0]["formulas"]
    assert formulas["B1"] == {"formula": "=A1*2", "cached_value": "10", "cache_present": True, "recalculated": False}
    assert formulas["C1"]["cached_value"] is None and not formulas["C1"]["cache_present"]
    assert "B1: =A1*2 [保存済み値: 10 / 未再計算]" in result.text
    assert "C1: =A1+7 [保存済み値なし / 未再計算]" in result.text
    assert result.status == "partial" and any("未再計算" in w for w in result.warnings)


def test_excel_sparse_format_only_range_does_not_create_empty_rows(tmp_path):
    from openpyxl.styles import PatternFill
    path = tmp_path / "疎な表.xlsx"
    book = openpyxl.Workbook()
    book.active["A1"] = "実際の内容"
    book.active["B99999"] = "離れた内容"
    book.active["XFD1048576"].fill = PatternFill("solid", fgColor="FF0000")
    book.save(path)
    result = ex.extract_document(path)
    assert result.status == "ok", result.as_dict()
    assert len(result.units) == 2
    assert "B99999: 離れた内容" in result.text
    assert result.metadata["sheet_scan"] == [{"sheet": "Sheet", "stored_rows": 3, "stored_cells": 3}]


def test_excel_shared_strings_dates_and_shared_formula(tmp_path):
    from datetime import datetime
    path = tmp_path / "表現.xlsx"
    book = openpyxl.Workbook()
    book.active.append([datetime(2026, 9, 24), "=C1*2", 3])
    book.active.append([None, "=C2*2", 4])
    book.save(path)
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
    for coord in ["B1", "B2"]:
        formula = root.find(f".//{{{ns}}}c[@r='{coord}']/{{{ns}}}f")
        formula.set("t", "shared")
        formula.set("si", "0")
        if coord == "B2":
            formula.text = None
    c = root.find(f".//{{{ns}}}c[@r='C1']")
    c.set("t", "s")
    c.find(f"{{{ns}}}v").text = "0"
    rewrite_zip(path, {"xl/worksheets/sheet1.xml": ET.tostring(root),
                       "xl/sharedStrings.xml": f'<sst xmlns="{ns}"><si><t>共通文字列</t></si></sst>'})
    result = ex.extract_document(path)
    assert "2026-09-24" in result.text and "共通文字列" in result.text
    assert "B2: =C2*2" in result.text


def test_powerpoint_structural_metadata_hidden_notes_and_tables(tmp_path):
    from lxml import etree
    path = tmp_path / "説明資料.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "安全規則"
    slide.element.set("show", "0")
    for role, text in [("ftr", "文書版 Rev.3"), ("dt", "2026/09/24"), ("sldNum", "1")]:
        box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(1))
        box.text = text
        ph = etree.SubElement(box.element.nvSpPr.nvPr, "{http://schemas.openxmlformats.org/presentationml/2006/main}ph")
        ph.set("type", role)
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(5), Inches(1)).table
    table.cell(0, 0).text = "測定値"
    table.cell(0, 1).text = "100"
    table.cell(1, 0).text = "測定値"
    table.cell(1, 1).text = "101"
    slide.notes_slide.notes_text_frame.text = "安全規則の補足。ノートを消さない。"
    deck.save(path)
    result = ex.extract_document(path)
    assert all(u["hidden"] for u in result.units)
    metadata = {u["role"]: u for u in result.units if u["kind"] == "metadata"}
    assert set(metadata) == {"footer", "date", "slide_number"}
    assert metadata["footer"]["text"] == "文書版 Rev.3"
    assert "文書版 Rev.3" in result.text  # 整理段階まで原文は残す。
    assert len([u for u in result.units if u["kind"] == "note"]) == 1
    assert result.text.count("測定値") == 2  # 表の本文二重抽出をしない。
    assert "100" in result.text and "101" in result.text


def test_same_image_bytes_ocr_once_with_every_parent_reference(tmp_path, monkeypatch):
    calls = mock_ocr(monkeypatch)
    path = tmp_path / "画像と本文.pptx"
    image = tmp_path / "same.png"
    Image.new("RGB", (30, 30), "white").save(image)
    deck = Presentation()
    for number in range(2):
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = "原文 100 / 条件A" if number == 0 else "原文 101 / 条件A"
        slide.shapes.add_picture(str(image), Inches(1), Inches(1))
    deck.save(path)
    # 別名で同じ画像が格納されていてもOCR計算のみ共有する。未参照部品も原文保持。
    rewrite_zip(path, {"ppt/media/unused-copy.png": image.read_bytes()})
    result = ex.extract_document(path)
    assert len(calls) == 1
    images = [u for u in result.units if u["kind"] == "image"]
    assert len(images) == 2 and all(u["image_hash"] == hashlib.sha256(image.read_bytes()).hexdigest() for u in images)
    referenced = next(u for u in images if "parent_locators" in u)
    assert referenced["parent_locators"] == ["スライド 1", "スライド 2"]
    assert any("スライド 1" in s for s in referenced["source_refs"])
    assert any("スライド 2" in s for s in referenced["source_refs"])
    unused = next(u for u in images if "unused-copy" in u["locator"])
    assert unused["candidate_reason"] == "unreferenced_asset"
    assert result.text.count("原文 100 / 条件A") == 3  # 本文とOCRを勝手に重複削除しない。
    assert next(u for u in result.units if u["kind"] == "body" and u["text"] == "原文 100 / 条件A")["candidate_reason"] == "native_and_image_text"
    assert "原文 101 / 条件A" in result.text
    assert result.metadata["ocr_units"] == 1 and result.metadata["ocr_cache_hits"] == 1


def test_embedded_excel_keeps_extra_values_and_parent_source(tmp_path):
    workbook = tmp_path / "inner.xlsx"
    book = openpyxl.Workbook()
    book.active.append(["親にもある", 100])
    book.create_sheet("補足").append(["埋め込みだけの条件", 900])
    book.save(workbook)
    path = tmp_path / "埋め込み.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "親にもある 100"
    deck.save(path)
    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    p_ns = "http://schemas.openxmlformats.org/presentationml/2006/main"
    r_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read("ppt/slides/slide1.xml"))
        rels = ET.fromstring(archive.read("ppt/slides/_rels/slide1.xml.rels"))
    ET.SubElement(root, f"{{{p_ns}}}oleObj", {f"{{{r_ns}}}id": "rIdEmbedded"})
    ET.SubElement(rels, f"{{{rel_ns}}}Relationship", {"Id": "rIdEmbedded", "Target": "../embeddings/inner.xlsx", "Type": f"{r_ns}/package"})
    rewrite_zip(path, {"ppt/slides/slide1.xml": ET.tostring(root),
                       "ppt/slides/_rels/slide1.xml.rels": ET.tostring(rels),
                       "ppt/embeddings/inner.xlsx": workbook.read_bytes()})
    result = ex.extract_document(path)
    assert "埋め込みだけの条件" in result.text and "900" in result.text
    embedded = [u for u in result.units if u["kind"] == "embedded"]
    assert len(embedded) == 2 and all(u["content_kind"] == "table_row" for u in embedded)
    assert all(any("スライド 1 / 埋め込み ppt/embeddings/inner.xlsx" in s for s in u["source_refs"]) for u in embedded)


def test_excel_image_has_sheet_source(tmp_path, monkeypatch):
    from openpyxl.drawing.image import Image as ExcelImage
    calls = mock_ocr(monkeypatch)
    image = tmp_path / "data.png"
    Image.new("RGB", (30, 30), "white").save(image)
    path = tmp_path / "図入り.xlsx"
    book = openpyxl.Workbook()
    book.active.title = "点検対象"
    book.active.add_image(ExcelImage(image), "C5")
    book.save(path)
    result = ex.extract_document(path)
    assert len(calls) == 1
    unit = next(u for u in result.units if u["kind"] == "image")
    assert any('シート「点検対象」' in s for s in unit["source_refs"])


def test_same_pdf_page_keeps_native_text_and_adds_image_ocr(tmp_path, monkeypatch):
    from pypdf import PdfReader, PdfWriter
    from test_extractors import make_text_pdf
    native, scan = tmp_path / "native.pdf", tmp_path / "scan.pdf"
    make_text_pdf(native)
    Image.new("RGB", (60, 60), "white").save(scan)
    writer = PdfWriter()
    writer.add_page(PdfReader(native).pages[0])
    page = writer.pages[0]
    page.merge_page(PdfReader(scan).pages[0])
    combined = tmp_path / "mixed_on_one_page.pdf"
    writer.write(combined)
    calls = mock_ocr(monkeypatch, "Production report ALPHA 101")
    result = ex.extract_document(combined)
    assert result.status == "ok", result.as_dict()
    assert len(calls) == 1 and calls[0].startswith("ページ 1 画像")
    assert result.units[0]["text"] == "Production report ALPHA"
    image_unit = next(u for u in result.units if u["kind"] == "image")
    assert image_unit["text"] == "Production report ALPHA 101"
    assert image_unit["parent_locator"] == "ページ 1"
    assert len(result.units) == 2


def test_common_metadata_and_unit_ids_are_stable(tmp_path):
    path = tmp_path / "test.txt"
    path.write_text("initial", encoding="utf-8")
    first = ex.extract_document(path)
    path.write_text("corrected", encoding="utf-8")
    second = ex.extract_document(path)
    assert first.units[0]["unit_id"] == second.units[0]["unit_id"]
    assert first.units[0]["source_refs"] == ["本文"]
    assert first.units[0]["kind"] == "body" and first.units[0]["hidden"] is False
    assert first.metadata["extractor_version"] == ex.EXTRACTOR_VERSION == "2"


def test_office_text_limit_cannot_be_bypassed_through_clean_text(tmp_path, monkeypatch):
    path = tmp_path / "limit.xlsx"
    book = openpyxl.Workbook()
    book.active.append(["header"])
    book.active.append(["content " * 100])
    book.active.add_table(Table(displayName="Limited", ref="A1:A2"))
    book.save(path)
    monkeypatch.setattr(ex, "MAX_TEXT_CHARS", 20)
    result = ex.extract_document(path)
    assert result.status == "partial"
    partial = next(u for u in result.units if u["status"] == "partial")
    assert not any(key in partial for key in ("values", "clean_text", "formulas", "table_headers"))


def test_hidden_powerpoint_shape_table_and_picture(tmp_path, monkeypatch):
    mock_ocr(monkeypatch)
    path, image = tmp_path / "shape.pptx", tmp_path / "image.png"
    Image.new("RGB", (20, 20), "white").save(image)
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "表示タイトル"
    group = slide.shapes.add_group_shape()
    box = group.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    box.text = "非表示グループの本文"
    group.element.nvGrpSpPr.cNvPr.set("hidden", "1")
    table = slide.shapes.add_table(1, 1, Inches(1), Inches(2), Inches(2), Inches(1))
    table.table.cell(0, 0).text = "非表示表"
    table.element.nvGraphicFramePr.cNvPr.set("hidden", "1")
    picture = slide.shapes.add_picture(str(image), Inches(1), Inches(3))
    picture.element.nvPicPr.cNvPr.set("hidden", "1")
    deck.save(path)
    result = ex.extract_document(path)
    assert next(u for u in result.units if u["text"] == "表示タイトル")["hidden"] is False
    assert all(u["hidden"] for u in result.units if u["text"] != "表示タイトル")


def test_unused_layout_image_does_not_gain_false_slide_reference(tmp_path, monkeypatch):
    mock_ocr(monkeypatch)
    path, image = tmp_path / "layouts.pptx", tmp_path / "unused.png"
    Image.new("RGB", (20, 20), "white").save(image)
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[5]).shapes.title.text = "layout6を使用"
    deck.save(path)
    p_ns = "http://schemas.openxmlformats.org/presentationml/2006/main"
    a_ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    r_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    with zipfile.ZipFile(path) as archive:
        layout = ET.fromstring(archive.read("ppt/slideLayouts/slideLayout1.xml"))
        rels = ET.fromstring(archive.read("ppt/slideLayouts/_rels/slideLayout1.xml.rels"))
    ET.SubElement(layout.find(f"{{{p_ns}}}cSld"), f"{{{a_ns}}}blip", {f"{{{r_ns}}}embed": "rUnused"})
    ET.SubElement(rels, f"{{{rel_ns}}}Relationship", {"Id": "rUnused", "Type": f"{r_ns}/image", "Target": "../media/unused.png"})
    rewrite_zip(path, {"ppt/slideLayouts/slideLayout1.xml": ET.tostring(layout),
                       "ppt/slideLayouts/_rels/slideLayout1.xml.rels": ET.tostring(rels),
                       "ppt/media/unused.png": image.read_bytes()})
    result = ex.extract_document(path)
    unit = next(u for u in result.units if u["kind"] == "image")
    assert unit["candidate_reason"] == "unreferenced_asset"
    assert not any("スライド 1" in ref for ref in unit["source_refs"])


def test_tiff_frames_are_distinct_and_ocr_partial_text_survives_limit(tmp_path, monkeypatch):
    path = tmp_path / "frames.tiff"
    Image.new("RGB", (20, 20), "white").save(path, save_all=True, append_images=[Image.new("RGB", (20, 20), "black")])
    calls = mock_ocr(monkeypatch, "同じOCR結果でも別フレーム")
    result = ex.extract_document(path)
    assert len(calls) == 2
    assert [u["image_frame"] for u in result.units] == [0, 1]
    monkeypatch.setattr(ex, "MAX_TEXT_CHARS", 4)
    limited = ex.extract_document(path)
    assert limited.units[0]["text"] == "同じOC"
    assert limited.units[0]["kind"] == "image" and limited.units[0]["status"] == "partial"
    assert limited.status == "partial"


def test_pdf_native_text_survives_missing_image_ocr(tmp_path, monkeypatch):
    from pypdf import PdfReader, PdfWriter
    from test_extractors import make_text_pdf
    native, scan = tmp_path / "native.pdf", tmp_path / "scan.pdf"
    make_text_pdf(native)
    Image.new("RGB", (30, 30), "white").save(scan)
    writer = PdfWriter()
    writer.add_page(PdfReader(native).pages[0])
    writer.pages[0].merge_page(PdfReader(scan).pages[0])
    mixed = tmp_path / "native_with_image.pdf"
    writer.write(mixed)
    monkeypatch.setattr(ex, "tesseract_path", lambda: None)
    result = ex.extract_document(mixed)
    assert "Production report ALPHA" in result.text
    assert result.status == "partial"
    assert any(u["kind"] == "image" and u["status"] == "needs_ocr" for u in result.units)
