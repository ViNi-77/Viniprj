"""登録先別の出力、非破壊整理、投入済み基準の差分を検証する。"""
import copy
import csv
import hashlib
import json
from pathlib import Path

import pytest

from contextgen_kai import exporting
from contextgen_kai.curation import prepare_document, render_unit, override_basis
from contextgen_kai.exporting import activate_export, build_export
from contextgen_kai.storage import Store, identifier, now


def raw_document(units, **overrides):
    original = "\n\n".join(render_unit(unit) for unit in units if unit.get("text"))
    return {"id": "doc-example", "units": units, "original_text": original, "effective_text": original,
            "edited_text": None, "conflict": False, "status": "ok", **overrides}


@pytest.fixture
def workspace(tmp_path):
    source = tmp_path / "日本語 資料"
    source.mkdir()
    store = Store(tmp_path / "state")
    library = store.add_library("業務資料", str(source))
    return store, library, source


def save_doc(store, library, source, name, text=None, units=None):
    if units is None:
        units = [{"unit_id": "body-1", "locator": "本文", "text": text, "kind": "body", "status": "ok"}]
    text = "\n\n".join(render_unit(unit) for unit in units if unit.get("text"))
    path = source / name
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    stat = path.stat()
    store.upsert_extraction(library["id"], name, path, stat.st_size, stat.st_mtime_ns,
                            hashlib.sha256(path.read_bytes()).hexdigest(),
                            {"text": text, "status": "ok", "warnings": [], "units": units, "metadata": {}})
    doc_id = store.one("SELECT id FROM documents WHERE library_id=? AND relative_path=?", (library["id"], name))["id"]
    return store.document(doc_id)


def save_set(store, library, **overrides):
    return store.save_collection({"name": "設備の保全手順", "library_id": library["id"], "purpose": "overview", **overrides})


def test_no_cleanup_is_lossless_and_does_not_mutate_original():
    units = [{"unit_id": str(n), "locator": f"ページ {n}", "text": "  同じ文。\n\n表  A  B\n", "kind": "body"} for n in range(3)]
    units += [{"unit_id": "note", "locator": "ノート", "text": "注意", "kind": "note", "hidden": True},
              {"unit_id": "embed", "locator": "埋め込み", "text": "記録", "kind": "embedded"}]
    doc = raw_document(units)
    before = copy.deepcopy(doc)
    result = prepare_document(doc, {"cleanup": "none"})
    assert result["text"] == doc["original_text"] == result["original_text"]
    assert result["changes"] == []
    assert doc == before
    assert len(result["units"]) == 5


def test_standard_cleanup_keeps_semantically_similar_body_and_all_exact_image_refs():
    units = [
        {"unit_id": "m1", "locator": "スライド 1 footer", "text": "共通注記", "kind": "metadata", "role": "footer"},
        {"unit_id": "m2", "locator": "スライド 2 footer", "text": "共通注記", "kind": "metadata", "role": "footer"},
        {"unit_id": "i1", "locator": "ページ 1 画像", "text": "図中の文字", "kind": "image", "image_hash": "image-A"},
        {"unit_id": "i2", "locator": "ページ 2 画像", "text": "図中の文字", "kind": "image", "image_hash": "image-A"},
        {"unit_id": "i3", "locator": "ページ 3 画像", "text": "図中の文字", "kind": "image", "image_hash": "image-B"},
        {"unit_id": "i4", "locator": "ページ 4 画像", "text": "図中の文字", "kind": "image", "image_hash": "image-A", "image_frame": 1},
        {"unit_id": "b1", "locator": "本文 1", "text": "毎朝、設備を点検する。", "kind": "body"},
        {"unit_id": "b2", "locator": "本文 2", "text": "毎朝設備を点検する。", "kind": "body"},
        {"unit_id": "b3", "locator": "本文 3", "text": "毎朝、設備を点検する。", "kind": "body"},
        {"unit_id": "t1", "locator": "表 A2", "text": "A2:軸 B2:0.02", "clean_text": "軸\t0.02", "kind": "table_row"},
    ]
    doc = raw_document(units)
    result = prepare_document(doc, {"cleanup": "standard"})
    assert len(result["units"]) == 8
    assert result["original_text"] == doc["original_text"]
    assert result["units"][0]["source_refs"] == ["スライド 1 footer", "スライド 2 footer"]
    assert result["units"][1]["source_unit_ids"] == ["doc-example:i1", "doc-example:i2"]
    assert "毎朝、設備を点検する。" in result["text"] and "毎朝設備を点検する。" in result["text"]
    assert result["units"][-1]["text"] == "軸\t0.02"
    assert result["units"][-1]["original_text"] == "A2:軸 B2:0.02"
    restored = prepare_document(doc, {"cleanup": "standard", "unit_overrides": {"doc-example:m2": "include", "doc-example:i2": "include"},
                                      "unit_override_bases": {f"doc-example:{unit['unit_id']}": override_basis(doc, unit) for unit in units if unit["unit_id"] in {"m2", "i2"}}})
    assert len(restored["units"]) == 10


def test_explicit_unit_choices_override_group_exclusion_and_full_edits_have_no_false_page():
    units = [{"unit_id": "n1", "locator": "ページ 2 ノート", "text": "補足", "kind": "note", "hidden": True},
             {"unit_id": "e1", "locator": "スライド 2 / 埋め込み", "text": "原価", "kind": "embedded"},
             {"unit_id": "b1", "locator": "本文", "text": "対象", "kind": "body"}]
    doc = raw_document(units)
    settings = {"include_hidden": False, "include_notes": False, "include_embedded": False,
                "unit_overrides": {"doc-example:n1": "include", "doc-example:b1": "exclude"},
                "unit_override_bases": {f"doc-example:{unit['unit_id']}": override_basis(doc, unit) for unit in units if unit["unit_id"] in {"n1", "b1"}}}
    result = prepare_document(doc, settings)
    assert [unit["unit_id"] for unit in result["units"]] == ["n1"]
    assert len(result["changes"]) == 2
    doc.update(edited_text="手動修正した本文\n", effective_text="手動修正した本文\n")
    result = prepare_document(doc, {})
    assert result["text"] == "手動修正した本文\n"
    assert result["before_chars"] == len(doc["original_text"])
    assert len(result["units"]) == 1
    assert "位置未特定" in result["units"][0]["locator"]
    assert "ページ 2" not in str(result["units"][0]["source_refs"])


def test_office_image_is_not_an_embedded_document_and_note_images_obey_notes_option():
    doc = raw_document([
        {"unit_id": "picture", "locator": "スライド 1 / 埋め込み ppt/media/image1.png / 画像", "text": "図中の数値", "kind": "image"},
        {"unit_id": "note-picture", "locator": "ノート画像", "text": "ノートの文字", "kind": "image", "note": True},
        {"unit_id": "office", "locator": "添付Excel", "text": "添付資料", "kind": "embedded"},
    ])
    prepared = prepare_document(doc, {"include_embedded": False, "include_notes": False})
    assert [unit["unit_id"] for unit in prepared["units"]] == ["picture"]
    assert len(prepared["raw_units"]) == 3
    assert prepared["raw_units"][1]["key"] == "doc-example:note-picture"


@pytest.mark.parametrize("target", ["builder", "studio"])
@pytest.mark.parametrize("purpose", list(exporting.PURPOSES))
def test_profiles_and_template_metadata_and_evaluation_csv(workspace, target, purpose):
    store, library, source = workspace
    save_doc(store, library, source, "点検.txt", "点検は始業前に実施する。")
    collection = save_set(store, library, target=target, purpose=purpose, audience="保全担当者", answer_scope="点検時期",
                          out_of_scope="設計変更", description="設備Aの保全記録", instructions="確認事項を最後に列挙",
                          evaluation_questions=[{"question": "いつ点検する？", "expected_response": "始業前", "source": "点検.txt 本文"}])
    result = build_export(store, collection["id"])
    assert result["state"] == "published"
    root = Path(result["path"])
    manifest = json.loads(result["manifest"])
    assert manifest["target"] == target
    assert len(manifest["knowledge_files"]) == 2
    suffix = ".txt" if target == "builder" else ".md"
    assert all(name.endswith(suffix) for name in manifest["knowledge_files"])
    instructions = (root / "instructions.txt").read_text(encoding="utf-8")
    for value in (exporting.PURPOSES[purpose], "保全担当者", "点検時期", "設計変更", "設備Aの保全記録", "確認事項を最後に列挙"):
        assert value in instructions
    assert "自動アップロードは行いません" in (root / "registration-guide.md").read_text(encoding="utf-8")
    with (root / "evaluation-questions.csv").open(encoding="utf-8-sig", newline="") as stream:
        generic = list(csv.reader(stream))
    assert generic == [["question", "expected_response", "source"], ["いつ点検する？", "始業前", "点検.txt 本文"]]
    with (root / "studio-evaluation.csv").open(encoding="utf-8-sig", newline="") as stream:
        assert list(csv.reader(stream)) == [["Question", "Expected response"], ["いつ点検する？", "始業前"]]
    assert manifest["collection_settings"]["audience"] == "保全担当者"
    assert manifest["source_snapshot"][0]["source_hash"]


def test_every_split_preserves_provenance_and_table_headers_without_losing_content(workspace):
    store, library, source = workspace
    units = [{"unit_id": "table-1", "locator": "シート 点検 / 行 2", "text": "A2:軸 B2:0.02\n" * 3000,
              "kind": "table_row", "table_headers": ["部位", "許容値"], "source_refs": ["シート 点検 / 行 2"],
              "formulas": {"B2": {"formula": "=1/50", "cached_value": 0.02, "cache_present": True, "recalculated": False}}},
             {"unit_id": "page-2", "locator": "ページ 2", "text": "次のページの重要な条件。" * 2000, "kind": "body"}]
    doc = save_doc(store, library, source, "点検.xlsx", units=units)
    collection = save_set(store, library, target="studio")
    result = build_export(store, collection["id"])
    root = Path(result["path"])
    entries = [json.loads(line) for line in (root / "context.jsonl").read_text(encoding="utf-8").splitlines()]
    assert len(entries) > 4
    assert "".join(entry["text"] for entry in entries) == doc["effective_text"]
    for entry in entries:
        body = (root / entry["knowledge_path"]).read_text(encoding="utf-8")
        assert f"位置: {entry['locator']}" in body
        assert entry["unit_key"] == f"{doc['id']}:{entry['unit_id']}"
        if entry["kind"] == "table_row":
            assert "表の見出し: 部位 | 許容値" in body
            assert "再計算は行っていません" in body
            assert entry["formulas"]["B2"]["cached_value"] == 0.02


def test_cross_library_selection_and_explicit_omission_are_recorded(workspace, tmp_path):
    store, library, source = workspace
    a = save_doc(store, library, source, "現場.txt", "設備Aの情報")
    other_source = tmp_path / "別フォルダ"
    other_source.mkdir()
    other_library = store.add_library("開発", str(other_source))
    b = save_doc(store, other_library, other_source, "設計.txt", "設備Bの情報")
    omitted = save_doc(store, other_library, other_source, "除外.txt", "収録しない情報")
    collection = save_set(store, library, library_ids=[library["id"], other_library["id"]],
                          excluded_document_ids=[omitted["id"]])
    result = build_export(store, collection["id"])
    assert result["document_count"] == 2
    manifest = json.loads(result["manifest"])
    assert manifest["omitted"] == 1
    assert {item["id"] for item in manifest["source_snapshot"]} == {a["id"], b["id"], omitted["id"]}
    refs = [json.loads(line) for line in (Path(result["path"]) / "sources.jsonl").read_text(encoding="utf-8").splitlines()]
    assert next(item for item in refs if item["id"] == omitted["id"])["omitted_reason"] == "資料セットの設定で除外"


def test_identical_names_in_different_libraries_keep_unambiguous_duplicate_sources(workspace, tmp_path):
    store, library, source = workspace
    a = save_doc(store, library, source, "手順.txt", "原文が完全一致する点検手順。")
    other_source = tmp_path / "別の登録元"
    other_source.mkdir()
    other_library = store.add_library("別部署", str(other_source))
    b = save_doc(store, other_library, other_source, "手順.txt", "原文が完全一致する点検手順。")
    collection = save_set(store, library, library_ids=[library["id"], other_library["id"]])
    result = build_export(store, collection["id"])
    root = Path(result["path"])
    refs = [json.loads(line) for line in (root / "sources.jsonl").read_text(encoding="utf-8").splitlines()]
    kept = next(item for item in refs if item["duplicate_of_id"] is None)
    alias = next(item for item in refs if item["duplicate_of_id"] is not None)
    assert {item["id"] for item in refs} == {a["id"], b["id"]}
    assert {item["library_id"] for item in refs} == {library["id"], other_library["id"]}
    assert alias["duplicate_of_id"] == kept["id"] != alias["id"]
    assert alias["duplicate_of_library_id"] == kept["library_id"] != alias["library_id"]
    index = (root / "M365/set_001/M365AgentContext_INDEX.txt").read_text(encoding="utf-8")
    assert all(item[key] in index for item in refs for key in ("id", "library_id"))
    assert result["document_count"] == 1


def test_diff_uses_recorded_handoff_not_most_recent_generated_version(workspace):
    store, library, source = workspace
    save_doc(store, library, source, "点検.txt", "初版の点検時期は毎朝。")
    collection = save_set(store, library)
    first = build_export(store, collection["id"])
    first_manifest = json.loads(first["manifest"])
    second_without_handoff = build_export(store, collection["id"])
    diff = json.loads(second_without_handoff["manifest"])["upload_changes"]
    assert set(diff["added"]) == set(first_manifest["knowledge_files"])
    assert diff["changed"] == [] and diff["baseline_handoff_id"] is None
    handoff_id = identifier()
    store.execute("INSERT INTO handoffs VALUES(?,?,?,?,?,?)", (handoff_id, collection["id"], "builder", first["id"], now(), json.dumps(first_manifest["knowledge_files"])))
    save_doc(store, library, source, "点検.txt", "改訂版の点検時期は毎朝と毎夕。")
    unsubmitted = build_export(store, collection["id"])
    latest = build_export(store, collection["id"])
    diff = json.loads(latest["manifest"])["upload_changes"]
    assert diff["baseline_generation_id"] == first["id"] != unsubmitted["id"]
    assert diff["changed"] and diff["added"] == []
    assert diff["baseline_handoff_id"] == handoff_id
    # 部分投入したファイルの状態を基準にする。別targetの記録は使わない。
    files = dict(first_manifest["knowledge_files"], **{"M365/old-file.txt": "oldhash"})
    store.execute("INSERT INTO handoffs VALUES(?,?,?,?,?,?)", (identifier(), collection["id"], "builder", unsubmitted["id"], now(), json.dumps(files)))
    diff = json.loads(build_export(store, collection["id"])["manifest"])["upload_changes"]
    assert diff["removed"] == ["M365/old-file.txt"]
    studio = save_set(store, library, target="studio")
    assert json.loads(build_export(store, studio["id"])["manifest"])["upload_changes"]["baseline_handoff_id"] is None


@pytest.mark.parametrize("limit", ["files", "bytes"])
def test_studio_capacity_holds_all_files_and_cannot_force_activate(workspace, monkeypatch, limit):
    store, library, source = workspace
    save_doc(store, library, source, "点検.txt", "常用の点検手順を記録する。")
    collection = save_set(store, library, target="studio")
    first = build_export(store, collection["id"])
    if limit == "files":
        monkeypatch.setattr(exporting, "STUDIO_MAX_FILES", 1)
    else:
        monkeypatch.setattr(exporting, "STUDIO_MAX_BYTES", 10)
    result = build_export(store, collection["id"], force=True)
    manifest = json.loads(result["manifest"])
    assert result["state"] == "held" and manifest["blocking_limits"]
    assert len(manifest["knowledge_files"]) == 2
    assert all((Path(result["path"]) / name).is_file() for name in manifest["knowledge_files"])
    with pytest.raises(ValueError, match="上限"):
        activate_export(store, result["id"])
    assert store.one("SELECT id FROM exports WHERE is_active=1")["id"] == first["id"]


def test_invalid_studio_evaluation_keeps_all_questions_in_generic_csv(workspace):
    store, library, source = workspace
    save_doc(store, library, source, "点検.txt", "点検資料")
    questions = [{"question": f"質問 {n}", "expected_response": "利用者の期待", "source": "本文"} for n in range(101)]
    collection = save_set(store, library, evaluation_questions=questions)
    result = build_export(store, collection["id"])
    root = Path(result["path"])
    assert not (root / "studio-evaluation.csv").exists()
    with (root / "evaluation-questions.csv").open(encoding="utf-8-sig", newline="") as stream:
        assert len(list(csv.reader(stream))) == 102
    assert "100問" in json.loads(result["manifest"])["evaluation_warnings"][0]


@pytest.mark.parametrize("kind", ["word", "powerpoint"])
@pytest.mark.parametrize("explicit_header", [True, False])
def test_office_explicit_headers_repeat_on_every_long_row_split(workspace, kind, explicit_header):
    from contextgen_kai.extractors import extract_document
    store, library, source = workspace
    long_text = "確認条件の記録。" * 2500
    if kind == "word":
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        document = Document()
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "対象"
        table.cell(0, 1).text = "確認事項"
        table.cell(1, 0).text = "設備A"
        table.cell(1, 1).text = long_text
        marker = OxmlElement("w:tblHeader")
        marker.set(qn("w:val"), "1" if explicit_header else "0")
        table.rows[0]._tr.get_or_add_trPr().append(marker)
        path = source / "手順.docx"
        document.save(path)
    else:
        from pptx import Presentation
        from pptx.util import Inches
        document = Presentation()
        slide = document.slides.add_slide(document.slide_layouts[6])
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(8), Inches(4)).table
        table.first_row = explicit_header
        table.cell(0, 0).text = "対象"
        table.cell(0, 1).text = "確認事項"
        table.cell(1, 0).text = "設備A"
        table.cell(1, 1).text = long_text
        path = source / "手順.pptx"
        document.save(path)
    extracted = extract_document(path, ocr=False)
    rows = [unit for unit in extracted.units if unit["kind"] == "table_row"]
    assert len(rows) == 2
    expected = ["対象", "確認事項"] if explicit_header else []
    assert all(unit["table_headers"] == expected for unit in rows)
    assert rows[1]["values"] == ["設備A", long_text]
    stat = path.stat()
    store.upsert_extraction(library["id"], path.name, path, stat.st_size, stat.st_mtime_ns,
                            hashlib.sha256(path.read_bytes()).hexdigest(), extracted.as_dict())
    collection = save_set(store, library, target="studio")
    result = build_export(store, collection["id"])
    root = Path(result["path"])
    entries = [json.loads(line) for line in (root / "context.jsonl").read_text(encoding="utf-8").splitlines()]
    assert len(entries) > 2
    assert "".join(entry["text"] for entry in entries) == extracted.text
    for entry in entries:
        assert entry["table_headers"] == expected
        body = (root / entry["knowledge_path"]).read_text(encoding="utf-8")
        assert f"位置: {entry['locator']}" in body
        if explicit_header:
            assert "表の見出し: 対象 | 確認事項" in body
        else:
            assert "表の見出し:" not in body
