"""ローカル資料の非破壊抽出。上限・欠落・OCRの失敗を必ず結果に残す。"""
from __future__ import annotations

from dataclasses import asdict, dataclass, field
from contextlib import contextmanager
from copy import deepcopy
import hashlib
import io
import os
from pathlib import Path, PurePosixPath
import posixpath
import re
import shutil
import stat
import subprocess
import sys
import tempfile
import time
from typing import Any
import zipfile

from defusedxml import ElementTree as ET

EXTRACTOR_VERSION = "2"
OFFICE_EXTENSIONS = {".docx", ".docm", ".xlsx", ".xlsm", ".pptx", ".pptm"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".tsv", ".json", ".jsonl", ".log"}
LEGACY_EXTENSIONS = {".doc", ".xls", ".ppt", ".rtf", ".odt", ".ods", ".odp"}
SUPPORTED_EXTENSIONS = OFFICE_EXTENSIONS | IMAGE_EXTENSIONS | TEXT_EXTENSIONS | LEGACY_EXTENSIONS | {".pdf", ".zip"}
MAX_FILE_BYTES = 256 * 1024 * 1024
MAX_MEMBER_BYTES = 64 * 1024 * 1024
MAX_ARCHIVE_BYTES = 512 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 5000
MAX_COMPRESSION_RATIO = 2000
MAX_TEXT_CHARS = 10_000_000
MAX_UNITS = 20_000
MAX_DEPTH = 3
MAX_IMAGE_PIXELS = 40_000_000
OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


@dataclass
class Extraction:
    text: str = ""
    status: str = "empty"
    warnings: list[str] = field(default_factory=list)
    units: list[dict[str, Any]] = field(default_factory=list)
    metadata: dict[str, Any] = field(default_factory=dict)

    def as_dict(self) -> dict[str, Any]:
        return asdict(self)


class _Limit(Exception):
    pass


@dataclass
class _Budget:
    deadline: float
    max_pages: int
    ocr: bool
    timeout: float
    expanded_bytes: int = 0
    entry_count: int = 0
    chars: int = 0
    units: int = 0
    image_cache: dict[str, Extraction] = field(default_factory=dict)

    def check(self):
        if time.monotonic() >= self.deadline:
            raise _Limit("処理時間の上限に達しました。未読部分があります。")

    def add(self, result: Extraction, locator: str, text: str, status: str = "ok", **attributes):
        self.check()
        text = text.strip()
        unit = {"locator": locator, "text": text, "status": status, "kind": "body",
                "hidden": False, "source_refs": [locator], **attributes}
        if self.units >= MAX_UNITS:
            raise _Limit(f"抽出単位の上限 {MAX_UNITS:,} 件に達しました。未読部分があります。")
        available = MAX_TEXT_CHARS - self.chars
        if len(text) > available:
            if available > 0:
                # 整理用情報にも切り詰め前の本文を残さず、原文上限を迂回させない。
                result.units.append({k: v for k, v in unit.items() if k not in {
                    "clean_text", "values", "formulas", "table_headers"}} | {"text": text[:available], "status": "partial"})
                self.chars += available
                self.units += 1
            raise _Limit(f"抽出文字数の上限 {MAX_TEXT_CHARS:,} 字に達しました。未読部分があります。")
        self.chars += len(text)
        self.units += 1
        result.units.append(unit)


def _warn(result: Extraction, text: str):
    if text not in result.warnings:
        result.warnings.append(text)


def _finalize(result: Extraction) -> Extraction:
    occurrences = {}
    for unit in result.units:
        unit.setdefault("kind", "body")
        unit.setdefault("hidden", False)
        unit.setdefault("source_refs", [unit["locator"]])
        key = f"{unit['locator']}\0{unit['kind']}"
        occurrences[key] = occurrences.get(key, 0) + 1
        # 内容訂正でIDが変わらないよう、本文ではなく出典と種別から生成する。
        unit["unit_id"] = hashlib.sha256(f"{key}\0{occurrences[key]}".encode()).hexdigest()[:24]
    result.text = "\n\n".join(f"[{u['locator']}]\n{u['text']}" for u in result.units if u["text"])
    result.metadata["unit_count"] = len(result.units)
    result.metadata["text_characters"] = sum(len(u["text"]) for u in result.units)
    result.metadata["extractor_version"] = EXTRACTOR_VERSION
    if result.text:
        result.status = "partial" if result.warnings or any(u["status"] not in {"ok", "ocr"} for u in result.units) else "ok"
    elif result.status not in {"protected", "needs_conversion", "needs_ocr", "error"}:
        result.status = "error" if result.warnings else "empty"
    return result


def extract_document(path: Path, *, ocr: bool = True, max_pages: int = 500, timeout: float = 60) -> Extraction:
    """本文・出典単位・警告を返す。原本への書込み、マクロ実行、外部通信は行わない。

    timeout は資料全体の協調的期限であり、OCR子プロセスには残り時間を強制適用。
    Office/PDFライブラリ内部の単一呼出しは中断できないため、バイト数上限も併用する。
    """
    if max_pages < 1 or timeout <= 0:
        raise ValueError("max_pages と timeout は正数で指定してください。")
    budget = _Budget(time.monotonic() + timeout, max_pages, ocr, timeout)
    return _extract(Path(path), budget, 0)


def _extract(path: Path, budget: _Budget, depth: int) -> Extraction:
    result = Extraction(metadata={"filename": path.name, "format": path.suffix.lower(), "ocr_requested": budget.ocr,
                                  "retryable": False})
    try:
        budget.check()
        if depth > MAX_DEPTH:
            raise _Limit(f"入れ子の上限 {MAX_DEPTH} 階層に達しました。内部資料を個別に登録してください。")
        size = path.stat().st_size
        result.metadata["source_bytes"] = size
        if size > MAX_FILE_BYTES:
            raise _Limit(f"ファイル容量の上限 {MAX_FILE_BYTES // 1024 // 1024} MB を超えています。資料を分割してください。")
        suffix = path.suffix.lower()
        if suffix in LEGACY_EXTENSIONS or suffix not in SUPPORTED_EXTENSIONS:
            result.status = "needs_conversion"
            _warn(result, "この形式は直接読み取れません。Officeの新形式（DOCX・XLSX・PPTX）またはPDF・TXTへ変換してください。")
        elif suffix in TEXT_EXTENSIONS:
            _plain(path, result, budget)
        elif suffix == ".pdf":
            _pdf(path, result, budget)
        elif suffix in IMAGE_EXTENSIONS:
            _image_file(path, result, budget, "画像")
        elif suffix == ".zip":
            _archive(path, result, budget, depth)
        else:
            with path.open("rb") as src:
                signature = src.read(8)
            if signature == OLE_MAGIC:
                result.status = "protected"
                _warn(result, "暗号化されたOffice文書、または拡張子と内容が異なる旧形式です。保護解除または新形式での保存が必要です。")
            else:
                _office(path, result, budget, depth)
    except _Limit as exc:
        _warn(result, str(exc))
    except PermissionError:
        result.status = "error"
        result.metadata["retryable"] = True
        _warn(result, "資料を読み取る権限がありません。アクセス権と他アプリでの使用状況を確認してください。")
    except FileNotFoundError:
        result.status = "error"
        result.metadata["retryable"] = True
        _warn(result, "資料が見つかりません。クラウド上の資料はローカルへ取得してから再実行してください。")
    except Exception as exc:
        result.status = "error"
        result.metadata["retryable"] = True
        # 本文・パス・外部ライブラリの生エラーをログやUIに漏らさない。
        _warn(result, f"読み取りに失敗しました（{type(exc).__name__}）。破損・保護・クラウド未取得を確認してください。")
    return _finalize(result)


def _plain(path: Path, result: Extraction, budget: _Budget):
    raw = path.read_bytes()
    if raw.startswith((b"\xff\xfe\x00\x00", b"\x00\x00\xfe\xff")):
        encodings = ["utf-32"]
    elif raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        encodings = ["utf-16"]
    else:
        encodings = ["utf-8-sig", "cp932"]
    text = None
    for encoding in encodings:
        try:
            text = raw.decode(encoding)
            result.metadata["encoding"] = encoding
            break
        except UnicodeDecodeError:
            pass
    if text is None:
        from charset_normalizer import from_bytes
        match = from_bytes(raw).best()
        if match is None:
            raise ValueError("文字コードを判定できません")
        text = str(match)
        result.metadata["encoding"] = match.encoding
        _warn(result, f"文字コードを {match.encoding} と推定しました。文字化けがないか確認してください。")
    if "\x00" in text:
        raise ValueError("バイナリデータを含むテキスト")
    budget.add(result, "本文", text)


def tesseract_path() -> str | None:
    """利用者指定、EXEと同梱のOCR、開発環境PATHの順で探す。"""
    configured = os.environ.get("CONTEXTGEN_TESSERACT")
    if configured:
        return str(Path(configured)) if Path(configured).is_file() else None
    # GitにはWindows用の実体も同梱する。macOS開発時はその.exeを実行しない。
    names = ("tesseract.exe", "tesseract") if sys.platform == "win32" else ("tesseract",)
    bases = [Path(sys.executable).parent / "ocr", Path(__file__).resolve().parents[1] / "ocr"]
    if getattr(sys, "_MEIPASS", None):
        bases.append(Path(sys._MEIPASS) / "ocr")
    for base in bases:
        for name in names:
            if (base / name).is_file():
                return str(base / name)
    return shutil.which("tesseract")


def _ocr_image(image, result: Extraction, budget: _Budget, locator: str):
    budget.check()
    if not budget.ocr:
        _warn(result, f"{locator}: OCRが無効のため画像内の文字は未読です。")
        result.status = "needs_ocr"
        budget.add(result, locator, "", "needs_ocr", kind="image")
        return
    executable = tesseract_path()
    if not executable:
        _warn(result, f"{locator}: Tesseractが見つからないため画像内の文字は未読です。OCR同梱版をご利用ください。")
        result.status = "needs_ocr"
        result.metadata["retryable"] = True
        budget.add(result, locator, "", "needs_ocr", kind="image")
        return
    if image.width * image.height > MAX_IMAGE_PIXELS:
        raise _Limit(f"{locator}: 画像が {MAX_IMAGE_PIXELS:,} 画素の上限を超えています。縮小して再登録してください。")
    with tempfile.TemporaryDirectory(prefix="contextgen-kai-ocr-") as tmp:
        img_path = Path(tmp) / "input.png"
        # DPI不明の画像には架空の値を与えず、Tesseractの推定を使う。
        options = {}
        dpi = image.info.get("dpi")
        if isinstance(dpi, (tuple, list)) and len(dpi) == 2 and all(70 <= d <= 2400 for d in dpi):
            options["dpi"] = dpi
        image.convert("RGB").save(img_path, **options)
        command = [executable, str(img_path), "stdout", "-l", os.environ.get("CONTEXTGEN_OCR_LANG", "jpn+eng")]
        tessdata = Path(executable).parent / "tessdata"
        if tessdata.is_dir():
            command.extend(["--tessdata-dir", str(tessdata)])
        flags = subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0
        remaining = min(30, max(0.01, budget.deadline - time.monotonic()))
        try:
            completed = subprocess.run(command, capture_output=True, timeout=remaining, creationflags=flags, check=False)
        except subprocess.TimeoutExpired:
            _warn(result, f"{locator}: OCRが制限時間を超えました。この部分は未読です。")
            budget.add(result, locator, "", "needs_ocr", kind="image")
            result.status = "needs_ocr"
            result.metadata["retryable"] = True
            return
        if completed.returncode:
            _warn(result, f"{locator}: OCRに失敗しました。日本語・英語の言語データと画像を確認してください。")
            budget.add(result, locator, "", "needs_ocr", kind="image")
            result.status = "needs_ocr"
            result.metadata["retryable"] = True
            return
        text = completed.stdout.decode("utf-8", errors="replace").strip()
        result.metadata["ocr_engine"] = "Tesseract"
        result.metadata["ocr_units"] = result.metadata.get("ocr_units", 0) + 1
        if not text:
            _warn(result, f"{locator}: OCRで文字を検出できませんでした。原本を確認してください。")
        budget.add(result, locator, text, "ocr" if text else "empty", kind="image")


def _image_file(path: Path, result: Extraction, budget: _Budget, locator: str):
    from PIL import Image, ImageOps
    with path.open("rb") as source:
        image_hash = hashlib.file_digest(source, "sha256").hexdigest()
    with Image.open(path) as image:
        frames = getattr(image, "n_frames", 1)
        if frames > budget.max_pages:
            _warn(result, f"画像は {frames} ページ中 {budget.max_pages} ページまで処理します。残りは未読です。")
        for i in range(min(frames, budget.max_pages)):
            budget.check()
            image.seek(i)
            if image.width * image.height > MAX_IMAGE_PIXELS:
                raise _Limit(f"{locator}: 画像サイズの上限を超えています。縮小して再登録してください。")
            with ImageOps.exif_transpose(image) as oriented:
                _cached_ocr(oriented, result, budget,
                            f"{locator} {i + 1}" if frames > 1 else locator, image_hash, i)


def _cached_ocr(image, result: Extraction, budget: _Budget, locator: str, image_hash: str, frame: int = 0):
    """画像バイト列が一致する場合だけOCR計算を共有し、各出典の原文単位は残す。"""
    key = f"{image_hash}:{frame}"
    cached = budget.image_cache.get(key)
    if cached is None:
        child = Extraction()
        try:
            _ocr_image(image, child, budget, locator)
        except _Limit:
            # 上限まで読み取った原文も失わない。
            for unit in child.units:
                unit.update(kind="image", image_hash=image_hash, image_frame=frame)
            result.units.extend(child.units)
            raise
        # 失敗・空結果は再試行可能にする。本文の一致や意味の類似では共有しない。
        if child.units and all(u["status"] == "ocr" for u in child.units):
            budget.image_cache[key] = deepcopy(child)
        result.units.extend(child.units)
        result.warnings.extend(w for w in child.warnings if w not in result.warnings)
        if child.status == "needs_ocr":
            result.status = child.status
        result.metadata["ocr_units"] = result.metadata.get("ocr_units", 0) + child.metadata.get("ocr_units", 0)
        if "ocr_engine" in child.metadata:
            result.metadata["ocr_engine"] = child.metadata["ocr_engine"]
        if child.metadata.get("retryable"):
            result.metadata["retryable"] = True
        units = child.units
    else:
        start = len(result.units)
        for unit in cached.units:
            budget.add(result, locator, unit["text"], unit["status"], kind="image", image_hash=image_hash, image_frame=frame)
        units = result.units[start:]
        result.metadata["ocr_cache_hits"] = result.metadata.get("ocr_cache_hits", 0) + 1
        result.metadata["ocr_engine"] = cached.metadata.get("ocr_engine", "Tesseract")
    for unit in units:
        unit.update(kind="image", image_hash=image_hash, image_frame=frame)


def _pdf_images(page, result: Extraction, budget: _Budget, locator: str):
    """文字と画像が同居するページでは画像部品だけを追加OCRする。本文は再OCRしない。"""
    try:
        for index, item in enumerate(page.images, 1):
            budget.check()
            image_locator = f"{locator} 画像 {index} ({item.name})"
            if not budget.ocr or not tesseract_path():
                before = len(result.units)
                _ocr_image(None, result, budget, image_locator)
                for unit in result.units[before:]:
                    unit.update(kind="image", parent_locator=locator)
                continue
            image_hash = hashlib.sha256(item.data).hexdigest()
            image = item.image
            try:
                before = len(result.units)
                _cached_ocr(image, result, budget, image_locator, image_hash)
                for unit in result.units[before:]:
                    unit["parent_locator"] = locator
                    unit["candidate_reason"] = "native_and_image_text"
            finally:
                image.close()
    except _Limit:
        raise
    except Exception as exc:
        result.metadata["retryable"] = True
        _warn(result, f"{locator}: 画像内文字を読み取れませんでした（{type(exc).__name__}）。本文は保持しています。")
        budget.add(result, f"{locator} 画像", "", "error", kind="image", parent_locator=locator)


def _pdf(path: Path, result: Extraction, budget: _Budget):
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            unlocked = reader.decrypt("")
        except Exception:
            unlocked = False
        if not unlocked:
            result.status = "protected"
            _warn(result, "パスワード保護されたPDFです。保護を解除してから登録してください。")
            return
    total = len(reader.pages)
    result.metadata["page_count"] = total
    if total > budget.max_pages:
        _warn(result, f"PDFは {total} ページ中 {budget.max_pages} ページまで処理します。残りは未読です。")
    rendered = None
    try:
        for number, page in enumerate(reader.pages[:budget.max_pages], 1):
            budget.check()
            locator = f"ページ {number}"
            try:
                text = page.extract_text() or ""
                if text.strip():
                    budget.add(result, locator, text)
                    _pdf_images(page, result, budget, locator)
                elif not budget.ocr or not tesseract_path():
                    _ocr_image(None, result, budget, locator)
                else:
                    if rendered is None:
                        import pypdfium2
                        rendered = pypdfium2.PdfDocument(str(path))
                    pdf_page = rendered[number - 1]
                    try:
                        width, height = pdf_page.get_size()
                        scale = min(200 / 72, (MAX_IMAGE_PIXELS / max(1, width * height)) ** 0.5)
                        if scale < 200 / 72:
                            _warn(result, f"{locator}: OCR画像を画素数の上限に合わせて縮小しました。小さい文字は原本で確認してください。")
                        bitmap = pdf_page.render(scale=scale)
                        try:
                            pil_image = bitmap.to_pil()
                            pil_image.info["dpi"] = (scale * 72, scale * 72)
                            _ocr_image(pil_image, result, budget, locator)
                            pil_image.close()
                        finally:
                            bitmap.close()
                    finally:
                        pdf_page.close()
            except _Limit:
                raise
            except Exception as exc:
                result.metadata["retryable"] = True
                _warn(result, f"{locator}: 読み取りに失敗しました（{type(exc).__name__}）。このページは未読です。")
                budget.add(result, locator, "", "error")
    finally:
        if rendered is not None:
            rendered.close()


def _safe_info(info: zipfile.ZipInfo, budget: _Budget):
    name = info.filename.replace("\\", "/")
    path = PurePosixPath(name)
    if path.is_absolute() or ".." in path.parts or re.match(r"^[A-Za-z]:", name) or "\x00" in name:
        raise ValueError("安全でない格納パス")
    if stat.S_ISLNK(info.external_attr >> 16):
        raise ValueError("シンボリックリンク")
    if info.flag_bits & 1:
        raise ValueError("パスワード保護")
    if info.file_size > MAX_MEMBER_BYTES:
        raise ValueError(f"単一項目が {MAX_MEMBER_BYTES // 1024 // 1024} MB を超過")
    if info.file_size / max(1, info.compress_size) > MAX_COMPRESSION_RATIO:
        raise ValueError("展開率の上限超過")
    budget.check()


def _read_member(zf: zipfile.ZipFile, info: zipfile.ZipInfo, budget: _Budget) -> bytes:
    _safe_info(info, budget)
    if budget.entry_count >= MAX_ARCHIVE_ENTRIES:
        raise _Limit(f"展開項目数が上限 {MAX_ARCHIVE_ENTRIES:,} 件に達しました。残りは未読です。")
    if budget.expanded_bytes + info.file_size > MAX_ARCHIVE_BYTES:
        raise _Limit("ZIP・Officeの累計展開容量が上限 512 MB に達しました。残りは未読です。")
    budget.entry_count += 1
    budget.expanded_bytes += info.file_size
    with zf.open(info) as src:
        data = src.read(MAX_MEMBER_BYTES + 1)
    if len(data) > MAX_MEMBER_BYTES:
        raise _Limit("ZIP項目が展開容量の上限を超えました。")
    return data


def _xml(zf: zipfile.ZipFile, name: str, budget: _Budget):
    return ET.fromstring(_read_member(zf, zf.getinfo(name), budget))


@contextmanager
def _xml_stream(zf: zipfile.ZipFile, name: str, budget: _Budget):
    """大きいシートは矩形の空セルを生成せず、格納されているXML行だけを読む。"""
    info = zf.getinfo(name)
    _safe_info(info, budget)
    if budget.entry_count >= MAX_ARCHIVE_ENTRIES or budget.expanded_bytes + info.file_size > MAX_ARCHIVE_BYTES:
        raise _Limit("Officeの累計展開容量・項目数の上限に達しました。残りは未読です。")
    budget.entry_count += 1
    budget.expanded_bytes += info.file_size
    with zf.open(info) as stream:
        yield ET.iterparse(stream, events=("start", "end"))


def _relations(zf: zipfile.ZipFile, part: str, budget: _Budget) -> dict[str, tuple[str, str]]:
    rel_name = str(PurePosixPath(part).parent / "_rels" / (PurePosixPath(part).name + ".rels"))
    if rel_name not in zf.namelist():
        return {}
    result = {}
    for item in _xml(zf, rel_name, budget):
        if item.get("TargetMode") == "External":
            continue
        target = item.get("Target", "").replace("\\", "/")
        target = posixpath.normpath(target.lstrip("/") if target.startswith("/") else posixpath.join(posixpath.dirname(part), target))
        if target in zf.namelist() and not target.startswith("../"):
            result[item.get("Id", "")] = (target, item.get("Type", "").rsplit("/", 1)[-1])
    return result


def _reference_ids(root) -> set[str]:
    return {value for node in root.iter() for key, value in node.attrib.items()
            if "relationships}" in key and _local(key) in {"id", "embed", "link"}}


def _shape_hidden(node) -> bool:
    if _local(node.tag) not in {"sp", "grpSp", "graphicFrame", "pic", "cxnSp"}:
        return False
    return any(_local(n.tag) == "cNvPr" and _truth(n.get("hidden"))
               for properties in node if _local(properties.tag).startswith("nv") for n in properties.iter())


def _reference_visibility(root) -> dict[str, bool]:
    result = {}
    pending = [(root, False)]
    while pending:
        node, hidden = pending.pop()
        hidden = hidden or _shape_hidden(node)
        for key, value in node.attrib.items():
            if "relationships}" in key and _local(key) in {"id", "embed", "link"}:
                result[value] = result.get(value, True) and hidden
        pending.extend((child, hidden) for child in node)
    return result


@dataclass
class _OfficeContext:
    roots: dict[str, list[tuple[str, bool]]] = field(default_factory=dict)
    used_ids: dict[str, set[str]] = field(default_factory=dict)
    hidden_refs: dict[str, dict[str, bool]] = field(default_factory=dict)
    relations: dict[str, dict[str, tuple[str, str]]] = field(default_factory=dict)

    def register(self, part: str, locator: str, hidden: bool, root=None):
        self.roots.setdefault(part, []).append((locator, hidden))
        if root is not None:
            self.used_ids[part] = _reference_ids(root)
            self.hidden_refs[part] = _reference_visibility(root)

    def links(self, zf, part, budget):
        if part not in self.relations:
            self.relations[part] = _relations(zf, part, budget)
        return self.relations[part]

    def asset_sources(self, zf, budget) -> dict[str, list[tuple[str, bool]]]:
        sources = {}
        # 親スライド→レイアウト→マスター、シート→drawing→画像、chart→埋込Excelを辿る。
        implicit = {"slideLayout", "slideMaster", "notesMaster"}
        for part, locations in self.roots.items():
            for locator, hidden in locations:
                pending, visited = [(part, 0, hidden)], set()
                while pending:
                    budget.check()
                    current, depth, path_hidden = pending.pop()
                    if (current, path_hidden) in visited:
                        continue
                    visited.add((current, path_hidden))
                    links = self.links(zf, current, budget)
                    if not links:
                        continue
                    if current not in self.used_ids:
                        root = _xml(zf, current, budget)
                        self.used_ids[current] = _reference_ids(root)
                        self.hidden_refs[current] = _reference_visibility(root)
                    for rid, (target, kind) in links.items():
                        # マスター内のレイアウト一覧は「このスライドが使うレイアウト」ではない。
                        if kind == "slideLayout" and not current.startswith("ppt/slides/"):
                            continue
                        # ノートは独立した出典rootから辿り、本文の画像と混同しない。
                        if kind == "notesSlide":
                            continue
                        if rid not in self.used_ids[current] and kind not in implicit:
                            continue
                        asset_hidden = path_hidden or self.hidden_refs.get(current, {}).get(rid, False)
                        if "/media/" in target or "/embeddings/" in target:
                            refs = sources.setdefault(target, [])
                            if (locator, asset_hidden) not in refs:
                                refs.append((locator, asset_hidden))
                        elif depth < 12 and target.endswith(".xml"):
                            pending.append((target, depth + 1, asset_hidden))
        return sources


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _paragraph_text(element) -> str:
    parts = []
    for node in element.iter():
        tag = _local(node.tag)
        if tag == "t":
            parts.append(node.text or "")
        elif tag == "tab":
            parts.append("\t")
        elif tag in {"br", "cr"}:
            parts.append("\n")
    return "".join(parts)


def _office(path: Path, result: Extraction, budget: _Budget, depth: int):
    with zipfile.ZipFile(path) as zf:
        context = _OfficeContext()
        if len(zf.infolist()) > MAX_ARCHIVE_ENTRIES:
            raise _Limit(f"Office内部の項目数が上限 {MAX_ARCHIVE_ENTRIES:,} 件を超えています。資料を分割してください。")
        # OfficeのXMLと埋め込み部品は共通のZIP上限を満たす文書だけ受け取る。
        for info in zf.infolist():
            _safe_info(info, budget)
        if sum(i.file_size for i in zf.infolist()) > MAX_ARCHIVE_BYTES:
            raise _Limit("Officeの展開容量が上限 512 MB を超えています。")
        if path.suffix.lower() in {".docx", ".docm"}:
            _word(zf, result, budget, context)
        elif path.suffix.lower() in {".xlsx", ".xlsm"}:
            _excel(zf, result, budget, context)
        else:
            _powerpoint(zf, result, budget, context)
        if any("vbaproject" in n.lower() for n in zf.namelist()):
            result.metadata["macros_present"] = True
            _warn(result, "マクロを含む資料です。マクロは実行せず本文のみ読み取っています。")
        _office_assets(zf, result, budget, depth, context)
        repeated = {}
        for unit in result.units:
            if unit["kind"] not in {"body", "note", "image"} or not unit["text"].strip():
                continue
            # 候補の提示だけ。数値・条件や表の別行を除去する判定には使わない。
            key = " ".join(unit["text"].split())
            repeated.setdefault(key, []).append(unit)
        for group in repeated.values():
            if len(group) > 1:
                reason = "native_and_image_text" if any(u["kind"] == "image" for u in group) else "repeated_text"
                for unit in group:
                    unit.setdefault("candidate_reason", reason)


def _word(zf: zipfile.ZipFile, result: Extraction, budget: _Budget, context: _OfficeContext):
    names = ["word/document.xml"]
    names += sorted(n for n in zf.namelist() if re.fullmatch(r"word/(header\d+|footer\d+|footnotes|endnotes)\.xml", n))
    for name in names:
        root = _xml(zf, name, budget)
        title = "本文" if name == "word/document.xml" else PurePosixPath(name).stem
        context.register(name, title, False, root)
        body = next((n for n in root if _local(n.tag) == "body"), root)
        paragraph = table = 0
        for node in body:
            budget.check()
            tag = _local(node.tag)
            if tag == "p":
                paragraph += 1
                text = _paragraph_text(node)
                if text:
                    budget.add(result, f"{title} 段落 {paragraph}", text)
            elif tag == "tbl":
                table += 1
                rows = []
                for row in node:
                    if _local(row.tag) == "tr":
                        values = [_paragraph_text(cell) for cell in row if _local(cell.tag) == "tc"]
                        properties = next((child for child in row if _local(child.tag) == "trPr"), [])
                        header = next((child for child in properties if _local(child.tag) == "tblHeader"), None)
                        header_value = next((value for key, value in header.attrib.items() if _local(key) == "val"), "1") if header is not None else "0"
                        rows.append((values, _truth(header_value)))
                # Wordで見出し行として明示された先頭行だけを用いる。太字や内容では推定しない。
                header_rows = []
                for values, is_header in rows:
                    if not is_header:
                        break
                    header_rows.append(values)
                headers = [" / ".join(row[column] for row in header_rows if column < len(row) and row[column])
                           for column in range(max((len(row) for row in header_rows), default=0))]
                for row_number, (values, is_header) in enumerate(rows, 1):
                    budget.add(result, f"{title} 表 {table} 行 {row_number}", "\t".join(values), kind="table_row",
                               table_id=f"{name}#table-{table}", table_headers=headers, table_header_row=is_header,
                               values=values, cells=[], parent_locator=f"{title} 表 {table}")
            elif tag in {"footnote", "endnote", "sdt", "customXml"}:
                paragraph += 1
                text = "\n".join(_paragraph_text(p) for p in node.iter() if _local(p.tag) == "p")
                if text:
                    budget.add(result, f"{title} 補足 {paragraph}", text)
    result.metadata["word_pagination"] = "Wordのレイアウト計算は行わず、段落・表番号で出典を示します。"


def _truth(value: str | None) -> bool:
    return value in {"1", "true", "True", "on"}


def _excel_strings(zf, budget):
    strings, chars = [], 0
    if "xl/sharedStrings.xml" in zf.namelist():
        with _xml_stream(zf, "xl/sharedStrings.xml", budget) as events:
            stack = []
            for event, node in events:
                if event == "start":
                    stack.append(node)
                    continue
                if _local(node.tag) == "si":
                    budget.check()
                    value = "".join(n.text or "" for n in node.iter() if _local(n.tag) == "t")
                    chars += len(value)
                    if chars > MAX_TEXT_CHARS or len(strings) >= 1_000_000:
                        raise _Limit("Excel共有文字列の上限に達しました。資料を分割してください。")
                    strings.append(value)
                    node.clear()
                    if len(stack) > 1:
                        stack[-2].remove(node)
                stack.pop()
    return strings


def _excel_date_styles(zf, budget):
    from openpyxl.styles.numbers import BUILTIN_FORMATS, is_date_format
    if "xl/styles.xml" not in zf.namelist():
        return set()
    root = _xml(zf, "xl/styles.xml", budget)
    formats = dict(BUILTIN_FORMATS)
    for node in root.iter():
        if _local(node.tag) == "numFmt":
            formats[int(node.get("numFmtId", "0"))] = node.get("formatCode", "")
    xfs = next((node for node in root if _local(node.tag) == "cellXfs"), [])
    return {i for i, xf in enumerate(xfs) if is_date_format(formats.get(int(xf.get("numFmtId", "0")), ""))}


def _excel_cell(node, strings, date_styles, epoch):
    from openpyxl.utils.datetime import from_excel
    kind = node.get("t", "n")
    cached = next((n for n in node if _local(n.tag) == "v"), None)
    raw = cached.text if cached is not None else None
    if kind == "inlineStr":
        value = "".join(n.text or "" for n in node.iter() if _local(n.tag) == "t")
    elif kind == "s" and raw is not None:
        value = strings[int(raw)]
    elif kind == "b" and raw is not None:
        value = "True" if raw == "1" else "False"
    elif raw is not None and kind == "n" and int(node.get("s", "0")) in date_styles:
        value = str(from_excel(float(raw), epoch=epoch))
    elif kind == "str" and cached is not None:
        value = raw or ""
    else:
        value = raw
    formula_node = next((n for n in node if _local(n.tag) == "f"), None)
    return value, formula_node


def _excel_tables(zf, part, context, budget):
    from openpyxl.utils.cell import range_boundaries
    tables = []
    for target, relation_type in context.links(zf, part, budget).values():
        if relation_type != "table":
            continue
        table = _xml(zf, target, budget)
        columns = [n.get("name", "") for n in table.iter() if _local(n.tag) == "tableColumn"]
        tables.append({"id": target, "name": table.get("displayName", table.get("name", "")),
                       "bounds": range_boundaries(table.get("ref", "A1")),
                       "headers": columns if table.get("headerRowCount", "1") != "0" else [],
                       "has_header": table.get("headerRowCount", "1") != "0"})
    return tables


def _excel(zf: zipfile.ZipFile, result: Extraction, budget: _Budget, context: _OfficeContext):
    from openpyxl.utils.cell import column_index_from_string, get_column_letter
    from openpyxl.utils.datetime import CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900
    from openpyxl.formula.translate import Translator
    workbook = _xml(zf, "xl/workbook.xml", budget)
    links = context.links(zf, "xl/workbook.xml", budget)
    sheets = [n for n in workbook.iter() if _local(n.tag) == "sheet"]
    epoch = CALENDAR_MAC_1904 if any(_local(n.tag) == "workbookPr" and _truth(n.get("date1904")) for n in workbook) else CALENDAR_WINDOWS_1900
    strings = _excel_strings(zf, budget)
    date_styles = _excel_date_styles(zf, budget)
    result.metadata["sheet_count"] = len(sheets)
    if len(sheets) > budget.max_pages:
        _warn(result, f"シート上限 {budget.max_pages} 件を超えたシートは未読です。")
    formula_present = False
    for sheet in sheets[:budget.max_pages]:
        budget.check()
        title = sheet.get("name", "無名")
        locator = f"シート「{title}」"
        rid = next((v for k, v in sheet.attrib.items() if _local(k) == "id"), "")
        if rid not in links:
            _warn(result, f"{locator}: シートの参照先が見つかりません。")
            continue
        part = links[rid][0]
        sheet_hidden = sheet.get("state", "visible") != "visible"
        context.register(part, locator, sheet_hidden)
        context.used_ids[part] = set()
        tables = _excel_tables(zf, part, context, budget)
        hidden_columns, shared_formulas = set(), {}
        rows_seen = cells_seen = 0
        with _xml_stream(zf, part, budget) as events:
            stack = []
            for event, node in events:
                tag = _local(node.tag)
                if event == "start":
                    stack.append(node)
                    for key, value in node.attrib.items():
                        if "relationships}" in key:
                            context.used_ids[part].add(value)
                    continue
                budget.check()
                if tag == "col" and _truth(node.get("hidden")):
                    hidden_columns.update(range(int(node.get("min", "1")), min(int(node.get("max", "1")), 256) + 1))
                if tag == "row":
                    rows_seen += 1
                    row_number = int(node.get("r", str(rows_seen)))
                    row_hidden = _truth(node.get("hidden"))
                    cells = []
                    next_column = 1
                    for cell in node:
                        if _local(cell.tag) != "c":
                            continue
                        cells_seen += 1
                        coordinate = cell.get("r", f"{get_column_letter(next_column)}{row_number}")
                        match = re.fullmatch(r"([A-Za-z]+)([1-9][0-9]*)", coordinate)
                        if not match:
                            _warn(result, f"{locator}: 不正なセル座標を読み飛ばしました。")
                            continue
                        column = column_index_from_string(match[1])
                        next_column = column + 1
                        value, formula = _excel_cell(cell, strings, date_styles, epoch)
                        if value is None and formula is None:
                            continue  # 書式だけのセルや空白範囲は本文・欠落に数えない。
                        if row_number > 100_000 or column > 256:
                            _warn(result, f"{locator}: 100,000 行・256 列を超える値のあるセルは未読です。")
                            continue
                        formulas = {}
                        if formula is not None:
                            formula_present = True
                            expression = "=" + (formula.text or "")
                            if formula.get("t") == "shared":
                                si = formula.get("si", "")
                                if formula.text:
                                    shared_formulas[si] = (coordinate, expression)
                                elif si in shared_formulas:
                                    origin, source_formula = shared_formulas[si]
                                    try:
                                        expression = Translator(source_formula, origin=origin).translate_formula(coordinate)
                                    except Exception:
                                        expression = f"[共有数式: {origin}] {source_formula}"
                                else:
                                    expression = "[共有数式の参照先を取得できません]"
                            formulas[coordinate] = {"formula": expression, "cached_value": value,
                                                    "cache_present": value is not None, "recalculated": False}
                            raw = expression + (f" [保存済み値: {value} / 未再計算]" if value is not None else " [保存済み値なし / 未再計算]")
                        else:
                            raw = str(value)
                        cells.append({"coordinate": coordinate, "column": column, "value": raw,
                                      "hidden": sheet_hidden or row_hidden or column in hidden_columns, "formulas": formulas})
                    # 実テーブルと範囲外、表示・非表示を別単位にし、出力時に安全に選べるようにする。
                    groups = {}
                    for cell in cells:
                        table = next((t for t in tables if t["bounds"][1] <= row_number <= t["bounds"][3]
                                      and t["bounds"][0] <= cell["column"] <= t["bounds"][2]), None)
                        key = (table["id"] if table else "", cell["hidden"])
                        groups.setdefault(key, (table, []))[1].append(cell)
                    for (table_id, hidden), (table, group) in groups.items():
                        coords = [c["coordinate"] for c in group]
                        values = [c["value"] for c in group]
                        text = "\t".join(f"{c['coordinate']}: {c['value']}" for c in group)
                        headers = [table["headers"][c["column"] - table["bounds"][0]]
                                   if c["column"] - table["bounds"][0] < len(table["headers"]) else "" for c in group] if table else []
                        attributes = {"kind": "table_row", "hidden": hidden, "cells": coords, "values": values,
                                      "hidden_sheet": sheet_hidden, "hidden_row": row_hidden,
                                      "hidden_columns": [c["coordinate"] for c in group if c["column"] in hidden_columns],
                                      "formulas": {k: v for c in group for k, v in c["formulas"].items()},
                                      "table_id": table_id, "table_headers": headers, "parent_locator": locator}
                        if table:
                            attributes.update(heading=table["name"], clean_text="\t".join(values),
                                              table_header_row=table["has_header"] and row_number == table["bounds"][1])
                        if hidden:
                            attributes["candidate_reason"] = "hidden_content"
                        suffix = (f" 表「{table['name']}」" if table else "") + (" 非表示" if hidden else "")
                        budget.add(result, f"{locator} 行 {row_number}{suffix}", text, **attributes)
                    node.clear()
                    if len(stack) > 1:
                        stack[-2].remove(node)
                stack.pop()
        result.metadata.setdefault("sheet_scan", []).append({"sheet": title, "stored_rows": rows_seen, "stored_cells": cells_seen})
    if formula_present:
        _warn(result, "数式と保存済み値を抽出しています。値は未再計算で、古い場合や保存されていない場合があります。表示結果はExcel原本で確認してください。")


def _natural_key(name: str):
    return [int(part) if part.isdigit() else part for part in re.split(r"(\d+)", name)]


def _powerpoint(zf: zipfile.ZipFile, result: Extraction, budget: _Budget, context: _OfficeContext):
    slides = sorted((n for n in zf.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)), key=_natural_key)
    # presentation.xmlの実際の表示順を優先し、ファイル名順に依存しない。
    if "ppt/presentation.xml" in zf.namelist() and "ppt/_rels/presentation.xml.rels" in zf.namelist():
        targets = {rid: target for rid, (target, _) in context.links(zf, "ppt/presentation.xml", budget).items()}
        actual = []
        for node in _xml(zf, "ppt/presentation.xml", budget).iter():
            if _local(node.tag) == "sldId":
                rid = next((v for k, v in node.attrib.items() if k.endswith("}id")), None)
                target = targets.get(rid, "")
                if target.startswith("ppt/ppt/"):
                    target = target[4:]
                if target in slides:
                    actual.append(target)
        if actual:
            slides = actual
    result.metadata["slide_count"] = len(slides)
    if len(slides) > budget.max_pages:
        _warn(result, f"スライドは {len(slides)} 枚中 {budget.max_pages} 枚まで読み取り、残りは未読です。")
    for number, name in enumerate(slides[:budget.max_pages], 1):
        root = _xml(zf, name, budget)
        locator = f"スライド {number}"
        hidden = root.get("show", "1") in {"0", "false", "off"}
        context.register(name, locator, hidden, root)
        parents = {id(child): parent for parent in root.iter() for child in parent}

        def is_hidden(node):
            while node is not None:
                if _shape_hidden(node):
                    return True
                node = parents.get(id(node))
            return hidden
        tables = [node for node in root.iter() if _local(node.tag) == "tbl"]
        table_paragraphs = {id(p) for table in tables for p in table.iter() if _local(p.tag) == "p"}
        handled = set(table_paragraphs)
        roles = {"sldNum": "slide_number", "dt": "date", "ftr": "footer", "hdr": "header"}
        headings = []
        for shape in root.iter():
            if _local(shape.tag) == "sp" and any(_local(n.tag) == "ph" and n.get("type") in {"title", "ctrTitle"} for n in shape.iter()):
                headings.append("\n".join(_paragraph_text(p) for p in shape.iter() if _local(p.tag) == "p"))
        heading = " / ".join(h for h in headings if h)
        shape_number = 0
        for shape in root.iter():
            if _local(shape.tag) != "sp":
                continue
            shape_number += 1
            paragraphs = [p for p in shape.iter() if _local(p.tag) == "p" and id(p) not in handled]
            handled.update(id(p) for p in paragraphs)
            text = "\n".join(_paragraph_text(p) for p in paragraphs if _paragraph_text(p))
            if not text:
                continue
            role = next((roles[n.get("type")] for n in shape.iter() if _local(n.tag) == "ph" and n.get("type") in roles), None)
            shape_hidden = is_hidden(shape)
            attributes = {"kind": "metadata" if role else "body", "hidden": shape_hidden,
                          "heading": heading, "parent_locator": locator}
            if role:
                attributes.update(role=role, candidate_reason="structural_metadata")
            elif shape_hidden:
                attributes["candidate_reason"] = "hidden_content"
            budget.add(result, locator if shape_number == 1 and not role else f"{locator} 図形 {shape_number}", text, **attributes)
        remaining = [_paragraph_text(p) for p in root.iter() if _local(p.tag) == "p" and id(p) not in handled]
        if any(remaining):
            budget.add(result, f"{locator} 本文", "\n".join(p for p in remaining if p), hidden=hidden, heading=heading)
        for table_number, table in enumerate(tables, 1):
            rows = [row for row in table if _local(row.tag) == "tr"]
            properties = next((child for child in table if _local(child.tag) == "tblPr"), None)
            has_header = properties is not None and _truth(properties.get("firstRow"))
            headers = [_paragraph_text(cell) for cell in rows[0] if _local(cell.tag) == "tc"] if has_header and rows else []
            for row_number, row in enumerate(rows, 1):
                values = [_paragraph_text(cell) for cell in row if _local(cell.tag) == "tc"]
                budget.add(result, f"{locator} 表 {table_number} 行 {row_number}", "\t".join(values),
                           kind="table_row", hidden=is_hidden(table), table_id=f"{name}#table-{table_number}",
                           table_headers=headers, table_header_row=has_header and row_number == 1,
                           values=values, cells=[], heading=heading, parent_locator=locator)
        for target, relation_type in context.links(zf, name, budget).values():
            if relation_type != "notesSlide":
                continue
            note = _xml(zf, target, budget)
            context.register(target, f"{locator} ノート", hidden, note)
            texts = []
            for shape in note.iter():
                if _local(shape.tag) == "sp":
                    # 従来のノート用テンプレート番号等の抑制は維持する。
                    if any(_local(n.tag) == "ph" and n.get("type") in {"sldNum", "dt", "ftr", "hdr"} for n in shape.iter()):
                        continue
                    texts.extend(_paragraph_text(p) for p in shape.iter() if _local(p.tag) == "p")
            if any(texts):
                budget.add(result, f"{locator} ノート", "\n".join(texts), kind="note", hidden=hidden,
                           heading=heading, parent_locator=locator, candidate_reason="speaker_notes")


def _merge_child(result: Extraction, child: Extraction, locator: str, *, sources=None, embedded=False):
    if child.metadata.get("retryable"):
        result.metadata["retryable"] = True
    for unit in child.units:
        merged = deepcopy(unit)
        merged["locator"] = f"{locator} / {unit['locator']}"
        if unit.get("table_id"):
            merged["table_id"] = f"{locator} / {unit['table_id']}"
        if unit.get("parent_locator"):
            merged["parent_locator"] = f"{locator} / {unit['parent_locator']}"
        references = unit.get("source_refs", [unit["locator"]])
        if sources:
            merged["source_refs"] = list(dict.fromkeys(f"{parent} / {locator} / {ref}" for parent, _ in sources for ref in references))
            merged["parent_locators"] = list(dict.fromkeys(parent for parent, _ in sources))
            merged["source_visibility"] = {parent: all(h for p, h in sources if p == parent) for parent, _ in sources}
            merged["hidden"] = unit.get("hidden", False) or all(is_hidden for _, is_hidden in sources)
            merged["note"] = unit.get("note", False) or unit.get("kind") == "note" or all(parent.endswith(" ノート") for parent, _ in sources)
        else:
            merged["source_refs"] = [f"{locator} / {ref}" for ref in references]
        if embedded:
            merged["content_kind"] = unit.get("kind", "body")
            merged["kind"] = "embedded"
            merged["candidate_reason"] = "embedded_document"
        if sources == []:
            merged["candidate_reason"] = "unreferenced_asset"
        result.units.append(merged)
    for warning in child.warnings:
        _warn(result, f"{locator}: {warning}")
    if child.status not in {"ok", "empty"} and not child.warnings:
        _warn(result, f"{locator}: 読み取り状態 {child.status}")


def _guess_office_suffix(data: bytes) -> str | None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as embedded:
            names = set(embedded.namelist())
            if "word/document.xml" in names:
                return ".docx"
            if "xl/workbook.xml" in names:
                return ".xlsx"
            if "ppt/presentation.xml" in names:
                return ".pptx"
    except zipfile.BadZipFile:
        pass
    return None


def _office_assets(zf: zipfile.ZipFile, result: Extraction, budget: _Budget, depth: int, context: _OfficeContext):
    images = embedded = 0
    sources = context.asset_sources(zf, budget)
    with tempfile.TemporaryDirectory(prefix="contextgen-kai-office-") as tmp:
        for index, info in enumerate(zf.infolist()):
            name = info.filename.replace("\\", "/")
            if info.is_dir() or ("/media/" not in name and "/embeddings/" not in name):
                continue
            budget.check()
            locator = f"埋め込み {name}"
            try:
                data = _read_member(zf, info, budget)
                suffix = PurePosixPath(name).suffix.lower()
                if "/media/" in name:
                    images += 1
                    if suffix not in IMAGE_EXTENSIONS:
                        _warn(result, f"{locator}: この画像形式のOCRは未対応です。原本を確認してください。")
                        continue
                else:
                    embedded += 1
                    if suffix == ".bin":
                        suffix = _guess_office_suffix(data) or suffix
                    if suffix not in SUPPORTED_EXTENSIONS:
                        _warn(result, f"{locator}: 旧OLE形式などのため抽出できません。元の資料をOffice新形式で保存して別途登録してください。")
                        continue
                # 元ZIP内のパスは使わず、管理下の一時ファイル名だけを利用。
                local = Path(tmp) / f"asset-{index}{suffix}"
                local.write_bytes(data)
                child = _extract(local, budget, depth + 1)
                _merge_child(result, child, locator, sources=sources.get(name, []), embedded="/embeddings/" in name)
                result.metadata["ocr_units"] = result.metadata.get("ocr_units", 0) + child.metadata.get("ocr_units", 0)
                result.metadata["ocr_cache_hits"] = result.metadata.get("ocr_cache_hits", 0) + child.metadata.get("ocr_cache_hits", 0)
            except _Limit:
                raise
            except Exception as exc:
                result.metadata["retryable"] = True
                _warn(result, f"{locator}: 読み取れませんでした（{type(exc).__name__}）。")
    result.metadata["embedded_images"] = images
    result.metadata["embedded_documents"] = embedded


def _archive(path: Path, result: Extraction, budget: _Budget, depth: int):
    with zipfile.ZipFile(path) as zf, tempfile.TemporaryDirectory(prefix="contextgen-kai-zip-") as tmp:
        result.metadata["archive_entries"] = len(zf.infolist())
        if len(zf.infolist()) > MAX_ARCHIVE_ENTRIES:
            _warn(result, f"ZIPの項目数上限 {MAX_ARCHIVE_ENTRIES:,} 件を超えた部分は未読です。")
        for index, info in enumerate(zf.infolist()[:MAX_ARCHIVE_ENTRIES]):
            budget.check()
            if info.is_dir():
                continue
            locator = f"ZIP内 {info.filename}"
            try:
                _safe_info(info, budget)
                suffix = PurePosixPath(info.filename.replace("\\", "/")).suffix.lower()
                if suffix not in SUPPORTED_EXTENSIONS:
                    _warn(result, f"{locator}: 非対応形式のため未読です。")
                    continue
                if suffix == ".zip" and depth >= MAX_DEPTH:
                    _warn(result, f"{locator}: 入れ子の上限 {MAX_DEPTH} 階層を超えたため未読です。")
                    continue
                data = _read_member(zf, info, budget)
                local = Path(tmp) / f"entry-{index}{suffix}"
                local.write_bytes(data)
                child = _extract(local, budget, depth + 1)
                _merge_child(result, child, locator)
            except _Limit:
                raise
            except ValueError as exc:
                _warn(result, f"{locator}: {exc}のため読み取りを省略しました。")
            except Exception as exc:
                result.metadata["retryable"] = True
                _warn(result, f"{locator}: 読み取りに失敗しました（{type(exc).__name__}）。")
