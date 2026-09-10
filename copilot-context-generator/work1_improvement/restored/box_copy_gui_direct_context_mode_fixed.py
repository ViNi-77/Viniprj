# =====================================================================
# 忠実復元版: box_copy_gui_direct_context_mode_fixed.pyw
# CopilotM365ContextGenerator_2026-05-28.exe (PyInstaller / Python 3.12)
# の逆アセンブル (analysis/box_copy_gui_disassembly.txt) から再構成。
# 動作・定数・文字列は原本と同一。このファイルは以後変更しない。
# =====================================================================
import os
import re
import json
import shutil
import threading
import queue
import tempfile
import zipfile
from pathlib import Path
from datetime import datetime
import tkinter as tk
from tkinter import ttk, messagebox, filedialog

try:
    import docx
except Exception:
    docx = None

try:
    from pypdf import PdfReader
except Exception:
    PdfReader = None

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except Exception:
    openpyxl = None
    get_column_letter = None

try:
    from pptx import Presentation
except Exception:
    Presentation = None

global SEARCH_ROOT, CONTEXT_OUTPUT_DIR, CONTEXT_MD, CONTEXT_JSONL
global CONTEXT_SOURCE_DIR, DIRECT_CONTEXT_FILES, worker_thread

SEARCH_ROOT_DEFAULT = Path(
    r"C:\BoxDrive\Box\第2ユニット生技部_全員\002_室運営\2026年\NZB00_アルミ開発室"
)

SEARCH_ROOT = SEARCH_ROOT_DEFAULT

CONTEXT_OUTPUT_DIR = Path(
    r"C:\BoxDrive\Box\第2ユニット生技部_全員\002_室運営\2026年\NZB00_アルミ開発室\2G\401_内製ソフト\ChatGPT用コンテキストファイル生成\生成コンテキスト"
)

BASE_DIR = Path(
    r"C:\BoxDrive\Box\第2ユニット生技部_全員\002_室運営\2026年\NZB00_アルミ開発室\2G\401_内製ソフト\ChatGPT用コンテキストファイル生成\管理"
)

LOG_PATH = BASE_DIR / "box_copy_gui_log.txt"
SETTINGS_PATH = BASE_DIR / "box_copy_gui_settings.json"

CONTEXT_MD = CONTEXT_OUTPUT_DIR / "_AI_CONTEXT_SUMMARY.md"
CONTEXT_JSONL = CONTEXT_OUTPUT_DIR / "_AI_CONTEXT_DATA.jsonl"


def refresh_context_output_paths():
    global CONTEXT_MD, CONTEXT_JSONL
    CONTEXT_MD = CONTEXT_OUTPUT_DIR / "_AI_CONTEXT_SUMMARY.md"
    CONTEXT_JSONL = CONTEXT_OUTPUT_DIR / "_AI_CONTEXT_DATA.jsonl"


KEYWORDS = []

CONTEXT_SOURCE_DIR = SEARCH_ROOT
DIRECT_CONTEXT_FILES = []

TARGET_EXTENSIONS = ['.doc', '.docx', '.txt', '.pdf', '.xlsx', '.xlsm', '.pptx', '.pptm', '.zip']
ZIP_READABLE_EXTENSIONS = [ext for ext in TARGET_EXTENSIONS if ext != '.zip']

EXCLUDE_DIRS = ['ChatGPT用コンテキストファイル生成', '収集データ', '生成コンテキスト', '管理']

MAX_TEXT_CHARS_PER_CONTEXT_RECORD = 12000
MAX_MARKDOWN_CHARS_PER_FILE = 120000
COPYRIGHT_TEXT = 'Copyright (c) 2026 第２ユニット生技部アルミ加工開発室. All rights reserved.'
M365_CONTEXT_BASENAME = 'M365AgentContext'
M365_CONTEXT_TARGET_CHARS_PER_FILE = 30000
M365_CONTEXT_MAX_BODY_FILES = 19
M365_CONTEXT_ARCHIVE_DIR_NAME = 'Archive'
MAX_SUMMARY_CHARS = 1200
MAX_ZIP_FILES = 300
MAX_ZIP_ENTRY_BYTES = 52428800
MAX_ZIP_TOTAL_BYTES = 314572800

stop_event = threading.Event()
log_queue = queue.Queue()
worker_thread = None


def now_text():
    return datetime.now().strftime('%Y-%m-%d %H:%M:%S')


def filename_timestamp():
    return datetime.now().strftime('%Y%m%d_%H%M%S')


def format_size(size):
    if size >= 1073741824:
        return f"{size / 1073741824:.2f} GB"
    if size >= 1048576:
        return f"{size / 1048576:.2f} MB"
    if size >= 1024:
        return f"{size / 1024:.2f} KB"
    return f"{size} B"


def clean_filename(name):
    return re.sub(r'[\\/:*?"<>|]', '_', name)


def is_same_or_child(path, parent):
    try:
        path.resolve().relative_to(parent.resolve())
        return True
    except ValueError:
        return False
    except Exception:
        return False


def load_settings():
    try:
        if SETTINGS_PATH.exists():
            with open(SETTINGS_PATH, 'r', encoding='utf-8') as f:
                data = json.load(f)
            return {
                'search_root': data.get('search_root', str(SEARCH_ROOT_DEFAULT)),
                'context_output_dir': data.get('context_output_dir', str(CONTEXT_OUTPUT_DIR)),
            }
    except Exception:
        pass
    return {
        'search_root': str(SEARCH_ROOT_DEFAULT),
        'context_output_dir': str(CONTEXT_OUTPUT_DIR),
    }


def save_settings(search_root, context_output_dir):
    try:
        BASE_DIR.mkdir(parents=True, exist_ok=True)
        data = {
            'search_root': str(search_root),
            'context_output_dir': str(context_output_dir),
            'saved_at': now_text(),
        }
        with open(SETTINGS_PATH, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
    except Exception as e:
        send_log(f"設定保存エラー: {e}")


def send_log(message):
    text = f"[{now_text()}] {message}"
    log_queue.put(('log', text))
    try:
        BASE_DIR.mkdir(parents=True, exist_ok=True)
        with open(LOG_PATH, 'a', encoding='utf-8') as f:
            f.write(text + '\n')
    except Exception:
        pass


def matched_file(path):
    if path.suffix.lower() not in TARGET_EXTENSIONS:
        return False
    if len(KEYWORDS) == 0:
        return True
    name = path.name.lower()
    return any(keyword.lower() in name for keyword in KEYWORDS)


def push_progress(counters):
    percent = counters['scan_checked'] % 100
    log_queue.put(('progress', {
        'percent': percent,
        'scan_checked': counters['scan_checked'],
        'target': counters['target'],
        'error': counters['error'],
        'referenced_bytes': counters['referenced_bytes'],
        'ext': counters['ext'].copy(),
    }))


def normalize_text(text):
    text = text.replace('\r', '\n')
    text = re.sub(r'\n{3,}', '\n\n', text)
    text = re.sub(r'[ \t]+', ' ', text)
    return text.strip()


def read_text_file(path):
    for enc in ('utf-8-sig', 'cp932', 'utf-8'):
        try:
            return path.read_text(encoding=enc, errors='ignore')
        except Exception:
            continue
    return ''


def read_docx_file(path):
    if docx is None:
        return ''
    try:
        document = docx.Document(str(path))
        lines = []
        for p in document.paragraphs:
            t = p.text.strip()
            if not t:
                continue
            lines.append(t)
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if not cells:
                    continue
                lines.append(' | '.join(cells))
        return '\n'.join(lines)
    except Exception:
        return ''


def read_pdf_file(path):
    if PdfReader is None:
        return ''
    try:
        reader = PdfReader(str(path))
        pages = []
        for page_index, page in enumerate(reader.pages, start=1):
            if stop_event.is_set():
                break
            text = page.extract_text() or ''
            if not text.strip():
                continue
            pages.append(f"# Page {page_index}")
            pages.append(text)
        return '\n'.join(pages)
    except Exception:
        return ''


def read_excel_file(path):
    if openpyxl is None:
        return ''
    try:
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        lines = []
        for ws in wb.worksheets:
            if stop_event.is_set():
                break
            lines.append(f"# Sheet: {ws.title}")
            for row_index, row in enumerate(ws.iter_rows(values_only=True), start=1):
                if stop_event.is_set():
                    break
                values = []
                for col_index, value in enumerate(row, start=1):
                    if value is None:
                        continue
                    cell_text = str(value).strip()
                    if not cell_text:
                        continue
                    if get_column_letter is not None:
                        cell_name = f"{get_column_letter(col_index)}{row_index}"
                    else:
                        cell_name = f"R{row_index}C{col_index}"
                    values.append(f"{cell_name}={cell_text}")
                if not values:
                    continue
                lines.append(f"Row {row_index}: " + ' | '.join(values))
        wb.close()
        return '\n'.join(lines)
    except Exception:
        return ''


def append_unique_text(lines, seen, text):
    text = normalize_text(text or '')
    if not text:
        return
    key = text.lower()
    if key in seen:
        return
    seen.add(key)
    lines.append(text)


def extract_ooxml_texts(element):
    texts = []
    try:
        for child in element.iter():
            tag = str(child.tag)
            if tag.endswith('}t') or tag.endswith('}v') or tag == 't' or tag == 'v':
                text = (child.text or '').strip()
                if not text:
                    continue
                texts.append(text)
    except Exception:
        pass
    return texts


def extract_shape_alt_texts(shape):
    texts = []
    try:
        for child in shape._element.iter():
            if not str(child.tag).endswith('}cNvPr'):
                continue
            for attr_name in ('descr', 'title'):
                text = (child.get(attr_name) or '').strip()
                if not text:
                    continue
                texts.append(text)
    except Exception:
        pass
    return texts


def collect_ppt_shape_text(shape, lines, seen):
    if hasattr(shape, 'text'):
        append_unique_text(lines, seen, shape.text)

    if hasattr(shape, 'table'):
        try:
            table = shape.table
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if not cells:
                    continue
                append_unique_text(lines, seen, ' | '.join(cells))
        except Exception:
            pass

    try:
        if getattr(shape, 'has_chart', False):
            chart = shape.chart

            try:
                append_unique_text(lines, seen, chart.chart_title.text_frame.text)
            except Exception:
                pass

            for axis_name in ('category_axis', 'value_axis'):
                try:
                    axis = getattr(chart, axis_name)
                    append_unique_text(lines, seen, axis.axis_title.text_frame.text)
                except Exception:
                    continue
    except Exception:
        pass

    for alt_text in extract_shape_alt_texts(shape):
        append_unique_text(lines, seen, alt_text)

    try:
        for child_shape in shape.shapes:
            collect_ppt_shape_text(child_shape, lines, seen)
    except Exception:
        pass


def read_pptx_file(path):
    if Presentation is None:
        return ''
    try:
        prs = Presentation(str(path))
        lines = []

        for slide_index, slide in enumerate(prs.slides, start=1):
            if stop_event.is_set():
                break

            lines.append(f"# Slide {slide_index}")
            seen = set()

            for shape in slide.shapes:
                if stop_event.is_set():
                    break

                collect_ppt_shape_text(shape, lines, seen)

            try:
                notes_text = slide.notes_slide.notes_text_frame.text.strip()
                if notes_text:
                    lines.append('## Notes')
                    append_unique_text(lines, seen, notes_text)
            except Exception:
                pass

            xml_texts = []
            for xml_text in extract_ooxml_texts(slide._element):
                normalized = normalize_text(xml_text)
                if normalized and normalized.lower() not in seen:
                    xml_texts.append(normalized)

            if not xml_texts:
                continue
            lines.append('## XML Text')
            for xml_text in xml_texts:
                append_unique_text(lines, seen, xml_text)

        return '\n'.join(lines)
    except Exception:
        return ''


def read_zip_file(path):
    try:
        if not zipfile.is_zipfile(path):
            return ''

        lines = [f"# ZIP: {path.name}"]
        read_count = 0
        total_bytes = 0

        with zipfile.ZipFile(path) as zf, tempfile.TemporaryDirectory(prefix='ai_context_zip_') as tmp_dir:
            for info in sorted(zf.infolist(), key=lambda item: item.filename.lower()):
                if stop_event.is_set():
                    break

                if info.is_dir():
                    continue

                inner_name = info.filename.replace('\\', '/')
                inner_path = Path(inner_name)

                if not inner_path.name:
                    continue

                if inner_path.name.startswith('_AI_CONTEXT_'):
                    continue

                ext = inner_path.suffix.lower()
                if ext not in ZIP_READABLE_EXTENSIONS:
                    continue

                if read_count >= MAX_ZIP_FILES:
                    lines.append(f"...（ZIP内ファイルは{MAX_ZIP_FILES}件で抽出打ち切り）")
                    break

                if info.file_size > MAX_ZIP_ENTRY_BYTES:
                    lines.append(f"## ZIP内ファイル: {inner_name}")
                    lines.append(f"...（{format_size(info.file_size)}のため抽出スキップ）")
                    continue

                if total_bytes + info.file_size > MAX_ZIP_TOTAL_BYTES:
                    lines.append(f"...（ZIP内の合計抽出サイズが{format_size(MAX_ZIP_TOTAL_BYTES)}を超えるため抽出打ち切り）")
                    break

                read_count += 1
                total_bytes += info.file_size

                temp_name = f"{read_count:04d}_{clean_filename(inner_path.name)}"
                temp_path = Path(tmp_dir) / temp_name

                with zf.open(info) as src, open(temp_path, 'wb') as dst:
                    shutil.copyfileobj(src, dst)

                inner_text = normalize_text(extract_text(temp_path))

                lines.append(f"## ZIP内ファイル {read_count}: {inner_name}")
                lines.append(f"- 種別: `{ext}`")
                lines.append(f"- サイズ: {format_size(info.file_size)}")
                lines.append('')

                if inner_text:
                    lines.append(inner_text)
                else:
                    lines.append('本文抽出不可。ZIP内ファイル名・パス・拡張子を参照。')

                lines.append('')

        if read_count == 0:
            lines.append('対応形式のファイルはZIP内に見つかりませんでした。')

        return '\n'.join(lines)
    except Exception:
        return ''


def extract_text(path):
    ext = path.suffix.lower()

    if ext == '.txt':
        return read_text_file(path)

    if ext == '.docx':
        return read_docx_file(path)

    if ext == '.pdf':
        return read_pdf_file(path)

    if ext in ('.xlsx', '.xlsm'):
        return read_excel_file(path)

    if ext in ('.pptx', '.pptm'):
        return read_pptx_file(path)

    if ext == '.zip':
        return read_zip_file(path)

    if ext == '.doc':
        return ''

    return ''


def simple_summary(text):
    text = normalize_text(text)

    if not text:
        return '本文抽出不可。ファイル名・パス・拡張子を参照。'

    lines = [line.strip() for line in text.splitlines() if line.strip()]
    head = '\n'.join(lines[:20])

    if len(head) > MAX_SUMMARY_CHARS:
        head = head[:MAX_SUMMARY_CHARS] + '...'

    return head


def split_text_for_context(text, max_chars=MAX_TEXT_CHARS_PER_CONTEXT_RECORD):
    text = normalize_text(text)

    if not text:
        return ['']

    chunks = []
    current = []
    current_len = 0

    for line in text.splitlines():
        line_len = len(line) + 1

        if line_len > max_chars:
            if current:
                chunks.append('\n'.join(current))
                current = []
                current_len = 0

            for start in range(0, len(line), max_chars):
                chunks.append(line[start:start + max_chars])

            continue

        if current and current_len + line_len > max_chars:
            chunks.append('\n'.join(current))
            current = []
            current_len = 0

        current.append(line)
        current_len += line_len

    if current:
        chunks.append('\n'.join(current))

    return chunks or ['']


def markdown_code_block(text):
    text = text or '本文抽出不可。ファイル名・パス・拡張子を参照。'
    max_backticks = 3

    for match in re.finditer('`+', text):
        max_backticks = max(max_backticks, len(match.group(0)) + 1)

    fence = '`' * max_backticks
    return [f"{fence}text", text, fence]


def context_markdown_header(generated_at, total, part_no):
    return [
        '# AI用コンテキスト要約',
        '',
        f"- 生成日時: {generated_at}",
        f"- 分割番号: {part_no}",
        f"- コンテキスト生成元フォルダ: `{CONTEXT_SOURCE_DIR}`",
        f"- 生成コンテキスト保存先: `{CONTEXT_OUTPUT_DIR}`",
        f"- 対象ファイル数: {total}",
        f"- 1ファイルあたり最大文字数: {MAX_MARKDOWN_CHARS_PER_FILE}",
        '',
        '## 全体インデックス',
        '',
    ]


def write_latest_markdown_manifest(generated_at, part_paths):
    lines = [
        '# AI用コンテキスト要約ファイル一覧',
        '',
        f"- 生成日時: {generated_at}",
        f"- 分割ファイル数: {len(part_paths)}",
        '',
        '## 分割ファイル',
        '',
    ]

    for path in part_paths:
        lines.append(f"- `{path.name}`")

    CONTEXT_MD.write_text('\n'.join(lines) + '\n', encoding='utf-8')


def m365_fixed_file_pattern():
    return re.compile(f"^{re.escape(M365_CONTEXT_BASENAME)}_(INDEX|\\d{{3}})\\.txt$")


def cleanup_m365_fixed_files():
    pattern = m365_fixed_file_pattern()

    for path in CONTEXT_OUTPUT_DIR.glob(f"{M365_CONTEXT_BASENAME}_*.txt"):
        if not path.is_file():
            continue
        if not pattern.match(path.name):
            continue
        try:
            path.unlink()
        except Exception as e:
            send_log(f"M365固定名ファイル削除エラー: {path}")
            send_log(str(e))


def m365_part_header(generated_at, generated_stamp, part_no, part_count, total_files, total_chunks):
    return [
        f"# M365 Agent Context {part_no:03d}",
        '',
        f"GeneratedAt: {generated_at}",
        f"GeneratedStamp: {generated_stamp}",
        f"Part: {part_no:03d} / {part_count:03d}",
        f"SourceScope: {CONTEXT_SOURCE_DIR}",
        'ContentType: measurement-context',
        'FileRole: body',
        f"SourceFileCount: {total_files}",
        f"SourceChunkCount: {total_chunks}",
        f"TargetCharactersPerFile: {M365_CONTEXT_TARGET_CHARS_PER_FILE}",
        '',
        '## Retrieval Notes',
        '- This file is generated for Microsoft 365 Copilot agent knowledge.',
        '- Use SourcePath, FileName, Page, Sheet, Row, Slide, and Chunk fields when answering.',
        '- Prefer precise source references when the answer depends on a specific file.',
        '',
        '## Content',
        '',
    ]


def build_m365_section(index, chunk_index, chunk_count, path, relative, size, text_length, summary, chunk_text):
    lines = [
        f"### Source {index}-{chunk_index}",
        '',
        f"FileName: {path.name}",
        f"FileType: {path.suffix.lower()}",
        f"SourcePath: {relative}",
        f"Size: {format_size(size)}",
        f"TextLength: {text_length}",
        f"Chunk: {chunk_index} / {chunk_count}",
        '',
        'Summary:',
        summary or '本文抽出不可。ファイル名・パス・拡張子を参照。',
        '',
        'Text:',
        chunk_text or '本文抽出不可。ファイル名・パス・拡張子を参照。',
        '',
        '---',
    ]

    return '\n'.join(lines)


def write_m365_agent_context_files(generated_at, generated_stamp, m365_parts, m365_entries, overflow_count, overflow_chars):
    cleanup_m365_fixed_files()

    archive_dir = CONTEXT_OUTPUT_DIR / M365_CONTEXT_ARCHIVE_DIR_NAME
    archive_dir.mkdir(parents=True, exist_ok=True)

    if not m365_parts:
        m365_parts = [['No content was extracted.']]

    part_count = len(m365_parts)
    fixed_paths = []
    archive_paths = []
    total_files = len({entry['relative_path'] for entry in m365_entries})
    total_chunks = len(m365_entries)

    for part_no, section_texts in enumerate(m365_parts, start=1):
        fixed_path = CONTEXT_OUTPUT_DIR / f"{M365_CONTEXT_BASENAME}_{part_no:03d}.txt"
        archive_path = archive_dir / f"{M365_CONTEXT_BASENAME}_{generated_stamp}_{part_no:03d}.txt"
        body_lines = m365_part_header(generated_at, generated_stamp, part_no, part_count, total_files, total_chunks)
        body_lines.extend(section_texts)
        text = '\n'.join(body_lines).rstrip() + '\n'

        fixed_path.write_text(text, encoding='utf-8')
        archive_path.write_text(text, encoding='utf-8')
        fixed_paths.append(fixed_path)
        archive_paths.append(archive_path)

    fixed_index_path = CONTEXT_OUTPUT_DIR / f"{M365_CONTEXT_BASENAME}_INDEX.txt"
    archive_index_path = archive_dir / f"{M365_CONTEXT_BASENAME}_{generated_stamp}_INDEX.txt"

    index_lines = [
        '# M365 Agent Context Index',
        '',
        f"GeneratedAt: {generated_at}",
        f"GeneratedStamp: {generated_stamp}",
        f"SourceScope: {CONTEXT_SOURCE_DIR}",
        'ContentType: measurement-context',
        'FileRole: index',
        f"BodyFileCount: {part_count}",
        f"TotalFileCountForUpload: {part_count + 1}",
        f"UploadLimitPolicy: 1 index file + up to {M365_CONTEXT_MAX_BODY_FILES} body files = 20 files",
        f"TargetCharactersPerBodyFile: {M365_CONTEXT_TARGET_CHARS_PER_FILE}",
        f"SourceChunkCount: {total_chunks}",
        f"OverflowChunkCount: {overflow_count}",
        f"OverflowCharacterCount: {overflow_chars}",
        '',
        '## How To Use',
        '- Add this INDEX file and all body files listed in RequiredFiles to the Microsoft 365 Copilot agent knowledge sources.',
        '- Keep the fixed file names unchanged so the agent can continue to reference the same files after regeneration.',
        '- The Archive folder keeps timestamped copies for traceability and normally does not need to be added to the agent.',
        '',
        '## RequiredFiles',
        f"- {fixed_index_path.name}",
    ]

    for path in fixed_paths:
        index_lines.append(f"- {path.name}")

    index_lines.extend([
        '',
        '## BodyFileMap',
    ])

    for part_no, path in enumerate(fixed_paths, start=1):
        entry_count = sum(1 for entry in m365_entries if entry.get('m365_part') == part_no)
        index_lines.append(f"- {path.name}: {entry_count} chunks")

    index_lines.extend([
        '',
        '## SourceChunkIndex',
    ])

    for entry in m365_entries:
        index_lines.append(
            f"- Part {entry['m365_part']:03d} | {entry['file_name']} | Chunk "
            f"{entry['chunk_index']}/{entry['chunk_count']} | {entry['relative_path']}"
        )

    if overflow_count:
        index_lines.extend([
            '',
            '## OverflowWarning',
            f"- {overflow_count} chunks were not included in the fixed 20-file M365 package.",
            '- The full extracted text remains available in _AI_CONTEXT_DATA.jsonl and the timestamped Markdown files.',
            '- Narrow the source folder, split by topic, or create multiple agents if all overflow content is required.',
        ])

    index_text = '\n'.join(index_lines).rstrip() + '\n'
    fixed_index_path.write_text(index_text, encoding='utf-8')
    archive_index_path.write_text(index_text, encoding='utf-8')

    for path in [fixed_index_path] + fixed_paths:
        send_log(f"M365 Copilot投入用TXT作成: {path}")

    send_log(f"M365 Copilot履歴保存先: {archive_dir}")

    return fixed_index_path, fixed_paths, archive_index_path, archive_paths


def build_ai_context(source_files=None):
    CONTEXT_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    send_log('AI用コンテキスト生成開始')

    files = []

    if source_files is not None:
        for path in source_files:
            if stop_event.is_set():
                send_log('AI用コンテキスト生成を停止しました')
                break

            path = Path(path)

            if not path.is_file():
                continue

            if path.name.startswith('_AI_CONTEXT_'):
                continue

            if path.suffix.lower() not in TARGET_EXTENSIONS:
                continue

            files.append(path)
    else:
        for path in CONTEXT_SOURCE_DIR.rglob('*'):
            if stop_event.is_set():
                send_log('AI用コンテキスト生成を停止しました')
                break

            if not path.is_file():
                continue

            if path.name.startswith('_AI_CONTEXT_'):
                continue

            if path.suffix.lower() not in TARGET_EXTENSIONS:
                continue

            files.append(path)

    files = sorted(files, key=lambda p: str(p).lower())
    total = len(files)

    send_log(f"AI用コンテキスト対象ファイル数: {total}")

    generated_at = now_text()
    generated_stamp = filename_timestamp()
    md_part_no = 1
    md_part_paths = []
    md_lines = context_markdown_header(generated_at, total, md_part_no)
    md_header_line_count = len(md_lines)
    md_char_count = len('\n'.join(md_lines)) + 1
    jsonl_record_count = 0
    m365_parts = [[]]
    m365_part_lengths = [0]
    m365_entries = []
    m365_overflow_count = 0
    m365_overflow_chars = 0

    def flush_markdown_part():
        nonlocal md_part_no, md_lines, md_char_count

        part_path = CONTEXT_OUTPUT_DIR / f"_AI_CONTEXT_SUMMARY_{generated_stamp}_{md_part_no:03d}.md"
        part_path.write_text('\n'.join(md_lines).rstrip() + '\n', encoding='utf-8')
        md_part_paths.append(part_path)

        md_part_no += 1
        md_lines = context_markdown_header(generated_at, total, md_part_no)
        md_char_count = len('\n'.join(md_lines)) + 1

    def append_markdown_section(section_lines):
        nonlocal md_char_count

        section_text = '\n'.join(section_lines) + '\n'

        if md_char_count + len(section_text) > MAX_MARKDOWN_CHARS_PER_FILE and len(md_lines) > md_header_line_count:
            flush_markdown_part()

        md_lines.extend(section_lines)
        md_lines.append('')
        md_char_count += len(section_text) + 1

    def append_m365_section(section_text, index_entry):
        nonlocal m365_overflow_count, m365_overflow_chars

        text = section_text.rstrip() + '\n\n'
        current_index = len(m365_parts) - 1

        if m365_part_lengths[current_index] > 0 and m365_part_lengths[current_index] + len(text) > M365_CONTEXT_TARGET_CHARS_PER_FILE:
            if len(m365_parts) < M365_CONTEXT_MAX_BODY_FILES:
                m365_parts.append([])
                m365_part_lengths.append(0)
                current_index += 1
            else:
                m365_overflow_count += 1
                m365_overflow_chars += len(text)
                return

        index_entry['m365_part'] = current_index + 1
        m365_parts[current_index].append(text)
        m365_part_lengths[current_index] += len(text)
        m365_entries.append(index_entry)

    with open(CONTEXT_JSONL, 'w', encoding='utf-8') as jsonl_file:
        for index, path in enumerate(files, start=1):
            if stop_event.is_set():
                break

            try:
                log_queue.put(('current', f"AIコンテキスト生成中... {index}/{total} {path.name}"))

                size = path.stat().st_size
                text = normalize_text(extract_text(path))
                summary = simple_summary(text)
                try:
                    relative = path.relative_to(CONTEXT_SOURCE_DIR)
                except Exception:
                    try:
                        relative = path.relative_to(SEARCH_ROOT)
                    except Exception:
                        relative = Path(path.name)

                chunks = split_text_for_context(text)
                chunk_count = len(chunks)

                for chunk_index, chunk_text in enumerate(chunks, start=1):
                    jsonl_record_count += 1

                    record = {
                        'index': index,
                        'chunk_index': chunk_index,
                        'chunk_count': chunk_count,
                        'file_name': path.name,
                        'relative_path': str(relative),
                        'extension': path.suffix.lower(),
                        'size_bytes': size,
                        'size_text': format_size(size),
                        'summary': summary,
                        'text_length': len(text),
                        'text_excerpt': chunk_text,
                        'text': chunk_text,
                    }

                    jsonl_file.write(json.dumps(record, ensure_ascii=False) + '\n')

                    section_title = f"### {index}. {path.name}"
                    if chunk_count > 1:
                        section_title = f"### {index}-{chunk_index}. {path.name}"

                    section_lines = [
                        section_title,
                        '',
                        f"- パス: `{relative}`",
                        f"- 種別: `{path.suffix.lower()}`",
                        f"- サイズ: {format_size(size)}",
                        f"- 抽出文字数: {len(text)}",
                        f"- 本文分割: {chunk_index}/{chunk_count}",
                        '',
                        '#### 要約 / 抜粋',
                        '',
                        summary,
                        '',
                        '#### 抽出本文',
                        '',
                    ]

                    section_lines.extend(markdown_code_block(chunk_text))
                    section_lines.extend(['', '---', ''])
                    append_markdown_section(section_lines)

                    m365_section = build_m365_section(
                        index,
                        chunk_index,
                        chunk_count,
                        path,
                        relative,
                        size,
                        len(text),
                        summary,
                        chunk_text,
                    )
                    append_m365_section(m365_section, {
                        'index': index,
                        'chunk_index': chunk_index,
                        'chunk_count': chunk_count,
                        'file_name': path.name,
                        'relative_path': str(relative),
                    })

                if index % 20 == 0:
                    send_log(f"AIコンテキスト生成中: {index}/{total}")
            except Exception as e:
                send_log(f"AIコンテキスト生成エラー: {path}")
                send_log(str(e))

    if len(md_lines) > md_header_line_count or not md_part_paths:
        flush_markdown_part()

    write_latest_markdown_manifest(generated_at, md_part_paths)
    write_m365_agent_context_files(
        generated_at,
        generated_stamp,
        m365_parts,
        m365_entries,
        m365_overflow_count,
        m365_overflow_chars,
    )

    for part_path in md_part_paths:
        send_log(f"AI用Markdown作成: {part_path}")

    send_log(f"AI用Markdown一覧作成: {CONTEXT_MD}")
    send_log(f"AI用JSONL作成: {CONTEXT_JSONL}")
    send_log(f"AI用JSONLレコード数: {jsonl_record_count}")
    if m365_overflow_count:
        send_log(f"M365 Copilot投入用TXTの20ファイル上限により未収録チャンク数: {m365_overflow_count}")
    send_log('AI用コンテキスト生成完了')


def backup_worker():
    global DIRECT_CONTEXT_FILES

    stop_event.clear()
    DIRECT_CONTEXT_FILES = []

    counters = {
        'scan_checked': 0,
        'target': 0,
        'error': 0,
        'referenced_bytes': 0,
        'ext': {
            '.doc': 0,
            '.docx': 0,
            '.txt': 0,
            '.pdf': 0,
            '.xlsx': 0,
            '.xlsm': 0,
            '.pptx': 0,
            '.pptm': 0,
            '.zip': 0,
        },
    }

    try:
        CONTEXT_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

        send_log('====================================')
        send_log('参照元直接読み取りコンテキスト生成開始')
        send_log('収集データへのコピーは行いません')
        send_log(f"検索元: {SEARCH_ROOT}")
        send_log(f"AIコンテキスト保存先: {CONTEXT_OUTPUT_DIR}")

        if not SEARCH_ROOT.exists():
            send_log(f"検索元フォルダが存在しません: {SEARCH_ROOT}")
            log_queue.put(('finish', None))
            return

        stack = [SEARCH_ROOT]

        while stack:
            if stop_event.is_set():
                send_log('停止要求を受けました')
                break

            current_dir = stack.pop()
            log_queue.put(('current_dir', str(current_dir)))

            try:
                with os.scandir(current_dir) as entries:
                    for entry in entries:
                        if stop_event.is_set():
                            send_log('停止要求を受けました')
                            break

                        counters['scan_checked'] += 1

                        try:
                            if entry.is_dir(follow_symlinks=False):
                                dir_name = Path(entry.path).name

                                if dir_name in EXCLUDE_DIRS:
                                    send_log(f"除外フォルダスキップ: {entry.path}")
                                    continue

                                stack.append(Path(entry.path))
                            elif entry.is_file(follow_symlinks=False):
                                path = Path(entry.path)

                                if matched_file(path):
                                    counters['target'] += 1

                                    try:
                                        counters['referenced_bytes'] += path.stat().st_size
                                    except Exception:
                                        pass

                                    ext = path.suffix.lower()
                                    if ext in counters['ext']:
                                        counters['ext'][ext] += 1

                                    DIRECT_CONTEXT_FILES.append(path)
                                    log_queue.put(('current', str(path)))
                                    send_log(f"参照元ファイル発見: {path}")

                                    push_progress(counters)

                            if counters['scan_checked'] % 20 == 0:
                                log_queue.put((
                                    'current',
                                    f"検索中... 確認:{counters['scan_checked']} / 参照元ファイル:{counters['target']}",
                                ))
                                push_progress(counters)
                        except Exception as e:
                            counters['error'] += 1
                            send_log(f"ファイル処理エラー: {entry.path}")
                            send_log(str(e))
                            push_progress(counters)
                            continue
            except Exception as e:
                counters['error'] += 1
                send_log(f"フォルダ読込エラー: {current_dir}")
                send_log(str(e))
                push_progress(counters)

        push_progress(counters)

        send_log('====================================')
        send_log(f"確認件数: {counters['scan_checked']}")
        send_log(f"参照元ファイル数: {counters['target']}")
        send_log(f"エラー数: {counters['error']}")
        send_log(f"参照元ファイル容量: {format_size(counters['referenced_bytes'])}")
        send_log('参照元検索処理終了')
        send_log('====================================')

        if not stop_event.is_set():
            send_log(f"コンテキスト生成対象数: {len(DIRECT_CONTEXT_FILES)}")
            build_ai_context(DIRECT_CONTEXT_FILES)

        log_queue.put(('finish', None))
    except Exception as e:
        send_log(f"致命的エラー: {e}")
        log_queue.put(('finish', None))


root = tk.Tk()
root.title('Copilot M365用コンテキスト生成 GUI')
root.geometry('1080x800')

main = ttk.Frame(root, padding=10)
main.pack(fill='both', expand=True)

ttk.Label(
    main,
    text='Copilot M365用コンテキスト生成',
    font=('Meiryo', 18, 'bold'),
).pack(anchor='w')

settings = load_settings()

SEARCH_ROOT = Path(settings['search_root'])
CONTEXT_OUTPUT_DIR = Path(settings['context_output_dir'])
refresh_context_output_paths()


def browse_folder(var, title, fallback):
    current = var.get().strip().strip('"')
    initial_dir = current if current and Path(current).exists() else str(fallback)
    selected = filedialog.askdirectory(
        title=title,
        initialdir=initial_dir,
        mustexist=True,
    )

    if selected:
        var.set(selected)


def browse_search_root():
    browse_folder(search_root_var, '参照元フォルダを選択', SEARCH_ROOT_DEFAULT)


def browse_context_output_dir():
    browse_folder(context_output_dir_var, 'コンテキスト出力先フォルダを選択', CONTEXT_OUTPUT_DIR)


path_frame = ttk.LabelFrame(main, text='フォルダ設定', padding=10)
path_frame.pack(fill='x', pady=(5, 10))

search_root_var = tk.StringVar(value=str(SEARCH_ROOT))
context_output_dir_var = tk.StringVar(value=str(CONTEXT_OUTPUT_DIR))


def add_folder_row(parent, label_text, variable, button_command):
    row = ttk.Frame(parent)
    row.pack(fill='x', pady=2)

    ttk.Label(row, text=label_text, font=('Meiryo', 9), width=18).pack(side='left')
    entry = ttk.Entry(row, textvariable=variable, font=('Meiryo', 9))
    entry.pack(side='left', fill='x', expand=True, padx=(0, 5))
    ttk.Button(row, text='参照', command=button_command).pack(side='left')

    return entry


search_entry = add_folder_row(path_frame, '参照元:', search_root_var, browse_search_root)
context_output_entry = add_folder_row(path_frame, 'コンテキスト出力先:', context_output_dir_var, browse_context_output_dir)

progress_frame = ttk.LabelFrame(main, text='動作状況', padding=10)
progress_frame.pack(fill='x', pady=10)

progress_var = tk.IntVar(value=0)

progress_bar = ttk.Progressbar(
    progress_frame,
    orient='horizontal',
    mode='determinate',
    variable=progress_var,
    maximum=100,
)

progress_bar.pack(fill='x')

percent_label = ttk.Label(progress_frame, text='動作待機中', font=('Meiryo', 16, 'bold'))
percent_label.pack(anchor='center', pady=5)

current_dir_var = tk.StringVar(value='現在検索中フォルダ: 待機中')

ttk.Label(
    progress_frame,
    textvariable=current_dir_var,
    font=('Meiryo', 9, 'bold'),
    foreground='blue',
    wraplength=1030,
).pack(anchor='w', pady=(5, 0))

current_file_var = tk.StringVar(value='現在処理中ファイル: 待機中')

ttk.Label(
    progress_frame,
    textvariable=current_file_var,
    font=('Meiryo', 9),
    wraplength=1030,
).pack(anchor='w', pady=(5, 0))

counter_frame = ttk.LabelFrame(main, text='カウンター', padding=10)
counter_frame.pack(fill='x', pady=10)

scan_var = tk.StringVar(value='確認: 0')
target_var = tk.StringVar(value='参照元ファイル: 0')
error_var = tk.StringVar(value='エラー: 0')
size_var = tk.StringVar(value='参照元容量: 0 B')

for i, var in enumerate([scan_var, target_var, error_var, size_var]):
    ttk.Label(counter_frame, textvariable=var, font=('Meiryo', 12, 'bold')).grid(row=0, column=i, padx=10)

ext_frame = ttk.LabelFrame(main, text='ファイル種別ごとの参照元ファイル数', padding=10)
ext_frame.pack(fill='x', pady=10)

doc_var = tk.StringVar(value='.doc: 0')
docx_var = tk.StringVar(value='.docx: 0')
txt_var = tk.StringVar(value='.txt: 0')
pdf_var = tk.StringVar(value='.pdf: 0')
xlsx_var = tk.StringVar(value='.xlsx: 0')
xlsm_var = tk.StringVar(value='.xlsm: 0')
pptx_var = tk.StringVar(value='.pptx: 0')
pptm_var = tk.StringVar(value='.pptm: 0')
zip_var = tk.StringVar(value='.zip: 0')

for i, var in enumerate([doc_var, docx_var, txt_var, pdf_var, xlsx_var, xlsm_var, pptx_var, pptm_var, zip_var]):
    ttk.Label(ext_frame, textvariable=var, font=('Meiryo', 11)).grid(row=i // 5, column=i % 5, padx=12, pady=2)

button_frame = ttk.Frame(main)
button_frame.pack(fill='x', pady=10)


def start_backup():
    global worker_thread, SEARCH_ROOT, CONTEXT_OUTPUT_DIR, CONTEXT_SOURCE_DIR

    if worker_thread and worker_thread.is_alive():
        messagebox.showinfo('実行中', 'すでに処理中です')
        return

    selected_search_text = search_root_var.get().strip().strip('"')
    selected_context_text = context_output_dir_var.get().strip().strip('"')

    if not selected_search_text:
        messagebox.showerror('参照元エラー', '参照元フォルダを指定してください。')
        return

    if not selected_context_text:
        messagebox.showerror('コンテキスト出力先エラー', 'コンテキスト出力先フォルダを指定してください。')
        return

    selected_root = Path(selected_search_text).expanduser()
    selected_context_output = Path(selected_context_text).expanduser()

    if not (selected_root.exists() and selected_root.is_dir()):
        messagebox.showerror('参照元エラー', f"参照元フォルダが存在しません。\n{selected_root}")
        return

    try:
        selected_context_output.mkdir(parents=True, exist_ok=True)
    except Exception as e:
        messagebox.showerror('出力先エラー', f"コンテキスト出力先を作成できません。\n{e}")
        return

    blocked_roots = [selected_context_output, BASE_DIR]

    for blocked_root in blocked_roots:
        if is_same_or_child(selected_root, blocked_root):
            messagebox.showerror(
                '参照元エラー',
                f"コンテキスト出力先・管理フォルダは参照元に指定できません。\n指定: {selected_root}"
                f"\n対象外: {blocked_root}",
            )
            return

    SEARCH_ROOT = selected_root
    CONTEXT_OUTPUT_DIR = selected_context_output
    CONTEXT_SOURCE_DIR = SEARCH_ROOT
    refresh_context_output_paths()

    search_root_var.set(str(SEARCH_ROOT))
    context_output_dir_var.set(str(CONTEXT_OUTPUT_DIR))

    save_settings(SEARCH_ROOT, CONTEXT_OUTPUT_DIR)

    progress_var.set(0)
    percent_label.config(text='検索・コンテキスト生成中')
    current_dir_var.set('現在検索中フォルダ: 開始準備中...')
    current_file_var.set('現在処理中ファイル: 開始準備中...')

    scan_var.set('確認: 0')
    target_var.set('参照元ファイル: 0')
    error_var.set('エラー: 0')
    size_var.set('参照元容量: 0 B')

    doc_var.set('.doc: 0')
    docx_var.set('.docx: 0')
    txt_var.set('.txt: 0')
    pdf_var.set('.pdf: 0')
    xlsx_var.set('.xlsx: 0')
    xlsm_var.set('.xlsm: 0')
    pptx_var.set('.pptx: 0')
    pptm_var.set('.pptm: 0')
    zip_var.set('.zip: 0')

    start_button.config(state='disabled')
    stop_button.config(state='normal')

    stop_event.clear()

    worker_thread = threading.Thread(target=backup_worker, daemon=True)
    worker_thread.start()


def stop_backup():
    stop_event.set()
    current_file_var.set('現在処理中ファイル: 停止要求中... 現在の処理が終わるまで待機')
    send_log('停止ボタンが押されました')


start_button = ttk.Button(button_frame, text='開始', command=start_backup)
start_button.pack(side='left', padx=5)

stop_button = ttk.Button(button_frame, text='停止', command=stop_backup, state='disabled')
stop_button.pack(side='left', padx=5)

copyright_label = ttk.Label(
    main,
    text=COPYRIGHT_TEXT,
    font=('Meiryo', 8),
    foreground='gray',
)

copyright_label.pack(side='bottom', anchor='center', fill='x', pady=(2, 0))

log_frame = ttk.LabelFrame(main, text='ログ', padding=10)
log_frame.pack(fill='both', expand=True, pady=10)

log_text = tk.Text(log_frame, height=18, font=('Consolas', 9), wrap='none')
log_text.pack(side='left', fill='both', expand=True)

scrollbar = ttk.Scrollbar(log_frame, orient='vertical', command=log_text.yview)
scrollbar.pack(side='right', fill='y')

log_text.configure(yscrollcommand=scrollbar.set)


def update_gui_from_queue():
    try:
        while True:
            event_type, data = log_queue.get_nowait()

            if event_type == 'log':
                log_text.insert('end', data + '\n')
                log_text.see('end')
            elif event_type == 'current':
                current_file_var.set(f"現在処理中ファイル: {data}")
            elif event_type == 'current_dir':
                current_dir_var.set(f"現在検索中フォルダ: {data}")
            elif event_type == 'progress':
                progress_var.set(data['percent'])
                percent_label.config(text='検索・コンテキスト生成中')

                scan_var.set(f"確認: {data['scan_checked']}")
                target_var.set(f"参照元ファイル: {data['target']}")
                error_var.set(f"エラー: {data['error']}")
                size_var.set(f"参照元容量: {format_size(data['referenced_bytes'])}")

                ext = data['ext']

                doc_var.set(f".doc: {ext.get('.doc', 0)}")
                docx_var.set(f".docx: {ext.get('.docx', 0)}")
                txt_var.set(f".txt: {ext.get('.txt', 0)}")
                pdf_var.set(f".pdf: {ext.get('.pdf', 0)}")
                xlsx_var.set(f".xlsx: {ext.get('.xlsx', 0)}")
                xlsm_var.set(f".xlsm: {ext.get('.xlsm', 0)}")
                pptx_var.set(f".pptx: {ext.get('.pptx', 0)}")
                pptm_var.set(f".pptm: {ext.get('.pptm', 0)}")
                zip_var.set(f".zip: {ext.get('.zip', 0)}")
            elif event_type == 'finish':
                start_button.config(state='normal')
                stop_button.config(state='disabled')

                if stop_event.is_set():
                    percent_label.config(text='停止')
                    current_file_var.set('現在処理中ファイル: 停止しました')
                else:
                    progress_var.set(100)
                    percent_label.config(text='完了')
                    current_file_var.set('現在処理中ファイル: 参照元直接読み取り + AIコンテキスト生成完了')
    except queue.Empty:
        pass

    root.after(200, update_gui_from_queue)


def on_close():
    if worker_thread and worker_thread.is_alive():
        stop_event.set()
        messagebox.showinfo('停止中', '処理中です。停止要求を出しました。数秒後に閉じてください。')
        return

    root.destroy()


root.protocol('WM_DELETE_WINDOW', on_close)

update_gui_from_queue()
root.mainloop()
