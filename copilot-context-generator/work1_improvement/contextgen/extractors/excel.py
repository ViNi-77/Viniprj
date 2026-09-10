"""Excel (.xlsx / .xlsm) 抽出.

原本と同一の整形（# Sheet: / Row N: A1=値 | B1=値）。
改善②: 暗号化ブック（OLE コンテナ）を protected として分類。
v3 柱1: シート単位で保護し、1シートの破損で全シートを失わない。
"""
from __future__ import annotations

import zipfile
from pathlib import Path

from . import ExtractResult, ExtractorContext, UnitCollector

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except Exception:  # pragma: no cover
    openpyxl = None
    get_column_letter = None

OLE_MAGIC = b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'


def _cell_name(row_index: int, col_index: int) -> str:
    if get_column_letter is not None:
        return f"{get_column_letter(col_index)}{row_index}"
    return f"R{row_index}C{col_index}"


def _collect_sheet(ws, collector: UnitCollector, ctx: ExtractorContext) -> None:
    collector.append(f"# Sheet: {ws.title}")
    for row_index, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if ctx.stopped():
            break
        values = []
        for col_index, value in enumerate(row, start=1):
            if value is None:
                continue
            cell_text = str(value).strip()
            if not cell_text:
                continue
            values.append(f"{_cell_name(row_index, col_index)}={cell_text}")
        if values:
            collector.append(f"Row {row_index}: " + ' | '.join(values))


def read_excel_file(path: Path, ctx: ExtractorContext) -> ExtractResult:
    if openpyxl is None:
        return ExtractResult(status='error', reason='openpyxl が利用できません')

    # ブックを開けない場合は文書全体の失敗として扱う
    try:
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    except zipfile.BadZipFile:
        try:
            with open(path, 'rb') as f:
                head = f.read(8)
            if head == OLE_MAGIC:
                return ExtractResult(status='protected',
                                     reason='暗号化ブック（またはxls形式）の可能性')
        except Exception:
            pass
        return ExtractResult(status='error', reason='zip 形式として読めないブック')
    except Exception as e:
        return ExtractResult(status='error', reason=f"Excel解析エラー: {e}")

    collector = UnitCollector()
    sheet_count = 0
    try:
        for ws in wb.worksheets:
            if ctx.stopped():
                break
            sheet_count += 1
            with collector.unit(f"シート「{getattr(ws, 'title', '?')}」"):
                _collect_sheet(ws, collector, ctx)
    finally:
        try:
            wb.close()
        except Exception:
            pass

    result = collector.build(empty_reason='セルに値がありません',
                             meta={'sheets': sheet_count})

    # シート見出しだけで実データが無い場合は empty 扱い（原本互換）
    if result.status == 'ok':
        has_data = any(not line.startswith('# Sheet: ') for line in collector.lines)
        if not has_data:
            result.status = 'empty'
            result.reason = 'セルに値がありません'
    return result
