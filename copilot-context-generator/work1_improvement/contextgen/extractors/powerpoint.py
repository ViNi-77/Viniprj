"""PowerPoint (.pptx / .pptm) 抽出.

原本と同一の整形（# Slide N / ## Notes / ## XML Text、図形・表・グラフ・
代替テキストの重複排除つき収集）。
"""
from __future__ import annotations

from pathlib import Path

from . import ExtractResult, ExtractorContext, UnitCollector
from ..textutil import normalize_text

try:
    from pptx import Presentation
except Exception:  # pragma: no cover
    Presentation = None


def append_unique_text(lines, seen, text):
    text = normalize_text(text or '')
    if not text:
        return
    key = text.lower()
    if key in seen:
        return
    seen.add(key)
    lines.append(text)


def extract_ooxml_texts(element):
    texts = []
    try:
        for child in element.iter():
            tag = str(child.tag)
            if tag.endswith('}t') or tag.endswith('}v') or tag == 't' or tag == 'v':
                text = (child.text or '').strip()
                if not text:
                    continue
                texts.append(text)
    except Exception:
        pass
    return texts


def extract_shape_alt_texts(shape):
    texts = []
    try:
        for child in shape._element.iter():
            if not str(child.tag).endswith('}cNvPr'):
                continue
            for attr_name in ('descr', 'title'):
                text = (child.get(attr_name) or '').strip()
                if not text:
                    continue
                texts.append(text)
    except Exception:
        pass
    return texts


def collect_ppt_shape_text(shape, lines, seen):
    try:
        text = getattr(shape, 'text', None)
        if text:
            append_unique_text(lines, seen, text)
    except Exception:
        pass

    # v3: hasattr(shape, 'table') は使ってはいけない。
    # python-pptx の GraphicFrame.table は「表でない場合 ValueError」を投げるが、
    # Python3 の hasattr が吸収するのは AttributeError のみ。そのため
    # グラフや埋め込みExcel（どちらも GraphicFrame）が1つあるだけで
    # 例外が呼び出し元まで伝播し、ファイル全体の抽出が失敗していた
    # （原本exeから継承していた不具合。実運用のレビューで発覚）。
    try:
        if getattr(shape, 'has_table', False):
            table = shape.table
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if not cells:
                    continue
                append_unique_text(lines, seen, ' | '.join(cells))
    except Exception:
        pass

    try:
        if getattr(shape, 'has_chart', False):
            chart = shape.chart

            try:
                append_unique_text(lines, seen, chart.chart_title.text_frame.text)
            except Exception:
                pass

            for axis_name in ('category_axis', 'value_axis'):
                try:
                    axis = getattr(chart, axis_name)
                    append_unique_text(lines, seen, axis.axis_title.text_frame.text)
                except Exception:
                    continue
    except Exception:
        pass

    for alt_text in extract_shape_alt_texts(shape):
        append_unique_text(lines, seen, alt_text)

    try:
        for child_shape in shape.shapes:
            collect_ppt_shape_text(child_shape, lines, seen)
    except Exception:
        pass


def _collect_slide(slide, slide_index: int, collector: UnitCollector, ctx) -> None:
    collector.append(f"# Slide {slide_index}")
    seen = set()

    for shape in slide.shapes:
        if ctx.stopped():
            break
        # 1つの図形の失敗でスライド全体を落とさない
        try:
            collect_ppt_shape_text(shape, collector.lines, seen)
        except Exception:
            continue

    try:
        notes_text = slide.notes_slide.notes_text_frame.text.strip()
        if notes_text:
            collector.append('## Notes')
            append_unique_text(collector.lines, seen, notes_text)
    except Exception:
        pass

    xml_texts = []
    for xml_text in extract_ooxml_texts(slide._element):
        normalized = normalize_text(xml_text)
        if normalized and normalized.lower() not in seen:
            xml_texts.append(normalized)

    if xml_texts:
        collector.append('## XML Text')
        for xml_text in xml_texts:
            append_unique_text(collector.lines, seen, xml_text)


def read_pptx_file(path: Path, ctx: ExtractorContext) -> ExtractResult:
    if Presentation is None:
        return ExtractResult(status='error', reason='python-pptx が利用できません')

    # ファイル自体を開けない場合は全体の失敗として扱う
    try:
        prs = Presentation(str(path))
        slides = list(prs.slides)
    except Exception as e:
        reason = str(e)
        if 'password' in reason.lower() or 'encrypted' in reason.lower():
            return ExtractResult(status='protected', reason=reason)
        return ExtractResult(status='error', reason=f"PowerPoint解析エラー: {reason}")

    # v3 柱1: スライド単位で保護する
    collector = UnitCollector()
    slide_count = 0

    for slide_index, slide in enumerate(slides, start=1):
        if ctx.stopped():
            break
        slide_count += 1
        with collector.unit(f"スライド{slide_index}"):
            _collect_slide(slide, slide_index, collector, ctx)

    meta = {'slides': slide_count}

    with collector.unit('埋め込みオブジェクト'):
        # 貼り付けられた埋め込みExcel等（python-pptx では本文にならない）
        from .embedded import extract_embedded_documents
        embedded_text, embedded_meta = extract_embedded_documents(path, ctx)
        if embedded_text:
            collector.append('# 埋め込みオブジェクト')
            collector.append(embedded_text)
            meta.update(embedded_meta)

    result = collector.build(empty_reason='テキストのないスライドのみ', meta=meta)

    # スライド見出しだけで中身が無い場合は empty 扱い（原本互換）
    if result.status == 'ok':
        has_content = any(
            not line.startswith('# Slide ') and line.strip() for line in collector.lines
        )
        if not has_content:
            result.status = 'empty'
            result.reason = 'テキストのないスライドのみ'
    return result
