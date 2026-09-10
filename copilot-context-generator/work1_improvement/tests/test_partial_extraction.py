"""v3 柱1: 部分的失敗の局所化（B3）の回帰テスト.

v2 では抽出器全体を1つの try/except で包んでいたため、
「1ページ／1シート／1エントリの失敗＝その文書が丸ごと0文字」だった。
Excel / PDF / Word / ZIP のすべてで全滅を実測で確認している。
無事な部分が必ず残ることを形式ごとに保証する。
"""
from __future__ import annotations

import threading
import zipfile

import pytest

from contextgen.config import RunConfig
from contextgen.events import NullEmitter
from contextgen.extractors import ExtractorContext, extract
from contextgen.worker import run_all
from conftest import make_text_pdf


@pytest.fixture
def ctx(tmp_path):
    config = RunConfig(search_root=tmp_path, context_output_dir=tmp_path / 'out',
                       base_dir=tmp_path / 'base')
    return ExtractorContext(config=config, stop_event=threading.Event())


# ---- Excel ----

def test_excel_broken_sheet_keeps_other_sheets(tmp_path, ctx, monkeypatch):
    import openpyxl

    import contextgen.extractors.excel as excel_mod

    wb = openpyxl.Workbook()
    for name, value in [('シートA', '値A'), ('壊れシート', '値B'), ('シートC', '値C')]:
        ws = wb.create_sheet(name)
        ws['A1'] = value
    wb.remove(wb['Sheet'])
    path = tmp_path / 'multi.xlsx'
    wb.save(str(path))

    real_load = excel_mod.openpyxl.load_workbook

    def broken_load(*args, **kwargs):
        book = real_load(*args, **kwargs)
        for ws in book.worksheets:
            if ws.title == '壊れシート':
                def boom(*a, **k):
                    raise RuntimeError('シート破損')
                ws.iter_rows = boom
        return book

    monkeypatch.setattr(excel_mod.openpyxl, 'load_workbook', broken_load)

    result = extract(path, ctx)
    assert result.status == 'partial'
    assert '値A' in result.text and '値C' in result.text
    assert result.units_ok == 2 and result.units_total == 3


# ---- PDF ----

def test_pdf_broken_page_keeps_other_pages(tmp_path, ctx, monkeypatch):
    import contextgen.extractors.pdfdoc as pdf_mod

    path = tmp_path / 'multi.pdf'
    path.write_bytes(make_text_pdf('PAGE-ONE-TEXT'))

    real_reader = pdf_mod.PdfReader

    class BrokenReader(real_reader):
        @property
        def pages(self):
            pages = list(super().pages)

            class BrokenPage:
                def extract_text(self):
                    raise RuntimeError('ページ破損')

            return pages + [BrokenPage()]

    monkeypatch.setattr(pdf_mod, 'PdfReader', BrokenReader)

    result = extract(path, ctx)
    assert result.status == 'partial'
    assert 'PAGE-ONE-TEXT' in result.text
    assert result.units_ok >= 1


# ---- Word ----

def test_docx_broken_table_keeps_body(tmp_path, ctx, monkeypatch):
    import docx as docx_lib

    import contextgen.extractors.worddoc as word_mod

    document = docx_lib.Document()
    document.add_paragraph('本文パラグラフ生存確認')
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = 'セル値'
    path = tmp_path / 'doc.docx'
    document.save(str(path))

    real_document = word_mod.docx.Document

    def broken_document(*args, **kwargs):
        doc = real_document(*args, **kwargs)

        class BrokenTables:
            def __iter__(self):
                raise RuntimeError('表破損')

        type(doc).tables = property(lambda self: BrokenTables())
        return doc

    monkeypatch.setattr(word_mod.docx, 'Document', broken_document)

    result = extract(path, ctx)
    assert result.status == 'partial'
    assert '本文パラグラフ生存確認' in result.text


# ---- ZIP ----

def test_zip_broken_entry_keeps_other_entries(tmp_path, ctx, monkeypatch):
    import contextgen.extractors.archive as archive_mod

    path = tmp_path / 'a.zip'
    with zipfile.ZipFile(path, 'w') as zf:
        zf.writestr('one.txt', 'ZIP内テキスト1')
        zf.writestr('two.txt', 'ZIP内テキスト2')

    real_copy = archive_mod.shutil.copyfileobj
    calls = {'n': 0}

    def broken_copy(src, dst):
        calls['n'] += 1
        if calls['n'] == 1:
            raise RuntimeError('展開失敗')
        return real_copy(src, dst)

    monkeypatch.setattr(archive_mod.shutil, 'copyfileobj', broken_copy)

    result = extract(path, ctx)
    assert result.status == 'partial'
    assert 'ZIP内テキスト2' in result.text


# ---- カバレッジの記録 ----

def test_coverage_survives_cache_roundtrip(tmp_path):
    """キャッシュヒットしたファイルでも構成単位の情報が失われないこと."""
    from contextgen.cache import ExtractCache
    from contextgen.extractors import ExtractResult

    db = tmp_path / 'cache.db'
    result = ExtractResult(text='本文', status='partial', reason='2単位中1単位',
                           units_total=2, units_ok=1, unit_failures=[('p.2', '破損')])

    with ExtractCache(db, lambda: '2026-01-01 00:00:00') as cache:
        cache.begin_run()
        cache.put(tmp_path / 'a.pdf', 10, 1.0, result)

    with ExtractCache(db, lambda: '2026-01-01 00:00:00') as cache:
        cache.begin_run()
        restored = cache.get(tmp_path / 'a.pdf', 10, 1.0)

    assert restored is not None
    assert restored.units_total == 2 and restored.units_ok == 1
    assert restored.unit_failures == [('p.2', '破損')]


def test_report_shows_partial_and_coverage(sample_tree, out_dirs, monkeypatch):
    """レポートに部分抽出とカバレッジが出ること."""
    import contextgen.extractors.excel as excel_mod

    real_load = excel_mod.openpyxl.load_workbook

    def broken_load(*args, **kwargs):
        book = real_load(*args, **kwargs)
        for ws in book.worksheets:
            def boom(*a, **k):
                raise RuntimeError('シート破損')
            ws.iter_rows = boom
        return book

    monkeypatch.setattr(excel_mod.openpyxl, 'load_workbook', broken_load)

    out, base = out_dirs
    run_all(RunConfig(search_root=sample_tree, context_output_dir=out, base_dir=base),
            NullEmitter(), threading.Event())

    report = (out / '_AI_CONTEXT_REPORT.md').read_text(encoding='utf-8')
    assert '## 抽出カバレッジ' in report
    assert '構成単位' in report
