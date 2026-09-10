"""テストデータ生成."""
from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest


def make_text_pdf(text: str) -> bytes:
    """テキスト1行入りの最小PDFを正しいxref付きで組み立てる."""
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode('latin-1')
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R"
        b" /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode())
        out.write(obj)
        out.write(b"\nendobj\n")
    xref_pos = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF\n".encode()
    )
    return out.getvalue()


def make_imageonly_pdf() -> bytes:
    """テキスト描画のない1ページPDF（OCR候補）."""
    content = b""
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R >>",
        b"<< /Length 0 >>\nstream\n" + content + b"\nendstream",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode())
        out.write(obj)
        out.write(b"\nendobj\n")
    xref_pos = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF\n".encode()
    )
    return out.getvalue()


def make_docx(path: Path, paragraphs: list[str], table: list[list[str]] | None = None):
    import docx
    document = docx.Document()
    for p in paragraphs:
        document.add_paragraph(p)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    document.save(str(path))


def make_xlsx(path: Path, sheets: dict[str, list[list]]):
    import openpyxl
    wb = openpyxl.Workbook()
    default = wb.active
    first = True
    for name, rows in sheets.items():
        ws = default if first else wb.create_sheet()
        ws.title = name
        first = False
        for row in rows:
            ws.append(row)
    wb.save(str(path))


def make_pptx(path: Path, slides: list[dict]):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    layout = prs.slide_layouts[1]  # Title and Content
    for spec in slides:
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = spec.get('title', '')
        body = spec.get('body')
        if body:
            slide.placeholders[1].text = body
        notes = spec.get('notes')
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    prs.save(str(path))


def make_sample_tree(root: Path) -> Path:
    """検証用の参照元フォルダを生成する."""
    src = root / 'source'
    (src / '設備').mkdir(parents=True)
    (src / '品質').mkdir(parents=True)
    (src / '管理').mkdir(parents=True)  # 除外対象フォルダ

    (src / 'メモ.txt').write_text(
        '生産ラインAの立ち上げメモ。\n担当: 田中\n進捗は順調。\n', encoding='utf-8'
    )
    (src / '設備' / '旧形式レポート.txt').write_text(
        'これはシフトJISで保存された設備点検記録です。異常なし。',
        encoding='cp932',
    )
    make_docx(
        src / '設備' / '点検手順書.docx',
        ['# 設備点検手順', '毎朝の始業前点検を実施すること。', '異常があれば保全へ連絡。'],
        table=[['項目', '基準'], ['油圧', '5MPa以上']],
    )
    make_xlsx(
        src / '品質' / '測定データ.xlsx',
        {'測定': [['品番', '寸法', '判定'], ['A-100', 25.01, 'OK'], ['A-101', 25.30, 'NG']]},
    )
    make_pptx(
        src / '品質' / '品質会議.pptx',
        [{'title': '月次品質報告', 'body': '不良率 0.3%', 'notes': '来月は目標0.2%'}],
    )
    (src / '設備' / '仕様書.pdf').write_bytes(make_text_pdf('Spec sheet for machine X-200'))
    (src / '品質' / 'スキャン報告書.pdf').write_bytes(make_imageonly_pdf())

    # ZIP（txt を内包）
    with zipfile.ZipFile(src / '過去資料.zip', 'w') as zf:
        zf.writestr('old/議事録2024.txt', '2024年の定例議事録。テーマ: 歩留まり改善。')

    # ロックファイル・対象外・除外フォルダ内
    (src / '設備' / '~$点検手順書.docx').write_bytes(b'\x00\x01lock')
    (src / '品質' / '写真.png').write_bytes(b'\x89PNG\r\n')
    (src / '管理' / '内部メモ.txt').write_text('これは収録されないはず', encoding='utf-8')

    # 中身が Word バイナリでない .doc（要変換扱いの確認用）
    (src / '古い報告書.doc').write_bytes(b'\xd0\xcf\x11\xe0old word file dummy')

    return src


@pytest.fixture
def sample_tree(tmp_path):
    return make_sample_tree(tmp_path)


@pytest.fixture
def out_dirs(tmp_path):
    out = tmp_path / 'out'
    base = tmp_path / 'base'
    out.mkdir()
    base.mkdir()
    return out, base
