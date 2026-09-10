"""v3 柱2: 人間向け読解キットのテスト.

「長いマニュアルを人間が読み解く」ための機能。LLM を使わずに
目次・鮮度・重複・参照切れを取り出せることを保証する。
"""
from __future__ import annotations

import threading

from contextgen.config import RunConfig
from contextgen.digest import (
    analyze,
    build_questions,
    extract_date_mentions,
    extract_headings,
    find_broken_references,
    find_duplicate_paragraphs,
)
from contextgen.events import NullEmitter
from contextgen.worker import run_all

MANUAL = """設備点検マニュアル

2019年4月版　Rev.3

第1章 目的と適用範囲
本マニュアルは設備点検の手順を定める。
1.1 適用設備
適用対象は加工機X-200とする。
1.2 適用範囲
別紙2参照のこと。

第2章 日常点検の手順
2.1 油圧の確認
油圧が5MPa以上であることを確認する。
3.2.1 詳細手順
細目はここに記す。

別表1 判定基準一覧
寸法公差は±0.05mmとする。
別表1による判定を行うこと。
"""


# ---- 目次 ----

def test_numbered_headings_keep_their_number():
    """「1.1」が「1.」にマッチして番号が壊れないこと（回帰）."""
    headings = extract_headings(MANUAL)
    titles = [h.title for h in headings]
    assert '1.1 適用設備' in titles
    assert '1.2 適用範囲' in titles
    assert '3.2.1 詳細手順' in titles
    assert '1 1 適用設備' not in titles


def test_heading_levels_are_nested():
    headings = {h.title: h.level for h in extract_headings(MANUAL)}
    assert headings['1 目的と適用範囲'] == 1
    assert headings['1.1 適用設備'] == 2
    assert headings['3.2.1 詳細手順'] == 3


def test_structure_markers_become_headings():
    """抽出器が付けた # Page / # Sheet: も目次になること."""
    headings = extract_headings('# Page 1\n本文\n## Notes\nメモ')
    titles = [h.title for h in headings]
    assert 'Page 1' in titles
    assert 'Notes' in titles


def test_body_sentences_are_not_headings():
    headings = extract_headings('これは通常の本文であり見出しではありません。\n')
    assert headings == []


# ---- 鮮度 ----

def test_date_mentions_and_newest_year():
    mentions, newest = extract_date_mentions(MANUAL)
    assert any('2019年4月' in m for m in mentions)
    assert newest == 2019


def test_wareki_is_converted():
    _, newest = extract_date_mentions('令和3年5月に改訂')
    assert newest == 2021


def test_no_date_returns_none():
    mentions, newest = extract_date_mentions('日付の記載がない文書')
    assert mentions == [] and newest is None


# ---- 重複 ----

def test_duplicate_paragraphs_detected():
    boiler = '本作業は関係部署と事前に調整のうえ実施し、実施記録を所定の様式に必ず残すこと。'
    text = f"章1\n{boiler}\n章2\n{boiler}\n章3\n{boiler}\n"
    duplicates = find_duplicate_paragraphs(text)
    assert duplicates and duplicates[0][0] == boiler
    assert duplicates[0][1] == 3


def test_short_lines_are_not_duplicates():
    assert find_duplicate_paragraphs('はい\nはい\nはい\n') == []


# ---- 参照切れ ----

def test_broken_reference_is_detected():
    """参照しているのに定義が無いものだけを挙げること."""
    broken = find_broken_references(MANUAL)
    assert '別紙2' in broken, '参照切れを検出できていない'
    assert '別表1' not in broken, '定義がある資料を誤って参照切れにしている'


def test_reference_line_is_not_mistaken_for_definition():
    """「別紙2参照のこと」を定義側と誤認しないこと（回帰）."""
    assert find_broken_references('別紙2参照のこと。') == ['別紙2']


# ---- 質問 ----

def test_questions_strip_numbers_and_skip_markers():
    digest = analyze(MANUAL)
    questions = [q for q, _ in build_questions(digest)]
    assert any('適用設備' in q for q in questions)
    assert not any(q.startswith('1.1') for q in questions)

    marker_digest = analyze('# Page 1\n本文\n')
    assert build_questions(marker_digest) == []


# ---- 出力 ----

def test_single_file_digest_via_cli(tmp_path, capsys):
    from contextgen.cli import main

    source = tmp_path / '設備点検マニュアル.txt'
    source.write_text(MANUAL, encoding='utf-8')
    out = tmp_path / 'out'

    code = main(['--digest-file', str(source), '--output', str(out)])
    assert code == 0

    digests = list((out / '_HUMAN_DIGEST').glob('*.md'))
    assert len(digests) == 1

    content = digests[0].read_text(encoding='utf-8')
    assert '## 30秒でわかる' in content
    assert '## 目次（自動抽出）' in content
    assert '1.1 適用設備' in content
    assert '2019' in content
    assert 'Copilotに聞くときのコピペ用' in content
    assert '別紙2' in content

    captured = capsys.readouterr().out
    assert '読み解きキットを作成しました' in captured


def test_digest_file_requires_existing_file(tmp_path, capsys):
    from contextgen.cli import main

    code = main(['--digest-file', str(tmp_path / 'missing.txt'), '--output', str(tmp_path / 'o')])
    assert code == 2
    assert 'ファイルが見つかりません' in capsys.readouterr().err


def test_source_is_required_without_digest_file(tmp_path, capsys):
    from contextgen.cli import main

    code = main(['--output', str(tmp_path / 'o')])
    assert code == 2
    assert '--source' in capsys.readouterr().err


def test_folder_digest_creates_index(sample_tree, out_dirs):
    out, base = out_dirs
    (sample_tree / '手順書.txt').write_text(MANUAL, encoding='utf-8')

    stats = run_all(
        RunConfig(search_root=sample_tree, context_output_dir=out, base_dir=base,
                  human_digest=True),
        NullEmitter(), threading.Event(),
    )

    assert stats.digest_paths
    index = out / '_HUMAN_DIGEST_INDEX.md'
    assert index.exists()

    content = index.read_text(encoding='utf-8')
    assert '読み解きキット一覧' in content
    assert '手順書' in content
    assert (out / '_HUMAN_DIGEST').is_dir()


def test_digest_off_by_default(sample_tree, out_dirs):
    out, base = out_dirs
    run_all(RunConfig(search_root=sample_tree, context_output_dir=out, base_dir=base),
            NullEmitter(), threading.Event())
    assert not (out / '_HUMAN_DIGEST').exists()


def test_docx_headings_become_markers(tmp_path):
    """Word の見出しスタイルが構造マーカーとして残ること."""
    import docx as docx_lib

    from contextgen.extractors import ExtractorContext, extract

    document = docx_lib.Document()
    document.add_heading('第1章 目的', level=1)
    document.add_paragraph('本文です。')
    document.add_heading('1.1 適用範囲', level=2)
    path = tmp_path / 'styled.docx'
    document.save(str(path))

    config = RunConfig(search_root=tmp_path, context_output_dir=tmp_path / 'out')
    ctx = ExtractorContext(config=config, stop_event=threading.Event())
    result = extract(path, ctx)

    assert '# 第1章 目的' in result.text
    assert '## 1.1 適用範囲' in result.text
