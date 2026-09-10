"""v3: 埋め込みオブジェクトの抽出と、図形1つでファイル全体が落ちない保証.

実運用のレビューで「週報PPTXに貼った埋め込みExcelが拾えない」
「そのファイルだけ本文が丸ごと取れない」という報告を受けて追加。
"""
from __future__ import annotations

import shutil
import threading
import zipfile
from pathlib import Path

import pytest

from contextgen.config import RunConfig
from contextgen.extractors import ExtractorContext, extract


@pytest.fixture
def ctx(tmp_path):
    config = RunConfig(search_root=tmp_path, context_output_dir=tmp_path / 'out',
                       base_dir=tmp_path / 'base')
    return ExtractorContext(config=config, stop_event=threading.Event())


def make_pptx_with_chart(path: Path):
    """グラフ入りPPTX。python-pptx がグラフ用に本物の埋め込みExcelを同梱するため、
    「埋め込みExcelを含むPPTX」の実物テストデータになる。"""
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = '週報 2026年6月'
    slide.placeholders[1].text = '設備稼働率は95%を達成しました'

    data = CategoryChartData()
    data.categories = ['4月', '5月']
    data.add_series('稼働率', (90, 95))
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(3), Inches(4), Inches(3), data
    )
    prs.save(str(path))
    return path


def embed_file_into(container: Path, embed_source: Path, inner_name: str, prefix: str):
    """既存のOffice文書zipに、埋め込みファイルを後から差し込む."""
    tmp = container.with_suffix('.tmp')
    with zipfile.ZipFile(container) as src, zipfile.ZipFile(tmp, 'w') as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item.filename))
        dst.write(embed_source, f"{prefix}/embeddings/{inner_name}")
    shutil.move(str(tmp), str(container))


def make_xlsx(path: Path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '実績'
    ws.append(['項目', '目標', '実績'])
    ws.append(['稼働率', '90%', '95%'])
    wb.save(str(path))
    return path


# --- 回帰: 図形1つでファイル全体が落ちない（原本exe由来の不具合） ---

def test_graphicframe_does_not_kill_whole_file(tmp_path, ctx):
    """グラフ/埋め込みExcel（GraphicFrame）があっても本文が失われないこと。

    修正前は hasattr(shape, 'table') が ValueError を投げ、
    タイトルも本文も含めて 0 文字になっていた。
    """
    pptx = make_pptx_with_chart(tmp_path / 'weekly_report.pptx')
    result = extract(pptx, ctx)

    assert result.status == 'ok', f"status={result.status} reason={result.reason}"
    assert '週報 2026年6月' in result.text
    assert '設備稼働率は95%を達成しました' in result.text


def test_broken_shape_is_skipped_not_fatal(tmp_path, ctx, monkeypatch):
    """想定外の壊れ方をする図形があっても、他の本文は残ること."""
    from contextgen.extractors import powerpoint

    original = powerpoint.collect_ppt_shape_text
    state = {'first': True}

    def flaky(shape, lines, seen):
        if state['first']:
            state['first'] = False
            raise RuntimeError('壊れた図形')
        return original(shape, lines, seen)

    monkeypatch.setattr(powerpoint, 'collect_ppt_shape_text', flaky)

    pptx = make_pptx_with_chart(tmp_path / 'flaky.pptx')
    result = extract(pptx, ctx)

    assert result.status == 'ok'
    assert result.text.strip()


# --- 新機能: 埋め込みファイルの本文抽出 ---

def test_pptx_embedded_excel_is_extracted(tmp_path, ctx):
    pptx = make_pptx_with_chart(tmp_path / 'with_embed.pptx')
    xlsx = make_xlsx(tmp_path / 'source.xlsx')
    embed_file_into(pptx, xlsx, '貼り付け実績表.xlsx', 'ppt')

    result = extract(pptx, ctx)

    assert result.status == 'ok'
    assert '# 埋め込みオブジェクト' in result.text
    assert '貼り付け実績表.xlsx' in result.text
    # 埋め込みExcelのセルの中身が本文として拾えていること
    assert '稼働率' in result.text
    assert '95%' in result.text
    assert result.meta.get('embedded_extracted', 0) >= 1


def test_docx_embedded_excel_is_extracted(tmp_path, ctx):
    import docx as docx_mod

    document = docx_mod.Document()
    document.add_paragraph('本文の説明です。')
    docx_path = tmp_path / 'manual.docx'
    document.save(str(docx_path))

    xlsx = make_xlsx(tmp_path / 'source2.xlsx')
    embed_file_into(docx_path, xlsx, '仕様一覧.xlsx', 'word')

    result = extract(docx_path, ctx)

    assert result.status == 'ok'
    assert '本文の説明です。' in result.text
    assert '仕様一覧.xlsx' in result.text
    assert '稼働率' in result.text


def test_legacy_ole_bin_is_reported_not_silent(tmp_path, ctx):
    """旧形式(.bin)は抽出できないが、存在は記録されること."""
    pptx = make_pptx_with_chart(tmp_path / 'legacy.pptx')
    ole = tmp_path / 'oleObject1.bin'
    ole.write_bytes(b'\xd0\xcf\x11\xe0' + b'dummy ole container')
    embed_file_into(pptx, ole, 'oleObject1.bin', 'ppt')

    result = extract(pptx, ctx)

    assert result.status == 'ok'
    assert '抽出できなかった埋め込みファイル' in result.text
    assert 'oleObject1.bin' in result.text
    assert result.meta.get('embedded_unreadable', 0) >= 1


def test_no_embeddings_keeps_output_unchanged(tmp_path, ctx):
    """埋め込みが無いファイルの出力に余計な見出しが増えないこと（既存互換）."""
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = 'タイトル'
    slide.placeholders[1].text = '本文'
    path = tmp_path / 'plain.pptx'
    prs.save(str(path))

    result = extract(path, ctx)

    assert result.status == 'ok'
    assert '埋め込みオブジェクト' not in result.text
